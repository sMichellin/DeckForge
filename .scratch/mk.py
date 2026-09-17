import sys
sys.path.insert(0,'.')
from pathlib import Path
from pptx import Presentation
prs = Presentation()
for _ in range(2):
    prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = "Пример из шаблона"
prs.save(".scratch/t.pptx")
print("ok")
