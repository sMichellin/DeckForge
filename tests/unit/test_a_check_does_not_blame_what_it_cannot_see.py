"""Проверка не обвиняет в том, чего не видит. Таски RG49 и RG50,
change `a-check-does-not-blame-what-it-cannot-see`.

После RG43/RG44/RG48 на трёх колодах осталось пять ошибок трёх видов, и замер каждой
показал, что **две из трёх — наши собственные проверки, обвиняющие нас зря**.

**Картинка «растянута».** У фигур Education s06 и s09 есть кадрирование
`l=14949 r=54601 t=7685 b=46639`: видно 0,3045 ширины и 0,45676 высоты, пропорция
видимой части 1,5 × 0,3045 / 0,45676 = 1,000 — ровно квадратная рамка. Кадр при этом
авторский: фигура скопирована из слайда-примера. Проверка сравнивала рамку с полным
исходником.

**«Q1 2026».** Разбор чисел увидел в номере квартала число 1 и потребовал подтвердить
его материалами. Годы, версии и даты он уже пропускал — квартал той же природы.

**«30 дней».** А вот это настоящий брак: срока в материалах нет, модель его придумала
в строке следующего шага. Правило 5а промпта композитора.

Сценарии — из дельт `openspec/changes/a-check-does-not-blame-what-it-cannot-see/`.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
import yaml
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

from deckforge.audit.deterministic.layout import _visible
from deckforge.parsing.content import extract_numbers
from deckforge.registry import get_prompt_registry

#: Числа кадра фигуры Education s06 и s09, символ в символ из файла колоды.
CROP = {"l": "14949", "r": "54601", "t": "7685", "b": "46639"}

#: Исходник 3:2 — та же пропорция, что у картинки шаблона.
NATIVE_CX, NATIVE_CY = 1500, 1000


def picture(tmp_path: Path, *, crop: dict[str, str] | None) -> Any:
    """Квадратная рамка с картинкой 3:2, с кадром или без."""
    from PIL import Image

    source = tmp_path / "picture.png"
    Image.new("RGB", (NATIVE_CX, NATIVE_CY), "white").save(source)
    prs = Presentation()
    slide = prs.slides.add_slide(next(la for la in prs.slide_layouts if la.name == "Blank"))
    shape = slide.shapes.add_picture(
        str(source), Emu(0), Emu(0), width=Emu(1_622_854), height=Emu(1_622_854)
    )
    if crop is not None:
        fill = shape._element.find(qn("p:blipFill"))
        rect = fill.makeelement(qn("a:srcRect"), crop)
        fill.insert(1, rect)
    return shape


# --- RG49: кадрированная картинка не растянута ------------------------------------


def test_a_cropped_picture_is_not_stretched(tmp_path: Path) -> None:
    """Нарушитель проверки: кадр делает видимую часть квадратной, как рамка."""
    shape = picture(tmp_path, crop=CROP)

    visible_cx, visible_cy = _visible(shape, NATIVE_CX, NATIVE_CY)

    assert abs(visible_cx / visible_cy - 1.0) < 0.01, "видимая часть не квадратная"


def test_an_uncropped_picture_keeps_its_own_aspect(tmp_path: Path) -> None:
    """Норма: кадра нет — видно весь файл, и пропорция прежняя 3:2."""
    shape = picture(tmp_path, crop=None)

    visible_cx, visible_cy = _visible(shape, NATIVE_CX, NATIVE_CY)

    assert (visible_cx, visible_cy) == (float(NATIVE_CX), float(NATIVE_CY))


# --- RG50: номер квартала не величина ---------------------------------------------


@pytest.mark.parametrize(
    "text", ["выделение команды на Q1 2026", "в 1 кв. 2026", "IV квартал 2026", "к1 2026"]
)
def test_a_quarter_is_not_a_number(text: str) -> None:
    """Нарушитель: номер квартала уезжал в фактчекинг как число."""
    assert extract_numbers(text) == [], f"в {text!r} нашлось число"


@pytest.mark.parametrize(
    ("text", "raw"), [("в течение 30 дней", "30"), ("рост 37 % за год", "37 %")]
)
def test_a_real_quantity_is_still_a_number(text: str, raw: str) -> None:
    """Норма: величина остаётся величиной — иначе ловушка съела бы проверку."""
    assert [number.raw for number in extract_numbers(text)] == [raw]


# --- RG50: закрывающий слайд не выдумывает сроков ---------------------------------


def test_the_composer_forbids_inventing_a_deadline() -> None:
    """Активная версия промпта называет соблазн строки следующего шага прямо."""
    active = yaml.safe_load(Path("prompts/registry.yaml").read_text())["skills"][
        "slide_composer"
    ]["active"]
    text = (Path("prompts/slide_composer") / active / "system.j2").read_text()

    assert "закрывающем слайде" in text
    assert "30 дней" in text, "правило не называет пример выдуманного срока"


def test_the_composer_response_schema_did_not_change() -> None:
    """Контракт ответа прежний: меняется просьба, а не схема."""
    bundle = get_prompt_registry().load("slide_composer")
    old = (Path("prompts/slide_composer/1.4.0/schema.json")).read_text()
    new = (Path("prompts/slide_composer") / bundle.version / "schema.json").read_text()

    assert old == new
