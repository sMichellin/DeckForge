"""Слайд по рецепту не несёт текста шаблона. Change `recipe-leaves-no-sample-text`.

Каталог честно делает зонами фигуры внутри групп (`p:grpSp`): подписи карточек, строки
таймлайна, имена в оргструктуре. Writer же искал зону только среди фигур верхнего уровня —
и молча пропускал её: наш текст терялся, а на слайде оставался текст автора шаблона.

Эталон «нашего» текста — `SlideIR`, а не файл шаблона: текст примера в тестах не зашит.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

from deckforge.designsystem import derive
from deckforge.designsystem.models import Recipe, RecipeKind, TypeLevel, Zone
from deckforge.domain.enums import TextRole
from deckforge.domain.slide import DeckIR, SlideIR, TextBlock
from deckforge.parsing import TemplateParser
from deckforge.rendering.recipe_slide import RecipeError, clone_recipe
from deckforge.rendering.writer import PptxWriter
from tests.case_templates import case_template

#: Колонтитулы и поля (номер слайда, дата) ставит макет, а не пример: это не текст автора.
FOOTERS = {"ftr", "dt", "sldNum", "hdr"}

#: Свой список, а не адресуемые фигуры writer'а: тест видит на слайде всё, у чего есть
#: `cNvPr id`, включая сами группы, — иначе он не заметил бы группу, оставшуюся от повтора.
SHAPES = (qn("p:sp"), qn("p:grpSp"), qn("p:pic"), qn("p:cxnSp"), qn("p:graphicFrame"))


def slide_ir(recipe: Recipe, zone_ids: list[str]) -> SlideIR:
    return SlideIR(
        slide_id="s01",
        layout_id="l01",
        variant="A",
        recipe_id=recipe.recipe_id,
        blocks=[
            TextBlock(
                block_id=f"b{n}", role=TextRole.BODY, text=f"Наш текст {zone_id}", zone_id=zone_id
            )
            for n, zone_id in enumerate(zone_ids)
        ],
    )


def ours(ir: SlideIR) -> set[str]:
    return {
        line
        for block in ir.blocks
        if isinstance(block, TextBlock)
        for line in block.text.splitlines()
    }


def paragraphs(slide) -> list[str]:
    """Непустые абзацы слайда на любой глубине групп, кроме колонтитулов и полей."""
    found: list[str] = []
    for shape in slide.shapes._spTree.iter(qn("p:sp"), qn("p:graphicFrame")):
        ph = shape.find(f"./*/{qn('p:nvPr')}/{qn('p:ph')}")
        if ph is not None and ph.get("type") in FOOTERS:
            continue
        for paragraph in shape.iter(qn("a:p")):
            if paragraph.find(qn("a:fld")) is not None:
                continue
            text = "".join(t.text or "" for t in paragraph.iter(qn("a:t"))).strip()
            if text:
                found.append(text)
    return found


def ids_on(slide) -> set[int]:
    ids: set[int] = set()
    for node in slide.shapes._spTree.iter(*SHAPES):
        props = node.find(f"./*/{qn('p:cNvPr')}")
        if props is not None:
            ids.add(int(props.get("id")))
    return ids


# --- синтетика: пример с группой ------------------------------------------------------


def grouped_example():
    """Пример из двух карточек-групп: в каждой плашка-подпись и текст."""
    prs = Presentation()
    source = prs.slides.add_slide(prs.slide_layouts[6])
    cards = []
    for n in range(2):
        group = source.shapes.add_group_shape()
        caption = group.shapes.add_textbox(Emu(100 + n * 3000), Emu(100), Emu(2000), Emu(500))
        caption.text_frame.text = f"Подпись примера {n}"
        body = group.shapes.add_textbox(Emu(100 + n * 3000), Emu(700), Emu(2000), Emu(500))
        body.text_frame.text = f"Текст примера {n}"
        cards.append((group, caption, body))
    part_name = str(source.part.partname)
    return prs, part_name, cards


def test_a_zone_inside_a_group_gets_our_text() -> None:
    """Зона внутри группы получает наш текст; пустая зона внутри группы не несёт текста примера."""
    prs, part_name, cards = grouped_example()
    (_, caption, body), _ = cards
    recipe = Recipe(
        recipe_id="ex001",
        example_index=1,
        part_name=part_name,
        kind=RecipeKind.TEXT,
        zones=[
            Zone(zone_id="zc", xml_id=caption.shape_id, role=TypeLevel.CAPTION, capacity_chars=40),
            Zone(zone_id="zb", xml_id=body.shape_id, role=TypeLevel.CAPTION, capacity_chars=40),
        ],
    )
    ir = slide_ir(recipe, ["zc"])

    slide = clone_recipe(prs, recipe, ir)

    texts = paragraphs(slide)
    assert "Наш текст zc" in texts, "наш текст потерян: зону внутри группы не нашли"
    assert set(texts) <= ours(ir) | {"Подпись примера 1", "Текст примера 1"}
    assert "Текст примера 0" not in texts, "пустая зона внутри группы сохранила текст примера"


def test_a_spare_repeat_inside_a_group_is_dropped() -> None:
    """Лишний повтор, чьи фигуры лежат в группе, уходит целиком — и пустая группа с ним."""
    prs, part_name, cards = grouped_example()
    recipe = spare_repeat_recipe(part_name, cards)
    rows = recipe.repeat_xml_ids
    ir = slide_ir(recipe, ["z0"])

    slide = clone_recipe(prs, recipe, ir)

    spare_group = cards[1][0].shape_id
    assert not ids_on(slide) & {*rows[1], spare_group}, "лишний повтор остался на слайде"
    assert "Наш текст z0" in paragraphs(slide)


def spare_repeat_recipe(part_name: str, cards) -> Recipe:
    rows = [[caption.shape_id, body.shape_id] for _, caption, body in cards]
    return Recipe(
        recipe_id="ex001",
        example_index=1,
        part_name=part_name,
        kind=RecipeKind.CARDS,
        repeats=2,
        repeat_xml_ids=rows,
        zones=[
            Zone(
                zone_id=f"z{n}", xml_id=row[0], role=TypeLevel.CAPTION, repeat=n, capacity_chars=40
            )
            for n, row in enumerate(rows)
        ],
    )


def test_a_group_keeps_what_the_writer_does_not_address() -> None:
    """Группа уходит, только когда в ней не осталось ничего, кроме её собственных свойств:
    рукописный фрагмент или `mc:AlternateContent` рядом с повтором — не наш, его не трогаем."""
    prs, part_name, cards = grouped_example()
    group = cards[1][0]
    etree.SubElement(
        group._element,
        "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent",
    )

    recipe = spare_repeat_recipe(part_name, cards)

    slide = clone_recipe(prs, recipe, slide_ir(recipe, ["z0"]))

    assert group.shape_id in ids_on(slide), "группа с чужим содержимым удалена вместе с повтором"


def test_an_ambiguous_address_is_an_error() -> None:
    """Два `cNvPr id` в одном примере: какую фигуру имел в виду каталог, не знает никто.
    Писать наугад нельзя — чужая фраза останется молча."""
    prs, part_name, cards = grouped_example()
    (_, caption, _), (_, twin, _) = cards
    twin._element.find(f"./*/{qn('p:cNvPr')}").set("id", str(caption.shape_id))
    recipe = Recipe(
        recipe_id="ex001",
        example_index=1,
        part_name=part_name,
        kind=RecipeKind.TEXT,
        zones=[
            Zone(zone_id="zc", xml_id=caption.shape_id, role=TypeLevel.CAPTION, capacity_chars=40)
        ],
    )

    with pytest.raises(RecipeError, match=str(caption.shape_id)):
        clone_recipe(prs, recipe, slide_ir(recipe, ["zc"]))


# --- настоящие шаблоны --------------------------------------------------------------


def write_recipe(name: str, recipe_id: str, tmp_path: Path, zones: int | None = None):
    path = case_template(name)
    manifest = TemplateParser().parse(path, use_cache=False)
    ds = derive(manifest)
    recipe = next((r for r in ds.recipes if r.recipe_id == recipe_id), None)
    if recipe is None:
        pytest.fail(f"в каталоге шаблона «{name}» нет рецепта {recipe_id}")
    example = next(e for e in manifest.examples if e.slide_index == recipe.example_index)
    ir = slide_ir(recipe, [z.zone_id for z in recipe.zones][:zones]).model_copy(
        update={"layout_id": example.layout_id or manifest.layouts[0].layout_id}
    )
    deck = DeckIR(deck_id="d01", template_id=manifest.template_id, variant="A", seed=1, slides=[ir])
    out = PptxWriter(path, manifest, design_system=ds).write(deck, tmp_path / "deck.pptx")
    return recipe, ir, Presentation(str(out)).slides[0]


@pytest.mark.parametrize(
    ("name", "recipe_id"),
    [("VK Tech шаблон.pptx", "ex013"), ("Шаблон презентации VK Education.pptx", "ex043")],
)
def test_filled_zones_inside_groups_leave_no_sample_text(
    name: str, recipe_id: str, tmp_path: Path
) -> None:
    """Все зоны заняты — на слайде нет абзаца, которого нет в нашем IR."""
    _recipe, ir, slide = write_recipe(name, recipe_id, tmp_path)

    left = [text for text in paragraphs(slide) if text not in ours(ir)]

    assert not left, f"текст шаблона на слайде по рецепту {recipe_id}: {left[:5]}"
    assert ours(ir) <= set(paragraphs(slide)), "часть нашего текста не дошла до слайда"


def test_spare_repeats_inside_groups_are_dropped_on_a_real_template(tmp_path: Path) -> None:
    """WorkSpace: картинка и подпись повтора лежат в группе рядом с его плашкой."""
    recipe, _ir, slide = write_recipe(
        "VK_WorkSpace_Клиентская_конференция_Шаблон_03.pptx", "ex005", tmp_path, zones=2
    )
    assert recipe.repeats > 1
    used = {z.repeat for z in recipe.zones[:2] if z.repeat is not None}
    spare = {
        xml_id for n, row in enumerate(recipe.repeat_xml_ids) if n not in used for xml_id in row
    }

    assert spare and not ids_on(slide) & spare, "фигуры лишних повторов остались на слайде"
