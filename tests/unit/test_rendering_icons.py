"""Иконки Lucide: SVG → нативная геометрия с цветом темы. Change (21) `smartart-icons`."""

from __future__ import annotations

import math
from itertools import pairwise

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

from deckforge.domain.enums import ColorRef
from deckforge.domain.slide import IconBlock
from deckforge.domain.template import TemplateManifest
from deckforge.domain.units import EMU_PER_CM
from deckforge.rendering.icons import (
    ICON_VIEWBOX,
    Close,
    Cubic,
    Line,
    Move,
    add_icon,
    icon_names,
    icon_nodes,
    node_segments,
    path_segments,
)


def ends(segments: list[object]) -> list[tuple[float, float]]:
    return [(s.x, s.y) for s in segments if not isinstance(s, Close)]  # type: ignore[attr-defined]


# --- разбор путей -------------------------------------------------------------------


def test_absolute_commands_and_close() -> None:
    assert path_segments("M4 4h16v16H4z") == [
        Move(4, 4), Line(20, 4), Line(20, 20), Line(4, 20), Close(),
    ]


def test_relative_move_repeats_as_line_to() -> None:
    assert path_segments("m9 12 2 2 4-4") == [Move(9, 12), Line(11, 14), Line(15, 10)]


def test_numbers_glued_without_separators() -> None:
    assert ends(path_segments("M1.17.2-.67-.01")) == [(1.17, 0.2), (-0.67, -0.01)]


def test_close_returns_to_the_start_of_the_subpath() -> None:
    segments = path_segments("M2 2h4v4zm1 1h1")
    assert segments[-2] == Move(3, 3)
    assert segments[-1] == Line(4, 3)


def test_drawing_after_close_starts_again_from_the_subpath_start() -> None:
    """В SVG `…zV7` продолжает от начала контура; в DrawingML после `close` нужен `moveTo`,
    иначе PowerPoint не знает, откуда вести линию (так устроены `scale`, `mop`)."""
    assert path_segments("M2 2h4v4zV7") == [
        Move(2, 2), Line(6, 2), Line(6, 6), Close(), Move(2, 2), Line(2, 7),
    ]


def test_every_close_in_lucide_is_followed_by_a_move() -> None:
    for name in icon_names():
        for tag, attrs in icon_nodes(name) or []:
            segments = node_segments(tag, attrs)
            for current, following in pairwise(segments):
                if isinstance(current, Close):
                    assert isinstance(following, Move), name


def test_arc_becomes_cubics_on_the_circle() -> None:
    segments = path_segments("M2 12a10 10 0 0 1 20 0")
    assert all(isinstance(s, Cubic) for s in segments[1:])
    assert ends(segments)[-1] == pytest.approx((22, 12))
    for x, y in ends(segments):
        assert math.hypot(x - 12, y - 12) == pytest.approx(10)
    # Флаг sweep=1 — по часовой стрелке в координатах SVG, то есть через верх (y = 2).
    assert min(y for _, y in ends(segments)) == pytest.approx(2)


def test_arc_flags_may_be_glued_to_the_next_number() -> None:
    assert ends(path_segments("M2 12a10 10 0 0120 0"))[-1] == pytest.approx((22, 12))


def test_smooth_cubic_reflects_the_previous_control_point() -> None:
    segments = path_segments("M0 0C0 10 10 10 10 0S20-10 20 0")
    assert segments[2] == Cubic(10, -10, 20, -10, 20, 0)


def test_quadratic_is_raised_to_a_cubic() -> None:
    (cubic,) = path_segments("M0 0Q10 10 20 0")[1:]
    controls = (cubic.x1, cubic.y1, cubic.x2, cubic.y2)
    assert controls == pytest.approx((20 / 3, 20 / 3, 40 / 3, 20 / 3))


def test_circle_rect_line_and_polyline_nodes() -> None:
    circle = node_segments("circle", {"cx": "12", "cy": "12", "r": "3"})
    assert isinstance(circle[-1], Close)
    assert all(math.hypot(x - 12, y - 12) == pytest.approx(3) for x, y in ends(circle))
    rect = node_segments("rect", {"x": "15", "y": "4", "width": "4", "height": "6", "ry": "2"})
    xs, ys = zip(*ends(rect), strict=True)
    assert (min(xs), max(xs), min(ys), max(ys)) == pytest.approx((15, 19, 4, 10))
    assert node_segments("line", {"x1": "1", "y1": "2", "x2": "3", "y2": "4"}) == [
        Move(1, 2), Line(3, 4),
    ]
    assert ends(node_segments("polygon", {"points": "1,1 5,1 3 4"})) == [(1, 1), (5, 1), (3, 4)]


def test_every_lucide_icon_converts() -> None:
    names = icon_names()
    assert len(names) > 1000
    for name in names:
        for tag, attrs in icon_nodes(name) or []:
            segments = node_segments(tag, attrs)
            assert segments and isinstance(segments[0], Move), name


# --- поиск ----------------------------------------------------------------------


@pytest.mark.parametrize(
    ("query", "name"),
    [("shield-check", "shield-check"), ("ShieldCheck", "shield-check"),
     ("shield_check", "shield-check"), (" Shield Check ", "shield-check"),
     ("Clock3", "clock-3"), ("Grid2x2", "grid-2x2"), ("CircleX", "circle-x")],
)
def test_icon_is_found_by_spelling_variants(query: str, name: str) -> None:
    assert icon_nodes(query) == icon_nodes(name)
    assert icon_nodes(query)


def test_unknown_icon_is_none() -> None:
    assert icon_nodes("нет-такой-иконки") is None


# --- вставка ----------------------------------------------------------------------


def icon_shape(query: str, color: ColorRef | None, manifest: TemplateManifest) -> object:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    block = IconBlock(block_id="i", query=query, color_ref=color,
                      x=EMU_PER_CM, y=EMU_PER_CM, cx=6 * EMU_PER_CM, cy=2 * EMU_PER_CM)
    return add_icon(slide, block)


def test_icon_is_a_square_vector_shape_centred_in_the_box(manifest: TemplateManifest) -> None:
    shape = icon_shape("shield-check", ColorRef.ACCENT2, manifest)
    assert (shape.width, shape.height) == (2 * EMU_PER_CM, 2 * EMU_PER_CM)  # type: ignore[attr-defined]
    assert shape.left == EMU_PER_CM + 2 * EMU_PER_CM  # type: ignore[attr-defined]
    xml = shape._element.xml  # type: ignore[attr-defined]
    assert "<a:custGeom>" in xml and "prstGeom" not in xml
    assert xml.count("<a:path ") == 2
    assert "blip" not in xml and "srgbClr" not in xml
    assert '<a:schemeClr val="accent2"/>' in xml
    # Толщина линии Lucide — 2 из 24: у иконки 2 см это 1/12 от стороны.
    line = shape._element.spPr.find(qn("a:ln"))  # type: ignore[attr-defined]
    assert int(line.get("w")) == round(2 * EMU_PER_CM * 2 / ICON_VIEWBOX)
    assert line.get("cap") == "rnd"


def test_stroke_only_paths_are_not_filled_but_dots_are(manifest: TemplateManifest) -> None:
    dotted = next(
        name for name in icon_names()
        if any("fill" in attrs for _, attrs in icon_nodes(name) or [])
    )
    xml = icon_shape(dotted, None, manifest)._element.xml  # type: ignore[attr-defined]
    assert 'fill="none"' in xml
    assert xml.count("<a:path ") > xml.count('fill="none"')
    assert '<a:schemeClr val="accent1"/>' in xml


def test_unknown_icon_cannot_be_added(manifest: TemplateManifest) -> None:
    with pytest.raises(KeyError, match="нет-такой"):
        icon_shape("нет-такой", None, manifest)
