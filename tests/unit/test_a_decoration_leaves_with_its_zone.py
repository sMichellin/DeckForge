"""Декор снятой зоны уходит вместе с ней. Change `a-decoration-leaves-with-its-zone`, таск RG52
(`docs/agents/tasks-26-09.md`).

Education s06, рецепт `ex013`: зоны схемы нашим текстом не заполнились, писатель их снял (RG40),
а стрелки, плашка и иконка остались — чертёж без единой подписи. Шов — `clone_recipe`
на синтетическом примере (python-pptx) и на настоящем шаблоне.

Сценарии — из дельты `openspec/changes/a-decoration-leaves-with-its-zone/specs/pptx-writer/`.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Emu

from deckforge.composition.recipe_binding import bind_to_recipe
from deckforge.designsystem import derive
from deckforge.designsystem.models import Recipe, RecipeKind, TypeLevel, Zone
from deckforge.domain.enums import TextRole
from deckforge.domain.slide import BulletItem, BulletsBlock, DeckIR, SlideIR, TextBlock
from deckforge.parsing import TemplateParser
from deckforge.rendering.recipe_slide import clone_recipe
from deckforge.rendering.writer import PptxWriter
from tests.case_templates import case_template
from tests.unit.test_recipe_leaves_no_sample_text import slide_ir

#: Допуск «примыкает» из спецификации: 1 % ширины слайда (калибровка — в proposal).
TOUCH_SHARE = 0.01


def _text_box(slide, x, y, cx, cy, text):
    box = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(cx), Emu(cy))
    box.text_frame.text = text
    return box


def _zone(zone_id, shape, role=TypeLevel.BODY, repeat=None) -> Zone:
    return Zone(
        zone_id=zone_id,
        xml_id=shape.shape_id,
        role=role,
        repeat=repeat,
        capacity_chars=60,
        x=int(shape.left),
        y=int(shape.top),
        cx=int(shape.width),
        cy=int(shape.height),
    )


def scheme_example():
    """Пример-схема на слайде 9144000 × 6858000 EMU (1 % ширины — 91 440 EMU).

    Фон во весь слайд; заголовок с чертой под ним, черта касается и зоны-подзаголовка S,
    которая не заполняется никогда; зона A на плашке; зона B на плашке с иконкой (иконка
    внутри плашки, но не у зоны); стрелка A → B; номер шага N на кружке; полоса во всю
    ширину касается только зоны N и кружка; член повтора и «картинка рецепта» вплотную
    к зоне A. Черта и полоса без своих защит (заголовок, `FULL_SPAN_SHARE`) ушли бы.
    """
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    shapes = s.shapes
    found = {
        "background": shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Emu(9144000), Emu(6858000)),
        "title": _text_box(s, 400000, 200000, 8000000, 600000, "Заголовок примера"),
        "title_rule": shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Emu(400000), Emu(820000), Emu(1400000), Emu(820000)
        ),
        "plate_a": shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(300000), Emu(1900000),
                                    Emu(2200000), Emu(1000000)),
        "plate_b": shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(3900000), Emu(1900000),
                                    Emu(2200000), Emu(1600000)),
        "icon_b": shapes.add_shape(MSO_SHAPE.OVAL, Emu(5600000), Emu(3050000),
                                   Emu(300000), Emu(300000)),
        "badge": shapes.add_shape(MSO_SHAPE.OVAL, Emu(6950000), Emu(3950000),
                                  Emu(500000), Emu(500000)),
        "member": shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(1000000), Emu(2850000),
                                   Emu(400000), Emu(300000)),
        "picture": shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(400000), Emu(2850000),
                                    Emu(400000), Emu(300000)),
        "zone_a": _text_box(s, 400000, 2000000, 2000000, 800000, "Текстовый блок A"),
        "zone_b": _text_box(s, 4000000, 2000000, 2000000, 800000, "Текстовый блок B"),
        "zone_n": _text_box(s, 7000000, 4000000, 400000, 400000, "01"),
        "zone_c": _text_box(s, 400000, 5000000, 2000000, 800000, "Карточка"),
        "zone_s": _text_box(s, 400000, 900000, 3000000, 400000, "Подзаголовок"),
        "band": shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Emu(4300000), Emu(9144000),
                                 Emu(500000)),
        "arrow": shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Emu(2400000), Emu(2400000), Emu(4000000), Emu(2400000)
        ),
    }
    return prs, str(s.part.partname), found


def scheme_recipe(part_name: str, f) -> Recipe:
    return Recipe(
        recipe_id="ex013",
        example_index=1,
        part_name=part_name,
        kind=RecipeKind.TEXT,
        zones=[
            _zone("zt", f["title"], TypeLevel.SLIDE_TITLE),
            _zone("za", f["zone_a"]),
            _zone("zb", f["zone_b"]),
            _zone("zn", f["zone_n"], TypeLevel.CARD_TITLE),
            _zone("zc", f["zone_c"], repeat=0),
            _zone("zs", f["zone_s"]),
        ],
        repeat_xml_ids=[[f["zone_c"].shape_id, f["member"].shape_id]],
        picture_xml_id=f["picture"].shape_id,
    )


def ids_left(slide) -> set[int]:
    tree = slide.shapes._spTree
    return {int(node.get("id")) for node in tree.iter() if node.tag.endswith("cNvPr")}


def cloned(filled: list[str]):
    prs, part_name, f = scheme_example()
    recipe = scheme_recipe(part_name, f)
    slide = clone_recipe(prs, recipe, slide_ir(recipe, filled))
    return ids_left(slide), f


def test_an_arrow_to_a_dropped_zone_leaves_with_it() -> None:
    """История 44: зона B не заполнена — стрелки A → B нет, зона A и наш текст на месте."""
    left, f = cloned(["zt", "za", "zc"])

    assert f["zone_b"].shape_id not in left
    assert f["arrow"].shape_id not in left, "стрелка ведёт к пустому месту"
    assert f["zone_a"].shape_id in left


def test_a_plate_of_a_dropped_zone_leaves_with_its_icon() -> None:
    """История 45: плашка вокруг снятой зоны без другого текста уходит, иконка на ней — тоже."""
    left, f = cloned(["zt", "za", "zc"])

    assert f["plate_b"].shape_id not in left, "пустая рамка карточки осталась"
    assert f["icon_b"].shape_id not in left, "иконка без подписи осталась"


def test_the_badge_of_a_dropped_step_number_leaves() -> None:
    """Сверка G2, п. 1: у снятой зоны-номера уходит и её подложка."""
    left, f = cloned(["zt", "za", "zb", "zc"])

    assert f["zone_n"].shape_id not in left
    assert f["badge"].shape_id not in left


def test_a_plate_that_keeps_our_text_stays() -> None:
    """Плашка, на которой остался наш текст, — фигура автора (RG40)."""
    left, f = cloned(["zt", "za", "zc"])

    assert f["plate_a"].shape_id in left


def test_the_slide_design_stays() -> None:
    """История 46: фон во всю ширину, декор у заголовка, член повтора, картинка рецепта.

    Черта под заголовком касается снятой зоны S, полоса во всю ширину — только снятой зоны N:
    обе остаются только благодаря своим защитам.
    """
    left, f = cloned(["zt", "zc"])

    assert f["zone_s"].shape_id not in left and f["zone_n"].shape_id not in left
    for name in ("background", "band", "title_rule", "member", "picture"):
        assert f[name].shape_id in left, f"{name}: оформление слайда снято"


def lines_example():
    """Развилка и цепочка на слайде 9144000 × 6858000 EMU.

    Развилка: ствол от заполненной зоны A к узлу J, из узла — ветка к снятой зоне B и ветка
    к заполненной зоне C. Цепочка: три линии от снятой зоны D и кружок без текста у конца
    последней — дальше допуска от любой зоны.
    """
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])

    def line(x0, y0, x1, y1):
        return s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(x0), Emu(y0), Emu(x1), Emu(y1))

    f = {
        "title": _text_box(s, 400000, 200000, 8000000, 600000, "Заголовок примера"),
        "zone_a": _text_box(s, 400000, 2000000, 2000000, 800000, "A"),
        "zone_b": _text_box(s, 4000000, 1000000, 2000000, 800000, "B"),
        "zone_c": _text_box(s, 4000000, 3400000, 2000000, 800000, "C"),
        "zone_d": _text_box(s, 6800000, 5000000, 1800000, 800000, "D"),
        "trunk": line(2400000, 2400000, 3200000, 2400000),
        "to_b": line(3200000, 2400000, 4000000, 1400000),
        "to_c": line(3200000, 2400000, 4000000, 3800000),
        "chain_1": line(6800000, 5400000, 5800000, 5400000),
        "chain_2": line(5800000, 5400000, 5800000, 6200000),
        "chain_3": line(5800000, 6200000, 4800000, 6200000),
        "tail": s.shapes.add_shape(MSO_SHAPE.OVAL, Emu(4500000), Emu(6050000),
                                   Emu(300000), Emu(300000)),
    }
    recipe = Recipe(
        recipe_id="ex013",
        example_index=1,
        part_name=str(s.part.partname),
        kind=RecipeKind.TEXT,
        zones=[
            _zone("zt", f["title"], TypeLevel.SLIDE_TITLE),
            *(_zone(f"z{key}", f[f"zone_{key}"]) for key in "abcd"),
        ],
    )
    slide = clone_recipe(prs, recipe, slide_ir(recipe, ["zt", "za", "zc"]))
    return ids_left(slide), f


def test_the_trunk_of_a_fork_between_filled_zones_stays() -> None:
    """Ревью D08: ветка к снятой зоне уходит, а ствол и ветка к заполненной зоне — нет,
    хотя касаются снятой линии: их другой конец у нашего текста."""
    left, f = lines_example()

    assert f["to_b"].shape_id not in left
    assert f["trunk"].shape_id in left, "ствол между заполненными зонами снят каскадом"
    assert f["to_c"].shape_id in left, "ветка к заполненной зоне снята каскадом"


def test_a_chain_to_a_dropped_zone_leaves_with_its_tail() -> None:
    """D08: цепочка из трёх линий к снятой зоне уходит целиком, кружок у её конца — тоже."""
    left, f = lines_example()

    for name in ("chain_1", "chain_2", "chain_3", "tail"):
        assert f[name].shape_id not in left, f"{name}: остался у пустого места"


# --- настоящий шаблон -----------------------------------------------------------------

EDUCATION = "Шаблон презентации VK Education.pptx"


def _box(shape) -> tuple[int, int, int, int]:
    return int(shape.left), int(shape.top), int(shape.width), int(shape.height)


def _gap(a: tuple[int, int, int, int], b: tuple[int, int, int, int]) -> int:
    dx = max(b[0] - (a[0] + a[2]), a[0] - (b[0] + b[2]), 0)
    dy = max(b[1] - (a[1] + a[3]), a[1] - (b[1] + b[3]), 0)
    return max(dx, dy)


def test_the_education_scheme_with_two_facts_keeps_no_arrow_to_nowhere(tmp_path: Path) -> None:
    """Education `ex013` с двумя фактами: ни одной линии, конец которой в снятой зоне."""
    path = case_template(EDUCATION)
    manifest = TemplateParser().parse(path, use_cache=False)
    ds = derive(manifest)
    recipe = next(r for r in ds.recipes if r.recipe_id == "ex013")
    example = next(e for e in manifest.examples if e.slide_index == recipe.example_index)
    source = SlideIR(
        slide_id="s01",
        layout_id=example.layout_id or manifest.layouts[0].layout_id,
        variant="A",
        blocks=[
            TextBlock(block_id="t", role=TextRole.TITLE, text="Как устроен конвейер"),
            BulletsBlock(
                block_id="b",
                items=[
                    BulletItem(text="Анализ шаблона и извлечение дизайн-системы"),
                    BulletItem(text="Композиция слайдов по рецептам автора шаблона"),
                ],
            ),
        ],
    )
    ir = bind_to_recipe(source, recipe, [])
    deck = DeckIR(deck_id="d01", template_id=manifest.template_id, variant="A", seed=1, slides=[ir])
    out = PptxWriter(path, manifest, design_system=ds).write(deck, tmp_path / "deck.pptx")
    prs = Presentation(str(out))
    slide = prs.slides[0]

    present = {shape.shape_id for shape in slide.shapes}
    dropped = [
        (z.zone_id, (z.x, z.y, z.cx, z.cy))
        for z in recipe.zones
        if z.has_frame and z.xml_id not in present
    ]
    assert dropped, "замер: у ex013 с двумя фактами зоны схемы снимаются"
    tolerance = TOUCH_SHARE * prs.slide_width
    leading = []
    for shape in slide.shapes:
        if not shape.element.tag.endswith("cxnSp"):
            continue
        x, y, cx, cy = _box(shape)
        for end in ((x, y, 0, 0), (x + cx, y + cy, 0, 0)):
            leading += [
                (shape.shape_id, zone_id) for zone_id, box in dropped if _gap(end, box) <= tolerance
            ]
    assert not leading, f"линии ведут к снятым зонам: {leading}"
