"""Слайды в файле идут в порядке колоды. Change `slide-order-follows-the-deck`.

Шов — `write()`: собрать смешанную колоду (обычные слайды и слайды по рецепту),
открыть файл и сравнить порядок. `slide_id` в файле не хранится, поэтому слайд
узнаётся по тексту заголовка, который writer пишет сам.

Синтетический шаблон строится в тесте и идёт везде, VK Tech — на машине с шаблонами кейса.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Pt

import deckforge.audit.deterministic.integrity  # noqa: F401  — регистрирует проверки
from deckforge.audit.context import AuditContext
from deckforge.audit.registry import REGISTRY, CheckUnavailable
from deckforge.designsystem import derive
from deckforge.designsystem.models import (
    DesignSystem,
    Recipe,
    RecipeKind,
    TypeLevel,
    Zone,
)
from deckforge.domain.enums import TextRole
from deckforge.domain.slide import DeckIR, FitResult, SlideIR, TextBlock
from deckforge.domain.template import LayoutSpec, TemplateManifest
from deckforge.parsing import TemplateParser
from deckforge.rendering.writer import PptxWriter
from tests.case_templates import case_template

Case = tuple[Path, TemplateManifest, DesignSystem]

#: Раскладки колоды: True — слайд по рецепту, False — обычный.
ARRANGEMENTS = {
    "plain-recipe-plain": (False, True, False),
    "recipe-first": (True, False, False),
    "recipe-last": (False, False, True),
    "two-recipes-in-a-row": (False, True, True, False),
}


@pytest.fixture(scope="module")
def synthetic(tmp_path_factory: pytest.TempPathFactory) -> Case:
    """Шаблон с двумя слайдами-примерами; у каждого заголовок — надпись вне плейсхолдера."""
    prs = Presentation()
    blank = next(layout for layout in prs.slide_layouts if layout.name == "Blank")
    examples: list[tuple[str, int, int]] = []
    for index in range(2):
        slide = prs.slides.add_slide(blank)
        title = slide.shapes.add_textbox(Pt(36), Pt(36), Pt(400), Pt(60))
        title.text_frame.text = f"Заголовок примера {index + 1}"
        body = slide.shapes.add_textbox(Pt(36), Pt(120), Pt(400), Pt(200))
        body.text_frame.text = f"Текст примера {index + 1}"
        examples.append((str(slide.part.partname).lstrip("/"), title.shape_id, body.shape_id))
    path = tmp_path_factory.mktemp("synthetic") / "template.pptx"
    prs.save(str(path))

    manifest = TemplateParser().parse(path, use_cache=False)
    recipes = [
        Recipe(
            recipe_id=f"ex{index + 1:03d}",
            example_index=index + 1,
            part_name=part_name,
            kind=RecipeKind.TEXT,
            zones=[
                Zone(zone_id="title", xml_id=title_id, role=TypeLevel.SLIDE_TITLE,
                     capacity_chars=60),
                Zone(zone_id="body", xml_id=body_id, role=TypeLevel.BODY, capacity_chars=200),
            ],
        )
        for index, (part_name, title_id, body_id) in enumerate(examples)
    ]
    return path, manifest, derive(manifest).model_copy(update={"recipes": recipes})


@pytest.fixture(scope="module")
def vk_tech() -> Case:
    path = case_template("VK Tech шаблон.pptx")
    manifest = TemplateParser().parse(path, use_cache=False)
    return path, manifest, derive(manifest)


def titled_recipes(ds: DesignSystem) -> list[tuple[Recipe, Zone, Zone]]:
    """Рецепты с зоной заголовка и зоной текста: по заголовку слайд узнаётся в файле,
    текст нужен, чтобы слайд не был «одним заголовком» для аудита."""
    out = []
    for recipe in ds.recipes:
        title = next(
            (z for z in recipe.zones if z.role is TypeLevel.SLIDE_TITLE and z.xml_id), None
        )
        body = next(
            (z for z in recipe.zones if z.role is not TypeLevel.SLIDE_TITLE and z.xml_id), None
        )
        if title is not None and body is not None:
            out.append((recipe, title, body))
    return out


def body_text(slide_id: str) -> str:
    """Текст у каждого слайда свой: одинаковый дал бы находку `integrity.duplicate_slides`."""
    return f"Содержание слайда {slide_id}"


def plain_slide(manifest: TemplateManifest, slide_id: str, title: str) -> SlideIR:
    def placeholder(layout: LayoutSpec, role: TextRole) -> int | None:
        return next((ph.idx for ph in layout.placeholders if ph.role is role), None)

    layout = next(
        lt
        for lt in manifest.layouts
        if placeholder(lt, TextRole.TITLE) is not None
        and placeholder(lt, TextRole.BODY) is not None
    )
    return SlideIR(
        slide_id=slide_id,
        layout_id=layout.layout_id,
        variant="A",
        blocks=[
            TextBlock(
                block_id="b0",
                role=TextRole.TITLE,
                text=title,
                placeholder_idx=placeholder(layout, TextRole.TITLE),
            ),
            TextBlock(
                block_id="b1",
                role=TextRole.BODY,
                text=body_text(slide_id),
                placeholder_idx=placeholder(layout, TextRole.BODY),
            ),
        ],
        # Кегли — со шкалы шаблона: вписывание здесь не проверяется.
        fit_report={
            "b0": FitResult(final_size_pt=manifest.size_ladder_pt[0]),
            "b1": FitResult(final_size_pt=manifest.size_ladder_pt[-1]),
        },
    )


def recipe_slide(
    manifest: TemplateManifest,
    slide_id: str,
    title: str,
    recipe: tuple[Recipe, Zone, Zone],
) -> SlideIR:
    chosen, title_zone, body_zone = recipe
    return SlideIR(
        slide_id=slide_id,
        layout_id=manifest.layouts[0].layout_id,
        variant="A",
        blocks=[
            TextBlock(block_id="b0", role=TextRole.TITLE, text=title,
                      zone_id=title_zone.zone_id),
            TextBlock(block_id="b1", role=TextRole.BODY, text=body_text(slide_id),
                      zone_id=body_zone.zone_id),
        ],
        recipe_id=chosen.recipe_id,
    )


def mixed_deck(case: Case, arrangement: tuple[bool, ...]) -> DeckIR:
    """Колода по раскладке; заголовок каждого слайда — его `slide_id`. Рецепты чередуются."""
    _path, manifest, ds = case
    recipes = titled_recipes(ds)
    if not recipes:
        pytest.fail("у шаблона нет рецепта с зоной заголовка и зоной текста: колоду не собрать")
    slides = []
    for number, by_recipe in enumerate(arrangement, start=1):
        slide_id = f"s{number:02d}"
        if by_recipe:
            slides.append(
                recipe_slide(manifest, slide_id, slide_id, recipes[number % len(recipes)])
            )
        else:
            slides.append(plain_slide(manifest, slide_id, slide_id))
    return DeckIR(
        deck_id="d01", template_id=manifest.template_id, variant="A", seed=1, slides=slides
    )


def write(case: Case, deck: DeckIR, out: Path) -> Path:
    path, manifest, ds = case
    return PptxWriter(path, manifest, design_system=ds).write(deck, out)


def order_in_file(path: Path, deck: DeckIR) -> list[str]:
    """`slide_id` каждого слайда файла по его заголовку; чужой слайд — «?»."""
    ids = {slide.slide_id for slide in deck.slides}
    out = []
    for slide in Presentation(str(path)).slides:
        texts = {
            shape.text_frame.text.strip() for shape in slide.shapes if shape.has_text_frame
        }
        found = sorted(texts & ids)
        out.append(found[0] if len(found) == 1 else "?")
    return out


@pytest.mark.parametrize("arrangement", ARRANGEMENTS.values(), ids=ARRANGEMENTS.keys())
def test_slides_follow_the_deck_on_a_synthetic_template(
    synthetic: Case, arrangement: tuple[bool, ...], tmp_path: Path
) -> None:
    deck = mixed_deck(synthetic, arrangement)

    out = write(synthetic, deck, tmp_path / "deck.pptx")

    assert order_in_file(out, deck) == [slide.slide_id for slide in deck.slides]


@pytest.mark.parametrize("arrangement", ARRANGEMENTS.values(), ids=ARRANGEMENTS.keys())
def test_slides_follow_the_deck_on_vk_tech(
    vk_tech: Case, arrangement: tuple[bool, ...], tmp_path: Path
) -> None:
    deck = mixed_deck(vk_tech, arrangement)

    out = write(vk_tech, deck, tmp_path / "deck.pptx")

    assert order_in_file(out, deck) == [slide.slide_id for slide in deck.slides]


def test_the_mixed_deck_on_vk_tech_opens_clean(vk_tech: Case, tmp_path: Path) -> None:
    """Файл открывается заново, примеров шаблона в нём нет, integrity-проверки без находок.

    Проверяются все `integrity.*` любого уровня; пропуск проверки — тоже провал: «не
    проверяли» не значит «чисто».
    """
    _path, manifest, _ds = vk_tech
    deck = mixed_deck(vk_tech, ARRANGEMENTS["plain-recipe-plain"])

    out = write(vk_tech, deck, tmp_path / "deck.pptx")

    assert len(Presentation(str(out)).slides) == len(deck.slides)
    ctx = AuditContext(manifest=manifest, deck=deck, deck_path=out)
    findings = []
    for item in REGISTRY.all():
        if not item.check_id.startswith("integrity."):
            continue
        try:
            findings += list(item.fn(ctx.with_params({})))
        except CheckUnavailable as skipped:
            pytest.fail(f"{item.check_id} пропущена: {skipped.reason}")
    assert findings == []
