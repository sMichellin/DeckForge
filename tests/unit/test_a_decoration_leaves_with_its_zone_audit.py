"""Аудит: стрелка или рамка примера ведёт к месту без текста. Change
`a-decoration-leaves-with-its-zone`, таск RG52 (`docs/agents/tasks-26-09.md`).

Education s06, рецепт `ex013`: зоны схемы писатель снял, а стрелки и плашка остались, и ни одна
проверка этого не видела. Проверка `template.decor_leads_nowhere` меряет **готовый файл**:
геометрия зон — из каталога, фигур — из слайда. Шов — функция проверки на `.pptx` и `SlideIR`,
как у `template.sample_text_left`. Правило 7 AGENTS.md: нарушитель и норма.

Сценарии — из дельты
`openspec/changes/a-decoration-leaves-with-its-zone/specs/audit-deterministic/`.
"""

from __future__ import annotations

from dataclasses import replace
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu

from deckforge.audit.deterministic.template import decor_leads_nowhere
from deckforge.audit.registry import CheckUnavailable
from deckforge.designsystem import derive
from deckforge.designsystem.models import Recipe, RecipeKind, TypeLevel, Zone
from deckforge.domain.enums import TextRole
from deckforge.domain.slide import SlideIR, TextBlock
from deckforge.domain.template import TemplateExample, TemplateManifest
from tests.unit._audit_builders import context_for, deck

CHECK = "template.decor_leads_nowhere"

#: Рамки зон примера на слайде 9 144 000 × 6 858 000 EMU (1 % ширины — 91 440 EMU).
TITLE = (400000, 200000, 8000000, 600000)
ZONE_A = (400000, 2000000, 2000000, 800000)
ZONE_B = (4000000, 2000000, 2000000, 800000)


def _zone(zone_id: str, xml_id: int, frame: tuple[int, int, int, int],
          role: TypeLevel = TypeLevel.BODY) -> Zone:
    x, y, cx, cy = frame
    return Zone(zone_id=zone_id, xml_id=xml_id, role=role, capacity_chars=60,
                x=x, y=y, cx=cx, cy=cy)


#: Зоны адресуют фигуры по `cNvPr id`; у зоны B фигуры в файле может не быть — её сняли.
RECIPE = Recipe(
    recipe_id="ex013",
    example_index=1,
    part_name="ppt/slides/slide1.xml",
    kind=RecipeKind.TEXT,
    zones=[
        _zone("zt", 900, TITLE, TypeLevel.SLIDE_TITLE),
        _zone("za", 901, ZONE_A),
        _zone("zb", 902, ZONE_B),
    ],
)


def _with_id(shape: object, xml_id: int) -> object:
    shape._element.find(f"./*/{qn('p:cNvPr')}").set("id", str(xml_id))  # type: ignore[attr-defined]
    return shape


def _text(slide: object, frame: tuple[int, int, int, int], text: str, xml_id: int) -> object:
    box = slide.shapes.add_textbox(*(Emu(v) for v in frame))  # type: ignore[attr-defined]
    box.text_frame.text = text
    return _with_id(box, xml_id)


def scheme_file(path: Path, *, zone_b: str | None = None, caption: bool = False) -> dict[str, int]:
    """Копия примера после писателя: заголовок и зона A с нашим текстом, стрелка A → B,
    плашка вокруг B. `zone_b` — текст зоны B (нет — зону сняли); `caption` — у конца стрелки
    лежит чужая для рецепта фигура с текстом."""
    prs = Presentation()
    page = prs.slides.add_slide(prs.slide_layouts[6])
    shapes = page.shapes
    plate = _with_id(shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(3900000), Emu(1900000),
                                      Emu(2200000), Emu(1000000)), 950)
    arrow = _with_id(shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(2400000), Emu(2400000),
                                          Emu(4000000), Emu(2400000)), 951)
    _text(page, TITLE, "Выручка выросла на треть", 900)
    _text(page, ZONE_A, "Первый факт", 901)
    if zone_b is not None:
        _text(page, ZONE_B, zone_b, 902)
    if caption:
        _text(page, ZONE_B, "Подпись этапа", 960)
    prs.save(str(path))
    return {"plate": plate.shape_id, "arrow": arrow.shape_id}  # type: ignore[attr-defined]


def recipe_slide(*zones: str, recipe_id: str | None = "ex013") -> SlideIR:
    texts = {"zt": "Выручка выросла на треть", "za": "Первый факт", "zb": "Второй факт"}
    return SlideIR(
        slide_id="s01",
        layout_id="L07",
        variant="A",
        recipe_id=recipe_id,
        blocks=[
            TextBlock(
                block_id=f"b{number}",
                role=TextRole.TITLE if zone == "zt" else TextRole.BODY,
                text=texts[zone],
                zone_id=zone if recipe_id else None,
            )
            for number, zone in enumerate(zones)
        ],
    )


def run(manifest: TemplateManifest, ir: SlideIR, path: Path | None) -> list:
    examples = manifest.model_copy(update={"examples": [TemplateExample(slide_index=1)]})
    design = derive(examples).model_copy(update={"recipes": [RECIPE]})
    ctx = replace(context_for(CHECK, deck(ir), examples, deck_path=path), design_system=design)
    return list(decor_leads_nowhere(ctx))


def test_an_arrow_and_a_plate_at_a_dropped_zone_are_findings(
    manifest: TemplateManifest, tmp_path: Path
) -> None:
    """Нарушитель: зону B сняли, стрелка к ней и плашка вокруг неё остались."""
    ids = scheme_file(tmp_path / "deck.pptx")

    findings = run(manifest, recipe_slide("zt", "za"), tmp_path / "deck.pptx")

    assert sorted((f.slide_id, f.evidence["xml_id"], f.evidence["zone_id"]) for f in findings) == [
        ("s01", str(ids["plate"]), "zb"),
        ("s01", str(ids["arrow"]), "zb"),
    ]


def test_a_filled_zone_is_not_a_finding(manifest: TemplateManifest, tmp_path: Path) -> None:
    """Норма: зона B заполнена нашим текстом — стрелка ведёт к нему, плашка его держит."""
    scheme_file(tmp_path / "deck.pptx", zone_b="Второй факт")

    assert run(manifest, recipe_slide("zt", "za", "zb"), tmp_path / "deck.pptx") == []


def test_an_arrow_to_a_shape_with_text_is_not_a_finding(
    manifest: TemplateManifest, tmp_path: Path
) -> None:
    """Норма: зоны B нет, но на её месте фигура с текстом — стрелка ведёт к подписи."""
    scheme_file(tmp_path / "deck.pptx", caption=True)

    assert run(manifest, recipe_slide("zt", "za"), tmp_path / "deck.pptx") == []


def test_nothing_to_check_is_skipped_not_passed(
    manifest: TemplateManifest, tmp_path: Path
) -> None:
    """Нет файла или нет слайдов по рецепту — проверка не прошла, а пропущена."""
    scheme_file(tmp_path / "deck.pptx")

    with pytest.raises(CheckUnavailable):
        run(manifest, recipe_slide("zt", "za"), None)
    with pytest.raises(CheckUnavailable):
        run(manifest, recipe_slide("zt", "za", recipe_id=None), tmp_path / "deck.pptx")
