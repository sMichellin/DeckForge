"""Таблица, диаграмма и SmartArt примера, которых мы не заполнили, со слайда уходят.
Change `an-unfilled-graphic-frame-leaves-the-slide`, таск RG45 (`docs/agents/tasks-26-09.md`),
он же Т9 из `docs/agents/requirements-from-notes-26-09.md`.

Зонами каталог делает только текстовые фигуры, и `p:graphicFrame` примера писатель копировал
целиком — с таблицей автора и её текстом. На титуле VK WorkSpace (`ex014`) так в колоду
уезжала таблица «Заголовок / Текст»; это была последняя ошибка аудита 26.09.

Сценарии — из дельты `openspec/changes/an-unfilled-graphic-frame-leaves-the-slide/specs/`.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu

from deckforge.designsystem.models import Recipe, RecipeKind, TypeLevel, Zone
from deckforge.rendering.recipe_slide import clone_recipe
from tests.unit.test_recipe_leaves_no_sample_text import (
    ours,
    paragraphs,
    slide_ir,
    write_recipe,
)

FRAME = qn("p:graphicFrame")


def example_with_frames():
    """Пример: заголовок, плашка-декор, таблица автора и его диаграмма."""
    prs = Presentation()
    source = prs.slides.add_slide(prs.slide_layouts[6])
    title = source.shapes.add_textbox(Emu(100), Emu(100), Emu(6000), Emu(600))
    title.text_frame.text = "Заголовок примера"
    plate = source.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(100), Emu(900), Emu(500), Emu(500))
    table = source.shapes.add_table(2, 2, Emu(100), Emu(1600), Emu(4000), Emu(800))
    table.table.cell(0, 0).text = "Заголовок"
    table.table.cell(1, 0).text = "Текст"
    data = CategoryChartData()
    data.categories = ["А", "Б"]
    data.add_series("Ряд", (1, 2))
    chart = source.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(4500), Emu(1600), Emu(3000), Emu(2000), data
    )
    return prs, str(source.part.partname), title, plate, table, chart


def recipe_for(part_name: str, title_id: int, addressed: int | None = None) -> Recipe:
    return Recipe(
        recipe_id="ex014",
        example_index=1,
        part_name=part_name,
        kind=RecipeKind.COVER,
        zones=[Zone(zone_id="zt", xml_id=title_id, role=TypeLevel.SLIDE_TITLE, capacity_chars=60)],
        picture_xml_id=addressed,
    )


def frames_on(slide) -> list:
    return [node for node in slide.shapes._spTree.iter(FRAME)]


def test_an_unfilled_table_and_chart_leave_the_slide() -> None:
    """Нарушитель до правки: таблица автора с текстом «Заголовок / Текст» на слайде колоды."""
    prs, part_name, title, plate, _table, _chart = example_with_frames()
    recipe = recipe_for(part_name, title.shape_id)
    ir = slide_ir(recipe, ["zt"])

    slide = clone_recipe(prs, recipe, ir)

    assert frames_on(slide) == [], "таблица или диаграмма примера осталась на слайде"
    assert set(paragraphs(slide)) <= ours(ir), "текст шаблона остался на слайде"
    assert any(shape.shape_id == plate.shape_id for shape in slide.shapes), "декор снят зря"


def test_the_relations_of_the_removed_frames_go_too(tmp_path: Path) -> None:
    """Диаграмма — отдельная часть пакета: без связи она не сохраняется в файл колоды."""
    prs, part_name, title, *_ = example_with_frames()
    recipe = recipe_for(part_name, title.shape_id)

    slide = clone_recipe(prs, recipe, slide_ir(recipe, ["zt"]))

    kinds = [rel.reltype for rel in slide.part.rels.values()]
    assert not [kind for kind in kinds if "/chart" in kind]
    assert any("/slideLayout" in kind for kind in kinds), "связь с макетом снимать нельзя"
    out = tmp_path / "deck.pptx"
    prs.save(str(out))
    assert Presentation(str(out)).slides[-1].slide_layout is not None


def test_a_frame_the_recipe_addresses_stays() -> None:
    """Норма: рамку, которую адресует рецепт, решает тот, кто её адресовал, — не писатель."""
    prs, part_name, title, _plate, table, _chart = example_with_frames()
    # Адрес через `picture_xml_id`: его не снимает ни одно другое правило писателя,
    # и тест проверяет ровно RG45, а не судьбу лишнего повтора.
    recipe = recipe_for(part_name, title.shape_id, addressed=table.shape_id)

    slide = clone_recipe(prs, recipe, slide_ir(recipe, ["zt"]))

    ids = {int(node.find(f"./*/{qn('p:cNvPr')}").get("id")) for node in frames_on(slide)}
    assert table.shape_id in ids


def test_the_title_of_workspace_carries_no_table_of_the_template(tmp_path: Path) -> None:
    """Тот самый слайд прогона 26.09: титул WorkSpace по `ex014`."""
    _recipe, ir, slide = write_recipe(
        "VK_WorkSpace_Клиентская_конференция_Шаблон_03.pptx", "ex014", tmp_path
    )

    assert frames_on(slide) == []
    left = [text for text in paragraphs(slide) if text not in ours(ir)]
    assert not left, f"текст шаблона на титуле: {left[:5]}"
