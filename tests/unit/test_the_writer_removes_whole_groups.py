"""Писатель снимает незаполненное группами паспорта. Change `the-writer-removes-whole-groups`
(план Б, 4; issue #244).

VK Tech `ex018`: ряд из четырёх карточек и широкая карточка под ним. Legacy-IR 28.09 кладёт текст
в 1-ю и 4-ю карточку ряда, 2-я, 3-я и широкая остаются плашками без текста — на шести слайдах.
С паспортом писатель не угадывает: карточка без нашего текста уходит целиком, а ряд
заполняется с начала. Без паспорта — прежний путь, байт в байт (эталон снят до правки кода).

Шов — `clone_recipe`: синтетический пример python-pptx и настоящие шаблоны кейса.
Сценарии — из дельты `openspec/changes/the-writer-removes-whole-groups/specs/pptx-writer/`.
"""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Emu

from deckforge.composition.passport import with_passports
from deckforge.designsystem.models import (
    ExamplePassport,
    Place,
    PlaceGroup,
    PlaceKind,
    Recipe,
    RecipeKind,
    TypeLevel,
    Zone,
)
from deckforge.layout.fonts import FontLibrary
from deckforge.pipeline.replay import from_fixture
from deckforge.rendering.recipe_slide import clone_recipe
from tests.case_templates import case_template
from tests.unit.test_recipe_leaves_no_sample_text import slide_ir

GOLDEN = Path(__file__).resolve().parents[1] / "fixtures" / "the-writer-removes-whole-groups"
RUNS = Path(__file__).resolve().parents[1] / "fixtures" / "runs" / "2026-09-28"
VK_TECH = "VK Tech шаблон.pptx"
WORKSPACE = "VK_WorkSpace_Клиентская_конференция_Шаблон_03.pptx"

#: Карточки ряда: левый край плашки. Четвёртая карточка собрана в `p:grpSp`.
CARDS = (300000, 2500000, 4700000, 6900000)


def _text_box(shapes, x, y, cx, cy, text):
    box = shapes.add_textbox(Emu(x), Emu(y), Emu(cx), Emu(cy))
    box.text_frame.text = text
    return box


def cards_example():
    """Пример «ряд карточек» на слайде 9144000 × 6858000 EMU.

    Заголовок; четыре карточки ряда — плашка, иконка на ней, подзаголовок и текст (четвёртая
    в группе `p:grpSp`); широкая карточка — плашка, акцентная черта слева от неё и текст;
    линия от третьей карточки вниз,
    которой нет в паспорте; таблица автора, которую рецепт не адресует.
    """
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shapes = slide.shapes
    found = {"title": _text_box(shapes, 400000, 200000, 8000000, 600000, "Заголовок примера")}
    for n, x in enumerate(CARDS, start=1):
        target = shapes.add_group_shape().shapes if n == 4 else shapes
        found[f"plate{n}"] = target.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(1200000), Emu(2000000), Emu(2400000)
        )
        found[f"icon{n}"] = target.add_shape(
            MSO_SHAPE.OVAL, Emu(x + 100000), Emu(1300000), Emu(300000), Emu(300000)
        )
        found[f"sub{n}"] = _text_box(target, x + 100000, 1700000, 1800000, 400000, "Подзаголовок")
        found[f"text{n}"] = _text_box(target, x + 100000, 2200000, 1800000, 1200000, "Текст")
    found["wide_plate"] = shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Emu(300000), Emu(3900000), Emu(8600000), Emu(1200000)
    )
    found["wide_bar"] = shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Emu(150000), Emu(3900000), Emu(100000), Emu(1200000)
    )
    found["wide_text"] = _text_box(shapes, 400000, 4000000, 8400000, 1000000, "Широкий текст")
    found["line"] = shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Emu(5700000), Emu(3400000), Emu(5700000), Emu(3900000)
    )
    found["table"] = shapes.add_table(2, 1, Emu(300000), Emu(5400000), Emu(3000000), Emu(800000))
    found["table"].table.cell(0, 0).text = "Заголовок"
    found["table"].table.cell(1, 0).text = "Текст"
    return prs, str(slide.part.partname), found


def _zone(zone_id, shape, role, repeat=None) -> Zone:
    return Zone(
        zone_id=zone_id,
        xml_id=shape.shape_id,
        role=role,
        repeat=repeat,
        capacity_chars=60,
        x=int(shape.left),
        y=int(shape.top),
        cx=int(shape.width),
        cy=int(shape.height),
    )


def cards_recipe(part_name: str, f, *, passport: bool) -> Recipe:
    """Рецепт примера: зоны `zt`, `zs1..4`/`zb1..4` (подзаголовок/текст карточки), `zw`.

    Паспорт: g01 — заголовок; g02–g05 — ряд `r1`, место 1 — подзаголовок, место 2 — текст,
    декор — плашка и иконка; g06 — широкая карточка с плашкой и чертой. Линии и таблицы
    в паспорте нет.
    """
    zones = [_zone("zt", f["title"], TypeLevel.SLIDE_TITLE)]
    for n in range(1, 5):
        zones.append(_zone(f"zs{n}", f[f"sub{n}"], TypeLevel.CARD_TITLE, repeat=n - 1))
        zones.append(_zone(f"zb{n}", f[f"text{n}"], TypeLevel.BODY, repeat=n - 1))
    zones.append(_zone("zw", f["wide_text"], TypeLevel.BODY))
    recipe = Recipe(
        recipe_id="ex018",
        example_index=1,
        part_name=part_name,
        kind=RecipeKind.CARDS,
        zones=zones,
        repeats=4,
        repeat_xml_ids=[
            [f[f"{name}{n}"].shape_id for name in ("plate", "sub", "text", "icon")]
            for n in range(1, 5)
        ],
    )
    if not passport:
        return recipe
    return recipe.model_copy(update={"passport": cards_passport(f)})


def _place(number: int, zone: Zone) -> Place:
    return Place(
        place_id=f"p{number:02d}", kind=PlaceKind.TEXT, zone_id=zone.zone_id,
        xml_id=zone.xml_id, role=zone.role, capacity_chars=60,
    )


def cards_passport(f) -> ExamplePassport:
    def zone(zone_id: str, shape, role) -> Zone:
        return _zone(zone_id, shape, role)

    groups = [PlaceGroup(group_id="g01", places=[_place(1, zone("zt", f["title"],
                                                                 TypeLevel.SLIDE_TITLE))])]
    for n in range(1, 5):
        groups.append(PlaceGroup(
            group_id=f"g{n + 1:02d}",
            places=[
                _place(2 * n, zone(f"zs{n}", f[f"sub{n}"], TypeLevel.CARD_TITLE)),
                _place(2 * n + 1, zone(f"zb{n}", f[f"text{n}"], TypeLevel.BODY)),
            ],
            decor_xml_ids=[f[f"plate{n}"].shape_id, f[f"icon{n}"].shape_id],
            row="r1",
        ))
    groups.append(PlaceGroup(
        group_id="g06",
        places=[_place(10, zone("zw", f["wide_text"], TypeLevel.BODY))],
        decor_xml_ids=[f["wide_plate"].shape_id, f["wide_bar"].shape_id],
    ))
    return ExamplePassport(groups=groups)


def slide_xml(slide) -> bytes:
    """XML части слайда — то, что сверяется с эталоном (без zip-меток времени)."""
    return bytes(etree.tostring(slide._element, encoding="utf-8"))


#: IR эталона: заголовок, 1-я карточка целиком и текст 4-й — как legacy-IR 28.09 на `ex018`.
LEGACY_FILL = ["zt", "zs1", "zb1", "zb4"]


def synthetic_legacy_xml() -> bytes:
    prs, part_name, f = cards_example()
    recipe = cards_recipe(part_name, f, passport=False)
    return slide_xml(clone_recipe(prs, recipe, slide_ir(recipe, LEGACY_FILL)))


def vk_tech_legacy_xml() -> bytes:
    """VK Tech `ex018` по первому слайду legacy-IR 28.09, который его берёт, без паспорта."""
    prs = Presentation(str(case_template(VK_TECH)))
    run = from_fixture(RUNS / "vk-tech")
    recipe = next(r for r in run.design_system.recipes if r.recipe_id == "ex018")
    slide = next(s for s in run.deck.slides if s.recipe_id == "ex018")
    return slide_xml(clone_recipe(prs, recipe.model_copy(update={"passport": None}), slide))


def test_without_a_passport_the_synthetic_slide_is_byte_for_byte() -> None:
    """История 4: без паспорта XML слайда равен эталону, снятому до правки."""
    assert synthetic_legacy_xml() == (GOLDEN / "synthetic-legacy.xml").read_bytes()


def test_without_a_passport_the_vk_tech_slide_is_byte_for_byte() -> None:
    """История 4 на настоящем шаблоне: VK Tech `ex018` без паспорта — как до правки."""
    assert vk_tech_legacy_xml() == (GOLDEN / "vk-tech-ex018-legacy.xml").read_bytes()


def ids_left(slide) -> set[int]:
    tree = slide.shapes._spTree
    return {int(node.get("id")) for node in tree.iter() if node.tag.endswith("cNvPr")}


def with_passport(filled: list[str]):
    prs, part_name, f = cards_example()
    recipe = cards_recipe(part_name, f, passport=True)
    slide = clone_recipe(prs, recipe, slide_ir(recipe, filled))
    return slide, {name: shape.shape_id for name, shape in f.items()}


def test_a_line_outside_the_passport_stays() -> None:
    """История 10: при паспорте каскад по линиям не вызывается — линии нет в группах, она остаётся.

    Без паспорта эта линия ушла бы: её конец у рамки снятой зоны 3-й карточки.
    """
    slide, ids = with_passport(["zt", "zs1", "zb1", "zs2", "zb2"])

    assert ids["line"] in ids_left(slide)


def test_a_group_leaves_with_its_plate_and_icon() -> None:
    """Истории 6, 12: карточки без нашего текста уходят целиком — места, плашка, иконка;
    опустевшая `p:grpSp` 4-й карточки уходит следом."""
    slide, ids = with_passport(["zt", "zs1", "zb1", "zs2", "zb2"])
    left = ids_left(slide)

    for n in (3, 4):
        for name in ("sub", "text", "plate", "icon"):
            assert ids[f"{name}{n}"] not in left, f"{name}{n} осталась"
    assert not slide.shapes._spTree.findall(".//{*}grpSp"), "пустая группа осталась"
    assert {ids["plate1"], ids["icon1"], ids["plate2"], ids["icon2"]} <= left


def test_a_single_group_without_text_leaves() -> None:
    """История 12: одиночная группа (широкая карточка) без текста уходит с плашкой и чертой."""
    slide, ids = with_passport(["zt", "zs1", "zb1"])
    left = ids_left(slide)

    assert not {ids["wide_text"], ids["wide_plate"], ids["wide_bar"]} & left


def _text_of(slide, xml_id: int) -> str:
    node = next(n for n in slide.shapes._spTree.iter() if n.tag.endswith("cNvPr")
                and n.get("id") == str(xml_id)).getparent().getparent()
    return "".join(t.text or "" for t in node.iter("{*}t"))


def test_a_row_of_four_with_two_items_keeps_the_first_two() -> None:
    """Истории 7, 12: IR заполнил 1-ю и 4-ю карточку ряда — остаются 1-я и 2-я, текст 4-й
    переехал во 2-ю место в место; 3-й и 4-й карточек нет."""
    slide, ids = with_passport(["zt", "zs1", "zb1", "zb4"])
    left = ids_left(slide)

    assert {ids["plate1"], ids["plate2"], ids["text1"], ids["text2"]} <= left
    assert _text_of(slide, ids["text1"]) == "Наш текст zb1"
    assert _text_of(slide, ids["text2"]) == "Наш текст zb4"
    for n in (3, 4):
        for name in ("sub", "text", "plate", "icon"):
            assert ids[f"{name}{n}"] not in left, f"{name}{n} осталась"


def test_an_empty_place_of_a_filled_card_leaves() -> None:
    """История 8: карточка заполнена текстом без подзаголовка — фигуры подзаголовка нет
    (текст шаблона «Подзаголовок» не остаётся), плашка карточки на месте."""
    slide, ids = with_passport(["zt", "zb1"])
    left = ids_left(slide)

    assert ids["sub1"] not in left
    assert {ids["plate1"], ids["icon1"], ids["text1"]} <= left


def test_an_empty_slide_title_stays_and_is_erased() -> None:
    """Истории 8, 9: блока для заголовка нет — рамка на месте, текста примера в ней нет."""
    slide, ids = with_passport(["zs1", "zb1"])

    assert ids["title"] in ids_left(slide)
    assert _text_of(slide, ids["title"]) == ""


def test_an_unfilled_author_table_leaves_with_a_passport() -> None:
    """История 11: таблица автора, которую рецепт не адресует, уходит и при паспорте."""
    slide, ids = with_passport(["zt", "zs1", "zb1"])

    assert ids["table"] not in ids_left(slide)


def test_the_workspace_cover_table_leaves_with_a_passport() -> None:
    """История 11 на настоящем шаблоне: титул WorkSpace `ex014` с паспортом — таблицы автора
    «Заголовок / Текст» (`graphicFrame` вне паспорта) на слайде нет, решение Насти 28.09."""
    prs = Presentation(str(case_template(WORKSPACE)))
    run = from_fixture(RUNS / "workspace")
    ds, _ = with_passports(run.design_system, run.manifest, FontLibrary.default())
    recipe = next(r for r in ds.recipes if r.recipe_id == "ex014")
    assert recipe.passport is not None, "у титула WorkSpace нет паспорта — тест не о том"
    slide = next(s for s in run.deck.slides if s.recipe_id == "ex014")

    written = clone_recipe(prs, recipe, slide)

    assert not written.shapes._spTree.findall(".//{*}graphicFrame")
