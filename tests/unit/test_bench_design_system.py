"""Замер дизайн-системы повторяется одной командой. Change `design-system-from-examples`.

Цифры из `docs/agents/tasks-design-system.md` приводятся в предложениях и отчётах.
Замер без теста через неделю расходится с кодом и становится преданием.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "scripts"))

from bench_design_system import REFERENCE_SLIDE_CX, measure
from tests.case_templates import case_template


@pytest.mark.parametrize(
    ("name", "examples", "outside", "font"),
    [
        ("VK Tech шаблон.pptx", 54, 0.90, "Play"),
        ("VK_WorkSpace_Клиентская_конференция_Шаблон_03.pptx", 29, 0.85, "Play"),
        ("Шаблон презентации VK Education.pptx", 55, 0.80, "Play"),
    ],
)
def test_the_table_of_the_task_is_reproduced(
    name: str, examples: int, outside: float, font: str
) -> None:
    row = measure(case_template(name))

    assert row.examples == examples
    assert row.outside_placeholders >= outside
    assert font in [family for family, _ in row.used_fonts], "гарнитура набора не найдена"
    assert row.theme_fonts == ["Arial"], "тема всех трёх шаблонов объявляет Arial"
    assert row.seconds < 10, "разбор шаблона дольше отведённых десяти секунд (D1)"


def test_sizes_are_brought_to_one_scale() -> None:
    """У VK Tech слайд 10″: без приведения его кегли несравнимы с остальными."""
    tech = measure(case_template("VK Tech шаблон.pptx"))
    education = measure(case_template("Шаблон презентации VK Education.pptx"))

    assert tech.slide_scale > 1.3, "масштаб слайда не замечен"
    assert education.slide_scale == 1.0
    assert any("кегли приведены" in note for note in tech.notes)
    assert max(tech.sizes_pt) > 20, "после приведения кегли VK Tech сравнимы с обычными"


def test_a_template_without_examples_says_so(tmp_path: Path) -> None:
    """Холодный шаблон (правило 10): примеров нет, и это не ошибка замера."""
    from pptx import Presentation

    path = tmp_path / "cold.pptx"
    Presentation().save(path)

    row = measure(path)

    assert row.examples == 0
    assert row.shapes == 0
    assert any("примеров нет" in note for note in row.notes)
    assert row.slide_scale == pytest.approx(REFERENCE_SLIDE_CX / 9144000, rel=0.01)
