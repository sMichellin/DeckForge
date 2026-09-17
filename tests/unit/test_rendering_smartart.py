"""Составные компоненты из автофигур и коннекторов. Change (21) `smartart-icons`."""

from __future__ import annotations

from pptx import Presentation
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from deckforge.domain.enums import ColorRef
from deckforge.domain.slide import SmartArtBlock
from deckforge.domain.template import TemplateManifest
from deckforge.domain.units import EMU_PER_CM, EMU_PER_PT
from deckforge.rendering.smartart import add_smartart, text_on


def render(
    manifest: TemplateManifest, pattern: str, items: list[str], **update: object
) -> object:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    block = SmartArtBlock(block_id="sa", pattern=pattern, items=items,
                          x=2 * EMU_PER_CM, y=5 * EMU_PER_CM, cx=24 * EMU_PER_CM,
                          cy=8 * EMU_PER_CM, **update)
    return add_smartart(slide, block, manifest, size_pt=18, text_color=ColorRef.DK2)


def parts(group: object) -> tuple[list[object], list[object]]:
    shapes = list(group.shapes)  # type: ignore[attr-defined]
    connectors = [s for s in shapes if s.shape_type == MSO_SHAPE_TYPE.LINE]
    return [s for s in shapes if s not in connectors], connectors


def geometry(shape: object) -> str:
    return shape._element.spPr.find(qn("a:prstGeom")).get("prst")  # type: ignore[attr-defined]


def test_process_is_a_group_of_editable_steps_and_arrows(manifest: TemplateManifest) -> None:
    group = render(manifest, "process", ["Сбор", "Анализ", "Решение"])
    assert group.shape_type == MSO_SHAPE_TYPE.GROUP  # type: ignore[attr-defined]
    nodes, arrows = parts(group)
    assert [n.text_frame.text for n in nodes] == ["Сбор", "Анализ", "Решение"]
    assert {geometry(n) for n in nodes} == {"roundRect"}
    assert len(arrows) == 2
    for arrow in arrows:
        tail = arrow._element.spPr.find(qn("a:ln")).find(qn("a:tailEnd"))
        assert tail is not None and tail.get("type") == "triangle"


def test_text_uses_the_fitted_size_and_theme_references(manifest: TemplateManifest) -> None:
    group = render(manifest, "process", ["Сбор", "Анализ"])
    xml = group._element.xml  # type: ignore[attr-defined]
    assert "srgbClr" not in xml
    assert "<p:style>" not in xml
    nodes, _ = parts(group)
    run = nodes[0].text_frame.paragraphs[0].runs[0]
    assert run.font.size == 18 * EMU_PER_PT
    assert run.font.name == "+mn-lt"
    assert nodes[0].text_frame.word_wrap is True


def test_nodes_take_color_refs_in_turn(manifest: TemplateManifest) -> None:
    group = render(manifest, "cycle", ["а", "б", "в"],
                   color_refs=[ColorRef.ACCENT1, ColorRef.ACCENT3])
    nodes, arrows = parts(group)
    fills = [n._element.spPr.find(qn("a:solidFill")).find(qn("a:schemeClr")).get("val")
             for n in nodes]
    assert fills == ["accent1", "accent3", "accent1"]
    assert {geometry(n) for n in nodes} == {"roundRect"}
    assert len(arrows) == 3


def test_text_on_a_fill_takes_the_more_contrasting_of_dark_and_light(
    manifest: TemplateManifest,
) -> None:
    assert text_on(ColorRef.ACCENT1, manifest) is ColorRef.LT1
    assert text_on(ColorRef.ACCENT3, manifest) is ColorRef.DK1
    group = render(manifest, "process", ["а", "б"], color_refs=[ColorRef.ACCENT3])
    nodes, _ = parts(group)
    run = nodes[0].text_frame.paragraphs[0].runs[0]
    assert run.font.color.theme_color == MSO_THEME_COLOR.DARK_1


def test_timeline_labels_sit_under_markers_in_the_layout_text_color(
    manifest: TemplateManifest,
) -> None:
    group = render(manifest, "timeline", ["2024", "2025", "2026"])
    shapes, lines = parts(group)
    markers = [s for s in shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]
    labels = [s for s in shapes if s.shape_type == MSO_SHAPE_TYPE.TEXT_BOX]
    assert len(markers) == len(labels) == 3
    assert all(not m.text_frame.text for m in markers)
    assert [label.text_frame.text for label in labels] == ["2024", "2025", "2026"]
    for marker, label in zip(markers, labels, strict=True):
        assert label.top >= marker.top + marker.height
    (axis,) = lines
    assert axis._element.spPr.find(qn("a:ln")).find(qn("a:tailEnd")) is None
    run = labels[0].text_frame.paragraphs[0].runs[0]
    assert run.font.color.theme_color == MSO_THEME_COLOR.DARK_2  # цвет текста макета
