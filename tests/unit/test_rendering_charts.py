"""Нативные диаграммы и таблицы с цветами темы. Change (14) `native-charts-tables`."""

from __future__ import annotations

import re

import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Emu

from deckforge.domain.content import Dataset, Series
from deckforge.domain.enums import ChartType, ColorRef, TextRole
from deckforge.domain.slide import ChartBlock, TableBlock
from deckforge.domain.template import ChartDefaults, TemplateManifest
from deckforge.rendering.charts import CHART_TYPES, add_chart, chart_problem
from deckforge.rendering.tables import add_table, template_table_style

BOX = {"x": 100_000, "y": 100_000, "cx": 6_000_000, "cy": 3_000_000}


def dataset(**update: object) -> Dataset:
    base = Dataset(
        dataset_id="d001", title="Выручка", categories=["2024", "2025", "2026"],
        series=[Series(name="Россия", values=[1.0, 2.0, 3.0]),
                Series(name="СНГ", values=[0.5, 0.7, None])],
        unit="млн ₽",
    )
    return base.model_copy(update=update)


def blank_slide() -> object:
    prs = Presentation()
    return prs.slides.add_slide(prs.slide_layouts[6])


def chart_xml(frame: object) -> str:
    return etree.tostring(frame.chart._chartSpace, encoding="unicode")  # type: ignore[attr-defined]


def scheme_colors(xml: str) -> list[str]:
    return re.findall(r'schemeClr val="(\w+)"', xml)


# --- chart_problem ---------------------------------------------------------------


def test_every_chart_type_has_a_native_counterpart() -> None:
    assert set(CHART_TYPES) == set(ChartType)


def test_regular_dataset_can_be_charted() -> None:
    assert chart_problem(ChartType.CLUSTERED_COLUMN, dataset()) is None


@pytest.mark.parametrize(
    ("chart_type", "update", "reason"),
    [
        (ChartType.CLUSTERED_BAR, {"categories": []}, "нет категорий"),
        (ChartType.CLUSTERED_BAR, {"series": []}, "нет серий"),
        (ChartType.LINE, {"series": [Series(name="а", values=[1.0])]}, "длин"),
        (ChartType.AREA, {"series": [Series(name="а", values=[None, None, None])]}, "значени"),
        (ChartType.PIE, {}, "одна серия"),
        (ChartType.DOUGHNUT, {"series": [Series(name="а", values=[1.0, -2.0, 3.0])]},
         "отрицательн"),
        (ChartType.SCATTER, {"categories": ["янв", "фев", "мар"]}, "числ"),
        (ChartType.SCATTER, {"categories": ["1", "nan", "3"]}, "числ"),
        (ChartType.CLUSTERED_COLUMN,
         {"series": [Series(name="а", values=[1.0, float("nan"), 2.0])]}, "конечн"),
        (ChartType.LINE, {"series": [Series(name="а", values=[1.0, float("inf"), 2.0])]},
         "конечн"),
    ],
)
def test_chart_problems_are_named(chart_type: ChartType, update: dict[str, object],
                                  reason: str) -> None:
    problem = chart_problem(chart_type, dataset(**update))
    assert problem is not None
    assert reason in problem


def test_scatter_with_numeric_categories_is_fine() -> None:
    assert chart_problem(ChartType.SCATTER, dataset(categories=["1", "2,5", "4"])) is None


# --- add_chart ---------------------------------------------------------------------


def test_series_are_painted_with_theme_accents(manifest: TemplateManifest) -> None:
    block = ChartBlock(block_id="c", chart_type=ChartType.CLUSTERED_COLUMN, dataset_ref="d001",
                       **BOX)
    xml = chart_xml(add_chart(blank_slide(), block, dataset(), manifest))
    assert "srgbClr" not in xml
    assert scheme_colors(xml)[:2] == ["accent1", "accent2"]


def test_block_colors_win_over_template_defaults(manifest: TemplateManifest) -> None:
    block = ChartBlock(block_id="c", chart_type=ChartType.STACKED_BAR, dataset_ref="d001",
                       series_color_refs=[ColorRef.ACCENT4, ColorRef.DK2], **BOX)
    xml = chart_xml(add_chart(blank_slide(), block, dataset(), manifest))
    assert scheme_colors(xml)[:2] == ["accent4", "dk2"]


def test_template_chart_defaults_are_used(manifest: TemplateManifest) -> None:
    with_defaults = manifest.model_copy(
        update={"chart_defaults": ChartDefaults(series_color_refs=[ColorRef.ACCENT3])}
    )
    block = ChartBlock(block_id="c", chart_type=ChartType.LINE_MARKERS, dataset_ref="d001", **BOX)
    xml = chart_xml(add_chart(blank_slide(), block, dataset(), with_defaults))
    assert set(scheme_colors(xml)) == {"accent3"}


def test_pie_paints_each_slice(manifest: TemplateManifest) -> None:
    single = dataset(series=[Series(name="Доля", values=[50.0, 30.0, 20.0])])
    block = ChartBlock(block_id="c", chart_type=ChartType.DOUGHNUT, dataset_ref="d001", **BOX)
    xml = chart_xml(add_chart(blank_slide(), block, single, manifest))
    assert scheme_colors(xml)[:3] == ["accent1", "accent2", "accent3"]


def test_unit_becomes_the_value_axis_title(manifest: TemplateManifest) -> None:
    """Без подписи единиц аудит выдаст `integrity.chart_labels_missing`."""
    block = ChartBlock(block_id="c", chart_type=ChartType.CLUSTERED_COLUMN, dataset_ref="d001",
                       **BOX)
    frame = add_chart(blank_slide(), block, dataset(), manifest)
    axis = frame.chart.value_axis  # type: ignore[attr-defined]
    assert axis.has_title
    assert axis.axis_title.text_frame.text == "млн ₽"


def test_legend_and_labels_follow_the_block(manifest: TemplateManifest) -> None:
    block = ChartBlock(block_id="c", chart_type=ChartType.CLUSTERED_COLUMN, dataset_ref="d001",
                       legend=False, data_labels=False, **BOX)
    chart = add_chart(blank_slide(), block, dataset(), manifest).chart  # type: ignore[attr-defined]
    assert chart.has_legend is False
    assert chart.plots[0].has_data_labels is False


def test_chart_text_uses_caption_size_and_theme_font(manifest: TemplateManifest) -> None:
    caption = manifest.typography(TextRole.CAPTION)
    assert caption is not None
    block = ChartBlock(block_id="c", chart_type=ChartType.CLUSTERED_COLUMN, dataset_ref="d001",
                       **BOX)
    chart = add_chart(blank_slide(), block, dataset(), manifest).chart  # type: ignore[attr-defined]
    assert chart.font.size.pt == caption.size_pt
    assert chart.font.name == "+mn-lt"


def test_chart_text_takes_the_layout_text_color(manifest: TemplateManifest) -> None:
    block = ChartBlock(block_id="c", chart_type=ChartType.CLUSTERED_COLUMN, dataset_ref="d001",
                       **BOX)
    frame = add_chart(blank_slide(), block, dataset(), manifest, text_color=ColorRef.LT1)
    text_properties = frame.chart._chartSpace.find(  # type: ignore[attr-defined]
        "{http://schemas.openxmlformats.org/drawingml/2006/chart}txPr"
    )
    assert '<a:schemeClr val="lt1"/>' in etree.tostring(text_properties, encoding="unicode")


def test_scatter_is_built_from_numeric_categories(manifest: TemplateManifest) -> None:
    data = dataset(categories=["1", "2,5", "4"])
    block = ChartBlock(block_id="c", chart_type=ChartType.SCATTER, dataset_ref="d001", **BOX)
    xml = chart_xml(add_chart(blank_slide(), block, data, manifest))
    assert "scatterChart" in xml
    assert "srgbClr" not in xml


# --- tables ------------------------------------------------------------------------


def test_table_cells_flags_and_size() -> None:
    block = TableBlock(block_id="t", header=["Год", "Выручка"], rows=[["2025", "2"]],
                       banding=False, **BOX)
    frame = add_table(blank_slide(), block, [["Год", "Выручка"], ["2025", "2"]],
                      size_pt=12, style_id=None)
    table = frame.table  # type: ignore[attr-defined]
    assert [[c.text for c in row.cells] for row in table.rows] == [["Год", "Выручка"],
                                                                   ["2025", "2"]]
    assert table.first_row is True
    assert table.horz_banding is False
    sizes = {r.font.size.pt for row in table.rows for c in row.cells
             for p in c.text_frame.paragraphs for r in p.runs}
    assert sizes == {12}
    assert Emu(frame.width) == BOX["cx"]  # type: ignore[attr-defined]


def test_table_rows_get_the_measured_heights_and_theme_font() -> None:
    block = TableBlock(block_id="t", header=["Год", ""], rows=[["2025", "2"]], **BOX)
    frame = add_table(blank_slide(), block, [["Год", ""], ["2025", "2"]], size_pt=12,
                      style_id=None, row_heights=[500_000, 300_000], font_token="+mn-lt")
    table = frame.table  # type: ignore[attr-defined]
    assert [row.height for row in table.rows] == [500_000, 300_000]
    xml = frame._element.xml  # type: ignore[attr-defined]
    assert 'typeface="+mn-lt"' in xml
    # Пустая ячейка без прогонов: без endParaRPr PowerPoint взял бы высоту строки от 18 pt.
    assert re.search(r'<a:endParaRPr[^>]*sz="1200"', xml)


def test_table_without_header_does_not_style_the_first_row() -> None:
    block = TableBlock(block_id="t", rows=[["а"], ["б"]], **BOX)
    frame = add_table(blank_slide(), block, [["а"], ["б"]], size_pt=12, style_id=None,
                      has_header=False)
    assert frame.table.first_row is False  # type: ignore[attr-defined]


def test_line_markers_are_outlined_in_the_series_color(manifest: TemplateManifest) -> None:
    block = ChartBlock(block_id="c", chart_type=ChartType.LINE_MARKERS, dataset_ref="d001",
                       series_color_refs=[ColorRef.ACCENT5], **BOX)
    xml = chart_xml(add_chart(blank_slide(), block, dataset(), manifest))
    marker = xml.split("<c:marker>", 1)[1].split("</c:marker>", 1)[0]
    assert marker.count('<a:schemeClr val="accent5"/>') == 2


def test_template_table_style_is_applied() -> None:
    style = "{073A0DAA-6AF3-43AB-8588-CEC1D06C72B9}"
    block = TableBlock(block_id="t", header=["а"], rows=[["б"]], **BOX)
    frame = add_table(blank_slide(), block, [["а"], ["б"]], size_pt=12, style_id=style)
    assert f"<a:tableStyleId>{style}</a:tableStyleId>" in frame._element.xml  # type: ignore[attr-defined]


def test_template_default_table_style_is_read_from_the_package() -> None:
    prs = Presentation()
    assert template_table_style(prs) is not None
