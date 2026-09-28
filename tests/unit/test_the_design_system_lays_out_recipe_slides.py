"""Дизайн-система подгоняет рецептный слайд. Change `the-design-system-lays-out-recipe-slides`
(план Б, 5а; issue #244).

После change 4 от ряда из четырёх карточек при двух пунктах остаются две — и дыра на месте
снятых. Теперь при паспорте оставшиеся карточки линии встают на всю ширину ряда с отступом шага
сетки ДС, наш кегль не ниже порога читаемости (подписи VK Tech `ex018` — 9 pt при пороге 10),
а цвет нашего текста на плашке читается. Без паспорта — всё как было, байт в байт.

Шов — `clone_recipe` (синтетический пример python-pptx и настоящий шаблон) и `PptxWriter.write`.
Сценарии — из дельты `openspec/changes/the-design-system-lays-out-recipe-slides/specs/pptx-writer/`.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from deckforge.composition.passport import probe_size, with_passports
from deckforge.designsystem import derive
from deckforge.designsystem.models import (
    ExamplePassport,
    Place,
    PlaceGroup,
    PlaceKind,
    Recipe,
    RecipeKind,
    TypeLevel,
    Zone,
)
from deckforge.domain.enums import ColorRef
from deckforge.domain.rules import contrast_ratio
from deckforge.domain.slide import DeckIR
from deckforge.domain.template import TemplateManifest
from deckforge.layout.by_design import DesignRules
from deckforge.layout.fonts import FontLibrary
from deckforge.pipeline.replay import from_fixture
from deckforge.rendering.recipe_slide import clone_recipe
from deckforge.rendering.writer import PptxWriter
from tests.case_templates import case_template
from tests.unit.test_recipe_leaves_no_sample_text import slide_ir
from tests.unit.test_the_writer_removes_whole_groups import (
    GOLDEN,
    RUNS,
    VK_TECH,
    cards_example,
    cards_recipe,
    slide_xml,
)

#: Карточка: плашка 2 000 000 × 2 400 000 EMU, шаг ряда 2 200 000, линии через 2 600 000.
LEFT, TOP, STEP, LINE_STEP = 300000, 1200000, 2200000, 2600000
PLATE_CX, PLATE_CY = 2000000, 2400000
#: Масштаб карточки в `p:grpSp`: координаты детей вдвое крупнее слайдовых.
SCALE = 2


def _text(shapes, x, y, cx, cy, text):
    box = shapes.add_textbox(Emu(x), Emu(y), Emu(cx), Emu(cy))
    box.text_frame.text = text
    return box


def _card(shapes, x, y, *, scale=1):
    """Плашка, иконка на ней, подзаголовок и текст; `x, y` — угол плашки в своих координатах."""
    def s(v):
        return v * scale

    return {
        "plate": shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(s(PLATE_CX)),
                                  Emu(s(PLATE_CY))),
        "icon": shapes.add_shape(MSO_SHAPE.OVAL, Emu(x + s(100000)), Emu(y + s(100000)),
                                 Emu(s(300000)), Emu(s(300000))),
        "sub": _text(shapes, x + s(100000), y + s(500000), s(1800000), s(400000), "Подзаголовок"),
        "text": _text(shapes, x + s(100000), y + s(1000000), s(1800000), s(1200000), "Текст"),
    }


def grid_example(lines=(4,), grouped=None):
    """Ряд карточек по линиям (`lines` — сколько карточек в линии), слайд 9144000 × 6858000.

    Карточка номер `grouped` собрана в `p:grpSp` с масштабом 1:`SCALE` — её фигуры записаны
    в координатах группы, а на слайде стоят там же, где стояли бы без группы.
    """
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shapes = slide.shapes
    found = {"title": _text(shapes, 400000, 200000, 8000000, 600000, "Заголовок примера")}
    cards = []
    for line, count in enumerate(lines):
        for column in range(count):
            n = len(cards) + 1
            x, y = LEFT + column * STEP, TOP + line * LINE_STEP
            if n != grouped:
                cards.append(_card(shapes, x, y))
                continue
            group = shapes.add_group_shape()
            cards.append(_card(group.shapes, 0, 0, scale=SCALE))
            xfrm = group._element.find(f"{qn('p:grpSpPr')}/{qn('a:xfrm')}")
            for tag, cx, cy in (("a:off", x, y), ("a:ext", PLATE_CX, PLATE_CY),
                                ("a:chOff", 0, 0), ("a:chExt", PLATE_CX * SCALE,
                                                    PLATE_CY * SCALE)):
                node = xfrm.find(qn(tag))
                keys = ("x", "y") if tag.endswith("Off") else ("cx", "cy")
                node.set(keys[0], str(cx))
                node.set(keys[1], str(cy))
    return prs, str(slide.part.partname), found, cards


def _zone(zone_id, shape, role) -> Zone:
    return Zone(zone_id=zone_id, xml_id=shape.shape_id, role=role, capacity_chars=60)


def _place(number: int, zone: Zone) -> Place:
    return Place(place_id=f"p{number:02d}", kind=PlaceKind.TEXT, zone_id=zone.zone_id,
                 xml_id=zone.xml_id, role=zone.role, capacity_chars=60)


def grid_recipe(part_name: str, found, cards, lines=(4,)) -> Recipe:
    """Рецепт с паспортом: g01 — заголовок; g02… — ряд `r1`, место 1 — подзаголовок,
    место 2 — текст, декор — плашка и иконка; рамка группы — рамка плашки на слайде."""
    title = _zone("zt", found["title"], TypeLevel.SLIDE_TITLE)
    zones = [title]
    groups = [PlaceGroup(group_id="g01", places=[_place(1, title)])]
    frames = [(LEFT + column * STEP, TOP + line * LINE_STEP)
              for line, count in enumerate(lines) for column in range(count)]
    for n, (card, (x, y)) in enumerate(zip(cards, frames, strict=True), start=1):
        sub = _zone(f"zs{n}", card["sub"], TypeLevel.CARD_TITLE)
        body = _zone(f"zb{n}", card["text"], TypeLevel.BODY)
        zones += [sub, body]
        groups.append(PlaceGroup(
            group_id=f"g{n + 1:02d}", places=[_place(2 * n, sub), _place(2 * n + 1, body)],
            decor_xml_ids=[card["plate"].shape_id, card["icon"].shape_id], row="r1",
            x=x, y=y, cx=PLATE_CX, cy=PLATE_CY,
        ))
    return Recipe(
        recipe_id="ex018", example_index=1, part_name=part_name, kind=RecipeKind.CARDS,
        zones=zones, passport=ExamplePassport(groups=groups),
    )


def test_design_without_a_passport_is_not_read(manifest: TemplateManifest) -> None:
    """История 15 / 4: `design` передан, паспорта нет — слайд равен эталону `legacy` change 4."""
    prs, part_name, f = cards_example()
    recipe = cards_recipe(part_name, f, passport=False)
    ir = slide_ir(recipe, ["zt", "zs1", "zb1", "zb4"])
    slide = clone_recipe(prs, recipe, ir, DesignRules(manifest))
    assert slide_xml(slide) == (GOLDEN / "synthetic-legacy.xml").read_bytes()


#: Шаг сетки ДС синтетического манифеста (`tests/conftest.py`): отступ колонок — полсантиметра.
GAP = 180000
#: Ряд синтетики: от левого края первой до правого края четвёртой карточки.
ROW_LEFT, ROW_RIGHT = LEFT, LEFT + 3 * STEP + PLATE_CX


def own_box(slide, shape) -> tuple[int, int, int, int]:
    """`a:off`/`a:ext` фигуры в её собственных координатах (у детей `p:grpSp` — группы)."""
    node = next(n for n in slide.shapes._spTree.iter(qn("p:cNvPr"))
                if int(n.get("id")) == shape.shape_id).getparent().getparent()
    xfrm = node.find(f"{qn('p:spPr')}/{qn('a:xfrm')}")
    off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
    return int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy"))


def test_two_cards_of_four_take_the_whole_row(manifest: TemplateManifest) -> None:
    """История 16: 2 из 4 — две карточки на всю ширину ряда, отступ — шаг сетки ДС.

    Вторая карточка — в `p:grpSp` с масштабом 1:2: её координаты пишутся в системе группы.
    """
    prs, part_name, found, cards = grid_example(grouped=2)
    recipe = grid_recipe(part_name, found, cards)
    slide = clone_recipe(prs, recipe, slide_ir(recipe, ["zt", "zs1", "zb1", "zb4"]),
                         DesignRules(manifest))
    width = (ROW_RIGHT - ROW_LEFT - GAP) // 2
    first, second = ROW_LEFT, ROW_LEFT + width + GAP
    assert own_box(slide, cards[0]["plate"]) == (first, TOP, width, PLATE_CY)
    assert own_box(slide, cards[0]["text"]) == (first + 100000, TOP + 1000000,
                                                width - 200000, 1200000)
    # Иконка — мелкий декор: размер тот же, отступ от левого края плашки тот же.
    assert own_box(slide, cards[0]["icon"]) == (first + 100000, TOP + 100000, 300000, 300000)
    # Карточка в группе: слайдовая координата x ↦ (x − 2 200 000 − LEFT) · 2 от `a:off` группы.
    origin = LEFT + STEP
    assert own_box(slide, cards[1]["plate"]) == (
        (second - origin) * SCALE, 0, width * SCALE, PLATE_CY * SCALE)
    # Подзаголовка у второй карточки нет: из 4-й переехал только текст (change 4).
    assert own_box(slide, cards[1]["text"]) == (
        (second + 100000 - origin) * SCALE, 1000000 * SCALE, (width - 200000) * SCALE,
        1200000 * SCALE)
    assert own_box(slide, cards[1]["icon"]) == (
        (second + 100000 - origin) * SCALE, 100000 * SCALE, 300000 * SCALE, 300000 * SCALE)


def _grid_3_2(manifest: TemplateManifest, fill: list[str]):
    prs, part_name, found, cards = grid_example(lines=(3, 2))
    recipe = grid_recipe(part_name, found, cards, lines=(3, 2))
    slide = clone_recipe(prs, recipe, slide_ir(recipe, fill), DesignRules(manifest))
    present = {int(n.get("id")) for n in slide.shapes._spTree.iter(qn("p:cNvPr"))}
    return slide, cards, present


#: Ряд сетки 3 + 2 — ширина первой линии: от левого края 1-й до правого края 3-й карточки.
GRID_RIGHT = LEFT + 2 * STEP + PLATE_CX


def test_grid_3_2_the_lonely_card_of_the_second_line_takes_the_whole_row(
    manifest: TemplateManifest,
) -> None:
    """История 17: четыре пункта на сетке 3 + 2 — первая линия прежняя, единственная
    карточка второй линии — на всю ширину ряда; пятой карточки нет."""
    slide, cards, present = _grid_3_2(manifest, ["zt", "zb1", "zb2", "zb3", "zb4"])
    for n in range(3):
        assert own_box(slide, cards[n]["plate"]) == (LEFT + n * STEP, TOP, PLATE_CX, PLATE_CY)
    assert own_box(slide, cards[3]["plate"]) == (
        LEFT, TOP + LINE_STEP, GRID_RIGHT - LEFT, PLATE_CY)
    assert cards[4]["plate"].shape_id not in present


def test_grid_3_2_two_points_leave_one_line(manifest: TemplateManifest) -> None:
    """История 17: два пункта — две карточки первой линии на всю ширину, второй линии нет."""
    slide, cards, present = _grid_3_2(manifest, ["zt", "zb1", "zb2"])
    width = (GRID_RIGHT - LEFT - GAP) // 2
    assert own_box(slide, cards[0]["plate"]) == (LEFT, TOP, width, PLATE_CY)
    assert own_box(slide, cards[1]["plate"]) == (LEFT + width + GAP, TOP, width, PLATE_CY)
    assert not {cards[n][k].shape_id for n in (2, 3, 4) for k in ("plate", "icon")} & present


def _sizes(slide, shape) -> list[str | None]:
    node = next(n for n in slide.shapes._spTree.iter(qn("p:cNvPr"))
                if int(n.get("id")) == shape.shape_id).getparent().getparent()
    return [props.get("sz") for props in node.iter(qn("a:rPr"))]


def _caption_9pt(manifest: TemplateManifest, floor: float):
    """Две карточки из четырёх; текст карточек у автора — 9 pt (как подписи VK Tech `ex018`)."""
    prs, part_name, found, cards = grid_example()
    for card in cards:
        card["text"].text_frame.paragraphs[0].runs[0].font.size = Pt(9)
    recipe = grid_recipe(part_name, found, cards)
    ir = slide_ir(recipe, ["zt", "zs1", "zb1", "zs2", "zb2"])
    slide = clone_recipe(prs, recipe, ir, DesignRules(manifest, reading_floor_pt=floor))
    return slide, cards, recipe


@pytest.mark.parametrize(("floor", "step"), [(10.0, "1200"), (13.0, "1800")])
def test_a_caption_below_the_floor_takes_the_lowest_step_above_it(
    manifest: TemplateManifest, floor: float, step: str
) -> None:
    """История 19: подпись 9 pt при пороге 10 → 12 pt — наименьшая ступень шкалы
    40/24/18/12 не ниже порога (правило 6); при пороге 13 — 18 pt."""
    slide, cards, recipe = _caption_9pt(manifest, floor)
    for card in cards[:2]:
        assert _sizes(slide, card["text"]) == [step]
    # Стык с A (#249): паспорт меряет ёмкость этой зоны тем же кеглем, до которого писатель
    # поднял подпись, — иначе текст «ровно под место» в место не встанет.
    zone = next(z for z in recipe.zones if z.zone_id == "zb1").model_copy(update={"size_pt": 9.0})
    assert probe_size(zone, manifest, floor) == int(step) / 100


#: Светлая плашка синтетики — цвет автора примера, не шаблона кейса.
PLATE = RGBColor(0xEE, 0xF1, 0xF6)


def _on_plate(manifest: TemplateManifest, ink: RGBColor, notes: list[str] | None = None):
    """Две карточки из четырёх; плашки светлые, текст карточек у автора — цветом `ink`."""
    prs, part_name, found, cards = grid_example()
    for card in cards:
        card["plate"].fill.solid()
        card["plate"].fill.fore_color.rgb = PLATE
        card["text"].text_frame.paragraphs[0].runs[0].font.color.rgb = ink
    recipe = grid_recipe(part_name, found, cards)
    source = etree.tostring(_run_fill(prs.slides[0], cards[0]["text"]))
    slide = clone_recipe(prs, recipe, slide_ir(recipe, ["zt", "zb1", "zb2"]),
                         DesignRules(manifest), notes=notes)
    return slide, cards, source


def _run_fill(slide, shape):
    node = next(n for n in slide.shapes._spTree.iter(qn("p:cNvPr"))
                if int(n.get("id")) == shape.shape_id).getparent().getparent()
    return node.find(f".//{qn('a:rPr')}/{qn('a:solidFill')}")


def test_a_readable_author_ink_stays_byte_for_byte(manifest: TemplateManifest) -> None:
    """История 18: тёмный текст автора на светлой плашке читается — цвет тот же, байт в байт."""
    slide, cards, source = _on_plate(manifest, RGBColor(0x10, 0x10, 0x14))
    for card in cards[:2]:
        assert etree.tostring(_run_fill(slide, card["text"])) == source


def test_an_unreadable_author_ink_takes_a_readable_theme_slot(
    manifest: TemplateManifest,
) -> None:
    """История 18: светлый текст на светлой плашке — ссылка на слот темы, читаемый на ней."""
    slide, cards, _ = _on_plate(manifest, RGBColor(0xC8, 0xC8, 0xC8))
    for card in cards[:2]:
        fill = _run_fill(slide, card["text"])
        assert [child.tag for child in fill] == [qn("a:schemeClr")]
        slot = ColorRef(fill[0].get("val"))
        assert contrast_ratio(manifest.theme.colors.get(slot), "#EEF1F6") >= 4.5


def vk_tech_ex018(design: bool = True):
    """VK Tech `ex018` по первому слайду IR 28.09, который его берёт, с паспортом и ДС шаблона."""
    prs = Presentation(str(case_template(VK_TECH)))
    run = from_fixture(RUNS / "vk-tech")
    ds, _ = with_passports(run.design_system, run.manifest, FontLibrary.default())
    recipe = next(r for r in ds.recipes if r.recipe_id == "ex018")
    ir = next(s for s in run.deck.slides if s.recipe_id == "ex018")
    rules = DesignRules(run.manifest, ds) if design else None
    return clone_recipe(prs, recipe, ir, rules), run.manifest


#: Слоты ссылок: цвет гиперссылки на подписи выглядит кликабельным — это не цвет текста.
LINK_SLOTS = {ColorRef.HLINK.value, ColorRef.FOL_HLINK.value}


def test_a_caption_on_a_plate_is_not_inked_as_a_link() -> None:
    """История 18 на VK Tech `ex018`: серый автора на белой плашке не читается, и подпись
    получает читаемый слот темы — но не `hlink`/`folHlink`, хотя `hlink` ближе всех по цвету."""
    slide, manifest = vk_tech_ex018()
    slots = [fill[0].get("val") for fill in slide._element.iter(qn("a:solidFill"))
             if fill.getparent().tag == qn("a:rPr") and fill[0].tag == qn("a:schemeClr")]
    assert slots, "подписи на плашке не получили слота темы"
    assert not set(slots) & LINK_SLOTS
    for slot in slots:
        assert contrast_ratio(manifest.theme.colors.get(ColorRef(slot)), "#FFFFFF") >= 4.5


def test_the_writer_hands_its_design_rules_to_the_recipe(
    manifest: TemplateManifest, tmp_path: Path
) -> None:
    """История 15 через `PptxWriter.write`: писатель передаёт в `clone_recipe` свои ответы ДС —
    синтетический шаблон с паспортом, 2 пункта из 4, и две карточки на всю ширину ряда.
    Что ДС судить не смогла (цвет текста на плашке унаследован), — в `degradations` писателя."""
    prs, part_name, found, cards = grid_example()
    for card in cards:
        card["plate"].fill.solid()
        card["plate"].fill.fore_color.rgb = PLATE
    # Макеты python-pptx под именами синтетического манифеста — писатель сверяет их по имени.
    for spec in manifest.layouts:
        prs.slide_layouts[spec.index]._element.cSld.set("name", spec.name)
    template = tmp_path / "template.pptx"
    prs.save(str(template))
    recipe = grid_recipe(part_name, found, cards)
    ds = derive(manifest).model_copy(update={"recipes": [recipe]})
    ir = slide_ir(recipe, ["zt", "zb1", "zb2"]).model_copy(
        update={"layout_id": manifest.layouts[0].layout_id})
    deck = DeckIR(deck_id="d01", template_id=manifest.template_id, variant="A", seed=1,
                  slides=[ir])
    writer = PptxWriter(template, manifest, design_system=ds)
    out = writer.write(deck, tmp_path / "deck.pptx")
    slide = Presentation(str(out)).slides[0]
    width = (ROW_RIGHT - ROW_LEFT - GAP) // 2
    assert own_box(slide, cards[0]["plate"]) == (ROW_LEFT, TOP, width, PLATE_CY)
    assert own_box(slide, cards[1]["plate"]) == (ROW_LEFT + width + GAP, TOP, width, PLATE_CY)
    assert [note.split(":")[0] for note in writer.degradations] == ["s01/zb1", "s01/zb2"]


def _plate_by_scheme(prs, cards, plate_slot: str) -> None:
    """Плашки карточек — заливкой `a:schemeClr` по имени из карты цветов (`bg1`, `tx1`…)."""
    for card in cards:
        card["plate"].fill.solid()
        fill = card["plate"]._element.spPr.find(qn("a:solidFill"))
        for child in list(fill):
            fill.remove(child)
        etree.SubElement(fill, qn("a:schemeClr"), val=plate_slot)


def test_the_plate_is_read_through_the_color_map_of_the_master(
    manifest: TemplateManifest,
) -> None:
    """История 18: `bg1` плашки — слот по `p:clrMap` мастера, а не по стандартной карте.

    Карта обращена (`bg1 → dk1`): плашка тёмная, светлый текст автора на ней читается —
    цвет тот же байт в байт. По стандартной карте плашка была бы белой, и цвет бы сменился.
    """
    prs, part_name, found, cards = grid_example()
    clr_map = prs.slide_masters[0]._element.find(qn("p:clrMap"))
    for role, slot in (("bg1", "dk1"), ("tx1", "lt1"), ("bg2", "dk2"), ("tx2", "lt2")):
        clr_map.set(role, slot)
    _plate_by_scheme(prs, cards, "bg1")
    for card in cards:
        card["text"].text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xC8, 0xC8, 0xC8)
    source = etree.tostring(_run_fill(prs.slides[0], cards[0]["text"]))
    recipe = grid_recipe(part_name, found, cards)
    slide = clone_recipe(prs, recipe, slide_ir(recipe, ["zt", "zb1", "zb2"]),
                         DesignRules(manifest))
    for card in cards[:2]:
        assert etree.tostring(_run_fill(slide, card["text"])) == source


def test_an_inherited_ink_on_a_plate_is_named_not_judged(manifest: TemplateManifest) -> None:
    """История 18: у прогона нет своего `a:solidFill` — цвет он наследует, и какой именно,
    писатель не знает. Цвет не трогается, но и не молча: заметка называет слайд и зону."""
    prs, part_name, found, cards = grid_example()
    for card in cards:
        card["plate"].fill.solid()
        card["plate"].fill.fore_color.rgb = PLATE
    recipe = grid_recipe(part_name, found, cards)
    notes: list[str] = []
    slide = clone_recipe(prs, recipe, slide_ir(recipe, ["zt", "zb1", "zb2"]),
                         DesignRules(manifest), notes=notes)
    assert _run_fill(slide, cards[0]["text"]) is None
    assert [note.split(":")[0] for note in notes] == ["s01/zb1", "s01/zb2"]
    assert all("унаследован" in note for note in notes)


def test_no_readable_slot_but_a_link_keeps_the_author_ink_and_names_it(
    manifest: TemplateManifest,
) -> None:
    """История 18: читается на плашке только `hlink` — ссылкой подпись не красится, цвет автора
    остаётся, но не молча: заметка называет слайд, зону и причину."""
    light = {ref.value: "#D0D4DA" for ref in ColorRef}
    colors = manifest.theme.colors.model_copy(update={**light, "hlink": "#101014"})
    theme = manifest.theme.model_copy(update={"colors": colors})
    dim = manifest.model_copy(update={"theme": theme})
    notes: list[str] = []
    slide, cards, source = _on_plate(dim, RGBColor(0xC8, 0xC8, 0xC8), notes)
    for card in cards[:2]:
        assert etree.tostring(_run_fill(slide, card["text"])) == source
    assert [note.split(":")[0] for note in notes] == ["s01/zb1", "s01/zb2"]
    assert all("читаемого слота нет" in note for note in notes)
