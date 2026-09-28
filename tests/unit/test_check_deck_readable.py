"""Замер читаемости готовой колоды. Change `acceptance-is-measured`.

Колоды в репозитории нет и быть не должно, поэтому каждый случай собирается здесь же:
`python-pptx` создаёт слайд с одной надписью, у которой заданы рамка, кегль и текст.
Так тест проверяет ровно то правило, ради которого написан, и не зависит ни от прогона,
ни от шаблона.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))

from check_deck_readable import check_deck  # noqa: E402
from deckforge.layout.fonts import FontLibrary  # noqa: E402

EMU_PER_INCH = 914400


@pytest.fixture(scope="module")
def library() -> FontLibrary:
    return FontLibrary.default()


def _deck(
    path: Path,
    text: str,
    *,
    left: int = EMU_PER_INCH,
    top: int = EMU_PER_INCH,
    width: int = 4 * EMU_PER_INCH,
    height: int = 2 * EMU_PER_INCH,
    size_pt: float | None = 18.0,
) -> Path:
    """Колода из одного слайда с одной надписью. `size_pt=None` — кегль не задан."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    if size_pt is not None:
        run.font.size = Pt(size_pt)
        run.font.name = "Arial"
    presentation.save(str(path))
    return path


def _kinds(path: Path, library: FontLibrary, **kwargs: float) -> list[str]:
    report = check_deck(path, size_floor_pt=kwargs.get("size_floor_pt", 10.0), library=library)
    return [finding.kind for finding in report.findings]


def test_shape_inside_the_slide_is_not_a_finding(tmp_path: Path, library: FontLibrary) -> None:
    deck = _deck(tmp_path / "ok.pptx", "Короткая строка")
    assert _kinds(deck, library) == []


def test_shape_outside_the_slide_is_found(tmp_path: Path, library: FontLibrary) -> None:
    deck = _deck(tmp_path / "out.pptx", "Текст", left=12 * EMU_PER_INCH)
    assert "за краем слайда" in _kinds(deck, library)


def test_word_wider_than_its_frame_is_found(tmp_path: Path, library: FontLibrary) -> None:
    deck = _deck(
        tmp_path / "wide.pptx", "Автоматизация", width=EMU_PER_INCH, size_pt=72.0
    )
    assert "слово шире рамки" in _kinds(deck, library)


def test_same_word_in_a_wide_frame_is_not_found(tmp_path: Path, library: FontLibrary) -> None:
    deck = _deck(
        tmp_path / "narrow.pptx", "Автоматизация", width=8 * EMU_PER_INCH, size_pt=12.0
    )
    assert "слово шире рамки" not in _kinds(deck, library)


def test_text_taller_than_its_frame_is_found(tmp_path: Path, library: FontLibrary) -> None:
    deck = _deck(
        tmp_path / "tall.pptx",
        "Строка текста " * 40,
        height=EMU_PER_INCH // 2,
        size_pt=36.0,
    )
    assert "текст выше рамки" in _kinds(deck, library)


def test_text_that_fits_its_frame_is_not_found(tmp_path: Path, library: FontLibrary) -> None:
    deck = _deck(tmp_path / "fits.pptx", "Одна строка", size_pt=12.0)
    assert "текст выше рамки" not in _kinds(deck, library)


def test_size_below_the_floor_is_found(tmp_path: Path, library: FontLibrary) -> None:
    deck = _deck(tmp_path / "small.pptx", "Мелко", size_pt=7.8)
    assert "кегль ниже порога" in _kinds(deck, library)


def test_size_above_the_floor_is_not_found(tmp_path: Path, library: FontLibrary) -> None:
    deck = _deck(tmp_path / "big.pptx", "Нормально", size_pt=12.0)
    assert "кегль ниже порога" not in _kinds(deck, library)


def test_inherited_size_is_counted_not_reported(tmp_path: Path, library: FontLibrary) -> None:
    """Кегль от макета из слайда не виден: такой прогон уходит в «не измерено».

    Замалчивать его нельзя — замер, молча пропускающий половину колоды, хуже
    отсутствующего.
    """
    deck = _deck(tmp_path / "inherited.pptx", "Текст без кегля", size_pt=None)
    report = check_deck(deck, size_floor_pt=10.0, library=library)
    assert report.unmeasured == 1
    assert report.findings == []


def test_speaker_notes_are_counted(tmp_path: Path, library: FontLibrary) -> None:
    deck = _deck(tmp_path / "notes.pptx", "Текст")
    presentation = Presentation(str(deck))
    presentation.slides[0].notes_slide.notes_text_frame.text = "Начать с вывода"
    presentation.save(str(deck))
    report = check_deck(deck, size_floor_pt=10.0, library=library)
    assert report.notes == 1


def test_left_edges_are_counted(tmp_path: Path, library: FontLibrary) -> None:
    deck = _deck(tmp_path / "edges.pptx", "Первый")
    presentation = Presentation(str(deck))
    slide = presentation.slides[0]
    second = slide.shapes.add_textbox(
        Emu(3 * EMU_PER_INCH), Emu(4 * EMU_PER_INCH), Emu(EMU_PER_INCH), Emu(EMU_PER_INCH)
    )
    second.text_frame.text = "Второй"
    presentation.save(str(deck))
    report = check_deck(deck, size_floor_pt=10.0, library=library)
    assert len(report.left_edges) == 2


# --- наложение меряется текстом (change `the-overlap-check-measures-the-text`, RG47) ---

SHORT = "Короткая строка"
TALL = "\n".join(["Строка тела слайда"] * 6)


def _two_boxes(
    path: Path,
    first_text: str,
    second_text: str,
    *,
    first_top: int,
    first_height: int,
    second_top: int,
    second_height: int,
) -> Path:
    """Колода из одного слайда с двумя надписями одной ширины, одна под другой."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for text, top, height in (
        (first_text, first_top, first_height),
        (second_text, second_top, second_height),
    ):
        box = slide.shapes.add_textbox(
            Emu(EMU_PER_INCH), Emu(top), Emu(4 * EMU_PER_INCH), Emu(height)
        )
        run = box.text_frame.paragraphs[0].add_run()
        run.text = text
        run.font.size = Pt(18.0)
        run.font.name = "Arial"
    presentation.save(str(path))
    return path


def _collisions_of(path: Path, library: FontLibrary) -> list[str]:
    report = check_deck(path, size_floor_pt=10.0, library=library)
    return [f.message for f in report.findings if f.kind == "тексты наложились"]


def test_overlapping_frames_with_separate_texts_are_not_a_finding(
    tmp_path: Path, library: FontLibrary
) -> None:
    """Нарушитель RG47: рамки перекрываются, тексты — нет.

    Рамки зон в шаблонах перекрываются у самих авторов, и автору это не мешало: его текст
    короче рамки. Прежняя формула брала большее из рамки и текста и обвиняла шаблон
    в наложении, которого на слайде нет (WorkSpace s05).
    """
    deck = _two_boxes(
        tmp_path / "frames.pptx",
        SHORT,
        SHORT,
        first_top=EMU_PER_INCH,
        first_height=3 * EMU_PER_INCH,
        second_top=2 * EMU_PER_INCH,
        second_height=2 * EMU_PER_INCH,
    )

    assert _collisions_of(deck, library) == []


def test_overlapping_texts_are_a_finding(tmp_path: Path, library: FontLibrary) -> None:
    """Норма: настоящее наложение по-прежнему находится.

    У верхней надписи шесть строк, и её текст заходит на место нижней больше чем
    на полстроки — буквы легли на буквы.
    """
    deck = _two_boxes(
        tmp_path / "texts.pptx",
        TALL,
        SHORT,
        first_top=EMU_PER_INCH,
        first_height=3 * EMU_PER_INCH,
        second_top=2 * EMU_PER_INCH,
        second_height=EMU_PER_INCH,
    )

    assert _collisions_of(deck, library), "настоящее наложение пропущено"


def test_a_shape_without_a_size_is_not_counted_for_overlap(
    tmp_path: Path, library: FontLibrary
) -> None:
    """Норма: у фигуры с унаследованным кеглем текст не померить.

    Такая фигура в счёт наложения не идёт и уходит в справку — замалчивать её нельзя,
    но и судить по её рамке значило бы вернуть ту же ошибку.
    """
    deck = _deck(tmp_path / "inherited-place.pptx", SHORT, size_pt=None)
    report = check_deck(deck, size_floor_pt=10.0, library=library)

    assert report.unmeasured_place == 1
    assert [f for f in report.findings if f.kind == "тексты наложились"] == []
