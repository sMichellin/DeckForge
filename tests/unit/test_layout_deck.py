"""Колода «по слайду на макет»: каждый слайд — на своём макете. Change (6), исправление C9.

LibreOffice здесь не нужен: проверяется сама колода, а не её картинки.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from deckforge.parsing import TemplateParser
from deckforge.rendering.layout_deck import build_layout_deck


def _template_with_an_empty_layout_first(path: Path) -> Path:
    """Первый макет без плейсхолдеров: парсер его пропускает, python-pptx — нет.

    Ровно так устроен VK Tech: индекс манифеста и позиция в списке python-pptx
    расходятся, и превью по индексу уезжали к чужим макетам (прогон 9ce69f5cbd39).
    """
    prs = Presentation()
    first = prs.slide_masters[0].slide_layouts[0]
    for placeholder in list(first.placeholders):
        element = placeholder._element
        element.getparent().remove(element)
    prs.save(str(path))
    return path


def test_each_slide_sits_on_the_layout_it_is_named_after(tmp_path: Path) -> None:
    template = _template_with_an_empty_layout_first(tmp_path / "tpl.pptx")
    manifest = TemplateParser().parse(template, use_cache=False)
    first_part = "ppt/slideLayouts/slideLayout1.xml"
    assert first_part not in {spec.part_name.lstrip("/") for spec in manifest.layouts}, (
        "предпосылка теста: пустой макет парсер пропускает"
    )

    out = tmp_path / "layouts.pptx"
    order = build_layout_deck(template, manifest, out)

    deck = Presentation(str(out))
    by_id = {spec.layout_id: spec for spec in manifest.layouts}
    for number, slide in enumerate(deck.slides):
        used = str(slide.slide_layout.part.partname).lstrip("/")
        assert used == by_id[order[number]].part_name.lstrip("/"), (
            f"слайд {number} назван {order[number]}, а стоит на другом макете"
        )
    assert len(order) == len(manifest.layouts)


def test_only_keeps_the_mapping_right(tmp_path: Path) -> None:
    template = _template_with_an_empty_layout_first(tmp_path / "tpl.pptx")
    manifest = TemplateParser().parse(template, use_cache=False)
    wanted = [manifest.layouts[-1].layout_id]

    out = tmp_path / "layouts.pptx"
    order = build_layout_deck(template, manifest, out, only=wanted)

    (slide,) = Presentation(str(out)).slides
    assert order == {0: wanted[0]}
    assert str(slide.slide_layout.part.partname).lstrip("/") == (
        manifest.layouts[-1].part_name.lstrip("/")
    )
