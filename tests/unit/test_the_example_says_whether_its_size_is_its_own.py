"""Кегль зоны — число автора или наша догадка. Таск RG42 (`docs/agents/tasks-25-09.md`).

Разбор RG39 упёрся в решение **D04**: прогон примера без `sz` писатель не трогает,
«настоящего числа он не знает». Осторожность верная, но причина была шире правды.

Источников кегля два, и оба авторские: свой кегль прогона (`a:rPr/@sz`) и кегль
плейсхолдера, заданный макетом. Второй парсер не доставал, каталог подставлял вместо
него ступень лестницы для роли — и писатель справедливо не верил числу. Заголовок
VK WorkSpace s05 (`<p:ph type="title"/>` без единого `sz`) из-за этого уезжал в файл
кеглем 36 pt, хотя вписывание опустило его до 23,4 и он в две строки ложился на тело.

Подстановка ступени при этом опасна по-настоящему: ею уже роняли VK Tech s03 с 16
до 7,8 pt и поднимали VK Education s04/s11 с 36 до 39. Поэтому различие, а не отмена:
своё число писатель ставит, вычисленное — не ставит.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

from deckforge.designsystem import derive
from deckforge.parsing import TemplateParser
from deckforge.rendering.recipe_slide import write_zone
from tests.case_templates import case_template

TITLE_PT = 36.0


#: Шаблон кейса: заголовок в нём — плейсхолдер без своего `sz`, кегль задаёт макет.
#: Синтетический шаблон этот случай не воспроизводит — мастер python-pptx подставляет
#: кегль сам, и наследования не остаётся. Без файла тест пропускается (правило CI).
CASE = "VK_WorkSpace_Клиентская_конференция_Шаблон_03.pptx"


def zones_of(path: Path) -> dict[str, Any]:
    ds = derive(TemplateParser().parse(path, use_cache=False))
    return {zone.zone_id: zone for recipe in ds.recipes for zone in recipe.zones}


def test_a_placeholders_size_comes_from_the_layout() -> None:
    """Нарушитель: у заголовка-плейсхолдера нет своего `sz`, и кегль знает макет.

    Раньше `ExampleShape.size_pt` был `None`, а каталог подставлял ступень лестницы —
    число совпадало случайно и помечено как своё не было.
    """
    manifest = TemplateParser().parse(case_template(CASE), use_cache=False)
    holders = [
        shape
        for example in manifest.examples
        for shape in example.shapes
        if shape.placeholder_idx is not None and shape.text_len
    ]

    assert holders, "в примерах шаблона нет ни одного текстового плейсхолдера"
    assert [shape for shape in holders if shape.size_pt is not None], (
        "кегль плейсхолдера не достали из макета"
    )


def test_a_zone_says_whether_the_size_is_its_own() -> None:
    """Зона отличает авторское число от подставленного — и таких зон большинство."""
    zones = zones_of(case_template(CASE))
    own = [zone for zone in zones.values() if zone.size_is_own]

    assert own, "ни одна зона не назвала кегль своим"
    assert all(zone.size_pt is not None for zone in own)
    assert len(own) > len(zones) / 2, "своих кеглей меньше половины — похоже на догадку"


# --- писатель ---------------------------------------------------------------------


def zone_shape(prs: Any, *, with_size: bool) -> Any:
    slide = prs.slides.add_slide(next(la for la in prs.slide_layouts if la.name == "Blank"))
    box = slide.shapes.add_textbox(Emu(400_000), Emu(300_000), Emu(6_000_000), Emu(600_000))
    box.text_frame.text = "Текст примера"
    run = box.element.find(f"{qn('p:txBody')}/{qn('a:p')}/{qn('a:r')}")
    props = run.find(qn("a:rPr"))
    if props is None:
        props = etree.SubElement(run, qn("a:rPr"))
        run.insert(0, props)
    if with_size:
        props.set("sz", "3600")
    else:
        props.attrib.pop("sz", None)
    return box.element


def written_size(shape: Any) -> str | None:
    run = shape.find(f"{qn('p:txBody')}/{qn('a:p')}/{qn('a:r')}")
    props = run.find(qn("a:rPr")) if run is not None else None
    return props.get("sz") if props is not None else None


@pytest.fixture
def prs() -> Any:
    return Presentation()


def test_the_writer_lowers_a_size_the_run_states(prs: Any) -> None:
    """Норма, как было: кегль стоит у прогона примера — писатель его опускает."""
    shape = zone_shape(prs, with_size=True)

    write_zone(shape, ["Наш текст"], 23.4, TITLE_PT)

    assert written_size(shape) == "2340"


def test_the_writer_lowers_a_size_the_layout_states(prs: Any) -> None:
    """Нарушитель RG39: у прогона кегля нет, но число авторское — из макета."""
    shape = zone_shape(prs, with_size=False)

    write_zone(shape, ["Наш текст"], 23.4, TITLE_PT)

    assert written_size(shape) == "2340"


def test_the_writer_stays_silent_when_the_size_is_a_guess(prs: Any) -> None:
    """Норма D04: число вычислили мы — писатель не трогает ничего.

    Именно подстановка ступени роняла VK Tech s03 с 16 до 7,8 pt.
    """
    shape = zone_shape(prs, with_size=False)

    write_zone(shape, ["Наш текст"], 7.8, None)

    assert written_size(shape) is None


def test_the_writer_never_raises_the_authors_size(prs: Any) -> None:
    """Норма D02: кегль автора писатель только опускает — VK Education 36 не станет 39."""
    shape = zone_shape(prs, with_size=False)

    write_zone(shape, ["Наш текст"], 39.0, TITLE_PT)

    assert written_size(shape) is None
