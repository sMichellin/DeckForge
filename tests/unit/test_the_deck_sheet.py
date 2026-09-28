"""Лист колоды: слайды рядом с примерами шаблона. Change `the-deck-sheet`, план Б, шаг 6 (#245).

Разбор 28.09 шёл глазами по PDF, а выбор примера был строкой текста. Лист показывает слайд
и пример, по которому он собран, называет повтор примера и пустые карточки — без модели
и без рендера. Сервис отдаёт превью примера из кэша шаблона, интерфейс собирает строки.

Сценарии — из дельты `openspec/changes/the-deck-sheet/specs/web-ui/`.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient
from frontend.sheet import EMPTY_GROUP, repeat_warnings, sheet_rows
from pptx import Presentation

from deckforge.api.app import create_app
from deckforge.api.examples import example_preview, shown_slide_parts
from deckforge.api.store import RunStore
from deckforge.parsing.template import template_id_of
from deckforge.pipeline.nodes.render import EXAMPLES_DIR

RUNS = Path(__file__).resolve().parents[1] / "fixtures" / "runs" / "2026-09-28"
PNG = b"\x89PNG\r\n\x1a\n" + b"0" * 32


# --- лист: строки и предупреждения --------------------------------------------------------


def _choice(slide_id: str, recipe_id: str | None, layout_id: str = "L07") -> dict[str, Any]:
    return {"slide_id": slide_id, "recipe_id": recipe_id, "layout_id": layout_id, "why": "так"}


def test_vk_tech_28_09_shows_ex018_on_six_slides() -> None:
    """Мерило: на прогоне 28.09 VK Tech видно, что `ex018` стоит на шести слайдах из десяти."""
    report = json.loads((RUNS / "vk-tech" / "run.json").read_text(encoding="utf-8"))

    rows = sheet_rows(report)

    assert [row.slide_id for row in rows if row.recipe_id == "ex018"] == [
        "s02", "s03", "s04", "s05", "s07", "s09",
    ]
    assert {row.uses for row in rows if row.recipe_id == "ex018"} == {6}
    assert "ex018 — на 6 слайдах из 10: s02, s03, s04, s05, s07, s09" in repeat_warnings(rows)
    assert "ex018 — на соседних s02 и s03" in repeat_warnings(rows)


def test_empty_cards_are_shown_at_their_slide() -> None:
    """Находки `integrity.empty_group` (#255) — у своего слайда, а не в общем списке."""
    report = {
        "slide_choices": [_choice("s01", "ex001"), _choice("s02", "ex018")],
        "findings_detail": [
            {"check_id": EMPTY_GROUP, "slide_id": "s02"},
            {"check_id": EMPTY_GROUP, "slide_id": "s02"},
            {"check_id": "template.sample_text_left", "slide_id": "s01"},
        ],
    }

    rows = sheet_rows(report)

    assert [row.empty_groups for row in rows] == [0, 2]
    assert "пустых карточек: 2" in rows[1].caption
    assert "пустых" not in rows[0].caption


def test_a_varied_deck_has_no_warnings() -> None:
    """Норма: каждый пример не больше двух раз и не подряд — предупреждать не о чем."""
    report = {"slide_choices": [
        _choice("s01", "ex001"), _choice("s02", "ex018"), _choice("s03", "ex009"),
        _choice("s04", "ex018"), _choice("s05", None),
    ]}

    assert repeat_warnings(sheet_rows(report)) == []


def test_the_caption_agrees_with_the_number() -> None:
    rows = sheet_rows({"slide_choices": [_choice("s01", "ex001"), _choice("s02", "ex018"),
                                         _choice("s03", "ex009"), _choice("s04", "ex009")]})

    assert "ex001 — на 1 слайде из 4" in rows[0].caption
    assert "ex009 — на 2 слайдах из 4" in rows[2].caption


def test_a_slide_by_layout_says_so() -> None:
    row = sheet_rows({"slide_choices": [_choice("s05", None, "L12")]})[0]

    assert row.recipe_id is None
    assert "по макету L12" in row.caption


def test_a_report_without_choices_gives_no_sheet() -> None:
    assert sheet_rows({}) == []


# --- сервис: превью примера из кэша шаблона ----------------------------------------------


def _template(path: Path, *, hide_second: bool) -> Path:
    """Шаблон из трёх слайдов: `slide1.xml`, `slide2.xml`, `slide3.xml`; второй можно скрыть."""
    prs = Presentation()
    for _ in range(3):
        prs.slides.add_slide(prs.slide_layouts[6])
    if hide_second:
        prs.slides[1]._element.set("show", "0")
    prs.save(str(path))
    return path


def _cache(cache_root: Path, template: Path, pages: int) -> Path:
    """Превью примеров, как их кладёт рендер: `<имя>-N.png` в порядке страниц PDF."""
    folder = cache_root / EXAMPLES_DIR / template_id_of(template).replace(":", "_")
    folder.mkdir(parents=True)
    for page in range(1, pages + 1):
        (folder / f"template-{page}.png").write_bytes(PNG + bytes([page]))
    return folder


def test_a_hidden_slide_shifts_the_pages(tmp_path: Path) -> None:
    """Страница — по порядку показа без скрытых: `slide3.xml` за скрытым вторым — страница 2."""
    template = _template(tmp_path / "t.pptx", hide_second=True)
    folder = _cache(tmp_path / "cache", template, pages=2)

    assert shown_slide_parts(template) == ["ppt/slides/slide1.xml", "ppt/slides/slide3.xml"]
    assert example_preview(template, "ex003", tmp_path / "cache") == folder / "template-2.png"
    assert example_preview(template, "ex002", tmp_path / "cache") is None


@pytest.fixture
def served(tmp_path: Path) -> tuple[TestClient, str]:
    store = RunStore(tmp_path / "runs")
    run_id = store.create(request={"variant": "A"})
    template = store.save_template(
        run_id, "t.pptx", _template(tmp_path / "t.pptx", hide_second=False).read_bytes()
    )
    _cache(tmp_path / "cache", template, pages=3)
    app = create_app(store=store, queue=_NoQueue(), cache_root=tmp_path / "cache")
    return TestClient(app), run_id


class _NoQueue:
    async def submit(self, job: str, run_id: str, *args: Any) -> None: ...
    async def close(self) -> None: ...


def test_the_service_gives_the_example_preview(served: tuple[TestClient, str]) -> None:
    client, run_id = served

    answer = client.get(f"/runs/{run_id}/examples/ex002")

    assert answer.status_code == 200
    assert answer.headers["content-type"] == "image/png"
    assert answer.content == PNG + bytes([2])


@pytest.mark.parametrize("recipe_id", ["ex009", "slide2", "..%2F..%2Fx", "ex"])
def test_no_such_example_is_404(served: tuple[TestClient, str], recipe_id: str) -> None:
    """Нет примера, чужой вид идентификатора или путь наружу — 404, файл не отдаётся."""
    client, run_id = served

    assert client.get(f"/runs/{run_id}/examples/{recipe_id}").status_code == 404


def test_a_cold_cache_is_404_not_a_crash(tmp_path: Path) -> None:
    """Примеры не рендерились (нет LibreOffice) — 404, и лист показывает слайд без примера."""
    store = RunStore(tmp_path / "runs")
    run_id = store.create(request={"variant": "A"})
    store.save_template(run_id, "t.pptx", _template(tmp_path / "t.pptx", hide_second=False)
                        .read_bytes())
    client = TestClient(create_app(store=store, queue=_NoQueue(), cache_root=tmp_path / "none"))

    assert client.get(f"/runs/{run_id}/examples/ex001").status_code == 404
