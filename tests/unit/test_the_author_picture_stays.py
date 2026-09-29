"""Иллюстрация автора остаётся. Change `the-author-picture-stays` (план Б, fix-forward к 4).

После #257 писатель при паспорте снимал группу из одной картинки как «незаполненную»: ассет
в рецептный слайд он не ставит, и такую группу заполнить нечем. На VK Tech уходил куб
иллюстрации, на Education — две картинки `ex045`/`ex046`. Правило: группа из одних картинок
остаётся; картинка в карточке с текстовыми местами уходит вместе с незаполненной карточкой.
"""

from __future__ import annotations

from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu

from deckforge.designsystem.models import Place, PlaceGroup, PlaceKind
from deckforge.rendering.recipe_slide import clone_recipe
from tests.unit.test_recipe_leaves_no_sample_text import slide_ir
from tests.unit.test_the_writer_removes_whole_groups import (
    _with_picture,
    cards_example,
    cards_recipe,
    ids_left,
)


def _with_illustration():
    """Пример «ряд карточек» и иллюстрация автора — отдельная группа паспорта из картинки."""
    prs, part_name, f = cards_example()
    slide = prs.slides[0]
    art = slide.shapes.add_shape(MSO_SHAPE.CUBE, Emu(7000000), Emu(200000), Emu(900000),
                                 Emu(900000))
    recipe = cards_recipe(part_name, f, passport=True)
    passport = recipe.passport
    assert passport is not None
    group = PlaceGroup(
        group_id="g07",
        places=[Place(place_id="p12", kind=PlaceKind.PICTURE, xml_id=art.shape_id)],
    )
    recipe = recipe.model_copy(
        update={"passport": passport.model_copy(update={"groups": [*passport.groups, group]})}
    )
    return prs, recipe, f, art.shape_id


def test_a_picture_only_group_stays_when_nothing_is_written() -> None:
    """Норма: группа из одной картинки остаётся даже при IR только с заголовком."""
    prs, recipe, _, art = _with_illustration()

    slide = clone_recipe(prs, recipe, slide_ir(recipe, ["zt"]))

    assert art in ids_left(slide)


def test_the_unfilled_cards_still_leave_beside_the_picture() -> None:
    """Нарушитель рядом: незаполненные карточки при этом уходят, как в change 4."""
    prs, recipe, f, art = _with_illustration()

    slide = clone_recipe(prs, recipe, slide_ir(recipe, ["zt", "zs1", "zb1"]))
    left = ids_left(slide)

    assert art in left
    assert f["plate3"].shape_id not in left and f["wide_plate"].shape_id not in left


def test_a_picture_inside_an_unfilled_card_leaves_with_it() -> None:
    """Нарушитель: картинка в карточке с текстовым местом уходит вместе с незаполненной картой."""
    prs, part_name, f = cards_example()
    recipe = cards_recipe(part_name, f, passport=True)
    assert recipe.passport is not None
    bar = f["wide_bar"].shape_id
    recipe = recipe.model_copy(update={"passport": _with_picture(recipe.passport, bar)})

    slide = clone_recipe(prs, recipe, slide_ir(recipe, ["zt", "zs1", "zb1"]))

    assert bar not in ids_left(slide)


def test_a_picture_inside_a_filled_card_stays() -> None:
    """Норма: в заполненной карточке картинка автора — её оформление, остаётся."""
    prs, part_name, f = cards_example()
    recipe = cards_recipe(part_name, f, passport=True)
    assert recipe.passport is not None
    bar = f["wide_bar"].shape_id
    recipe = recipe.model_copy(update={"passport": _with_picture(recipe.passport, bar)})

    slide = clone_recipe(prs, recipe, slide_ir(recipe, ["zt", "zw"]))

    assert bar in ids_left(slide)


def test_a_shape_shared_with_a_leaving_card_stays_with_the_picture() -> None:
    """Нарушитель паспорта (Education `ex045`, фигура 1001): адрес иллюстрации есть и в декоре
    незаполненной карточки. Остающуюся фигуру писатель не снимает, карточка уходит без неё."""
    prs, recipe, f, art = _with_illustration()
    passport = recipe.passport
    assert passport is not None
    groups = [
        g.model_copy(update={"decor_xml_ids": [*g.decor_xml_ids, art]}) if g.group_id == "g05"
        else g
        for g in passport.groups
    ]
    recipe = recipe.model_copy(update={"passport": passport.model_copy(update={"groups": groups})})

    slide = clone_recipe(prs, recipe, slide_ir(recipe, ["zt", "zs1", "zb1"]))
    left = ids_left(slide)

    assert art in left
    assert f["plate4"].shape_id not in left
