"""Вёрстка по дизайн-системе шаблона. Change `draw-by-the-design-system` (DG3, сторона B).

Что проверяется, по требованиям `compose-by-the-design-system` и задачи DG3:

* акцент — роль `accent` из `palette_roles`, а не `accent1` по порядку; в IR и в файле —
  имя слота темы (правило 5);
* кегль числа — `number_sizes`, приведённый к шкале (правило 6), и не мельче тела
  рядом (страховка от регрессии B13, #139);
* цитата и callout рисуются нативными объектами pptx (C3) и в html, полоса и отбивка —
  шагом дизайн-системы;
* холодный шаблон (правило 10): ролей и кеглей нет — поведение прежнее, и это названо.

Шаблоны — синтетический манифест `conftest` и стандартный шаблон python-pptx: ни один
не похож на шаблоны кейса (C6).
"""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from deckforge.designsystem import DesignSystem, Origin, derive
from deckforge.designsystem.models import (
    ColorRoleKind,
    ComponentCard,
    NumberScale,
    PaletteRole,
)
from deckforge.domain.base import BBox
from deckforge.domain.enums import CalloutTone, ColorRef, SmartArtPattern, TextRole
from deckforge.domain.slide import (
    CalloutBlock,
    DeckIR,
    KpiBlock,
    KpiItem,
    QuoteBlock,
    SlideIR,
    TextBlock,
)
from deckforge.domain.template import ComponentKind, ComponentSpec, ShapeKind, TemplateManifest
from deckforge.domain.units import EMU_PER_CM
from deckforge.export.html import export_html
from deckforge.layout.by_design import MAX_INSET_SHARE, DesignRules
from deckforge.layout.diagram import diagram_geometry
from deckforge.layout.fitting import fit_kpi, fit_slide
from deckforge.layout.fonts import FontLibrary
from deckforge.parsing import TemplateParser
from deckforge.pipeline.nodes.fit import _fit_shortening
from deckforge.rendering.writer import PptxWriter
from tests.integration.test_native_objects import build_template
from tests.unit.test_layout_fonts import make_font

TALL = (12 * EMU_PER_CM, 14 * EMU_PER_CM)


@pytest.fixture
def fonts(tmp_path: Path, manifest: TemplateManifest) -> FontLibrary:
    directory = tmp_path / "fonts"
    directory.mkdir()
    make_font(directory, manifest.theme.fonts.major_latin, advance=500, bold=True)
    make_font(directory, manifest.theme.fonts.minor_latin, advance=500)
    make_font(directory, manifest.theme.fonts.minor_latin, advance=550, bold=True)
    return FontLibrary([directory])


def accent_role(manifest: TemplateManifest, ref: ColorRef, share: float) -> PaletteRole:
    return PaletteRole(
        role=ColorRoleKind.ACCENT,
        color_hex=manifest.theme.colors.get(ref),
        nearest_ref=ref,
        share=share,
    )


def with_roles(manifest: TemplateManifest, *roles: PaletteRole) -> DesignSystem:
    return derive(manifest).model_copy(update={"palette_roles": list(roles)})


def with_numbers(manifest: TemplateManifest, large: float) -> DesignSystem:
    """Дизайн-система, измерившая кегль числа `large` на KPI-примере шаблона."""
    ds = derive(manifest)
    return ds.model_copy(
        update={"number_sizes": NumberScale(large_pt=large, medium_pt=None, small_pt=None,
                                            origin=Origin.MEASURED)}
    )


# --- акцент по роли -------------------------------------------------------------


def test_the_accent_is_the_one_the_template_uses(manifest: TemplateManifest) -> None:
    """Нарушитель прежнего правила: шаблон ставит акцентом третий слот, а не первый."""
    rules = DesignRules(
        manifest,
        with_roles(
            manifest,
            accent_role(manifest, ColorRef.ACCENT1, 0.01),
            accent_role(manifest, ColorRef.ACCENT3, 0.2),
        ),
    )

    assert rules.roles_measured
    assert rules.accents()[:2] == [ColorRef.ACCENT3, ColorRef.ACCENT1]
    assert rules.block_accent(None) is ColorRef.ACCENT3
    assert rules.callout_accent(CalloutTone.INSIGHT, None) is ColorRef.ACCENT3
    assert rules.callout_accent(CalloutTone.RISK, None) is ColorRef.ACCENT1


def test_a_color_shared_by_text_and_accent_slots_is_named_as_accent(
    manifest: TemplateManifest,
) -> None:
    """Один синий в `dk2` и `accent1` — акцентом он записывается слотом акцента."""
    blue = manifest.theme.colors.get(ColorRef.ACCENT1)
    colors = manifest.theme.colors.model_copy(update={"dk2": blue})
    same_blue = manifest.model_copy(
        update={"theme": manifest.theme.model_copy(update={"colors": colors})}
    )
    role = PaletteRole(role=ColorRoleKind.ACCENT, color_hex=blue, nearest_ref=ColorRef.DK2,
                       share=0.2)

    rules = DesignRules(same_blue, with_roles(same_blue, role))

    assert rules.accents()[0] is ColorRef.ACCENT1


def test_an_accent_invisible_on_the_background_yields(manifest: TemplateManifest) -> None:
    """Акцент, неразличимый на фоне слайда (жёлтый на белом — ниже 3:1), уступает
    следующему, а не рисуется невидимым."""
    rules = DesignRules(
        manifest,
        with_roles(
            manifest,
            accent_role(manifest, ColorRef.ACCENT3, 0.2),
            accent_role(manifest, ColorRef.ACCENT5, 0.1),
        ),
    )
    background = manifest.theme.colors.get(ColorRef.LT1)

    assert rules.accent(None) is ColorRef.ACCENT3, "без фона выбирать не по чему"
    assert rules.accent(background) is ColorRef.ACCENT5


def test_a_cold_template_keeps_accent1_and_says_so(manifest: TemplateManifest) -> None:
    """Правило 10: примеров нет — роли не измерены, акцент прежний, заметка есть."""
    rules = DesignRules(manifest)

    assert manifest.examples == []
    assert not rules.roles_measured
    assert rules.block_accent(None) is ColorRef.ACCENT1
    assert any("роли цветов" in note for note in rules.cold_notes())


# --- кегль числа ----------------------------------------------------------------


def kpi(*values: str) -> KpiBlock:
    return KpiBlock(
        block_id="k",
        items=[KpiItem(value=value, label="за год") for value in values],
        x=0, y=0, cx=TALL[0], cy=TALL[1],
    )


def test_the_number_takes_the_template_size_on_the_ladder(
    manifest: TemplateManifest, fonts: FontLibrary
) -> None:
    """Шаблон набрал число 26 pt — значение встаёт на ближайшую ступень шкалы, подпись
    — на мелкий кегль числа, оба со шкалы шаблона."""
    rules = DesignRules(manifest, with_numbers(manifest, 26.0).model_copy(
        update={"number_sizes": NumberScale(large_pt=26.0, medium_pt=24.0, small_pt=18.0,
                                            origin=Origin.MEASURED)}
    ))
    sizes = rules.kpi_sizes()
    assert sizes is not None
    block = kpi("37 %", "×2,3")
    assert block.bbox is not None

    result = fit_kpi(block, block.bbox, manifest, fonts=fonts, sizes=sizes)

    assert sizes.value_pt == 24.0
    assert result.final_size_pt == 24.0, "значение ушло от кегля шаблона"
    assert result.final_size_pt in manifest.size_ladder_pt
    assert sizes.label_for(result.final_size_pt, manifest) == 18.0
    assert sizes.label_for(result.final_size_pt, manifest) < result.final_size_pt


def test_the_old_rule_grows_where_the_template_did_not_answer(
    manifest: TemplateManifest, fonts: FontLibrary
) -> None:
    """Норма (#87): без ответа ДС короткое число в высокой рамке растёт до потолка —
    на ответе ДС стоит на своём кегле. Разница и есть «кегль из шаблона»."""
    block = kpi("37 %", "×2,3")
    assert block.bbox is not None

    old = fit_kpi(block, block.bbox, manifest, fonts=fonts)

    assert old.final_size_pt == max(manifest.size_ladder_pt)


def test_a_number_not_above_the_body_is_rejected(manifest: TemplateManifest) -> None:
    """Страховка от B13: измеренный кегль ложится на тело — показатель по прежнему
    правилу, а не вровень со списком рядом, и заметка это называет."""
    body = manifest.typography(TextRole.BODY)
    assert body is not None
    rules = DesignRules(manifest, with_numbers(manifest, body.size_pt + 0.5))

    assert rules.kpi_sizes() is None
    assert any("B13" in note for note in rules.cold_notes())


def test_a_phrase_value_still_wraps_by_words_on_the_design_path(
    manifest: TemplateManifest, fonts: FontLibrary
) -> None:
    """B13 (#139) не возвращается: фраза в значении не тянет блок ниже тела."""
    rules = DesignRules(manifest, with_numbers(manifest, 26.0))
    sizes = rules.kpi_sizes()
    assert sizes is not None
    block = kpi("часы → минуты", "100%")
    assert block.bbox is not None

    result = fit_kpi(block, block.bbox, manifest, fonts=fonts, sizes=sizes)

    body = manifest.typography(TextRole.BODY)
    assert body is not None
    assert result.overflow is False
    assert result.final_size_pt > body.size_pt


def test_a_cold_template_keeps_the_old_number_rule(manifest: TemplateManifest) -> None:
    """Правило 10: показателя в шаблоне нет — кегли не измерены, правило прежнее."""
    rules = DesignRules(manifest)

    assert rules.ds.number_sizes.origin is Origin.DERIVED
    assert rules.kpi_sizes() is None
    assert any("кегль числа не измерен" in note for note in rules.cold_notes())


# --- пропорции ------------------------------------------------------------------


def test_the_tile_comes_from_the_design_system_catalog(manifest: TemplateManifest) -> None:
    """Плитка ДС и та же плитка манифеста дают одну и ту же сетку."""
    spec = ComponentSpec(
        kind=ComponentKind.TILE, repeats=3, axis="row", width_share=0.3, height_share=0.15,
        gap_share=0.33, parts=[ShapeKind.TEXT], seen_on=[1, 2],
    )
    card = ComponentCard(
        kind=ComponentKind.TILE, repeats=3, axis="row", width_share=0.3, height_share=0.15,
        gap_share=0.33, seen_on=[1, 2],
    )
    rules = DesignRules(manifest, derive(manifest).model_copy(update={"components": [card]}))
    box = BBox(x=0, y=0, cx=24 * EMU_PER_CM, cy=10 * EMU_PER_CM)

    assert rules.tile() == card
    from_card = diagram_geometry(SmartArtPattern.MATRIX, 4, box, rules.tile())
    from_spec = diagram_geometry(SmartArtPattern.MATRIX, 4, box, spec)
    assert from_card.nodes == from_spec.nodes


def test_an_oversized_grid_step_is_halved_into_the_block(manifest: TemplateManifest) -> None:
    """Шаг сетки в шестую часть слайда отбивкой не становится: делится пополам,
    пока не встанет в предел, а полоса тоньшает вместе с ним."""
    style = DesignRules(manifest).quote_style()
    huge = style.__class__(**{**style.__dict__, "spacing_emu": 6 * EMU_PER_CM,
                              "line_emu": 3 * EMU_PER_CM})
    box_cx = 12 * EMU_PER_CM

    inset = huge.inset(box_cx)

    assert inset.pad_emu <= box_cx * MAX_INSET_SHARE
    assert inset.bar_emu * 2 == pytest.approx(inset.pad_emu, abs=1)


def test_the_quote_step_is_the_design_system_step(manifest: TemplateManifest) -> None:
    """Норма: шаг в предел укладывается — отбивка и полоса ровно из дизайн-системы."""
    rules = DesignRules(manifest)
    quote = next(item for item in rules.ds.synthesized if item.kind == "quote")
    style = rules.quote_style()

    inset = style.inset(30 * EMU_PER_CM)

    assert inset.pad_emu == quote.spacing_emu
    assert inset.bar_emu == quote.line_emu
    assert style.text_pt == quote.size_pt


# --- цитата и callout: вписывание -----------------------------------------------


def slide_with(manifest: TemplateManifest, *blocks: object, box: BBox | None = None) -> SlideIR:
    area = box or manifest.content_bbox
    half = area.cx // 2
    placed = [
        block.model_copy(  # type: ignore[attr-defined]
            update={"x": area.x + i * half, "y": area.y + area.cy // 2, "cx": half,
                    "cy": area.cy // 2}
        )
        for i, block in enumerate(blocks)
    ]
    return SlideIR(slide_id="s1", layout_id=manifest.layouts[-1].layout_id, variant="A",
                   blocks=[TextBlock(block_id="t", placeholder_idx=0, role=TextRole.TITLE,
                                     text="Вывод слайда"), *placed])


def quote() -> QuoteBlock:
    return QuoteBlock(block_id="q", text="Шаблон задаёт язык колоды", author="Бриф")


def callout(tone: CalloutTone = CalloutTone.RISK) -> CalloutBlock:
    return CalloutBlock(block_id="c", text="Срок сдачи под угрозой", tone=tone)


def test_quote_and_callout_are_fitted_on_the_ladder(
    manifest: TemplateManifest, fonts: FontLibrary
) -> None:
    fitted = fit_slide(slide_with(manifest, quote(), callout()), manifest, fonts=fonts)

    rules = DesignRules(manifest)
    for block_id, start in (("q", rules.quote_style().text_pt),
                            ("c", rules.callout_style(CalloutTone.RISK).text_pt)):
        fit = fitted.fit_report[block_id]
        assert fit.overflow is False
        assert fit.final_size_pt in manifest.size_ladder_pt
        assert fit.final_size_pt == start, "кегль ушёл от ступени дизайн-системы без нужды"


def test_a_quote_that_does_not_fit_is_dropped_with_a_note(
    manifest: TemplateManifest, fonts: FontLibrary
) -> None:
    """Нарушитель: цитата в полосе ниже строки не роняет запись колоды — снимается,
    и заметка это называет."""
    tiny = BBox(x=manifest.content_bbox.x, y=manifest.content_bbox.y,
                cx=4 * EMU_PER_CM, cy=EMU_PER_CM // 10)
    long_quote = quote().model_copy(update={"text": "очень длинная цитата " * 12})
    slide = slide_with(manifest, long_quote, box=tiny)

    fitted, notes = _fit_shortening(slide, manifest, fonts, _content())

    assert "q" not in {block.block_id for block in fitted.blocks}
    assert any(note.startswith("s1/q: блок quote") and "снят" in note for note in notes)


def _content() -> object:
    from deckforge.domain.content import Brief, ContentPackage

    return ContentPackage(brief=Brief(purpose="product", audience="правление", target_slides=6))


# --- цитата и callout: рисунок --------------------------------------------------


@pytest.fixture
def real(tmp_path: Path) -> tuple[Path, TemplateManifest]:
    template = build_template(tmp_path / "template.pptx")
    return template, TemplateParser(cache_dir=tmp_path / "cache").parse(template)


def written(
    tmp_path: Path, template: Path, manifest: TemplateManifest, slide: SlideIR,
    design: DesignSystem | None = None,
) -> tuple[Path, SlideIR]:
    fitted = fit_slide(slide, manifest, design=DesignRules(manifest, design))
    deck = DeckIR(deck_id="d", variant="A", template_id=manifest.template_id, seed=1,
                  slides=[fitted])
    out = PptxWriter(template, manifest, design_system=design).write(deck, tmp_path / "d.pptx")
    return out, fitted


def groups(path: Path) -> dict[str, object]:
    slide = Presentation(str(path)).slides[0]
    return {
        shape.name: shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.GROUP
    }


def scheme_colors(shape: object) -> list[str]:
    element = shape._element  # type: ignore[attr-defined]
    return [node.get("val") for node in element.iter() if node.tag.endswith("}schemeClr")]


def test_quote_and_callout_are_native_shapes(
    tmp_path: Path, real: tuple[Path, TemplateManifest]
) -> None:
    """C3: полоса — автофигура, текст — текстовая рамка, оба редактируемы; цвет —
    ссылка на слот темы, а не RGB (правило 5)."""
    template, manifest = real
    path, _ = written(tmp_path, template, manifest, slide_with(manifest, quote(), callout()))

    drawn = groups(path)
    assert {"Цитата", "Callout risk"} <= set(drawn)
    quote_group = drawn["Цитата"]
    kinds = [shape.shape_type for shape in quote_group.shapes]  # type: ignore[attr-defined]
    assert MSO_SHAPE_TYPE.AUTO_SHAPE in kinds and MSO_SHAPE_TYPE.TEXT_BOX in kinds
    text = "\n".join(
        shape.text_frame.text for shape in quote_group.shapes  # type: ignore[attr-defined]
        if shape.has_text_frame and shape.text_frame.text
    )
    assert "Шаблон задаёт язык колоды" in text and "— Бриф" in text
    assert "srgbClr" not in quote_group._element.xml  # type: ignore[attr-defined]
    risk_label = DesignRules(manifest).callout_style(CalloutTone.RISK).label
    callout_text = "".join(
        shape.text_frame.text for shape in drawn["Callout risk"].shapes  # type: ignore[attr-defined]
        if shape.has_text_frame
    )
    assert risk_label and risk_label in callout_text


def test_the_bar_takes_the_role_accent_when_ir_names_none(
    tmp_path: Path, real: tuple[Path, TemplateManifest]
) -> None:
    """`accent_ref` не задан — полоса акцентом по роли ДС; задан — им."""
    template, manifest = real
    design = with_roles(manifest, accent_role(manifest, ColorRef.ACCENT4, 0.3))
    path, _ = written(tmp_path, template, manifest, slide_with(manifest, quote()), design)
    assert "accent4" in scheme_colors(groups(path)["Цитата"])

    named = quote().model_copy(update={"accent_ref": ColorRef.ACCENT2})
    path, _ = written(tmp_path, template, manifest, slide_with(manifest, named), design)
    colors = scheme_colors(groups(path)["Цитата"])
    assert "accent2" in colors and "accent4" not in colors


def test_the_kpi_value_takes_the_role_accent(
    tmp_path: Path, real: tuple[Path, TemplateManifest]
) -> None:
    """Значение показателя без слота в IR — акцент по роли, а не `accent1`."""
    template, manifest = real
    design = with_roles(manifest, accent_role(manifest, ColorRef.ACCENT4, 0.3))
    block = kpi("37 %")
    path, _ = written(tmp_path, template, manifest, slide_with(manifest, block), design)

    slide = Presentation(str(path)).slides[0]
    values = [
        node.get("val") for shape in slide.shapes if shape.has_text_frame
        for node in shape._element.iter() if node.tag.endswith("}schemeClr")
        and "37 %" in shape.text_frame.text
    ]
    assert "accent4" in values


def test_html_draws_the_same_quote(
    tmp_path: Path, real: tuple[Path, TemplateManifest]
) -> None:
    _template, manifest = real
    design = with_roles(manifest, accent_role(manifest, ColorRef.ACCENT4, 0.3))
    fitted = fit_slide(slide_with(manifest, quote(), callout(CalloutTone.INSIGHT)), manifest,
                       design=DesignRules(manifest, design))
    deck = DeckIR(deck_id="d", variant="A", template_id=manifest.template_id, seed=1,
                  slides=[fitted])

    html = export_html(deck, manifest, tmp_path / "d.html", design_system=design).read_text(
        encoding="utf-8"
    )

    assert 'class="block boxed quote"' in html and 'class="block boxed callout"' in html
    assert "Шаблон задаёт язык колоды" in html and "— Бриф" in html
    assert "background: var(--accent4)" in html
    assert "#" not in html.split("</style>")[1].split("data:")[0].replace("#arrow", ""), (
        "литеральный цвет в разметке слайда"
    )


# --- конвейер -------------------------------------------------------------------


@pytest.mark.cold
async def test_the_run_names_what_the_design_system_did_not_measure(tmp_path: Path) -> None:
    """Правило 10 на всём конвейере: у шаблона python-pptx нет показателя на примерах —
    прогон проходит, а заметка называет, где осталось прежнее правило. Одна на прогон,
    а не на слайд."""
    from tests.integration.test_pipeline_end_to_end import run_pipeline
    from tests.unit.test_ir_callout_and_quote import workspace

    result, _ = await run_pipeline(workspace(tmp_path))

    notes = result.state.get("notes") or []
    about = [note for note in notes if note.startswith("дизайн-система:")]
    assert any("кегль числа не измерен" in note for note in about), notes
    assert len(about) == len(set(about))
    assert not [note for note in notes if "не нарисован" in note]


@pytest.mark.cold
async def test_a_fix_round_does_not_repeat_the_note(tmp_path: Path) -> None:
    """Виток починки проходит `fit` заново — заметки копятся, и вторая была бы шумом."""
    from langgraph.runtime import Runtime

    from deckforge.pipeline.nodes.fit import fit_node
    from tests.integration.test_pipeline_end_to_end import run_pipeline
    from tests.unit.test_ir_callout_and_quote import workspace

    result, deps = await run_pipeline(workspace(tmp_path))
    state = {**result.state, "fix_round": 1}

    update = await fit_node(state, Runtime(context=deps))  # type: ignore[arg-type]

    assert not [note for note in update["notes"] if note.startswith("дизайн-система:")]
