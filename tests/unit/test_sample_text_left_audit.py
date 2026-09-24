"""Аудит: текст шаблона остался на слайде по рецепту. Change `sample-text-left-audit`.

Слайд по рецепту — копия слайда-примера шаблона вместе с его текстом. Writer заменяет
текст зон нашим или удаляет зону; всё, что осталось на таком слайде и не написано нами,
пришло из примера. Шов — функция проверки на готовом `.pptx` и `SlideIR` колоды.

Синтетические тесты идут везде, включая CI: файл собирается python-pptx, манифест —
синтетический из `tests/conftest.py`. Тесты на шаблоне кейса проверяют то же на настоящей
копии примера и пропускаются, если шаблонов в чекауте нет.
"""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Cm

from deckforge.audit.deterministic.template import sample_text_left
from deckforge.audit.registry import CheckUnavailable
from deckforge.designsystem import derive
from deckforge.designsystem.models import DesignSystem, Recipe
from deckforge.domain.content import Brief, ContentPackage, Dataset, Series
from deckforge.domain.enums import TextRole
from deckforge.domain.slide import (
    BulletItem,
    BulletsBlock,
    DeckIR,
    KpiBlock,
    KpiItem,
    SlideIR,
    TableBlock,
    TextBlock,
)
from deckforge.domain.template import TemplateExample, TemplateManifest
from deckforge.parsing import TemplateParser
from deckforge.rendering.writer import PptxWriter
from tests.case_templates import case_template
from tests.unit._audit_builders import context_for, deck

CHECK = "template.sample_text_left"

#: Текст, который в синтетическом «примере» написал автор шаблона.
SAMPLE = "Описание первого преимущества"


def with_examples(manifest: TemplateManifest) -> TemplateManifest:
    """Шаблон с одним слайдом-примером: рецептам есть откуда взяться."""
    return manifest.model_copy(update={"examples": [TemplateExample(slide_index=1)]})


def by_recipe(*texts: str, slide_id: str = "s01", recipe_id: str | None = "ex001") -> SlideIR:
    """Слайд колоды: каждый текст — блок в своей зоне рецепта."""
    return SlideIR(
        slide_id=slide_id,
        layout_id="L07",
        variant="A",
        recipe_id=recipe_id,
        blocks=[
            TextBlock(
                block_id=f"b{index}",
                role=TextRole.TITLE if index == 0 else TextRole.BODY,
                text=text,
                zone_id=f"z{index}" if recipe_id else None,
            )
            for index, text in enumerate(texts)
        ],
    )


def pptx_with(path: Path, *slides: list[str]) -> Path:
    """Файл колоды: на каждом слайде по надписи на каждый текст."""
    prs = Presentation()
    for texts in slides:
        page = prs.slides.add_slide(prs.slide_layouts[6])
        for number, text in enumerate(texts):
            box = page.shapes.add_textbox(Cm(2), Cm(2 + 3 * number), Cm(10), Cm(2))
            box.text_frame.text = text
    prs.save(str(path))
    return path


def test_sample_text_left_on_a_recipe_slide_is_a_finding(
    manifest: TemplateManifest, tmp_path: Path
) -> None:
    """Нарушитель: зона заголовка заменена, а надпись примера осталась как была."""
    path = pptx_with(tmp_path / "deck.pptx", ["Выручка выросла на треть", SAMPLE])
    colony = deck(by_recipe("Выручка выросла на треть"))

    findings = list(
        sample_text_left(context_for(CHECK, colony, with_examples(manifest), deck_path=path))
    )

    assert [f.slide_id for f in findings] == ["s01"]
    assert SAMPLE in findings[0].evidence["text"]


def test_zones_replaced_or_dropped_is_not_a_finding(
    manifest: TemplateManifest, tmp_path: Path
) -> None:
    """Норма: одна зона заменена нашим текстом, вторая — нашим списком, третьей нет.

    Пункты списка стоят в зоне отдельными абзацами — так их пишет writer.
    """
    path = pptx_with(tmp_path / "deck.pptx", ["Выручка выросла на треть", "Первый\nВторой"])
    slide = by_recipe("Выручка выросла на треть")
    bullets = BulletsBlock(
        block_id="b1",
        items=[BulletItem(text="Первый"), BulletItem(text="Второй")],
        zone_id="z1",
    )
    colony = deck(slide.model_copy(update={"blocks": [*slide.blocks, bullets]}))
    assert len(Presentation(str(path)).slides[0].shapes[1].text_frame.paragraphs) == 2

    context = context_for(CHECK, colony, with_examples(manifest), deck_path=path)

    assert list(sample_text_left(context)) == []


def test_a_soft_break_inside_our_paragraph_is_not_a_finding(
    manifest: TemplateManifest, tmp_path: Path
) -> None:
    """Мягкий перенос (`a:br`) делит абзац на строки, как перевод строки в нашем тексте."""
    path = tmp_path / "deck.pptx"
    prs = Presentation()
    page = prs.slides.add_slide(prs.slide_layouts[6])
    box = page.shapes.add_textbox(Cm(2), Cm(2), Cm(10), Cm(2))
    box.text_frame.paragraphs[0].text = "Выручка выросла\vна треть"
    prs.save(str(path))
    colony = deck(by_recipe("Выручка выросла\nна треть"))

    context = context_for(CHECK, colony, with_examples(manifest), deck_path=path)

    assert list(sample_text_left(context)) == []


@pytest.mark.parametrize("kind", ["ftr", "dt", "sldNum", "hdr"])
def test_text_typed_into_a_footer_placeholder_is_not_a_finding(
    manifest: TemplateManifest, tmp_path: Path, kind: str
) -> None:
    """Колонтитул — не текст примера, даже набранный вручную: «Конфиденциально»."""
    path = pptx_with(tmp_path / "deck.pptx", ["Выручка выросла на треть", "Конфиденциально"])
    prs = Presentation(str(path))
    footer = prs.slides[0].shapes[1]._element
    footer.find(f"./{qn('p:nvSpPr')}/{qn('p:nvPr')}").append(
        footer.makeelement(qn("p:ph"), {"type": kind})
    )
    prs.save(str(path))
    colony = deck(by_recipe("Выручка выросла на треть"))

    context = context_for(CHECK, colony, with_examples(manifest), deck_path=path)

    assert list(sample_text_left(context)) == []


def test_kpi_and_table_text_written_by_the_deck_is_not_a_finding(
    manifest: TemplateManifest, tmp_path: Path
) -> None:
    """Показатель и ячейки таблиц writer пишет отдельными абзацами, а числа датасета —
    по-русски («2,5»). Это наш текст, где бы он ни оказался в файле.

    Страница — как у writer'а в смешанной колоде: копия рецепта впереди, слайд с цифрами
    за ней, и страница слайда `s02` по номеру — слайд с цифрами.
    """
    path = pptx_with(tmp_path / "deck.pptx", ["План на год"], ["37 %", "рост"])
    prs = Presentation(str(path))
    for cells in (["Год", "Выручка"], ["2025", "2,5"]):
        table = prs.slides[1].shapes.add_table(1, 2, Cm(2), Cm(12), Cm(10), Cm(2)).table
        for column, text in enumerate(cells):
            table.cell(0, column).text = text
    prs.save(str(path))
    numbers = SlideIR(
        slide_id="s01",
        layout_id="L07",
        variant="A",
        blocks=[
            KpiBlock(block_id="k", items=[KpiItem(value="37 %", label="рост")]),
            TableBlock(block_id="t1", header=["Год"], rows=[["2025"]]),
            TableBlock(block_id="t2", dataset_ref="d001"),
        ],
    )
    content = ContentPackage(
        brief=Brief(purpose="report", audience="команда", target_slides=2),
        datasets=[
            Dataset(
                dataset_id="d001",
                title="Выручка",
                categories=["2025"],
                series=[Series(name="Выручка", values=[2.5])],
            )
        ],
    )
    colony = deck(numbers, by_recipe("План на год", slide_id="s02"))

    context = context_for(
        CHECK, colony, with_examples(manifest), content=content, deck_path=path
    )

    assert list(sample_text_left(context)) == []


def test_a_slide_without_a_recipe_is_not_checked(
    manifest: TemplateManifest, tmp_path: Path
) -> None:
    """Слайд без рецепта строится из макета: текста примера на нём быть не может.

    Незнакомая строка на нём — забота других проверок, а не этой. Слайд по рецепту
    рядом с ним проверяется как обычно.
    """
    path = pptx_with(
        tmp_path / "deck.pptx", ["Выручка выросла на треть"], ["Вставлено вручную"]
    )
    colony = deck(
        by_recipe("Выручка выросла на треть"),
        by_recipe("План на год", slide_id="s02", recipe_id=None),
    )

    context = context_for(CHECK, colony, with_examples(manifest), deck_path=path)

    assert list(sample_text_left(context)) == []


def test_nothing_to_check_is_skipped_not_passed(
    manifest: TemplateManifest, tmp_path: Path
) -> None:
    """Нет рецептов, нет примеров, нет файла — «не проверяли», а не «прошла» и не падение."""
    path = pptx_with(tmp_path / "deck.pptx", [SAMPLE])
    plain = deck(by_recipe("План на год", recipe_id=None))
    recipe = deck(by_recipe("План на год"))
    cases = {
        "без слайдов по рецепту": context_for(
            CHECK, plain, with_examples(manifest), deck_path=path
        ),
        "холодный шаблон без примеров": context_for(CHECK, recipe, manifest, deck_path=path),
        "без файла колоды": context_for(CHECK, recipe, with_examples(manifest)),
    }

    skipped = []
    for why, context in cases.items():
        try:
            list(sample_text_left(context))
        except CheckUnavailable:
            skipped.append(why)
    assert skipped == list(cases)


def test_a_deck_file_that_does_not_match_the_ir_is_skipped(
    manifest: TemplateManifest, tmp_path: Path
) -> None:
    """Слайдов в файле меньше, чем в IR, или файл испорчен — «не проверяли», а не «прошла»."""
    short = pptx_with(tmp_path / "short.pptx", ["Выручка выросла на треть"])
    broken = tmp_path / "broken.pptx"
    broken.write_bytes(b"not a zip")
    colony = deck(
        by_recipe("Выручка выросла на треть"),
        by_recipe("План на год", slide_id="s02"),
    )
    cases = {
        "слайдов меньше, чем в IR": context_for(
            CHECK, colony, with_examples(manifest), deck_path=short
        ),
        "испорченный файл": context_for(
            CHECK, colony, with_examples(manifest), deck_path=broken
        ),
    }

    skipped = []
    for why, context in cases.items():
        try:
            list(sample_text_left(context))
        except CheckUnavailable:
            skipped.append(why)
    assert skipped == list(cases)


# --- на шаблоне кейса: настоящая копия примера из writer ---

TEMPLATE = "VK Tech шаблон.pptx"


@pytest.fixture(scope="module")
def case() -> tuple[Path, TemplateManifest, DesignSystem]:
    path = case_template(TEMPLATE)
    manifest = TemplateParser().parse(path, use_cache=False)
    return path, manifest, derive(manifest)


def richest(ds: DesignSystem) -> Recipe:
    """Рецепт с наибольшим числом повторов: на нём writer и заменяет, и удаляет зоны."""

    def rows(recipe: Recipe) -> int:
        return len({zone.repeat for zone in recipe.zones if zone.repeat is not None})

    return max(ds.recipes, key=lambda r: (rows(r), len(r.zones)))


def written_by_recipe(
    case: tuple[Path, TemplateManifest, DesignSystem], out: Path
) -> tuple[Path, DeckIR, int]:
    """Колода из одного слайда по рецепту, записанная writer'ом; и адрес занятой зоны."""
    path, manifest, ds = case
    recipe = richest(ds)
    zone = next(z for z in recipe.zones if z.repeat == 0 and z.xml_id is not None)
    colony = deck(
        SlideIR(
            slide_id="s01",
            layout_id=manifest.layouts[0].layout_id,
            variant="A",
            recipe_id=recipe.recipe_id,
            blocks=[
                TextBlock(block_id="b1", role=TextRole.BODY, text="Первый", zone_id=zone.zone_id)
            ],
        )
    ).model_copy(update={"template_id": manifest.template_id})
    PptxWriter(path, manifest, design_system=ds).write(colony, out)
    assert zone.xml_id is not None
    return out, colony, zone.xml_id


def shape_by_xml_id(slide: object, xml_id: int) -> object:
    tree = slide.shapes._spTree  # type: ignore[attr-defined]
    return next(
        node
        for node in tree.iter(qn("p:sp"))
        if node.find(f"./*/{qn('p:cNvPr')}").get("id") == str(xml_id)
    )


def test_text_of_the_example_restored_in_a_zone_is_a_finding(
    case: tuple[Path, TemplateManifest, DesignSystem], tmp_path: Path
) -> None:
    """Нарушитель на настоящей копии: в зону вернулся исходный текст примера."""
    template, manifest, ds = case
    out, colony, xml_id = written_by_recipe(case, tmp_path / "deck.pptx")
    recipe = richest(ds)
    source = next(
        slide
        for slide in Presentation(str(template)).slides
        if str(slide.part.partname).lstrip("/") == (recipe.part_name or "").lstrip("/")
    )
    original = shape_by_xml_id(source, xml_id).find(qn("p:txBody"))  # type: ignore[attr-defined]
    prs = Presentation(str(out))
    zone = shape_by_xml_id(prs.slides[0], xml_id)
    zone.replace(zone.find(qn("p:txBody")), deepcopy(original))  # type: ignore[attr-defined]
    prs.save(str(out))

    findings = list(sample_text_left(context_for(CHECK, colony, manifest, deck_path=out)))

    assert [(f.slide_id, f.evidence["xml_id"]) for f in findings] == [("s01", str(xml_id))]


def test_a_slide_written_by_the_recipe_is_not_a_finding(
    case: tuple[Path, TemplateManifest, DesignSystem], tmp_path: Path
) -> None:
    """Норма на настоящей копии: writer заменил занятую зону, остальные убрал или стёр."""
    _template, manifest, _ds = case
    out, colony, _xml_id = written_by_recipe(case, tmp_path / "deck.pptx")

    assert list(sample_text_left(context_for(CHECK, colony, manifest, deck_path=out))) == []
