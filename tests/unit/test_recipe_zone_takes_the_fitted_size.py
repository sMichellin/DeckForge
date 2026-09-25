"""Кегль вписывания доходит до файла слайда по рецепту. Change `recipe-zone-takes-the-fitted-size`.

Вписывание (RG29, `fitting-does-not-skip-the-recipe`) находит зоне кегль, при котором
самое длинное слово уже рамки, а текст встаёт в неё целиком. Писатель же копировал `rPr`
первого прогона примера и `fit_report` не читал: в файл уходил кегль автора примера —
144 pt у титула, 54 pt у «извлечения» — и PowerPoint рвал слова посередине.

Эталон кегля — число из `fit_report`, эталон оформления автора — `rPr` фигуры примера,
прочитанный из файла шаблона напрямую, а не через писатель.
"""

from __future__ import annotations

import copy
from pathlib import Path
from typing import Any

import pytest
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

from deckforge.designsystem import derive
from deckforge.designsystem.models import Recipe, RecipeKind, TypeLevel, Zone
from deckforge.domain.enums import TextRole
from deckforge.domain.slide import (
    BulletItem,
    BulletsBlock,
    DeckIR,
    FitResult,
    SlideIR,
    TextBlock,
)
from deckforge.parsing import TemplateParser
from deckforge.rendering.recipe_slide import clone_recipe
from deckforge.rendering.writer import PptxWriter
from tests.case_templates import case_template


def shape_by_id(slide: Any, xml_id: int) -> Any:
    """Фигура слайда по `cNvPr id` на любой глубине групп."""
    for node in slide.shapes._spTree.iter(qn("p:sp")):
        props = node.find(f"./*/{qn('p:cNvPr')}")
        if props is not None and props.get("id") == str(xml_id):
            return node
    raise AssertionError(f"фигуры {xml_id} на слайде нет")


def runs(shape: Any) -> list[Any]:
    return list(shape.iter(qn("a:r")))


def props_of(run: Any) -> Any | None:
    return run.find(qn("a:rPr"))


def canonical(node: Any) -> bytes:
    """XML узла без объявлений пространств имён, унаследованных от своего документа:
    у слайда шаблона и у записанного слайда они разные, а оформление — то же."""
    return bytes(etree.tostring(node, method="c14n", exclusive=True))


# --- синтетика: пример с двумя текстовыми фигурами -----------------------------------


#: `rPr` и `pPr` автора у второй надписи примера — они же эталон того, что зона без записи
#: о вписывании сохраняет.
AUTHORS_RPR = (
    '<a:rPr sz="2400" b="1"><a:solidFill><a:schemeClr val="accent1"/></a:solidFill>'
    '<a:latin typeface="+mn-lt"/></a:rPr>'
)
AUTHORS_PPR = '<a:pPr algn="ctr"/>'
A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def fragment(xml: str) -> Any:
    return etree.fromstring(f'<a:root xmlns:a="{A}">{xml}</a:root>')[0]


def example(size_pt: float | None = None):
    """Пример из двух надписей: у первой прогон без `rPr`, у второй — `pPr` и `rPr` автора
    и два абзаца. `size_pt` — кегль обеих зон в каталоге, с него вписывание стартует."""
    prs = Presentation()
    source = prs.slides.add_slide(prs.slide_layouts[6])
    bare = source.shapes.add_textbox(Emu(100), Emu(100), Emu(3000), Emu(1000))
    bare.text_frame.text = "Текст примера"
    styled = source.shapes.add_textbox(Emu(100), Emu(1300), Emu(3000), Emu(1000))
    styled.text_frame.text = "Подпись примера\nВторая строка примера"
    for paragraph in styled.text_frame.paragraphs:
        paragraph._p.insert(0, fragment(AUTHORS_PPR))
        paragraph.runs[0]._r.insert(0, fragment(AUTHORS_RPR))
    recipe = Recipe(
        recipe_id="ex001",
        example_index=1,
        part_name=str(source.part.partname),
        kind=RecipeKind.TEXT,
        zones=[
            Zone(
                zone_id="zb",
                xml_id=bare.shape_id,
                role=TypeLevel.BODY,
                capacity_chars=80,
                size_pt=size_pt,
            ),
            Zone(
                zone_id="zs",
                xml_id=styled.shape_id,
                role=TypeLevel.CAPTION,
                capacity_chars=80,
                size_pt=size_pt,
            ),
        ],
    )
    return prs, recipe, bare.shape_id, styled.shape_id


def slide_ir(recipe: Recipe, fit_report: dict[str, FitResult]) -> SlideIR:
    return SlideIR(
        slide_id="s01",
        layout_id="l01",
        variant="A",
        recipe_id=recipe.recipe_id,
        blocks=[
            BulletsBlock(
                block_id="b0",
                role=TextRole.BODY,
                items=[BulletItem(text="Первый пункт"), BulletItem(text="Второй пункт")],
                zone_id="zb",
            ),
            TextBlock(
                block_id="b1",
                role=TextRole.CAPTION,
                text="Наша подпись\nВторая строка",
                zone_id="zs",
            ),
        ],
        fit_report=fit_report,
    )


def test_every_run_of_the_zone_takes_the_fitted_size() -> None:
    """Кегль опущен — у каждого прогона зоны `sz` из `fit_report`, в сотых пункта;
    оформление автора не теряется. Прогон без `sz` в примере не трогается (D04) —
    это `test_a_run_without_the_authors_size_is_left_alone`."""
    prs, recipe, _bare, styled = example()
    ir = slide_ir(recipe, {"b1": FitResult(final_size_pt=13.5, strategy="shrink")})

    slide = clone_recipe(prs, recipe, ir)

    captions = runs(shape_by_id(slide, styled))
    assert len(captions) == 2
    for caption in captions:
        props = props_of(caption)
        assert props is not None and props.get("sz") == "1350"
        assert props.get("b") == "1", "полужирный автора потерян"
        assert props.find(f"{qn('a:solidFill')}/{qn('a:schemeClr')}").get("val") == "accent1"
        assert props.find(qn("a:latin")).get("typeface") == "+mn-lt"
        assert caption.index(props) == 0, "rPr обязан стоять перед a:t"


def expected_body(source: Any, paragraphs: list[str]) -> bytes:
    """`p:txBody` зоны, каким его писал код до правки: `bodyPr` и `lstStyle` примера,
    а абзацы — выписанные вручную, по абзацу на строку, с `pPr`/`rPr` автора."""
    body = copy.deepcopy(source.find(qn("p:txBody")))
    for paragraph in body.findall(qn("a:p")):
        body.remove(paragraph)
    for xml in paragraphs:
        body.append(fragment(xml))
    return canonical(body)


def authors_bodies(prs: Any, bare: int, styled: int) -> dict[int, bytes]:
    """Эталон зоны с оформлением автора: `p:txBody`, каким его писал код до правки RG29,
    выписанный вручную по абзацу на строку, с `pPr`/`rPr` автора."""
    source = prs.slides[0]
    return {
        bare: expected_body(
            shape_by_id(source, bare),
            [
                f"<a:p><a:r><a:t>{line}</a:t></a:r></a:p>"
                for line in ("Первый пункт", "Второй пункт")
            ],
        ),
        styled: expected_body(
            shape_by_id(source, styled),
            [
                f"<a:p>{AUTHORS_PPR}<a:r>{AUTHORS_RPR}<a:t>{line}</a:t></a:r></a:p>"
                for line in ("Наша подпись", "Вторая строка")
            ],
        ),
    }


def test_without_a_record_the_zone_is_written_as_before() -> None:
    """Записи нет — `p:txBody` зоны целиком тот же, что писал код до правки: число абзацев,
    `pPr`, `rPr` каждого прогона. Страж прежнего поведения: до правки он тоже зелёный."""
    prs, recipe, bare, styled = example()
    expected = authors_bodies(prs, bare, styled)

    slide = clone_recipe(prs, recipe, slide_ir(recipe, {}))

    for xml_id, body in expected.items():
        assert canonical(shape_by_id(slide, xml_id).find(qn("p:txBody"))) == body


@pytest.mark.parametrize(
    ("final_size_pt", "strategy"),
    [(60, "as_is"), (24, "as_is"), (24, "shrink")],
    ids=["as_is-above-the-zone", "as_is-equal", "shrink-equal"],
)
def test_a_zone_the_fitting_did_not_lower_keeps_the_authors_look(
    final_size_pt: float, strategy: str
) -> None:
    """D02/D03: кегль не опущен — `final_size_pt` не ниже кегля зоны в каталоге (24 pt),
    с какой бы `strategy` ни пришла запись. Зона та же, что без записи, байт в байт:
    у прогона без `rPr` в примере его нет и в файле. Так заголовки VK Education
    выросли 36 → 60 pt: `Zone.size_pt` каталога не равен кеглю прогона примера."""
    prs, recipe, bare, styled = example(size_pt=24)
    expected = authors_bodies(prs, bare, styled)
    record = FitResult(final_size_pt=final_size_pt, strategy=strategy)

    slide = clone_recipe(prs, recipe, slide_ir(recipe, {"b0": record, "b1": record}))

    for xml_id, body in expected.items():
        assert canonical(shape_by_id(slide, xml_id).find(qn("p:txBody"))) == body, (
            "кегль не опущен, а оформление автора изменилось"
        )


def test_the_fitted_size_never_exceeds_the_authors() -> None:
    """D02: вписывание опустило кегль (30 pt при зоне каталога в 60 pt), но не ниже
    авторского — у прогона остаётся `sz` автора, 24 pt, а не 30."""
    prs, recipe, bare, styled = example(size_pt=60)
    expected = authors_bodies(prs, bare, styled)[styled]
    shrunk = FitResult(final_size_pt=30, strategy="shrink")

    slide = clone_recipe(prs, recipe, slide_ir(recipe, {"b1": shrunk}))

    assert canonical(shape_by_id(slide, styled).find(qn("p:txBody"))) == expected, (
        "кегль вписывания больше авторского попал в файл"
    )


def test_a_title_that_yielded_its_size_is_written_lowered() -> None:
    """D03: заголовок уступил кегль (`_titles_yield_size` ставит блоку `size_pt` и вписывает
    заново), и запись приходит как `as_is` — 20 pt при зоне каталога в 60 pt. Кегль опущен,
    значит у прогонов `sz` 20, а не авторские 60 (VK Education ex015/ex017/ex030/ex033)."""
    prs, recipe, _bare, styled = example(size_pt=60)
    for run in runs(shape_by_id(prs.slides[0], styled)):
        props_of(run).set("sz", "6000")
    yielded = FitResult(final_size_pt=20, strategy="as_is")

    slide = clone_recipe(prs, recipe, slide_ir(recipe, {"b1": yielded}))

    assert [
        None if props_of(run) is None else props_of(run).get("sz")
        for run in runs(shape_by_id(slide, styled))
    ] == ["2000", "2000"], "опущенный кегль с записью as_is не дошёл до файла"
    assert all(props_of(run).get("b") == "1" for run in runs(shape_by_id(slide, styled)))


@pytest.mark.parametrize("strategy", ["shrink", "as_is"])
def test_a_run_without_the_authors_size_is_left_alone(strategy: str) -> None:
    """D04: у прогона примера нет `sz` — кегль унаследован из макета, и настоящего числа
    писатель не знает (`Zone.size_pt` с ним не совпадает). Кегль опущен (20 pt при зоне
    в 60 pt), а зона всё равно та же, что у автора: `sz` не пишется, `rPr` не создаётся.
    Иначе VK Tech s03 падал 16 → 7,8 pt, а VK Education s04/s11 рос 36 → 39."""
    prs, recipe, bare, styled = example(size_pt=60)
    expected = authors_bodies(prs, bare, styled)[bare]
    lowered = FitResult(final_size_pt=20, strategy=strategy)

    slide = clone_recipe(prs, recipe, slide_ir(recipe, {"b0": lowered}))

    assert canonical(shape_by_id(slide, bare).find(qn("p:txBody"))) == expected, (
        "прогону без sz автора писатель поставил кегль вслепую"
    )


# --- настоящий шаблон ---------------------------------------------------------------

TEMPLATE = "VK Tech шаблон.pptx"


def authors_props(path: Path, recipe: Recipe, xml_id: int) -> Any | None:
    """`rPr` первого прогона фигуры примера — прямо из файла шаблона, мимо писателя."""
    for slide in Presentation(str(path)).slides:
        if str(slide.part.partname).lstrip("/") == (recipe.part_name or "").lstrip("/"):
            return shape_by_id(slide, xml_id).find(f".//{qn('a:r')}/{qn('a:rPr')}")
    raise AssertionError(f"примера {recipe.part_name} в шаблоне нет")


def test_the_written_file_carries_the_fitted_size(tmp_path: Path) -> None:
    """Сквозной `write()`: зона с записью — кегль вписывания у всех прогонов, зона без
    записи — `rPr` автора примера целиком. Файл открывается python-pptx."""
    path = case_template(TEMPLATE)
    manifest = TemplateParser().parse(path, use_cache=False)
    ds = derive(manifest)

    def sized(recipe: Recipe) -> list[Zone]:
        """Зоны, у которых автор поставил кегль прогону, — на них видно и замену, и сохранение."""
        found = []
        for zone in recipe.zones:
            if zone.xml_id is None:
                continue
            props = authors_props(path, recipe, zone.xml_id)
            if props is not None and props.get("sz"):
                found.append(zone)
        return found

    found = next(((r, z) for r in ds.recipes if r.repeats <= 1 and len(z := sized(r)) >= 2), None)
    if found is None:
        pytest.fail(
            f"в каталоге «{TEMPLATE}» нет рецепта без повторов с двумя зонами, у которых "
            "автор поставил кегль прогону: сквозную проверку кегля ставить не на чем"
        )
    recipe, zones = found
    fitted, kept = zones[0], zones[1]
    authors_size = int(authors_props(path, recipe, fitted.xml_id).get("sz"))
    #: Ниже кегля автора и кегля зоны в каталоге: иначе кегль не опущен (D02, D03).
    size_pt = min(authors_size / 100, fitted.size_pt or authors_size / 100) / 2

    example = next(e for e in manifest.examples if e.slide_index == recipe.example_index)
    ir = SlideIR(
        slide_id="s01",
        layout_id=example.layout_id or manifest.layouts[0].layout_id,
        variant="A",
        recipe_id=recipe.recipe_id,
        blocks=[
            TextBlock(
                block_id=f"b{n}", role=TextRole.BODY, text=f"Наш текст {n}", zone_id=z.zone_id
            )
            for n, z in enumerate([fitted, kept])
        ],
        fit_report={"b0": FitResult(final_size_pt=size_pt, strategy="shrink")},
    )
    deck = DeckIR(deck_id="d01", template_id=manifest.template_id, variant="A", seed=1, slides=[ir])
    out = PptxWriter(path, manifest, design_system=ds).write(deck, tmp_path / "deck.pptx")

    slide = Presentation(str(out)).slides[0]

    fitted_runs = runs(shape_by_id(slide, fitted.xml_id))
    assert fitted_runs
    assert {props_of(run).get("sz") for run in fitted_runs} == {str(round(size_pt * 100))}

    expected = canonical(authors_props(path, recipe, kept.xml_id))
    kept_runs = runs(shape_by_id(slide, kept.xml_id))
    assert kept_runs
    assert all(canonical(props_of(run)) == expected for run in kept_runs), (
        "зона без записи потеряла оформление автора"
    )
