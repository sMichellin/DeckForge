"""Короткое слово не остаётся в конце строки. Change `no-hanging-prepositions`, Т4
(`docs/agents/requirements-from-notes-26-09.md`).

Заметка 26.09: «…при создании презентаций в / фирменном стиле» — предлог «в» остался
в конце строки. Правило одно на замер, pptx и html: замер, который переносит строку
не там, где её перенесёт PowerPoint, врёт о вписывании.

Сценарии — из дельты `openspec/changes/no-hanging-prepositions/specs/`.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation

from deckforge.domain.enums import TextRole
from deckforge.domain.slide import SlideIR, TextBlock
from deckforge.domain.template import TemplateManifest
from deckforge.layout.metrics import longest_word_em, split_paragraphs
from deckforge.layout.nonbreaking import NBSP, bind
from deckforge.parsing import TemplateParser
from tests.integration.test_native_objects import build_template
from tests.unit.test_compose_by_the_design_system import body_layout, written
from tests.unit.test_export_html import half, html_of

SCREENSHOT = "Генеративный ИИ не заменяет дизайнеров при создании презентаций в фирменном стиле"


def dotted(text: str) -> str:
    return text.replace(NBSP, "·")


# --- правило -----------------------------------------------------------------------------


def test_a_preposition_is_bound_to_the_next_word() -> None:
    """Нарушитель со скриншота: «в» на краю строки."""
    assert dotted(bind(SCREENSHOT)) == (
        "Генеративный ИИ не·заменяет дизайнеров при·создании презентаций в·фирменном стиле"
    )


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("Выручка 1 480 млн ₽", "Выручка 1·480·млн·₽"),
        ("Доступность 99,97 % за год", "Доступность 99,97·% за·год"),
        ("Срок — 14 месяцев", "Срок·— 14 месяцев"),
        ("Проблема не в содержании, а в стиле", "Проблема не·в·содержании, а·в·стиле"),
    ],
    ids=["разряды и единица", "процент", "тире", "подряд идущие короткие слова"],
)
def test_numbers_units_and_dashes(text: str, expected: str) -> None:
    assert dotted(bind(text)) == expected


@pytest.mark.parametrize(
    "text",
    ["из-за погоды", "Work in progress, and a plan", "Дизайнеры тратят часы", ""],
    ids=["дефис", "латиница", "без коротких слов", "пусто"],
)
def test_what_the_rule_does_not_touch(text: str) -> None:
    """Норма: слово внутри дефисного, латинский текст и текст без коротких слов."""
    assert bind(text) == text


def test_the_rule_is_idempotent() -> None:
    once = bind(SCREENSHOT)
    assert bind(once) == once


# --- замер -------------------------------------------------------------------------------


def test_measurement_sees_the_bound_pair_as_one_word() -> None:
    """Вписывание обязано мерить то, что PowerPoint не разорвёт: «в фирменном» — одно слово."""
    assert split_paragraphs("в фирменном") == [f"в{NBSP}фирменном"]
    pair = longest_word_em("в фирменном", font_family="Arial")
    alone = longest_word_em("фирменном", font_family="Arial")
    assert pair > alone


# --- файлы -------------------------------------------------------------------------------


def test_the_pptx_gets_the_same_rule(tmp_path: Path) -> None:
    template = build_template(tmp_path / "template.pptx")
    manifest = TemplateParser(cache_dir=tmp_path / "cache").parse(template)
    layout_id, idx = body_layout(manifest)
    block = TextBlock(block_id="b", role=TextRole.BODY, text=SCREENSHOT, placeholder_idx=idx)
    slide = SlideIR(slide_id="s01", layout_id=layout_id, variant="A", blocks=[block])

    out, fitted, _ = written(tmp_path, template, manifest, slide)

    texts = [
        shape.text_frame.text
        for shape in Presentation(str(out)).slides[0].shapes
        if shape.has_text_frame
    ]
    assert bind(SCREENSHOT) in texts
    assert fitted.blocks[0].text == SCREENSHOT, "IR правило не трогает — только файл"


def test_the_html_gets_the_same_rule(manifest: TemplateManifest, tmp_path: Path) -> None:
    block = TextBlock(block_id="b", role=TextRole.BODY, text=SCREENSHOT, **half(manifest, 0))

    html = html_of(manifest, tmp_path, block)

    assert f"в{NBSP}фирменном" in html
