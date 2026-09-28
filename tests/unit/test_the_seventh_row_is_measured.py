"""Строка 7 мерила плана Б по файлу колоды. Change `the-seventh-row-is-measured`.

Блоки ниже порога читаемости считает `scripts/check_deck_readable.py` — тот же замер,
что у приёмки итерации 25.09. Мерило только сводит его: блок — пара «слайд, фигура»,
а без файла колоды строка остаётся прочерком, а не нулём.
"""

from __future__ import annotations

import importlib
import sys
from pathlib import Path
from types import ModuleType

import pytest
from pptx import Presentation
from pptx.util import Cm, Pt

ROOT = Path(__file__).resolve().parents[2]


@pytest.fixture(scope="module")
def plan_b() -> ModuleType:
    scripts = ROOT / "scripts"
    sys.path.insert(0, str(scripts))
    try:
        return importlib.import_module("plan_b_metrics")
    finally:
        sys.path.remove(str(scripts))


def deck_with(run_dir: Path, sizes: list[float]) -> None:
    """Колода из одного слайда: по текстовой рамке на кегль, в каждой — два прогона текста."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for index, size in enumerate(sizes):
        box = slide.shapes.add_textbox(Cm(1), Cm(1 + index * 3), Cm(20), Cm(2))
        paragraph = box.text_frame.paragraphs[0]
        for word in ("Выручка", "выросла"):
            run = paragraph.add_run()
            run.text = f"{word} "
            run.font.size = Pt(size)
    (run_dir / "out").mkdir(parents=True)
    presentation.save(str(run_dir / "out" / "deck.pptx"))


def test_a_block_below_the_floor_counts_once(plan_b: ModuleType, tmp_path: Path) -> None:
    """Нарушитель: рамка 9 pt при пороге 10 — один блок, хоть прогонов в ней два."""
    deck_with(tmp_path, [9, 12])

    assert plan_b.below_floor(tmp_path) == 1


def test_a_readable_deck_is_zero(plan_b: ModuleType, tmp_path: Path) -> None:
    """Норма: всё не ниже порога — ноль."""
    deck_with(tmp_path, [12, 16])

    assert plan_b.below_floor(tmp_path) == 0


def test_without_the_deck_file_the_row_is_not_measured(
    plan_b: ModuleType, tmp_path: Path
) -> None:
    """Норма: файла колоды нет (фикстура) — `None`, в таблице прочерк."""
    assert plan_b.below_floor(tmp_path) is None
