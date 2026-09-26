"""Писатель не меняет геометрию фигур примера. Change `a-shape-outside-the-slide`.

Education s02: две фигуры колоды за правым краем. Замер до правки показал наследство
шаблона: слайд-пример сам несёт их с тем же вылетом, а писатель копирует фигуры глубокой
копией и `a:xfrm` не трогает. Правки писателя нет — здесь закреплено поведение, которое
делает находку «не нашей»: геометрия каждой уцелевшей фигуры и группы в колоде та же,
что в шаблоне, в том числе у фигуры за краем и после удаления лишнего.

Шов — `write()`: эталон геометрии читается из файла шаблона напрямую, мимо писателя.
Шаблон синтетический и строится в тесте — идёт везде, в том числе в CI.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu

from deckforge.designsystem import derive
from deckforge.designsystem.models import DesignSystem, Recipe, RecipeKind, TypeLevel, Zone
from deckforge.domain.enums import TextRole
from deckforge.domain.slide import DeckIR, SlideIR, TextBlock
from deckforge.domain.template import TemplateManifest
from deckforge.parsing import TemplateParser
from deckforge.rendering.writer import PptxWriter

Case = tuple[Path, TemplateManifest, DesignSystem, dict[str, int]]

#: Вылет декора за правый край, заданный при сборке шаблона, — известная величина,
#: а не посчитанная тем же способом, что и код.
OVERHANG = 1_337_000
DECOR_WIDTH = 1_837_000
DECOR_ROTATION = 15.0

#: Фигуры с геометрией — всё, что человек видит на слайде, включая сами группы.
SHAPES = (qn("p:sp"), qn("p:grpSp"), qn("p:pic"), qn("p:cxnSp"), qn("p:graphicFrame"))


def xfrm_of(node: Any) -> Any | None:
    if node.tag == qn("p:graphicFrame"):
        return node.find(qn("p:xfrm"))
    if node.tag == qn("p:grpSp"):
        return node.find(f"{qn('p:grpSpPr')}/{qn('a:xfrm')}")
    return node.find(f"{qn('p:spPr')}/{qn('a:xfrm')}")


def geometry(slide: Any) -> dict[int, tuple]:
    """`cNvPr id` → (`rot`, отражения, `off`, `ext`, `chOff`, `chExt`) на любой глубине групп."""
    out: dict[int, tuple] = {}
    for node in slide.shapes._spTree.iter(*SHAPES):
        props = node.find(f"./*/{qn('p:cNvPr')}")
        xfrm = xfrm_of(node)
        if props is None or xfrm is None:
            continue

        def point(tag: str, xfrm: Any = xfrm) -> tuple | None:
            child = xfrm.find(qn(tag))
            return tuple(sorted(child.attrib.items())) if child is not None else None

        out[int(props.get("id"))] = (
            xfrm.get("rot"),
            xfrm.get("flipH"),
            xfrm.get("flipV"),
            point("a:off"),
            point("a:ext"),
            point("a:chOff"),
            point("a:chExt"),
        )
    return out


def example_geometry(case: Case, recipe: Recipe) -> dict[int, tuple]:
    """Геометрия слайда-примера — из файла шаблона, мимо писателя."""
    path = case[0]
    for slide in Presentation(str(path)).slides:
        if str(slide.part.partname).lstrip("/") == recipe.part_name:
            return geometry(slide)
    raise AssertionError(f"в шаблоне нет части {recipe.part_name}")


@pytest.fixture(scope="module")
def synthetic(tmp_path_factory: pytest.TempPathFactory) -> Case:
    """Слайд-пример: заголовок, декор за правым краем (не зона, повёрнут), ряд из двух
    карточек-повторов в группе (плашка + надпись) и отдельная зона текста."""
    prs = Presentation()
    blank = next(layout for layout in prs.slide_layouts if layout.name == "Blank")
    slide = prs.slides.add_slide(blank)
    width = int(prs.slide_width)

    title = slide.shapes.add_textbox(Emu(400_000), Emu(300_000), Emu(6_000_000), Emu(600_000))
    title.text_frame.text = "Заголовок примера"
    decor = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Emu(width + OVERHANG - DECOR_WIDTH),
        Emu(1_200_000),
        Emu(DECOR_WIDTH),
        Emu(900_000),
    )
    decor.rotation = DECOR_ROTATION
    row = slide.shapes.add_group_shape()
    cards = []
    for n in range(2):
        left = 400_000 + n * 2_600_000
        plate = row.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Emu(left), Emu(2_400_000), Emu(2_400_000), Emu(1_500_000)
        )
        label = row.shapes.add_textbox(
            Emu(left + 100_000), Emu(2_500_000), Emu(2_200_000), Emu(600_000)
        )
        label.text_frame.text = f"Карточка примера {n + 1}"
        cards.append((plate.shape_id, label.shape_id))
    body = slide.shapes.add_textbox(Emu(400_000), Emu(4_400_000), Emu(6_000_000), Emu(1_200_000))
    body.text_frame.text = "Текст примера"

    path = tmp_path_factory.mktemp("synthetic") / "template.pptx"
    prs.save(str(path))
    manifest = TemplateParser().parse(path, use_cache=False)
    recipe = Recipe(
        recipe_id="ex001",
        example_index=1,
        part_name=str(slide.part.partname).lstrip("/"),
        kind=RecipeKind.TEXT,
        zones=[
            Zone(zone_id="title", xml_id=title.shape_id, role=TypeLevel.SLIDE_TITLE,
                 capacity_chars=60),
            Zone(zone_id="card1", xml_id=cards[0][1], role=TypeLevel.BODY, repeat=0,
                 capacity_chars=40),
            Zone(zone_id="card2", xml_id=cards[1][1], role=TypeLevel.BODY, repeat=1,
                 capacity_chars=40),
            Zone(zone_id="body", xml_id=body.shape_id, role=TypeLevel.BODY, capacity_chars=200),
        ],
        repeat_xml_ids=[list(card) for card in cards],
    )
    ds = derive(manifest).model_copy(update={"recipes": [recipe]})
    ids = {"decor": decor.shape_id, "row": row.shape_id, "spare_plate": cards[1][0],
           "body": body.shape_id}
    return path, manifest, ds, ids


def written(case: Case, zone_ids: list[str], tmp_path: Path) -> Any:
    path, manifest, ds, _ids = case
    recipe = ds.recipes[0]
    slide = SlideIR(
        slide_id="s01",
        layout_id=manifest.layouts[0].layout_id,
        variant="A",
        recipe_id=recipe.recipe_id,
        blocks=[
            TextBlock(block_id=f"b{n}", role=TextRole.BODY, text=f"Наш текст {zone_id}",
                      zone_id=zone_id)
            for n, zone_id in enumerate(zone_ids)
        ],
    )
    deck = DeckIR(
        deck_id="d01", template_id=manifest.template_id, variant="A", seed=1, slides=[slide]
    )
    out = PptxWriter(path, manifest, design_system=ds).write(deck, tmp_path / "deck.pptx")
    prs = Presentation(str(out))
    assert len(prs.slides) == 1
    return prs


def test_a_shape_outside_the_slide_keeps_the_templates_geometry(
    synthetic: Case, tmp_path: Path
) -> None:
    """Фигура примера за правым краем — в колоде тот же `off/ext/rot`, тот же вылет."""
    recipe, decor = synthetic[2].recipes[0], synthetic[3]["decor"]

    prs = written(synthetic, ["title", "card1", "card2", "body"], tmp_path)

    deck = geometry(prs.slides[0])
    assert decor in deck, "декор примера пропал из колоды"
    assert deck[decor] == example_geometry(synthetic, recipe)[decor]
    xfrm = prs.slides[0].shapes._spTree.xpath(f".//p:cNvPr[@id='{decor}']/../../p:spPr/a:xfrm")[0]
    off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
    right = int(off.get("x")) + int(ext.get("cx"))
    assert right - int(prs.slide_width) == OVERHANG
    assert xfrm.get("rot") == str(int(DECOR_ROTATION * 60_000))


def test_dropping_the_unfilled_does_not_move_what_is_left(
    synthetic: Case, tmp_path: Path
) -> None:
    """Удаление незаполненного повтора и незаполненной зоны не сдвигает и не растягивает
    уцелевшие фигуры, в том числе группу, из которой ушёл повтор."""
    recipe, ids = synthetic[2].recipes[0], synthetic[3]

    prs = written(synthetic, ["title", "card1"], tmp_path)

    deck, example = geometry(prs.slides[0]), example_geometry(synthetic, recipe)
    assert ids["spare_plate"] not in deck and ids["body"] not in deck, "лишнее не удалено"
    assert {ids["decor"], ids["row"]} <= deck.keys()
    assert deck == {xml_id: example[xml_id] for xml_id in deck}
