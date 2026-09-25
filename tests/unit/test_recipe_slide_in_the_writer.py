"""Слайд по рецепту: копия слайда-примера шаблона. Change `recipe-slide-in-the-writer`.

Главный шов — `write()`: всё, что видит человек, проверяется на открытом файле, а не
на промежуточных структурах. Шаблон кейса нужен настоящий: копирование фигур, связей
и `rPr` на синтетическом манифесте не проверить.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation

from deckforge.designsystem import derive
from deckforge.designsystem.models import DesignSystem, Recipe
from deckforge.domain.enums import TextRole
from deckforge.domain.slide import DeckIR, FitResult, SlideIR, TextBlock
from deckforge.domain.template import TemplateManifest
from deckforge.parsing import TemplateParser
from deckforge.rendering.writer import PptxWriter
from tests.case_templates import case_template

TEMPLATE = "VK Tech шаблон.pptx"


@pytest.fixture(scope="module")
def case() -> tuple[Path, TemplateManifest, DesignSystem]:
    path = case_template(TEMPLATE)
    manifest = TemplateParser().parse(path, use_cache=False)
    return path, manifest, derive(manifest)


def richest(ds: DesignSystem) -> Recipe:
    """Рецепт с наибольшим числом занятых повторов: на нём видно и замену, и удаление."""

    def rows(recipe: Recipe) -> int:
        return len({zone.repeat for zone in recipe.zones if zone.repeat is not None})

    return max(ds.recipes, key=lambda r: (rows(r), len(r.zones)))


def deck_by_recipe(manifest: TemplateManifest, recipe: Recipe, texts: list[str]) -> DeckIR:
    """Колода из одного слайда: текст раскладывается по зонам повторов, по одному на повтор."""
    first_of_repeat: dict[int, str] = {}
    for zone in recipe.zones:
        if zone.repeat is not None:
            first_of_repeat.setdefault(zone.repeat, zone.zone_id)
    blocks = [
        TextBlock(
            block_id=f"b{index}",
            role=TextRole.BODY,
            text=text,
            zone_id=first_of_repeat[index],
        )
        for index, text in enumerate(texts)
        if index in first_of_repeat
    ]
    slide = SlideIR(
        slide_id="s01",
        layout_id=manifest.layouts[0].layout_id,
        variant="A",
        blocks=blocks,
        recipe_id=recipe.recipe_id,
    )
    return DeckIR(
        deck_id="d01", template_id=manifest.template_id, variant="A", seed=1, slides=[slide]
    )


def written(case: tuple[Path, TemplateManifest, DesignSystem], deck: DeckIR, tmp_path: Path):
    path, manifest, ds = case
    out = PptxWriter(path, manifest, design_system=ds).write(deck, tmp_path / "deck.pptx")
    return Presentation(str(out))


def all_shapes(container) -> list:
    """Все фигуры слайда, включая вложенные в группы.

    Ряд повторов у шаблона бывает сгруппирован (у VK Tech сетка `ex013` — одна группа),
    и обход только верхнего уровня такой слайд видит пустым: ни надписей, ни удалённых
    повторов. Считать надо то же, что видит человек.
    """
    out = []
    for shape in container.shapes:
        out.append(shape)
        if getattr(shape, "shapes", None) is not None:
            out.extend(all_shapes(shape))
    return out


def texts_of(slide) -> list[str]:
    return [
        shape.text_frame.text.strip()
        for shape in all_shapes(slide)
        if shape.has_text_frame and shape.text_frame.text.strip()
    ]


def test_a_slide_by_recipe_carries_the_shapes_of_the_example(
    case: tuple[Path, TemplateManifest, DesignSystem], tmp_path: Path
) -> None:
    """Оформление шаблона живёт в примере: вне плейсхолдеров стоит до 96 % его содержимого."""
    _path, manifest, ds = case
    recipe = richest(ds)

    prs = written(case, deck_by_recipe(manifest, recipe, ["Первый", "Второй"]), tmp_path)

    assert len(prs.slides) == 1, "слайды-примеры шаблона в колоду не попали"
    slide = prs.slides[0]
    decor = [shape for shape in all_shapes(slide) if not shape.has_text_frame]
    assert decor, "на слайде нет ни одной нетекстовой фигуры — оформление примера потеряно"
    assert len(all_shapes(slide)) > len(texts_of(slide)), "фигур не больше, чем надписей"
    assert {"Первый", "Второй"} <= set(texts_of(slide))


def test_no_text_of_the_template_is_left_on_the_slide(
    case: tuple[Path, TemplateManifest, DesignSystem], tmp_path: Path
) -> None:
    """Нарушитель `template.sample_text_left`: чужая фраза в колоде хуже пустой рамки."""
    _path, manifest, ds = case
    recipe = richest(ds)
    ours = ["Первый", "Второй"]

    prs = written(case, deck_by_recipe(manifest, recipe, ours), tmp_path)

    assert set(texts_of(prs.slides[0])) <= set(ours)


def test_spare_repeats_are_dropped_whole(
    case: tuple[Path, TemplateManifest, DesignSystem], tmp_path: Path
) -> None:
    """Лишний повтор уходит целиком: плашка вместе со своим текстом, а не одна надпись."""
    _path, manifest, ds = case
    recipe = richest(ds)

    full = written(case, deck_by_recipe(manifest, recipe, ["A", "B", "C"]), tmp_path / "full")
    short = written(case, deck_by_recipe(manifest, recipe, ["A"]), tmp_path / "short")

    assert len(all_shapes(short.slides[0])) < len(all_shapes(full.slides[0]))


def test_a_slide_without_a_recipe_is_built_as_before(
    case: tuple[Path, TemplateManifest, DesignSystem], tmp_path: Path
) -> None:
    """Норма: колода без рецептов собирается прежним путём, примеры удаляются."""
    _path, manifest, _ds = case
    layout = next(lt for lt in manifest.layouts if lt.capacity.max_chars_title > 0)
    title = next(ph for ph in layout.placeholders if ph.role is TextRole.TITLE)
    slide = SlideIR(
        slide_id="s01",
        layout_id=layout.layout_id,
        variant="A",
        blocks=[
            TextBlock(
                block_id="b0",
                role=TextRole.TITLE,
                text="Заголовок",
                placeholder_idx=title.idx,
            )
        ],
        fit_report={"b0": FitResult(final_size_pt=24.0)},
    )
    deck = DeckIR(
        deck_id="d01", template_id=manifest.template_id, variant="A", seed=1, slides=[slide]
    )

    prs = written(case, deck, tmp_path)

    assert len(prs.slides) == 1
    assert "Заголовок" in texts_of(prs.slides[0])
