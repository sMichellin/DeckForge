"""Инварианты рендерера до записи файла и привязка к теме. Change (13) `pptx-writer`."""

from __future__ import annotations

from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.dml import MSO_THEME_COLOR

from deckforge.domain.content import Asset, Brief, ContentPackage, Dataset, Series
from deckforge.domain.enums import ColorRef, FontRef, ImageSource, TextRole
from deckforge.domain.slide import (
    BulletItem,
    BulletsBlock,
    ChartBlock,
    DeckIR,
    FitResult,
    IconBlock,
    ImageBlock,
    KpiBlock,
    KpiItem,
    SlideIR,
    SmartArtBlock,
    TableBlock,
    TextBlock,
)
from deckforge.domain.template import PlaceholderSpec, TemplateManifest
from deckforge.layout.fonts import FontLibrary
from deckforge.parsing import TemplateParser
from deckforge.rendering.images import contain_box, cover_crop
from deckforge.rendering.theme_binding import (
    THEME_COLORS,
    apply_theme_color,
    resolve_font,
    theme_font_token,
)
from deckforge.rendering.writer import (
    PptxWriter,
    SlideDegrader,
    SlideValidator,
    WriterError,
)
from tests.case_templates import case_template
from tests.unit.test_layout_fonts import make_font

# --- theme_binding -----------------------------------------------------------


@pytest.mark.parametrize("ref", list(ColorRef))
def test_every_color_ref_maps_to_the_same_theme_slot(ref: ColorRef) -> None:
    """Слот в XML обязан называться так же, как ссылка в IR: иначе смена темы перекрасит
    объект не в тот цвет."""
    assert MSO_THEME_COLOR.to_xml(THEME_COLORS[ref]) == ref.value


def test_theme_color_is_written_as_scheme_reference_not_rgb() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, 1, 1)
    run = box.text_frame.paragraphs[0].add_run()
    apply_theme_color(run.font, ColorRef.ACCENT2)
    apply_theme_color(box.fill, ColorRef.LT1)
    xml = box._element.xml
    assert '<a:schemeClr val="accent2"/>' in xml
    assert '<a:schemeClr val="lt1"/>' in xml
    assert "srgbClr" not in xml


def test_font_tokens_point_to_the_theme() -> None:
    assert theme_font_token(FontRef.MAJOR_LATIN) == "+mj-lt"
    assert theme_font_token(FontRef.MINOR_LATIN) == "+mn-lt"
    assert theme_font_token(FontRef.MAJOR_CS) == "+mj-cs"
    assert theme_font_token(FontRef.MINOR_CS) == "+mn-cs"


def test_resolve_font_takes_the_family_from_the_manifest(manifest: TemplateManifest) -> None:
    assert resolve_font(FontRef.MAJOR_LATIN, manifest) == manifest.theme.fonts.major_latin
    # Сложной письменности в синтетической теме нет — берётся латинская пара.
    assert resolve_font(FontRef.MINOR_CS, manifest) == manifest.theme.fonts.minor_latin


# --- images: геометрия cover / contain ----------------------------------------


def test_cover_crops_the_long_side_evenly() -> None:
    """Картинка 2:1 в квадрат: срезаем по четверти слева и справа, пропорции целы."""
    assert cover_crop(200, 100, 1000, 1000) == pytest.approx((0.25, 0.0, 0.25, 0.0))


def test_cover_of_a_tall_image_crops_top_and_bottom() -> None:
    assert cover_crop(100, 400, 1000, 1000) == pytest.approx((0.0, 0.375, 0.0, 0.375))


def test_cover_with_matching_aspect_does_not_crop() -> None:
    assert cover_crop(160, 90, 1600, 900) == pytest.approx((0.0, 0.0, 0.0, 0.0))


def test_contain_fits_inside_and_centers() -> None:
    x, y, cx, cy = contain_box(200, 100, 100, 50, 1000, 1000)
    assert (cx, cy) == (1000, 500)
    assert (x, y) == (100, 50 + 250)


# --- validate ----------------------------------------------------------------


def fit(size: float = 18, overflow: bool = False) -> FitResult:
    return FitResult(final_size_pt=size, overflow=overflow)


def title(text: str = "Выручка выросла на 37 %") -> TextBlock:
    return TextBlock(block_id="t", placeholder_idx=0, role=TextRole.TITLE, text=text)


def ok_slide(**update: object) -> SlideIR:
    base = SlideIR(
        slide_id="s01",
        layout_id="L07",
        variant="A",
        blocks=[
            title(),
            BulletsBlock(block_id="b", placeholder_idx=1, items=[BulletItem(text="пункт")]),
        ],
        fit_report={"t": fit(40), "b": fit(18)},
    )
    return base.model_copy(update=update)


@pytest.fixture
def writer(manifest: TemplateManifest, tmp_path: Path) -> PptxWriter:
    return PptxWriter(tmp_path / "нет-файла.pptx", manifest)


def problems(writer: PptxWriter, slide: SlideIR) -> str:
    with pytest.raises(WriterError) as exc:
        writer.validate(slide)
    return str(exc.value)


def test_valid_slide_passes(writer: PptxWriter) -> None:
    writer.validate(ok_slide())


def test_unknown_layout_is_rejected(writer: PptxWriter) -> None:
    assert "L99" in problems(writer, ok_slide(layout_id="L99"))


def test_empty_slide_is_rejected(writer: PptxWriter) -> None:
    """Слайд без блоков — это пустой фон шаблона в колоде, а не содержание."""
    assert "без блоков" in problems(writer, ok_slide(blocks=[], fit_report={}))


def test_missing_placeholder_is_rejected(writer: PptxWriter) -> None:
    block = TextBlock(block_id="t", placeholder_idx=7, role=TextRole.TITLE, text="а")
    assert "idx=7" in problems(writer, ok_slide(blocks=[block], fit_report={"t": fit(40)}))


def test_text_into_picture_placeholder_is_rejected(
    manifest: TemplateManifest, tmp_path: Path
) -> None:
    layout = manifest.layout("L07")
    assert layout is not None
    pic = PlaceholderSpec(idx=9, ph_type="PIC", x=0, y=0, cx=100, cy=100)
    with_pic = manifest.model_copy(
        update={
            "layouts": [
                layout.model_copy(update={"placeholders": [*layout.placeholders, pic]})
                if lay.layout_id == "L07" else lay
                for lay in manifest.layouts
            ]
        }
    )
    block = TextBlock(block_id="x", placeholder_idx=9, role=TextRole.BODY, text="а")
    slide = ok_slide(blocks=[block], fit_report={"x": fit(18)})
    assert "PIC" in problems(PptxWriter(tmp_path / "t.pptx", with_pic), slide)


def test_text_block_without_placeholder_or_coordinates_is_rejected(writer: PptxWriter) -> None:
    block = TextBlock(block_id="x", role=TextRole.BODY, text="а")
    assert "x" in problems(writer, ok_slide(blocks=[block], fit_report={"x": fit(18)}))


def test_coordinates_outside_margins_are_rejected(
    writer: PptxWriter, manifest: TemplateManifest
) -> None:
    block = TextBlock(
        block_id="x", role=TextRole.BODY, text="а",
        x=0, y=0, cx=manifest.slide_size.cx_emu, cy=manifest.slide_size.cy_emu,
    )
    assert "пол" in problems(writer, ok_slide(blocks=[block], fit_report={"x": fit(18)}))


def content_with(tmp_path: Path) -> ContentPackage:
    image = tmp_path / "photo.png"
    Image.new("RGB", (16, 9)).save(image)
    return ContentPackage(
        brief=Brief(purpose="report", audience="команда", target_slides=1),
        assets=[Asset(asset_id="a001", kind="image", path=str(image))],
    )


def test_full_bleed_image_is_allowed(
    writer: PptxWriter, manifest: TemplateManifest, tmp_path: Path
) -> None:
    image = ImageBlock(
        block_id="i", source=ImageSource.ASSET, asset_ref="a001",
        x=0, y=0, cx=manifest.slide_size.cx_emu, cy=manifest.slide_size.cy_emu,
    )
    writer.validate(
        ok_slide(blocks=[title(), image], fit_report={"t": fit(40)}), content_with(tmp_path)
    )


def test_image_without_asset_is_rejected(writer: PptxWriter, manifest: TemplateManifest) -> None:
    content = manifest.content_bbox
    image = ImageBlock(block_id="i", source=ImageSource.ASSET, asset_ref="a001",
                       x=content.x, y=content.y, cx=content.cx, cy=content.cy)
    assert "a001" in problems(writer, ok_slide(blocks=[image], fit_report={}))


def test_generated_image_names_the_change_that_adds_it(
    writer: PptxWriter, manifest: TemplateManifest
) -> None:
    content = manifest.content_bbox
    image = ImageBlock(block_id="i", source=ImageSource.GENERATED, prompt="город",
                       x=content.x, y=content.y, cx=content.cx, cy=content.cy)
    assert "(24)" in problems(writer, ok_slide(blocks=[image], fit_report={}))


def test_image_outside_the_slide_is_rejected(
    writer: PptxWriter, manifest: TemplateManifest, tmp_path: Path
) -> None:
    image = ImageBlock(
        block_id="i", source=ImageSource.ASSET, asset_ref="a001",
        x=1, y=0, cx=manifest.slide_size.cx_emu, cy=manifest.slide_size.cy_emu,
    )
    slide = ok_slide(blocks=[title(), image], fit_report={"t": fit(40)})
    with pytest.raises(WriterError) as exc:
        writer.validate(slide, content_with(tmp_path))
    assert "выходит за слайд" in str(exc.value)


def test_asset_that_is_not_an_image_is_rejected(
    writer: PptxWriter, manifest: TemplateManifest, tmp_path: Path
) -> None:
    content = content_with(tmp_path)
    Path(content.assets[0].path).write_bytes(b"not an image")
    box = manifest.content_bbox
    image = ImageBlock(block_id="i", source=ImageSource.ASSET, asset_ref="a001",
                       x=box.x, y=box.y, cx=box.cx, cy=box.cy)
    slide = ok_slide(blocks=[title(), image], fit_report={"t": fit(40)})
    with pytest.raises(WriterError, match="не картинка"):
        writer.validate(slide, content)


def test_slide_of_a_single_picture_is_rejected(
    writer: PptxWriter, manifest: TemplateManifest, tmp_path: Path
) -> None:
    """C3: слайд из одной картинки — растр, а не редактируемые объекты."""
    box = manifest.content_bbox
    image = ImageBlock(block_id="i", source=ImageSource.ASSET, asset_ref="a001",
                       x=box.x, y=box.y, cx=box.cx, cy=box.cy)
    with pytest.raises(WriterError, match="C3"):
        writer.validate(ok_slide(blocks=[image], fit_report={}), content_with(tmp_path))


def test_empty_text_is_rejected(writer: PptxWriter) -> None:
    """Пустой текст оставил бы на слайде «Введите заголовок»."""
    block = TextBlock(block_id="t", placeholder_idx=0, role=TextRole.TITLE, text="   ")
    assert "пуст" in problems(writer, ok_slide(blocks=[block], fit_report={"t": fit(40)}))


def test_empty_bullets_are_rejected(writer: PptxWriter) -> None:
    block = BulletsBlock(block_id="b", placeholder_idx=1, items=[BulletItem(text=" ")])
    slide = ok_slide(blocks=[title(), block], fit_report={"t": fit(40), "b": fit(18)})
    assert "пуст" in problems(writer, slide)


def test_two_blocks_in_one_placeholder_are_rejected(writer: PptxWriter) -> None:
    second = TextBlock(block_id="t2", placeholder_idx=0, role=TextRole.TITLE, text="Другой")
    slide = ok_slide(blocks=[title(), second], fit_report={"t": fit(40), "t2": fit(40)})
    assert "idx=0" in problems(writer, slide)


def test_placeholder_and_coordinates_together_are_rejected(
    writer: PptxWriter, manifest: TemplateManifest
) -> None:
    """Координаты рядом с плейсхолдером были бы молча проигнорированы."""
    block = TextBlock(
        block_id="t", placeholder_idx=0, role=TextRole.TITLE, text="а",
        x=0, y=0, cx=manifest.slide_size.cx_emu, cy=manifest.slide_size.cy_emu,
    )
    assert "координаты" in problems(writer, ok_slide(blocks=[block], fit_report={"t": fit(40)}))


def test_text_without_a_measurement_is_not_rejected(writer: PptxWriter) -> None:
    """Норма к `test_overflow_is_rejected_before_writing`: замера нет — слайд проходит.

    Прежде здесь был отказ: писатель не писал слайд, про который не известно, влезает ли
    текст. Требование «Отсутствие замера не роняет колоду» это решение отменило — блок
    пишется кеглем своей роли, подмена называется в `degradations`. Отказ делал подмену
    в писателе недостижимой: до неё слайд не доживал.
    """
    writer.validate(ok_slide(fit_report={"t": fit(40)}))


def test_overflow_is_rejected_before_writing(writer: PptxWriter) -> None:
    assert "переполн" in problems(
        writer, ok_slide(fit_report={"t": fit(40), "b": fit(12, overflow=True)})
    )


def test_size_outside_the_template_scale_is_rejected(writer: PptxWriter) -> None:
    assert "13" in problems(writer, ok_slide(fit_report={"t": fit(40), "b": fit(13)}))


# --- change (21): составные компоненты и иконки ----------------------------------------


def smartart(
    manifest: TemplateManifest, pattern: str = "process", **update: object
) -> SmartArtBlock:
    fields: dict[str, object] = {"items": ["Сбор", "Анализ"], **box_of(manifest), **update}
    return SmartArtBlock(block_id="s", pattern=pattern, **fields)  # type: ignore[arg-type]


def icon(manifest: TemplateManifest, query: str = "shield-check") -> IconBlock:
    b = manifest.content_bbox
    return IconBlock(block_id="i", query=query, x=b.x, y=b.y, cx=b.cy // 4, cy=b.cy // 4)


def test_smartart_with_coordinates_and_fit_passes(
    writer: PptxWriter, manifest: TemplateManifest
) -> None:
    slide = ok_slide(blocks=[title(), smartart(manifest)], fit_report={"t": fit(40), "s": fit()})
    writer.validate(slide)


def test_smartart_needs_coordinates(writer: PptxWriter) -> None:
    """Координаты компоненту нужны по-прежнему: подменить можно кегль, но не место."""
    block = SmartArtBlock(block_id="s", pattern="cycle", items=["а", "б"])
    message = problems(writer, ok_slide(blocks=[title(), block], fit_report={"t": fit(40)}))
    assert "без координат" in message


def test_smartart_with_an_empty_item_is_rejected(
    writer: PptxWriter, manifest: TemplateManifest
) -> None:
    block = smartart(manifest, items=["Сбор", " "])
    slide = ok_slide(blocks=[title(), block], fit_report={"t": fit(40), "s": fit()})
    assert "пуст" in problems(writer, slide)


def test_unsupported_pattern_left_undegraded_is_rejected(
    writer: PptxWriter, manifest: TemplateManifest
) -> None:
    # matrix с этой правки рисуется плитками; неподдержанной осталась пирамида.
    slide = ok_slide(blocks=[title(), smartart(manifest, "pyramid")], fit_report={"t": fit(40)})
    assert "pyramid" in problems(writer, slide)


def test_icon_with_coordinates_passes(writer: PptxWriter, manifest: TemplateManifest) -> None:
    writer.validate(ok_slide(blocks=[title(), icon(manifest)], fit_report={"t": fit(40)}))


def test_unknown_icon_and_icon_without_coordinates_are_rejected(
    writer: PptxWriter, manifest: TemplateManifest
) -> None:
    lost = IconBlock(block_id="j", query="shield-check")
    slide = ok_slide(blocks=[title(), icon(manifest, "нет-такой"), lost],
                     fit_report={"t": fit(40)})
    message = problems(writer, slide)
    assert "нет-такой" in message
    assert "без координат" in message


@pytest.fixture
def fonts(tmp_path: Path, manifest: TemplateManifest) -> FontLibrary:
    make_font(tmp_path, manifest.theme.fonts.major_latin, advance=500, bold=True)
    make_font(tmp_path, manifest.theme.fonts.minor_latin, advance=500)
    return FontLibrary([tmp_path])


def test_unsupported_pattern_degrades_to_fitted_bullets(
    manifest: TemplateManifest, fonts: FontLibrary
) -> None:
    degrader = SlideDegrader(manifest, fonts)
    slide = degrader.degrade(
        ok_slide(blocks=[title(), smartart(manifest, "hierarchy")], fit_report={"t": fit(40)}),
        None,
    )
    bullets = slide.blocks[1]
    assert isinstance(bullets, BulletsBlock)
    assert [item.text for item in bullets.items] == ["Сбор", "Анализ"]
    assert bullets.bbox == smartart(manifest).bbox
    assert not slide.fit_report["s"].overflow
    assert degrader.degradations == [
        "s01/s: smartart hierarchy → буллеты (паттерн не поддерживается)"
    ]


def test_smartart_that_does_not_fit_degrades_to_bullets(
    manifest: TemplateManifest, fonts: FontLibrary
) -> None:
    degrader = SlideDegrader(manifest, fonts)
    slide = ok_slide(blocks=[title(), smartart(manifest)],
                     fit_report={"t": fit(40), "s": fit(12, overflow=True)})
    degraded = degrader.degrade(slide, None)
    assert isinstance(degraded.blocks[1], BulletsBlock)
    assert degrader.degradations == ["s01/s: smartart process → буллеты (не влез)"]


def test_unknown_icon_is_dropped_with_a_record(
    manifest: TemplateManifest, fonts: FontLibrary
) -> None:
    degrader = SlideDegrader(manifest, fonts)
    slide = ok_slide(blocks=[title(), icon(manifest, "нет-такой")], fit_report={"t": fit(40)})
    degraded = degrader.degrade(slide, None)
    assert [b.block_id for b in degraded.blocks] == ["t"]
    assert degrader.degradations == ["s01/i: иконки «нет-такой» нет в Lucide — блок убран"]


def data_content(**dataset_update: object) -> ContentPackage:
    dataset = Dataset(dataset_id="d001", title="Выручка", categories=["2024", "2025"],
                      series=[Series(name="Россия", values=[1.0, 2.0])])
    return ContentPackage(
        brief=Brief(purpose="report", audience="команда", target_slides=1),
        datasets=[dataset.model_copy(update=dataset_update)],
    )


def box_of(manifest: TemplateManifest) -> dict[str, int]:
    b = manifest.content_bbox
    return {"x": b.x, "y": b.y, "cx": b.cx, "cy": b.cy // 2}


def test_chart_with_dataset_and_coordinates_passes(
    writer: PptxWriter, manifest: TemplateManifest
) -> None:
    chart = ChartBlock(block_id="c", chart_type="clustered_bar", dataset_ref="d001",
                       **box_of(manifest))
    writer.validate(ok_slide(blocks=[title(), chart], fit_report={"t": fit(40)}), data_content())


def test_chart_without_dataset_is_rejected(
    writer: PptxWriter, manifest: TemplateManifest
) -> None:
    chart = ChartBlock(block_id="c", chart_type="clustered_bar", dataset_ref="d404",
                       **box_of(manifest))
    slide = ok_slide(blocks=[title(), chart], fit_report={"t": fit(40)})
    with pytest.raises(WriterError, match="d404"):
        writer.validate(slide, data_content())


def test_chart_without_coordinates_is_rejected(writer: PptxWriter) -> None:
    chart = ChartBlock(block_id="c", chart_type="clustered_bar", dataset_ref="d001")
    slide = ok_slide(blocks=[title(), chart], fit_report={"t": fit(40)})
    with pytest.raises(WriterError, match="без координат"):
        writer.validate(slide, data_content())


def test_impossible_chart_is_named_in_validation(
    writer: PptxWriter, manifest: TemplateManifest
) -> None:
    """`validate` не деградирует сам: это делает `write`, и там подмена записывается."""
    chart = ChartBlock(block_id="c", chart_type="pie", dataset_ref="d001", **box_of(manifest))
    content = data_content(series=[Series(name="а", values=[1.0, 2.0]),
                                   Series(name="б", values=[3.0, 4.0])])
    slide = ok_slide(blocks=[title(), chart], fit_report={"t": fit(40)})
    with pytest.raises(WriterError, match="одна серия"):
        writer.validate(slide, content)


def test_a_table_without_a_measurement_is_not_rejected(
    writer: PptxWriter, manifest: TemplateManifest
) -> None:
    """Таблица с координатами, но без замера, пишется кеглем роли, а не отвергается."""
    table = TableBlock(block_id="tb", header=["а"], rows=[["б"]], **box_of(manifest))
    writer.validate(ok_slide(blocks=[title(), table], fit_report={"t": fit(40)}))


def test_table_without_rows_or_dataset_is_rejected(
    writer: PptxWriter, manifest: TemplateManifest
) -> None:
    table = TableBlock(block_id="tb", dataset_ref="d404", **box_of(manifest))
    slide = ok_slide(blocks=[title(), table], fit_report={"t": fit(40), "tb": fit(18)})
    assert "d404" in problems(writer, slide)


def test_kpi_needs_coordinates(writer: PptxWriter) -> None:
    kpi = KpiBlock(block_id="k", items=[KpiItem(value="37 %", label="рост")])
    message = problems(writer, ok_slide(blocks=[title(), kpi], fit_report={"t": fit(40)}))
    assert "без координат" in message


def test_all_problems_are_reported_at_once(writer: PptxWriter) -> None:
    message = problems(writer, ok_slide(fit_report={"t": fit(40), "b": fit(13, overflow=True)}))
    assert "13" in message
    assert "переполн" in message


def test_deck_for_another_template_is_rejected_without_creating_a_file(
    writer: PptxWriter, tmp_path: Path
) -> None:
    deck = DeckIR(deck_id="d", variant="A", template_id="sha256:" + "f" * 64, seed=1,
                  slides=[ok_slide()])
    out = tmp_path / "out.pptx"
    with pytest.raises(WriterError, match="template_id"):
        writer.write(deck, out)
    assert not out.exists()


def test_invalid_slide_anywhere_prevents_the_file(
    writer: PptxWriter, manifest: TemplateManifest, tmp_path: Path
) -> None:
    deck = DeckIR(deck_id="d", variant="A", template_id=manifest.template_id, seed=1,
                  slides=[ok_slide(), ok_slide(slide_id="s02", layout_id="L99")])
    out = tmp_path / "out.pptx"
    with pytest.raises(WriterError, match="s02"):
        writer.write(deck, out)
    assert not out.exists()


def test_overflowing_table_degrades_to_bullets_without_content(
    manifest: TemplateManifest, fonts: FontLibrary
) -> None:
    """Строки таблицы лежат в самом IR — пакет контента для подмены не нужен."""
    table = TableBlock(block_id="tb", header=["Год", "Выручка"], rows=[["2025", "2"]],
                       **box_of(manifest))
    degrader = SlideDegrader(manifest, fonts)
    slide = ok_slide(blocks=[title(), table],
                     fit_report={"t": fit(40), "tb": fit(12, overflow=True)})
    bullets = degrader.degrade(slide, None).blocks[1]
    assert isinstance(bullets, BulletsBlock)
    assert [item.text for item in bullets.items] == ["Год: 2025 — Выручка: 2"]


# --- RG5: писатель не падает (change `recipe-is-not-the-models-word`) ----------


def zoned(block_id: str, text: str, zone_id: str | None) -> TextBlock:
    return TextBlock(block_id=block_id, role=TextRole.BODY, text=text, zone_id=zone_id)


@pytest.fixture(scope="module")
def case() -> tuple[Path, TemplateManifest]:
    """Настоящий шаблон: подмена кегля проверяется на записанном файле, а не на замысле."""
    path = case_template("VK Tech шаблон.pptx")
    return path, TemplateParser().parse(path, use_cache=False)


def slide_on(manifest: TemplateManifest, **update: object) -> SlideIR:
    layout = next(lt for lt in manifest.layouts if lt.capacity.max_chars_title > 0)
    ph = next(p for p in layout.placeholders if p.role is TextRole.TITLE)
    base = SlideIR(
        slide_id="s01",
        layout_id=layout.layout_id,
        variant="A",
        blocks=[
            TextBlock(block_id="t", placeholder_idx=ph.idx, role=TextRole.TITLE, text="Итоги")
        ],
    )
    return base.model_copy(update=update)


def test_a_deck_is_written_when_a_measurement_is_missing(
    case: tuple[Path, TemplateManifest], tmp_path: Path
) -> None:
    """Нарушитель: записи о вписывании нет ни для одного блока.

    Колода из десяти слайдов лучше отсутствия колоды: блок пишется кеглем своей роли
    из типошкалы шаблона, подмена называется. Прежде здесь был `KeyError` на последней
    стадии прогона, из которого причина не читалась. Проверяется записью файла целиком:
    подмена, до которой слайд не доживает, ничего не стоит.
    """
    path, manifest = case
    writer = PptxWriter(path, manifest)
    slide = slide_on(manifest, fit_report={})
    deck = DeckIR(
        deck_id="d", variant="A", template_id=manifest.template_id, seed=1, slides=[slide]
    )
    out = tmp_path / "deck.pptx"

    writer.write(deck, out)

    assert out.is_file(), "колода не написана"
    assert any("s01/t" in line for line in writer.degradations), "подмена не названа"
    title_pt = manifest.typography(TextRole.TITLE).size_pt
    assert f"{title_pt:g}" in " ".join(writer.degradations)
    assert title_pt in manifest.size_ladder_pt, "кегль обязан быть из шкалы шаблона"


def test_a_measured_deck_keeps_its_measurements(
    case: tuple[Path, TemplateManifest], tmp_path: Path
) -> None:
    """Норма: замер есть — берётся он, подмены не случается."""
    path, manifest = case
    writer = PptxWriter(path, manifest)
    size = manifest.size_ladder_pt[0]
    slide = slide_on(manifest, fit_report={"t": FitResult(final_size_pt=size)})
    deck = DeckIR(
        deck_id="d", variant="A", template_id=manifest.template_id, seed=1, slides=[slide]
    )

    writer.write(deck, tmp_path / "deck.pptx")

    assert writer.degradations == []


def test_a_recipe_outside_the_catalogue_is_a_named_error(
    manifest: TemplateManifest,
) -> None:
    """Нарушитель: слайд назван собранным по рецепту, которого у шаблона нет.

    Выдуманный моделью рецепт проходил весь конвейер и падал в писателе `KeyError`
    через 138 секунд прогона. Утечка контракта обязана называть себя.
    """
    validator = SlideValidator(manifest, ["ex001", "ex002"])

    out = validator.problems(ok_slide(recipe_id="L12_15_title_closing"), None)

    assert out, "рецепт вне каталога прошёл молча"
    assert "s01" in out[0]
    assert "L12_15_title_closing" in out[0]


def test_a_recipe_from_the_catalogue_passes(manifest: TemplateManifest) -> None:
    """Норма: рецепт в каталоге есть — проверка о нём молчит."""
    validator = SlideValidator(manifest, ["ex001", "ex002"])

    out = validator.problems(ok_slide(recipe_id="ex001"), None)

    assert out == []


def test_a_catalogue_that_is_not_given_says_nothing_about_recipes(
    manifest: TemplateManifest,
) -> None:
    """Норма: без каталога о композициях не судят — html зовёт валидатор именно так."""
    assert SlideValidator(manifest).problems(ok_slide(recipe_id="что-угодно"), None) == []


def test_a_half_zoned_slide_is_not_a_slide_by_recipe(manifest: TemplateManifest) -> None:
    """Нарушитель: часть блоков стоит в зонах шаблона, часть — нет.

    Такой слайд не собрать ни одним из двух способов, и его блоки вне зон обязаны
    пройти обычные проверки. Прежде писатель судил по каждому блоку отдельно,
    и блок вне зоны уезжал в запись без координат.
    """
    validator = SlideValidator(manifest, ["ex001"])
    slide = ok_slide(
        recipe_id="ex001",
        blocks=[zoned("a", "в зоне", "z1"), zoned("b", "без зоны", None)],
        fit_report={},
    )

    out = validator.problems(slide, None)

    assert any("s01/b" in line for line in out), "блок вне зоны прошёл как блок рецепта"


def test_a_fully_zoned_slide_is_a_slide_by_recipe(manifest: TemplateManifest) -> None:
    """Норма: все блоки в зонах — слайд собран по рецепту, координат с него не спрашивают."""
    validator = SlideValidator(manifest, ["ex001"])
    slide = ok_slide(
        recipe_id="ex001",
        blocks=[zoned("a", "раз", "z1"), zoned("b", "два", "z2")],
        fit_report={},
    )

    assert validator.problems(slide, None) == []
