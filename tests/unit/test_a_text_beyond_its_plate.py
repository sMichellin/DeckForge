"""Текст вышел за свою плашку — находка. Change `a-text-beyond-its-plate`, план Б, круг 3,
задача C4 (#245) — мерило для B2.

Прогон `a3a8f3a2319b`, WorkSpace s03: карточка — сама текстовая фигура с заливкой и верхним
отступом под иконку; вписывание намерило строки без отступа — «влезло», а последняя строка
ушла за край карточки. Проверка `layout.text_beyond_plate` меряет готовый файл. Правило 7:
нарушитель и норма. Шрифт синтетический — замер одинаков на любой машине.

Сценарии — из дельты `openspec/changes/a-text-beyond-its-plate/specs/audit-deterministic/`.
"""

from __future__ import annotations

from dataclasses import replace
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from deckforge.audit.deterministic.layout import text_beyond_plate
from deckforge.designsystem import derive
from deckforge.designsystem.models import Recipe, RecipeKind, TypeLevel, Zone
from deckforge.domain.enums import TextRole
from deckforge.domain.slide import DeckIR, SlideIR, TextBlock
from deckforge.domain.template import TemplateManifest
from deckforge.layout import fonts as font_module
from deckforge.layout.fonts import FontLibrary
from tests.unit._audit_builders import context_for, deck
from tests.unit.test_layout_fonts import make_font

CHECK = "layout.text_beyond_plate"
CARD = (4_000_000, 900_000, 4_500_000, 1_700_000)
TOP_INSET = 756_000
LONG = " ".join(["Инструменты генерируют содержание, но хуже держат структуру"] * 3)
SHORT = "Короткая строка"
RUNS = Path(__file__).resolve().parents[1] / "fixtures" / "runs"


@pytest.fixture(autouse=True)
def synthetic_font(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    folder = tmp_path / "fonts"
    folder.mkdir()
    make_font(folder, "Deck Sans", advance=500)
    monkeypatch.setattr(font_module.FontLibrary, "default", classmethod(
        lambda cls: FontLibrary([folder])
    ))


def _id(shape: object, xml_id: int) -> None:
    shape._element.find(f"./*/{qn('p:cNvPr')}").set("id", str(xml_id))  # type: ignore[attr-defined]


def _text(shape: object, text: str) -> None:
    frame = shape.text_frame  # type: ignore[attr-defined]
    frame.text = text
    for run in frame.paragraphs[0].runs:
        run.font.size = Pt(18)
        run.font.name = "Deck Sans"


def cards_file(path: Path, text: str, *, filled: bool = True, plate_behind: bool = False) -> Path:
    """Страница: карточка 901 — текстовая фигура с отступом под иконку; `filled` — сама
    залита; `plate_behind` — вместо этого залитая плашка 950 под ней."""
    prs = Presentation()
    page = prs.slides.add_slide(prs.slide_layouts[6])
    x, y, cx, cy = (Emu(v) for v in CARD)
    if plate_behind:
        plate = page.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, cx, cy)
        _id(plate, 950)
    card = page.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, cx, cy)
    _id(card, 901)
    if not filled:
        card.fill.background()
    body = card.text_frame._txBody.find(qn("a:bodyPr"))
    body.set("tIns", str(TOP_INSET))
    body.set("bIns", "0")
    body.set("anchor", "t")
    _text(card, text)
    prs.save(str(path))
    return path


def run(manifest: TemplateManifest, path: Path, *, zone_xml: int = 901) -> list:
    recipe = Recipe(recipe_id="ex005", example_index=1, kind=RecipeKind.CARDS, zones=[
        Zone(zone_id="z1", xml_id=zone_xml, role=TypeLevel.BODY, capacity_chars=200),
    ])
    slide = SlideIR(slide_id="s03", layout_id="L01", variant="A", recipe_id="ex005", blocks=[
        TextBlock(block_id="b1", role=TextRole.BODY, text="x", zone_id="z1"),
    ])
    design = derive(manifest).model_copy(update={"recipes": [recipe]})
    ctx = context_for(CHECK, deck(slide), manifest, deck_path=path)
    return list(text_beyond_plate(replace(ctx, design_system=design)))


def test_a_long_text_under_an_icon_inset_leaves_its_card(
    manifest: TemplateManifest, tmp_path: Path
) -> None:
    """Нарушитель (WorkSpace s03): строки плюс отступ под иконку выше карточки."""
    findings = run(manifest, cards_file(tmp_path / "deck.pptx", LONG))

    assert [(f.slide_id, f.evidence["shape_id"]) for f in findings] == [("s03", "901")]
    assert int(findings[0].evidence["beyond_emu"]) > 0


def test_a_short_text_stays_on_its_card(manifest: TemplateManifest, tmp_path: Path) -> None:
    """Норма: одна строка под отступом умещается в карточку."""
    assert run(manifest, cards_file(tmp_path / "deck.pptx", SHORT)) == []


def test_a_plate_behind_the_text_is_the_plate(manifest: TemplateManifest, tmp_path: Path) -> None:
    """Нарушитель: текстовая фигура без заливки, плашка — отдельная фигура под ней."""
    path = cards_file(tmp_path / "deck.pptx", LONG, filled=False, plate_behind=True)

    assert [f.evidence["shape_id"] for f in run(manifest, path)] == ["901"]


def test_a_text_without_a_plate_is_not_this_finding(
    manifest: TemplateManifest, tmp_path: Path
) -> None:
    """Норма: текст на фоне слайда — плашки нет, вылет судят проверки переполнения."""
    assert run(manifest, cards_file(tmp_path / "deck.pptx", LONG, filled=False)) == []


def test_the_text_of_the_author_is_not_judged(manifest: TemplateManifest, tmp_path: Path) -> None:
    """Норма: фигура не наша зона — надпись автора, его композиция."""
    assert run(manifest, cards_file(tmp_path / "deck.pptx", LONG), zone_xml=777) == []


def test_without_the_file_the_check_is_skipped(manifest: TemplateManifest) -> None:
    from deckforge.audit.registry import CheckUnavailable

    slide = SlideIR(slide_id="s03", layout_id="L01", variant="A", blocks=[
        TextBlock(block_id="b1", role=TextRole.TITLE, text="x"),
    ])
    with pytest.raises(CheckUnavailable):
        list(text_beyond_plate(context_for(CHECK, DeckIR(
            deck_id="d", variant="A", template_id=manifest.template_id, seed=1, slides=[slide]
        ), manifest)))


def test_the_card_of_workspace_s03_leaves_its_plate(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Мерило C4: WorkSpace 29.09 путём `by_example` — s03, карточка 442, как в `a3a8f3a2319b`.

    Замер настоящими шрифтами шаблона: без точной Arial (Linux без msttcorefonts) строки
    переносятся иначе, и граница находки сдвигается — тогда тест пропускается, а не врёт."""
    monkeypatch.undo()
    folder = RUNS / "2026-09-29" / "workspace"
    template = (Path(__file__).resolve().parents[1] / "fixtures" / "templates"
                / "VK_WorkSpace_Клиентская_конференция_Шаблон_03.pptx")
    if not folder.is_dir() or not template.is_file():
        pytest.skip("нет фикстуры WorkSpace 29.09 или шаблона")
    if not FontLibrary.default().resolve("Arial", bold=False).exact:
        pytest.skip("Arial не найдена точно: замер строк шаблона на этой машине другой")
    from deckforge.export.pptx import export_pptx
    from deckforge.pipeline.replay import from_fixture

    snapshot = from_fixture(folder)
    path = export_pptx(snapshot.deck, snapshot.manifest, template, tmp_path / "deck.pptx",
                       content=snapshot.content, design_system=snapshot.design_system,
                       by_example=True)
    ctx = context_for(CHECK, snapshot.deck, snapshot.manifest, deck_path=path)

    findings = list(text_beyond_plate(replace(ctx, design_system=snapshot.design_system)))

    assert [(f.slide_id, f.evidence["shape_id"]) for f in findings] == [("s03", "442")]
