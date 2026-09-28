"""Замер читаемости: фигура за краем сверяется с шаблоном.
Change `the-edge-check-knows-the-template` (RG58).

Прогон `96ef159` дал на трёх колодах единственную находку края — и она оказалась не наша:
фигура стоит за краем и в слайде-примере шаблона, ровно на столько же. Писатель копирует
пример целиком и фигур не двигает. Фигура за краем **по нашей вине** остаётся браком,
поэтому глушить проверку нельзя — её надо научить сравнивать.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))

from check_deck_readable import check_deck, outside_in_template  # noqa: E402
from deckforge.layout.fonts import FontLibrary  # noqa: E402

EMU_PER_INCH = 914400
#: Имя фигуры — ключ сверки: писатель копирует его вместе с фигурой.
NAME = "Google Shape;448;p50"
#: Вылет фигуры автора за правый край — число из примера Education (слайд 18).
AUTHOR_OVERHANG = 414587


@pytest.fixture(scope="module")
def library() -> FontLibrary:
    return FontLibrary.default()


def _file(path: Path, overhang: int, *, name: str = NAME) -> Path:
    """Файл из одного слайда с надписью, вылезающей за правый край на `overhang`."""
    presentation = Presentation()
    width = int(presentation.slide_width or 0)
    box_width = 2 * EMU_PER_INCH
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(
        Emu(width - box_width + overhang), Emu(EMU_PER_INCH), Emu(box_width), Emu(EMU_PER_INCH)
    )
    box.name = name
    box.text_frame.text = "Автоматизация вёрстки"
    presentation.save(str(path))
    return path


def _kinds(path: Path, library: FontLibrary, template: Path | None) -> list[str]:
    report = check_deck(
        path,
        size_floor_pt=10.0,
        library=library,
        outside=outside_in_template(template),
    )
    return [finding.kind for finding in report.findings]


def test_the_same_overhang_as_in_the_template_is_not_a_finding(
    tmp_path: Path, library: FontLibrary
) -> None:
    """Норма: фигура вылезает на столько же, на сколько у автора."""
    template = _file(tmp_path / "template.pptx", AUTHOR_OVERHANG)
    deck = _file(tmp_path / "deck.pptx", AUTHOR_OVERHANG)

    report = check_deck(
        deck, size_floor_pt=10.0, library=library, outside=outside_in_template(template)
    )

    assert [f.kind for f in report.findings] == []
    assert report.outside_in_template == 1


def test_moving_the_shape_further_out_is_a_finding(tmp_path: Path, library: FontLibrary) -> None:
    """Нарушитель: в колоде фигура вылезла дальше, чем в шаблоне."""
    template = _file(tmp_path / "template.pptx", AUTHOR_OVERHANG)
    deck = _file(tmp_path / "deck.pptx", 900000)

    report = check_deck(
        deck, size_floor_pt=10.0, library=library, outside=outside_in_template(template)
    )

    assert [f.kind for f in report.findings] == ["за краем слайда"]
    assert "900000" in report.findings[0].message


def test_a_shape_the_template_does_not_have_is_a_finding(
    tmp_path: Path, library: FontLibrary
) -> None:
    """Нарушитель: фигуры с таким именем в шаблоне нет вовсе."""
    template = _file(tmp_path / "template.pptx", AUTHOR_OVERHANG, name="Google Shape;1;p1")
    deck = _file(tmp_path / "deck.pptx", AUTHOR_OVERHANG)

    assert _kinds(deck, library, template) == ["за краем слайда"]


def test_without_a_template_any_overhang_is_a_finding(
    tmp_path: Path, library: FontLibrary
) -> None:
    """Норма прежнего поведения: сравнивать не с чем — вылет остаётся находкой."""
    deck = _file(tmp_path / "deck.pptx", AUTHOR_OVERHANG)

    assert _kinds(deck, library, None) == ["за краем слайда"]
