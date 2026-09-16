"""Живой рендер через LibreOffice. Change (6) `template-preview-render`.

Пропускаются, если LibreOffice недоступен: CI ставит только Python, а образ с офисом
собирается отдельно (`docker/Dockerfile.libreoffice`).

    podman build -f docker/Dockerfile.libreoffice -t deckforge:soffice .
"""

from __future__ import annotations

from pathlib import Path

import pytest

from deckforge.parsing import TemplateParser
from deckforge.rendering.layout_deck import build_layout_deck
from deckforge.rendering.layout_preview import CompositePreview, SofficePreview
from deckforge.rendering.soffice import SofficeRenderer, SofficeUnavailableError

pytestmark = [
    pytest.mark.needs_libreoffice,
    pytest.mark.slow,
    pytest.mark.skipif(
        not SofficeRenderer().available, reason="нет LibreOffice или pdftoppm"
    ),
]


@pytest.fixture(scope="module")
def template(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """Настоящий .pptx, а не наша синтетическая фикстура.

    `tests/integration/synthetic_template.py` собран так, чтобы гонять **наш** парсер:
    он читает пакет как zip и к объявлениям типов содержимого нетребователен. LibreOffice
    и python-pptx требовательны и такой файл отвергают — и правильно делают.

    Здесь нужен файл, который откроется в настоящем редакторе, поэтому берётся шаблон
    по умолчанию из python-pptx. Для наших целей он ровно так же «незнакомый», как любой
    другой: код о нём ничего не знает.
    """
    from pptx import Presentation

    path = tmp_path_factory.mktemp("tpl") / "default.pptx"
    Presentation().save(str(path))
    return path


def test_cyrillic_is_rendered_not_substituted() -> None:
    """Гейт дня 1 из плана: без шрифтов с кириллицей весь VLM-аудит врёт молча."""
    import subprocess
    import sys

    root = Path(__file__).resolve().parents[2]
    result = subprocess.run(
        [sys.executable, str(root / "scripts" / "check_cyrillic_render.py")],
        capture_output=True,
        text=True,
    )
    assert result.returncode == 0, result.stderr


def test_deck_converts_to_one_png_per_slide(tmp_path: Path) -> None:
    from pptx import Presentation

    prs = Presentation()
    for _ in range(3):
        prs.slides.add_slide(prs.slide_layouts[6])
    deck = tmp_path / "deck.pptx"
    prs.save(str(deck))

    pages = SofficeRenderer().to_pngs(deck, tmp_path / "out")
    assert len(pages) == 3
    assert all(p.read_bytes().startswith(b"\x89PNG") for p in pages)


def test_layout_deck_has_a_slide_per_layout(template: Path, tmp_path: Path) -> None:
    manifest = TemplateParser().parse(template, use_cache=False)
    order = build_layout_deck(template, manifest, tmp_path / "layouts.pptx")

    assert len(order) == len(manifest.layouts)
    assert list(order.values()) == [layout.layout_id for layout in manifest.layouts], (
        "порядок обязан совпадать: иначе превью уедут не к тем макетам"
    )


def test_rendered_layouts_alone_do_not_show_structure(template: Path, tmp_path: Path) -> None:
    """Пустые плейсхолдеры в PowerPoint не видны: у шаблона без декора все страницы белые.

    Это не дефект рендера, а причина, по которой он не может быть единственным
    поставщиком превью для классификатора.
    """
    manifest = TemplateParser().parse(template, use_cache=False)
    preview = SofficePreview(template, manifest, tmp_path)

    images = [
        preview.render(layout, manifest.slide_size.cx_emu, manifest.slide_size.cy_emu)
        for layout in manifest.layouts
    ]
    assert all(blob.startswith(b"\x89PNG") for blob in images)
    assert len(set(images)) == 1, "шаблон без оформления даёт одинаковые пустые страницы"


def test_composite_preview_shows_both_decor_and_structure(
    template: Path, tmp_path: Path
) -> None:
    """Схема поверх рендера: видно и оформление шаблона, и места под контент."""
    manifest = TemplateParser().parse(template, use_cache=False)
    composite = CompositePreview(template, manifest, tmp_path)

    images = [
        composite.render(layout, manifest.slide_size.cx_emu, manifest.slide_size.cy_emu)
        for layout in manifest.layouts
    ]
    assert all(blob.startswith(b"\x89PNG") for blob in images)
    assert len(set(images)) > 1, "разные макеты обязаны давать разные картинки"
    assert composite.fallbacks == 0


def test_composite_falls_back_to_the_schematic(template: Path, tmp_path: Path) -> None:
    """Без LibreOffice превью без декора хуже полного, но лучше отсутствующего."""
    manifest = TemplateParser().parse(template, use_cache=False)
    composite = CompositePreview(
        template, manifest, tmp_path, renderer=SofficeRenderer(binary="нет-такого")
    )
    blob = composite.render(
        manifest.layouts[0], manifest.slide_size.cx_emu, manifest.slide_size.cy_emu
    )
    assert blob.startswith(b"\x89PNG")
    assert composite.fallbacks == 1


def test_whole_deck_is_converted_once(template: Path, tmp_path: Path) -> None:
    """Запуск LibreOffice стоит секунды: 37 макетов нельзя конвертировать по одному."""
    manifest = TemplateParser().parse(template, use_cache=False)
    preview = SofficePreview(template, manifest, tmp_path)

    calls = 0
    original = preview._render_all

    def counted() -> dict[str, bytes]:
        nonlocal calls
        calls += 1
        return original()

    preview._render_all = counted  # type: ignore[method-assign]
    for layout in manifest.layouts:
        preview.render(layout, manifest.slide_size.cx_emu, manifest.slide_size.cy_emu)
    assert calls == 1


def test_previews_can_be_saved_for_the_ui(template: Path, tmp_path: Path) -> None:
    manifest = TemplateParser().parse(template, use_cache=False)
    saved = SofficePreview(template, manifest, tmp_path).save_previews(tmp_path / "png")

    assert set(saved) == {layout.layout_id for layout in manifest.layouts}
    assert all(path.exists() and path.stat().st_size > 0 for path in saved.values())


def test_unknown_layout_is_an_error(template: Path, tmp_path: Path) -> None:
    manifest = TemplateParser().parse(template, use_cache=False)
    preview = SofficePreview(template, manifest, tmp_path)
    stranger = manifest.layouts[0].model_copy(update={"layout_id": "НЕТ ТАКОГО"})

    with pytest.raises(SofficeUnavailableError, match="НЕТ ТАКОГО"):
        preview.render(stranger, manifest.slide_size.cx_emu, manifest.slide_size.cy_emu)


def test_only_the_requested_layouts_are_rendered(template: Path, tmp_path: Path) -> None:
    """Конвертация всего шаблона занимает 23–53 с при бюджете парсинга 25 с (§12).

    Спорных макетов, ради которых зовут модель, обычно втрое меньше — их и рендерим.
    """
    manifest = TemplateParser().parse(template, use_cache=False)
    wanted = [manifest.layouts[0].layout_id, manifest.layouts[-1].layout_id]

    preview = SofficePreview(template, manifest, tmp_path, only=wanted)
    saved = preview.save_previews(tmp_path / "png")
    assert set(saved) == set(wanted)

    with pytest.raises(SofficeUnavailableError):
        preview.render(manifest.layouts[1], manifest.slide_size.cx_emu, manifest.slide_size.cy_emu)
