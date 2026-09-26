"""Фигура за краем на Education — наследство шаблона. Change `a-shape-outside-the-slide`.

Находка прогона 25.09 (Education s02, раньше s03): две фигуры за правым краем на
~1 337 000 EMU. Слайд-пример 18 шаблона (рецепт `ex018`, вид `metrics`) несёт фигуры
445 и 447 с тем же вылетом — 1 338 791 и 1 337 452 EMU (замер до правки, proposal).
Тест держит это на рецепте, а не на номере слайда колоды: колода по `ex018` — геометрия
каждой уцелевшей фигуры та же, что в шаблоне, 445 и 447 на месте с тем же вылетом.

Эталон геометрии читается из файла шаблона напрямую, мимо писателя. Шаблона кейса
в чекауте нет — тест пропускается (`tests/case_templates.py`).
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

from deckforge.designsystem import derive
from deckforge.designsystem.models import DesignSystem, Recipe
from deckforge.domain.enums import TextRole
from deckforge.domain.slide import DeckIR, SlideIR, TextBlock
from deckforge.domain.template import TemplateManifest
from deckforge.parsing import TemplateParser
from deckforge.rendering.writer import PptxWriter
from tests.case_templates import case_template
from tests.unit.test_a_shape_outside_the_slide import geometry

Case = tuple[Path, TemplateManifest, DesignSystem]

TEMPLATE = "Шаблон презентации VK Education.pptx"
RECIPE = "ex018"

#: Вылет за правый край в самом шаблоне, EMU, — замер разведки 25.09 (proposal).
OVERHANG = {445: 1_338_791, 447: 1_337_452}


@pytest.fixture(scope="module")
def education() -> Case:
    path = case_template(TEMPLATE)
    manifest = TemplateParser().parse(path, use_cache=False)
    return path, manifest, derive(manifest)


def ex018(case: Case) -> Recipe:
    recipe = next((r for r in case[2].recipes if r.recipe_id == RECIPE), None)
    if recipe is None:
        pytest.fail(f"в каталоге Education нет рецепта {RECIPE}: замер proposal устарел")
    return recipe


def example_geometry(case: Case, recipe: Recipe) -> dict[int, tuple]:
    for slide in Presentation(str(case[0])).slides:
        if str(slide.part.partname).lstrip("/") == recipe.part_name:
            return geometry(slide)
    raise AssertionError(f"в шаблоне нет части {recipe.part_name}")


def written(case: Case, recipe: Recipe, zone_ids: list[str], tmp_path: Path) -> Any:
    path, manifest, ds = case
    slide = SlideIR(
        slide_id="s01",
        layout_id=manifest.layouts[0].layout_id,
        variant="A",
        recipe_id=recipe.recipe_id,
        blocks=[
            TextBlock(block_id=f"b{n}", role=TextRole.BODY, text=f"Текст {n}", zone_id=zone_id)
            for n, zone_id in enumerate(zone_ids)
        ],
    )
    deck = DeckIR(
        deck_id="d01", template_id=manifest.template_id, variant="A", seed=1, slides=[slide]
    )
    out = PptxWriter(path, manifest, design_system=ds).write(deck, tmp_path / "deck.pptx")
    return Presentation(str(out))


def overhang(prs: Any, xml_id: int) -> int:
    xfrm = prs.slides[0].shapes._spTree.xpath(f".//p:cNvPr[@id='{xml_id}']/../../p:spPr/a:xfrm")[0]
    right = int(xfrm.find(qn("a:off")).get("x")) + int(xfrm.find(qn("a:ext")).get("cx"))
    return right - int(prs.slide_width)


def test_ex018_keeps_the_shapes_outside_the_slide_as_the_template_does(
    education: Case, tmp_path: Path
) -> None:
    """Все зоны заполнены: геометрия всех фигур — шаблонная, 445 и 447 с тем же вылетом."""
    recipe = ex018(education)

    prs = written(education, recipe, [zone.zone_id for zone in recipe.zones], tmp_path)

    deck, example = geometry(prs.slides[0]), example_geometry(education, recipe)
    assert set(OVERHANG) <= deck.keys(), "фигуры 445/447 пропали из колоды"
    assert deck == {xml_id: example[xml_id] for xml_id in deck}
    assert {xml_id: overhang(prs, xml_id) for xml_id in OVERHANG} == OVERHANG


def test_ex018_with_one_repeat_moves_nothing_that_is_left(
    education: Case, tmp_path: Path
) -> None:
    """Один повтор из двух: лишний удалён, уцелевшие фигуры стоят там же, где в шаблоне."""
    recipe = ex018(education)
    first = [z.zone_id for z in recipe.zones if z.repeat in (None, 0)]

    prs = written(education, recipe, first, tmp_path)

    deck, example = geometry(prs.slides[0]), example_geometry(education, recipe)
    assert len(deck) < len(example), "лишний повтор не удалён"
    assert deck == {xml_id: example[xml_id] for xml_id in deck}
    assert overhang(prs, 447) == OVERHANG[447]
