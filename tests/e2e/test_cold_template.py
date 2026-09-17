"""C6: колода собирается на шаблоне, который не использовался при разработке.

Фикстура `tests/fixtures/templates/cold/` пополняется только тем, что команда не открывала
при написании парсера. Правило PLAN.md §8: любая правка парсера или композитора
проверяется на холодном шаблоне.

Композиции (поток A) здесь нет: колода собирается детерминированно — заголовок и текст
в плейсхолдеры каждого подходящего макета. Проверяется путь «манифест → вписывание →
нативный .pptx → .pdf» на незнакомом файле, а не качество содержания.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from deckforge.domain.enums import TextRole
from deckforge.domain.slide import BulletItem, BulletsBlock, DeckIR, SlideIR, TextBlock
from deckforge.domain.template import LayoutSpec, TemplateManifest
from deckforge.export.pdf import export_pdf
from deckforge.export.pptx import export_pptx
from deckforge.layout.fitting import fit_slide
from deckforge.parsing import TemplateParser
from deckforge.rendering.soffice import SofficeRenderer
from tests.e2e.cold_corpus import cold_templates

COLD_TEMPLATES = cold_templates()
TEXT_TYPES = {"TITLE", "CTRTITLE", "SUBTITLE", "BODY", "OBJ"}
MAX_SLIDES = 12

pytestmark = [pytest.mark.cold, pytest.mark.slow]


def _blocks(layout: LayoutSpec, number: int) -> list[TextBlock | BulletsBlock]:
    blocks: list[TextBlock | BulletsBlock] = []
    for ph in layout.placeholders:
        if ph.ph_type not in TEXT_TYPES or ph.role is None:
            continue
        block_id = f"p{ph.idx}"
        if ph.role is TextRole.TITLE:
            blocks.append(TextBlock(block_id=block_id, placeholder_idx=ph.idx, role=ph.role,
                                    text=f"Вывод {number}"))
        elif ph.role is TextRole.BODY:
            blocks.append(BulletsBlock(block_id=block_id, placeholder_idx=ph.idx,
                                       items=[BulletItem(text="Первый довод"),
                                              BulletItem(text="Второй довод")]))
        else:
            blocks.append(TextBlock(block_id=block_id, placeholder_idx=ph.idx, role=ph.role,
                                    text="Итоги"))
    return blocks


def cold_deck(manifest: TemplateManifest) -> DeckIR:
    """По слайду на макет с текстовыми плейсхолдерами; не влезшие блоки убираются —
    так поступила бы композиция, получив `shorten`."""
    slides = []
    for layout in manifest.layouts:
        blocks = _blocks(layout, len(slides) + 1)
        if not blocks:
            continue
        slide = fit_slide(SlideIR(slide_id=f"s{len(slides) + 1:02d}", layout_id=layout.layout_id,
                                  variant="A", blocks=blocks), manifest)
        kept = [b for b in slide.blocks if not slide.fit_report[b.block_id].overflow]
        if not kept:
            continue
        slides.append(slide.model_copy(update={
            "blocks": kept,
            "fit_report": {b.block_id: slide.fit_report[b.block_id] for b in kept},
        }))
        if len(slides) == MAX_SLIDES:
            break
    return DeckIR(deck_id="cold", variant="A", template_id=manifest.template_id, seed=1,
                  slides=slides)


@pytest.mark.skipif(
    not COLD_TEMPLATES, reason="нет холодных шаблонов: ни в cold/, ни от LibreOffice"
)
@pytest.mark.parametrize("template", COLD_TEMPLATES, ids=lambda p: p.name)
def test_cold_template_end_to_end(template: Path, tmp_path: Path) -> None:
    manifest = TemplateParser().parse(template, use_cache=False)
    deck = cold_deck(manifest)
    assert deck.slides, "ни одного слайда не удалось вписать в незнакомый шаблон"

    pptx = export_pptx(deck, manifest, template, tmp_path / "cold.pptx")
    prs = Presentation(str(pptx))
    assert len(prs.slides) == len(deck.slides)
    for slide in prs.slides:
        shapes = list(slide.shapes)
        assert shapes
        assert not all(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in shapes), "C3"

    if SofficeRenderer().available:
        pdf = export_pdf(pptx, tmp_path / "cold.pdf")
        assert pdf.read_bytes().startswith(b"%PDF")
