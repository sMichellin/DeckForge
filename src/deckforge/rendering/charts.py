"""Нативные диаграммы. Change (14) `native-charts-tables`.

Серии красятся в цвета **темы** ссылками (`schemeClr`), поэтому палитра меняется вместе
с шаблоном без единой правки кода. Цепочка деградации: диаграмма → таблица → буллеты (§15);
`chart_problem` говорит, почему диаграмму не построить, а подмену выполняет `PptxWriter`.
"""

from __future__ import annotations

import math
from itertools import cycle

from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Emu, Pt

from deckforge.domain.content import Dataset
from deckforge.domain.enums import ChartType, ColorRef, FontRef, TextRole
from deckforge.domain.slide import ChartBlock
from deckforge.domain.template import TemplateManifest
from deckforge.rendering.theme_binding import THEME_COLORS, apply_theme_color, theme_font_token

CHART_TYPES: dict[ChartType, XL_CHART_TYPE] = {
    ChartType.CLUSTERED_BAR: XL_CHART_TYPE.BAR_CLUSTERED,
    ChartType.STACKED_BAR: XL_CHART_TYPE.BAR_STACKED,
    ChartType.CLUSTERED_COLUMN: XL_CHART_TYPE.COLUMN_CLUSTERED,
    ChartType.STACKED_COLUMN: XL_CHART_TYPE.COLUMN_STACKED,
    ChartType.LINE: XL_CHART_TYPE.LINE,
    ChartType.LINE_MARKERS: XL_CHART_TYPE.LINE_MARKERS,
    ChartType.PIE: XL_CHART_TYPE.PIE,
    ChartType.DOUGHNUT: XL_CHART_TYPE.DOUGHNUT,
    ChartType.SCATTER: XL_CHART_TYPE.XY_SCATTER,
    ChartType.AREA: XL_CHART_TYPE.AREA,
}

_ROUND = frozenset({ChartType.PIE, ChartType.DOUGHNUT})
_LINES = frozenset({ChartType.LINE, ChartType.LINE_MARKERS})
#: Порядок акцентов, если ни блок, ни шаблон не задали свой.
_ACCENTS = [
    ColorRef.ACCENT1, ColorRef.ACCENT2, ColorRef.ACCENT3,
    ColorRef.ACCENT4, ColorRef.ACCENT5, ColorRef.ACCENT6,
]


def _as_number(text: str) -> float | None:
    try:
        number = float(text.replace(",", ".").replace(" ", ""))
    except ValueError:
        return None
    return number if math.isfinite(number) else None


def chart_problem(chart_type: ChartType, dataset: Dataset) -> str | None:
    """Почему эту диаграмму из этих данных не построить; `None` — строится."""
    if not dataset.categories:
        return "нет категорий"
    if not dataset.series:
        return "нет серий"
    if any(len(s.values) != len(dataset.categories) for s in dataset.series):
        return "длины серий не совпадают с числом категорий"
    values = [v for s in dataset.series for v in s.values if v is not None]
    if not values:
        return "нет ни одного значения"
    if not all(math.isfinite(v) for v in values):
        # NaN и ∞ проходят валидацию модели, но роняют запись данных диаграммы.
        return "значения должны быть конечными числами"
    if chart_type in _ROUND:
        if len(dataset.series) != 1:
            return "у круговой диаграммы должна быть одна серия"
        if any(v < 0 for v in values):
            return "у круговой диаграммы отрицательные доли"
        if not any(values):
            return "у круговой диаграммы все доли нулевые"
    if chart_type is ChartType.SCATTER and any(_as_number(c) is None for c in dataset.categories):
        return "у точечной диаграммы категории должны быть числами"
    return None


def _chart_data(chart_type: ChartType, dataset: Dataset) -> CategoryChartData | XyChartData:
    if chart_type is ChartType.SCATTER:
        xy = XyChartData()  # type: ignore[no-untyped-call]
        xs = [_as_number(c) for c in dataset.categories]
        for s in dataset.series:
            series = xy.add_series(s.name)  # type: ignore[no-untyped-call]
            for x, y in zip(xs, s.values, strict=True):
                if x is not None and y is not None:
                    series.add_data_point(x, y)
        return xy
    data = CategoryChartData()  # type: ignore[no-untyped-call]
    data.categories = dataset.categories
    for s in dataset.series:
        data.add_series(s.name, s.values)  # type: ignore[no-untyped-call]
    return data


def _paint_marker(series: object, ref: ColorRef) -> None:
    """Маркер целиком — и заливка, и обводка, иначе обводка останется автоматического цвета."""
    marker = series.marker  # type: ignore[attr-defined]
    apply_theme_color(marker.format.fill, ref)
    marker.format.line.color.theme_color = THEME_COLORS[ref]


def add_chart(
    slide: object,
    block: ChartBlock,
    dataset: Dataset,
    manifest: TemplateManifest,
    *,
    text_color: ColorRef | None = None,
) -> object:
    """Диаграмма в рамке блока. `text_color` — цвет текста макета: на тёмном фоне подписи
    осей и легенды иначе остались бы тёмными."""
    box = block.bbox
    if box is None:
        raise ValueError(f"диаграмма {block.block_id} без координат")
    frame = slide.shapes.add_chart(  # type: ignore[attr-defined]
        CHART_TYPES[block.chart_type], Emu(box.x), Emu(box.y), Emu(box.cx), Emu(box.cy),
        _chart_data(block.chart_type, dataset),
    )
    chart = frame.chart
    refs = cycle(block.series_color_refs or manifest.chart_defaults.series_color_refs or _ACCENTS)
    plot = chart.plots[0]

    if block.chart_type in _ROUND:
        for point in plot.series[0].points:
            apply_theme_color(point.format.fill, next(refs))
    else:
        for series in plot.series:
            ref = next(refs)
            if block.chart_type in _LINES:
                # У линии цвет — у обводки, а не у заливки.
                series.format.line.color.theme_color = THEME_COLORS[ref]
                if block.chart_type is ChartType.LINE_MARKERS:
                    _paint_marker(series, ref)
            elif block.chart_type is ChartType.SCATTER:
                _paint_marker(series, ref)
            else:
                apply_theme_color(series.format.fill, ref)

    chart.has_legend = block.legend
    if block.legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    if block.chart_type is not ChartType.SCATTER:
        # У точечной диаграммы python-pptx подписей данных на уровне графика не знает.
        plot.has_data_labels = block.data_labels

    if block.chart_type not in _ROUND:
        value_title = block.axis_titles.get("value") or dataset.unit
        if value_title:
            chart.value_axis.has_title = True
            chart.value_axis.axis_title.text_frame.text = value_title
        category_title = block.axis_titles.get("category")
        if category_title:
            chart.category_axis.has_title = True
            chart.category_axis.axis_title.text_frame.text = category_title

    # Текст диаграммы — кеглем подписи из шкалы шаблона и шрифтом темы.
    caption = manifest.typography(TextRole.CAPTION) or manifest.typography(TextRole.BODY)
    if caption is not None:
        chart.font.size = Pt(caption.size_pt)
    chart.font.name = theme_font_token(FontRef.MINOR_LATIN)
    if text_color is not None:
        chart.font.color.theme_color = THEME_COLORS[text_color]
    return frame
