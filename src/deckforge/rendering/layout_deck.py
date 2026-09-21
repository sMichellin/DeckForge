"""Колода «по слайду на макет». Change (6) `template-preview-render`.

Чтобы увидеть, как выглядит макет, его надо на чём-то показать: сам по себе
`slideLayout` не рендерится. Собираем из шаблона временную колоду, где каждый слайд
создан на своём макете, и конвертируем её одной командой.

Плейсхолдеры **не заполняются**: классификатору нужна структура макета, а вписанный
текст-рыба только сбил бы его с толку. Зато виден весь декор шаблона — фон, логотип,
плашки, — которого нет в схеме.
"""

from __future__ import annotations

from collections.abc import Collection
from pathlib import Path

from pptx import Presentation
from pptx.presentation import Presentation as PresentationPart

from deckforge.domain.template import TemplateManifest


def build_layout_deck(
    template: Path,
    manifest: TemplateManifest,
    out: Path,
    only: Collection[str] | None = None,
) -> dict[int, str]:
    """Собрать колоду из макетов шаблона.

    Возвращает соответствие «номер слайда (с нуля) → layout_id»: порядок макетов
    в манифесте и в колоде обязан совпадать, иначе превью уедут не к тем макетам.

    `only` ограничивает набор. Это не микрооптимизация: конвертация всех макетов
    настоящего шаблона занимает 23–53 с при бюджете парсинга в 25 с (§12), а спорных
    макетов, ради которых всё затевается, обычно втрое меньше.
    """
    prs = Presentation(str(template))

    # В .potx и в шаблонах с примерами слайды уже есть — убираем, чтобы нумерация
    # страниц pdf совпадала с порядком макетов.
    _drop_existing_slides(prs)

    # Макет ищется по имени части, а не по `spec.index`. Индекс манифеста считает
    # только макеты с плейсхолдерами (парсер пропускает пустые), а список python-pptx —
    # все подряд. На шаблоне с пустыми макетами индексы расходятся, и превью уезжали
    # к чужим макетам: у VK Tech вместо тёмного фона L16 рисовались светлые столбцы
    # (прогон 9ce69f5cbd39) — и классификатор, и метрики C9 смотрели не на тот макет.
    # Писатель колоды ищет макет так же (`rendering/writer.py`).
    by_part = {
        str(layout.part.partname).lstrip("/"): layout
        for master in prs.slide_masters
        for layout in master.slide_layouts
    }
    order: dict[int, str] = {}

    wanted = set(only) if only is not None else None
    for spec in manifest.layouts:
        layout = by_part.get(spec.part_name.lstrip("/"))
        if layout is None:
            continue
        if wanted is not None and spec.layout_id not in wanted:
            continue
        order[len(order)] = spec.layout_id
        prs.slides.add_slide(layout)

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return order


def _drop_existing_slides(prs: PresentationPart) -> None:
    """Удалить все слайды презентации, не трогая макеты и мастера."""
    id_list = prs.slides._sldIdLst  # приватный: публичного API удаления слайдов нет
    for slide_id in list(id_list):
        rel_id = slide_id.get(
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
        )
        if rel_id:
            prs.part.drop_rel(rel_id)
        id_list.remove(slide_id)
