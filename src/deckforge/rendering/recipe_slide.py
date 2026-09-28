"""Слайд по рецепту: копия слайда-примера шаблона. Change `recipe-slide-in-the-writer`.

Раньше writer строил каждый слайд из пустого макета, а слайды-примеры удалял первым делом.
Но оформление шаблона живёт не в макетах: вне плейсхолдеров стоит от 87 до 96 % содержимого
примеров. Пустой макет — это белый фон с рамками под текст, и «вёрстка по дизайн-системе»
поверх него плашек автора не вернёт.

Здесь пример копируется целиком — с фоном, плашками, декором, картинками и порядком
наложения, — а наш текст занимает его текстовые зоны. Модель файла по-прежнему не видит
(ADR-003, пересмотр 23.09): она пишет только `SlideIR` с `zone_id`, а копирует writer.

Фигуры остаются нативными объектами PowerPoint: ни картинки слайда, ни растра —
правило C3 цело.
"""

from __future__ import annotations

import copy
from typing import Any

from lxml import etree
from pptx.oxml.ns import qn

from deckforge.designsystem.models import (
    ExamplePassport,
    Place,
    PlaceGroup,
    PlaceKind,
    Recipe,
    TypeLevel,
    Zone,
)
from deckforge.domain.slide import Block, BulletsBlock, FitResult, SlideIR, TextBlock
from deckforge.layout.nonbreaking import bind as nonbreaking
from deckforge.rendering.units import size_hundredths

#: Атрибуты, которыми фигура ссылается на связь своей части: картинка, диаграмма, ссылка.
#: При копировании они переписываются на связи нового слайда, иначе PowerPoint предложит
#: «восстановить» файл.
REL_ATTRS = (qn("r:embed"), qn("r:id"), qn("r:link"))

#: Узлы дерева фигур, которые у нового слайда уже свои.
OWN_NODES = (qn("p:nvGrpSpPr"), qn("p:grpSpPr"))

#: Фигуры, которые адресует каталог рецептов по `cNvPr id`. Тот же набор, что берёт разбор
#: примеров (`parsing/ooxml/examples._collect`): группы раскрываются, сами не адресуются.
SHAPE_TAGS = (qn("p:sp"), qn("p:pic"), qn("p:cxnSp"), qn("p:graphicFrame"))
GROUP = qn("p:grpSp")
GRAPHIC_FRAME = qn("p:graphicFrame")

#: Пространство имён связей: `r:id`, `r:embed`, а у SmartArt ещё `r:dm`, `r:lo`, `r:qs`,
#: `r:cs`. Снятая рамка уносит свои связи — любые из них.
_R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

#: Части, на которые ссылается рамка: диаграмма, схема SmartArt, встроенный объект.
#: Снимается только такая связь — связь с макетом или заметками слайда трогать нельзя,
#: а идентификаторы `rId` у копии и у примера совпадают по случайности.
_FRAME_PARTS = ("/chart", "/diagram", "/package", "/oleObject")


class RecipeError(ValueError):
    """Рецепт не удалось применить: примера нет в файле или зоны не нашлись."""


def _part_name(slide: Any) -> str:
    return str(slide.part.partname).lstrip("/")


def source_slide(prs: Any, recipe: Recipe) -> Any:
    """Слайд-пример шаблона по имени части. Порядок в `sldIdLst` ненадёжен, имя — нет."""
    for slide in prs.slides:
        if _part_name(slide) == (recipe.part_name or "").lstrip("/"):
            return slide
    raise RecipeError(f"рецепт {recipe.recipe_id}: части {recipe.part_name} нет в шаблоне")


def _shapes_tree(slide: Any) -> Any:
    return slide.shapes._spTree


def _xml_id(node: Any) -> int | None:
    name = node.find(f"./*/{qn('p:cNvPr')}")
    if name is None:
        return None
    raw = name.get("id")
    return int(raw) if raw and raw.isdigit() else None


def _shapes_in(parent: Any) -> Any:
    """Фигуры в порядке документа, группы раскрыты — так же, как их видит разбор примеров."""
    for node in parent:
        if node.tag == GROUP:
            yield from _shapes_in(node)
        elif node.tag in SHAPE_TAGS:
            yield node


def _addressed(recipe: Recipe) -> set[int]:
    ids = {zone.xml_id for zone in recipe.zones if zone.xml_id is not None}
    ids.update(xml_id for row in recipe.repeat_xml_ids for xml_id in row)
    if recipe.picture_xml_id is not None:
        ids.add(recipe.picture_xml_id)
    return ids


def _by_xml_id(tree: Any, recipe: Recipe) -> dict[int, Any]:
    """Фигуры слайда по их `cNvPr id` на любой глубине групп.

    Каталог делает зонами и фигуры внутри `p:grpSp` — подписи карточек, строки таймлайна.
    Ищи writer только верхний уровень, такая зона молча пропускалась бы: наш текст
    терялся, а текст шаблона оставался.

    Если id, который адресует рецепт, в примере повторяется, какую фигуру имел в виду
    каталог, неизвестно: писать наугад — значит молча оставить чужую фразу. Это ошибка.
    """
    found: dict[int, Any] = {}
    twins: set[int] = set()
    for node in _shapes_in(tree):
        identifier = _xml_id(node)
        if identifier is None:
            continue
        if identifier in found:
            twins.add(identifier)
        found.setdefault(identifier, node)
    ambiguous = sorted(twins & _addressed(recipe))
    if ambiguous:
        raise RecipeError(
            f"рецепт {recipe.recipe_id}: cNvPr id {ambiguous} повторяется в примере, "
            "фигуру зоны не определить"
        )
    return found


def _relink(source: Any, target: Any) -> None:
    """Перевесить связи скопированных фигур на новый слайд.

    Картинка и диаграмма живут отдельными частями пакета; копия ссылается на них тем же
    `rId`, которого у нового слайда нет. Части при этом переиспользуются, а не дублируются.
    """
    rels = source.part.rels
    mapping: dict[str, str] = {}
    for node in _shapes_tree(target).iter():
        for attr in REL_ATTRS:
            old = node.get(attr)
            if not old:
                continue
            if old not in mapping:
                rel = rels[old]
                mapping[old] = (
                    target.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
                    if rel.is_external
                    else target.part.relate_to(rel.target_part, rel.reltype)
                )
            node.set(attr, mapping[old])


def clone_slide(prs: Any, recipe: Recipe) -> Any:
    """Копия слайда-примера: тот же макет, те же фигуры, свои связи."""
    source = source_slide(prs, recipe)
    target = prs.slides.add_slide(source.slide_layout)

    tree = _shapes_tree(target)
    for node in list(tree):
        if node.tag not in OWN_NODES:
            tree.remove(node)
    for node in _shapes_tree(source):
        if node.tag not in OWN_NODES:
            tree.append(copy.deepcopy(node))

    _relink(source, target)
    return target


def _first_run_props(shape: Any) -> Any | None:
    """Свойства первого прогона зоны: гарнитура, кегль и цвет, которые поставил автор."""
    body = shape.find(qn("p:txBody"))
    if body is None:
        return None
    props = body.find(f"{qn('a:p')}/{qn('a:r')}/{qn('a:rPr')}")
    return copy.deepcopy(props) if props is not None else None


def _paragraph_props(shape: Any) -> Any | None:
    body = shape.find(qn("p:txBody"))
    if body is None:
        return None
    props = body.find(f"{qn('a:p')}/{qn('a:pPr')}")
    return copy.deepcopy(props) if props is not None else None


def _lines_of(block: Block) -> list[str]:
    if isinstance(block, BulletsBlock):
        return [item.text for item in block.items if item.text]
    if isinstance(block, TextBlock):
        return [line for line in block.text.splitlines() if line]
    return []


def write_zone(
    shape: Any,
    lines: list[str],
    size_pt: float | None = None,
    author_pt: float | None = None,
) -> None:
    """Заменить текст зоны нашим, сохранив оформление автора.

    Берутся свойства первого прогона и первого абзаца: гарнитура и цвет остаются теми,
    что в шаблоне (правило 5 — в IR их нет вовсе). Кегль — тоже автора, пока вписывание
    не нашло свой (`size_pt`, change `recipe-zone-takes-the-fitted-size`): кегль примера
    стоял под короткое слово примера, и наш текст на нём рвёт слова и выходит за рамку.
    Кегль автора писатель только опускает (D02): каждому прогону
    `sz = min(size_pt, sz прогона примера)`.

    **D04 уточнён (RG42).** Прогон примера без `sz` не трогался вовсе: «настоящего числа
    писатель не знает». Осторожность верная, но причина была шире правды: у плейсхолдера
    число знает макет, и парсер его теперь достаёт (`ExampleShape.size_pt`), а каталог
    помечает, своё оно или подставлено ступенью лестницы (`Zone.size_is_own`). Своё
    приходит сюда как `author_pt` — и от него кегль можно опустить.

    Молчание осталось там, где ему место: число вычислили мы (`author_pt` пуст) —
    писатель не трогает ничего. Именно подстановка ступени роняла VK Tech s03
    с 16 до 7,8 pt и поднимала VK Education s04/s11 с 36 до 39.

    Цена прежнего поведения: у VK WorkSpace заголовок s05 — плейсхолдер без `sz`,
    вписывание опустило его 36 → 23,4 pt, писатель это проигнорировал, и заголовок
    в две строки лёг на первую строку тела.

    Без `size_pt` зона пишется как раньше, байт в байт.
    """
    body = shape.find(qn("p:txBody"))
    if body is None:
        return
    run_props, par_props = _first_run_props(shape), _paragraph_props(shape)
    if size_pt is not None and run_props is not None:
        fitted = size_hundredths(size_pt)
        own = run_props.get("sz")
        authors = int(own) if own is not None else (
            int(size_hundredths(author_pt)) if author_pt else None
        )
        if authors is not None and int(fitted) < authors:
            run_props.set("sz", fitted)
    for paragraph in body.findall(qn("a:p")):
        body.remove(paragraph)
    for line in lines or [""]:
        paragraph = etree.SubElement(body, qn("a:p"))
        if par_props is not None:
            paragraph.append(copy.deepcopy(par_props))
        run = etree.SubElement(paragraph, qn("a:r"))
        if run_props is not None:
            run.append(copy.deepcopy(run_props))
        text = etree.SubElement(run, qn("a:t"))
        text.text = nonbreaking(line)


def _used_repeats(recipe: Recipe, blocks: list[Block]) -> int:
    """Сколько повторов занято нашим текстом. Ноль — ряд не используется вовсе."""
    zones = {zone.zone_id: zone for zone in recipe.zones}
    used: list[int] = []
    for block in blocks:
        zone_id = getattr(block, "zone_id", None)
        zone = zones.get(zone_id) if zone_id else None
        if zone is not None and zone.repeat is not None:
            used.append(zone.repeat)
    return max(used) + 1 if used else 0


def _lowered_size(zone: Zone, fitted: FitResult | None) -> float | None:
    """Кегль вписывания по рамке зоны (RG29) — если вписывание его опустило, иначе `None`.

    Записи нет — зона без рамки или старый чекпойнт: остаётся кегль автора примера.
    «Опустило» значит ниже кегля зоны в каталоге, с которого вписывание стартует, а не
    `strategy`: заголовок, уступивший кегль, вписывается заново и приходит как `as_is`
    (D03). Не ниже — оформление автора байт в байт: старт каталога не всегда равен кеглю
    прогона примера, и писать его значило бы поднять заголовок 36 → 60 pt (D02).
    """
    if fitted is None:
        return None
    if zone.size_pt is not None and fitted.final_size_pt >= zone.size_pt:
        return None
    return fitted.final_size_pt


def clone_recipe(prs: Any, recipe: Recipe, slide_ir: SlideIR) -> Any:
    """Слайд колоды по рецепту: копия примера, наш текст по зонам, лишнее удалено.

    Развилка одна (change `the-writer-removes-whole-groups`, план Б, 4): у рецепта с паспортом
    лишнее снимается группами паспорта (`_fill_by_passport`), без паспорта — прежним угадыванием,
    байт в байт. Оба пути живут до приёмки плана Б, и расхождение должно быть видно здесь.
    """
    slide = clone_slide(prs, recipe)
    shapes = _by_xml_id(_shapes_tree(slide), recipe)
    if recipe.passport is not None:
        _fill_by_passport(recipe, recipe.passport, slide_ir, slide, shapes)
        # Таблица и диаграмма автора в паспорт не попадают — ни местом, ни декором. Решение
        # Насти 28.09: незаполненную чужую таблицу «Убирать» и при паспорте (титул WorkSpace
        # `ex014`, «Заголовок / Текст»). Отступление от буквы issue — в proposal.
        _drop_unfilled_frames(recipe, slide)
        return slide

    filled: set[str] = set()
    for block in slide_ir.blocks:
        zone_id = getattr(block, "zone_id", None)
        if not zone_id:
            continue
        zone = next((z for z in recipe.zones if z.zone_id == zone_id), None)
        if zone is None or zone.xml_id is None or zone.xml_id not in shapes:
            continue
        write_zone(
            shapes[zone.xml_id],
            _lines_of(block),
            _lowered_size(zone, slide_ir.fit_report.get(block.block_id)),
            zone.author_size_pt,
        )
        filled.add(zone_id)

    _drop_spare_repeats(recipe, _used_repeats(recipe, list(slide_ir.blocks)), shapes)
    _drop_empty_zones(recipe, filled, shapes)
    _drop_orphan_decor(recipe, slide, int(prs.slide_width), int(prs.slide_height))
    _drop_unfilled_frames(recipe, slide)
    return slide


def _texted_zones(slide_ir: SlideIR) -> list[tuple[str, Block, list[str]]]:
    """Блоки IR с текстом для зоны: (зона, блок, строки). Пустой блок место не заполняет."""
    out: list[tuple[str, Block, list[str]]] = []
    for block in slide_ir.blocks:
        zone_id = getattr(block, "zone_id", None)
        lines = _lines_of(block)
        if zone_id and lines:
            out.append((zone_id, block, lines))
    return out


def _row_moves(passport: ExamplePassport, texted: set[str]) -> dict[str, str]:
    """Куда переезжает текст в рядах паспорта: зона блока → зона цели.

    Группы ряда взаимозаменяемы (форма у них одна — это инвариант паспорта), поэтому
    заполненные группы по порядку ряда встают в первые k групп, место в место по номеру
    места в группе. IR при этом не меняется: он принадлежит композиции, писатель только
    решает, куда положить текст. Legacy-IR 28.09 на VK Tech `ex018` кладёт текст в 1-ю и 4-ю
    карточку — без переноса между ними осталась бы дыра из двух снятых карточек.
    """
    moves: dict[str, str] = {}
    for groups in passport.rows.values():
        taken = [g for g in groups if any(p.zone_id in texted for p in g.places)]
        for source, target in zip(taken, groups, strict=False):
            for mine, theirs in zip(source.places, target.places, strict=True):
                if mine.zone_id is not None and theirs.zone_id is not None:
                    moves[mine.zone_id] = theirs.zone_id
    return moves


def _node(tree: Any, shapes: dict[int, Any], xml_id: int) -> Any | None:
    """Фигура по `cNvPr id`, а если это группа `p:grpSp` декора — сама группа."""
    if xml_id in shapes:
        return shapes[xml_id]
    return next((g for g in tree.iter(GROUP) if _xml_id(g) == xml_id), None)


def _fill_by_passport(
    recipe: Recipe, passport: ExamplePassport, slide_ir: SlideIR, slide: Any,
    shapes: dict[int, Any],
) -> None:
    """Наш текст по зонам, лишнее — группами паспорта, без угадывания.

    1. Перенос в ряду (`_row_moves`) — до записи, кегль вписывания тот же (`fit_report` блока):
       группы ряда одного размера.
    2. Группа заполнена, если хоть одно её место получило наш непустой текст. Незаполненная
       уходит целиком: фигуры мест и `decor_xml_ids` — плашка, иконка, линии; опустевшая
       `p:grpSp` — следом (`_remove`).

    `_drop_spare_repeats`, `_drop_empty_zones` и `_drop_orphan_decor` здесь не вызываются:
    паспорт знает, что живёт и уходит вместе, а угадывание по номеру повтора и по геометрии
    линий снимало бы то, чего в группах нет. Удалять их нельзя — на них живёт `legacy`.
    """
    zones = {zone.zone_id: zone for zone in recipe.zones}
    texted = _texted_zones(slide_ir)
    moves = _row_moves(passport, {zone_id for zone_id, _, _ in texted})
    filled: set[str] = set()
    for zone_id, block, lines in texted:
        zone = zones.get(moves.get(zone_id, zone_id))
        if zone is None or zone.xml_id not in shapes:
            continue
        write_zone(
            shapes[zone.xml_id],
            lines,
            _lowered_size(zone, slide_ir.fit_report.get(block.block_id)),
            zone.author_size_pt,
        )
        filled.add(zone.zone_id)

    tree = _shapes_tree(slide)
    for group in passport.groups:
        titled = any(place.role is TypeLevel.SLIDE_TITLE for place in group.places)
        if titled or any(place.zone_id in filled for place in group.places):
            _drop_empty_places(group, filled, zones, shapes)
            continue
        for xml_id in [_place_xml_id(place, zones) for place in group.places] + group.decor_xml_ids:
            if xml_id is not None:
                _remove(_node(tree, shapes, xml_id))


def _drop_empty_places(
    group: PlaceGroup, filled: set[str], zones: dict[str, Zone], shapes: dict[int, Any]
) -> None:
    """Пустое текстовое место группы, которая остаётся, уходит — иначе в колоде текст шаблона.

    Заголовок слайда не уходит никогда (правило `_drop_empty_zones`): слайд без заголовка
    читать нечем, поэтому его группа остаётся даже пустой, а чужая фраза в нём стирается.
    Место-картинка остаётся: ассет в рецептный слайд писатель не ставит, и картинка автора
    в заполненной карточке — её оформление.
    """
    for place in group.places:
        if place.kind is PlaceKind.PICTURE or place.zone_id in filled:
            continue
        shape = shapes.get(_place_xml_id(place, zones) or -1)
        if shape is None:
            continue
        if place.role is TypeLevel.SLIDE_TITLE:
            write_zone(shape, [])
        else:
            _remove(shape)


def _place_xml_id(place: Place, zones: dict[str, Zone]) -> int | None:
    """Адрес фигуры места: свой у картинки, у текстового — тот же, что у его зоны."""
    if place.xml_id is not None:
        return place.xml_id
    zone = zones.get(place.zone_id or "")
    return zone.xml_id if zone is not None else None


def _remove(node: Any) -> None:
    """Удалить фигуру у её собственного родителя. Группа, в которой не осталось ничего,
    кроме её собственных свойств, уходит следом: пустая рамка группы на слайде ни к чему.
    Всё прочее в группе (`p:contentPart`, `mc:AlternateContent`) не наше — группа остаётся."""
    parent = node.getparent() if node is not None else None
    if parent is None:
        return
    parent.remove(node)
    if parent.tag == GROUP and all(child.tag in OWN_NODES for child in parent):
        _remove(parent)


def _drop_spare_repeats(recipe: Recipe, used: int, shapes: dict[int, Any]) -> None:
    """Лишние повторы удаляются целиком и с конца ряда: плашка вместе со своим текстом."""
    if not recipe.repeat_xml_ids:
        return
    for row in recipe.repeat_xml_ids[max(used, 0) :]:
        for xml_id in row:
            _remove(shapes.get(xml_id))


def _drop_empty_zones(recipe: Recipe, filled: set[str], shapes: dict[int, Any]) -> None:
    """Зона без нашего текста удаляется — иначе в колоде останется текст шаблона.

    Заголовок не удаляется никогда: слайд без заголовка читать нечем, и пустую рамку
    заметит человек, а чужую фразу — нет. Чужую фразу в нём стираем (проверка
    `template.sample_text_left`).

    Зона внутри уцелевшего повтора раньше оставалась пустой рамкой — тоже удаляется
    (RG40). Повтор у этих шаблонов не равен карточке: у VK Tech рецепт `ex018` держит
    ряд из десяти зон в двух повторах, и два факта оставляли восемь пустых рамок.
    В готовой колоде таких рамок было 48 на три колоды. Плашка карточки — отдельная
    фигура автора, она остаётся: её геометрию писатель не трогает (RG36).
    """
    for zone in recipe.zones:
        if zone.zone_id in filled or zone.xml_id is None or zone.xml_id not in shapes:
            continue
        if zone.role is TypeLevel.SLIDE_TITLE:
            write_zone(shapes[zone.xml_id], [])
            continue
        _remove(shapes[zone.xml_id])


def _relation_ids(node: Any) -> set[str]:
    """Все идентификаторы связей, на которые ссылается узел и его потомки."""
    return {
        value
        for element in node.iter()
        for attr, value in element.attrib.items()
        if attr.startswith(_R_NS)
    }


def _drop_unfilled_frames(recipe: Recipe, slide: Any) -> None:
    """Таблица, диаграмма и SmartArt примера, которых мы ничем не заполнили, уходят (RG45).

    Зонами каталог делает только текстовые фигуры, поэтому `p:graphicFrame` примера писатель
    копирует целиком — вместе с таблицей автора и её текстом. На титуле VK WorkSpace (`ex014`)
    так в колоду уезжала таблица «Заголовок / Текст», и это была последняя ошибка аудита
    (`template.sample_text_left`). Правило то же, что у пустой зоны: незаполненное — чужое.
    Рамка, которую рецепт адресует, остаётся — её судьбу решает тот, кто её адресовал.

    Связи снятых рамок снимаются следом, если на них больше никто не ссылается: иначе
    в пакете осталась бы диаграмма автора, которую PowerPoint не покажет, но сохранит.
    """
    addressed = _addressed(recipe)
    frames = [
        node
        for node in _shapes_in(_shapes_tree(slide))
        if node.tag == GRAPHIC_FRAME and _xml_id(node) not in addressed
    ]
    if not frames:
        return
    released = set().union(*(_relation_ids(frame) for frame in frames))
    for frame in frames:
        _remove(frame)
    still_used = _relation_ids(slide.part._element)
    rels = slide.part.rels
    for rid in sorted(released - still_used):
        rel = rels.get(rid)
        if rel is not None and any(kind in rel.reltype for kind in _FRAME_PARTS):
            rels.pop(rid)


#: Допуск «примыкает» — доля ширины слайда (RG52). Политика вёрстки, а не свойство шаблона:
#: линии схемы Education `ex013` лежат в 1–2 % от рамок зон. 2 % уже снимают оформление
#: слайда — точки-пагинатор раздела VK Tech `ex003` рядом с подзаголовком; 1 % — нет.
#: То же число у проверки `template.decor_leads_nowhere` (`configs/audit_checks.yaml`).
TOUCH_SHARE = 0.01

#: Линия — `p:cxnSp` любой формы или фигура, у которой меньшая сторона не больше этой доли
#: большей: у линий схем одна сторона нулевая, у градиентных черт VK Tech — около 1 %.
LINE_SHARE = 0.02

#: Фон и полосы во всю ширину или высоту слайда — оформление слайда, не декор зоны.
FULL_SPAN_SHARE = 0.9

#: Колонтитулы — не текст слайда (как `FOOTER_PLACEHOLDERS` аудита).
_FOOTERS = frozenset({"ftr", "dt", "sldNum", "hdr"})

_Box = tuple[float, float, float, float]


def _xfrm(node: Any) -> Any | None:
    if node.tag == GROUP:
        return node.find(f"{qn('p:grpSpPr')}/{qn('a:xfrm')}")
    found = node.find(f"{qn('p:spPr')}/{qn('a:xfrm')}")
    return found if found is not None else node.find(qn("p:xfrm"))


def _pair(xfrm: Any | None, off: str, ext: str) -> _Box | None:
    """Рамка из пары узлов `off/ext` (или `chOff/chExt` у группы)."""
    start = xfrm.find(qn(off)) if xfrm is not None else None
    size = xfrm.find(qn(ext)) if xfrm is not None else None
    if start is None or size is None:
        return None
    return (float(start.get("x", 0)), float(start.get("y", 0)),
            float(size.get("cx", 0)), float(size.get("cy", 0)))


def _placed(parent: Any, frame: _Box = (0, 0, 1, 1), origin: tuple[float, float] = (0, 0)) -> Any:
    """Фигуры с рамкой в координатах слайда (или `None`), группы раскрыты с их масштабом.

    Геометрию линий разбор примеров не хранит (`cx = 0` он отбрасывает), поэтому она
    берётся из самого слайда. `frame` — сдвиг и масштаб родителя, `origin` — `chOff`.
    """
    ox, oy, sx, sy = frame
    for node in parent:
        if node.tag not in (GROUP, *SHAPE_TAGS):
            continue
        xfrm = _xfrm(node)
        own = _pair(xfrm, "a:off", "a:ext")
        box = None if own is None else (
            ox + (own[0] - origin[0]) * sx, oy + (own[1] - origin[1]) * sy,
            own[2] * sx, own[3] * sy,
        )
        if node.tag != GROUP:
            yield node, box
            continue
        child = _pair(xfrm, "a:chOff", "a:chExt")
        if box is None or child is None or not child[2] or not child[3]:
            yield from _placed(node, frame, origin)
            continue
        yield from _placed(
            node, (box[0], box[1], box[2] / child[2], box[3] / child[3]), (child[0], child[1])
        )


def _gap(a: _Box, b: _Box) -> float:
    """Зазор между рамками; 0 — касаются или пересекаются."""
    dx = max(b[0] - (a[0] + a[2]), a[0] - (b[0] + b[2]), 0)
    dy = max(b[1] - (a[1] + a[3]), a[1] - (b[1] + b[3]), 0)
    return max(dx, dy)


def _within(inner: _Box, outer: _Box) -> bool:
    return (inner[0] >= outer[0] and inner[1] >= outer[1]
            and inner[0] + inner[2] <= outer[0] + outer[2]
            and inner[1] + inner[3] <= outer[1] + outer[3])


def _is_line(node: Any, box: _Box) -> bool:
    if node.tag == qn("p:cxnSp"):
        return True
    small, big = sorted((box[2], box[3]))
    return big > 0 and small <= LINE_SHARE * big


def _ends(node: Any, box: _Box) -> tuple[_Box, _Box]:
    """Концы линии — углы рамки с учётом отражения `flipH`/`flipV`."""
    xfrm = _xfrm(node)
    x0, y0, x1, y1 = box[0], box[1], box[0] + box[2], box[1] + box[3]
    if xfrm is not None and xfrm.get("flipH") == "1":
        x0, x1 = x1, x0
    if xfrm is not None and xfrm.get("flipV") == "1":
        y0, y1 = y1, y0
    return (x0, y0, 0, 0), (x1, y1, 0, 0)


def _drop_orphan_decor(recipe: Recipe, slide: Any, width: int, height: int) -> None:
    """Декор снятой зоны уходит вместе с ней (RG52).

    Education `ex013` — схема из стрелок: зоны нашим текстом не заполнились и ушли
    (`_drop_empty_zones`), а линии, плашка и иконка остались чертежом без подписей.
    Привязок у линий нет (`a:stCxn`/`a:endCxn` — 2 на 45 линий трёх шаблонов), поэтому
    декор зоны отличается только геометрией, в допуске `TOUCH_SHARE`:

    1. Линия уходит, если её конец у рамки снятой зоны, или каскадом — если конец у уже снятой
       линии, а ни один конец не у текста (ствол развилки между заполненными зонами остаётся);
       разветвители схемы (`ex013`: цепочка в три шага) уходят до конца.
    2. Фигура без текста уходит, если касается снятой зоны или конца снятой линии
       **и** не касается ни одной фигуры с текстом: плашка, на которой остался наш
       текст, — фигура автора (RG40).
    3. Фигура без текста целиком внутри снятой плашки уходит с ней: иконка карточки.

    Не трогаются никогда: фон и полосы во всю ширину или высоту, плейсхолдеры, всё, что
    адресует рецепт (зоны, члены повторов — их правило `_drop_spare_repeats`, картинка),
    таблицы и диаграммы (`_drop_unfilled_frames`), заголовок и всё, что касается его.
    Замер на трёх шаблонах — в proposal change `a-decoration-leaves-with-its-zone`.
    """
    tree = _shapes_tree(slide)
    placed = list(_placed(tree))
    present = {_xml_id(node) for node, _ in placed}
    framed: dict[int, _Box] = {
        zone.xml_id: (float(zone.x), float(zone.y), float(zone.cx), float(zone.cy))
        for zone in recipe.zones
        if zone.xml_id is not None and zone.x is not None and zone.y is not None
        and zone.cx is not None and zone.cy is not None
    }
    titles = {zone.xml_id for zone in recipe.zones if zone.role is TypeLevel.SLIDE_TITLE}
    dropped = [box for xml_id, box in framed.items() if xml_id not in present | titles]
    if not dropped:
        return
    tolerance = TOUCH_SHARE * width
    addressed = _addressed(recipe)

    def touches(box: _Box, targets: list[_Box]) -> bool:
        return any(_gap(box, target) <= tolerance for target in targets)

    content: list[_Box] = []
    title_boxes: list[_Box] = []
    decor: list[tuple[Any, _Box]] = []
    for node, own in placed:
        xml_id = _xml_id(node)
        mark = node.find(f"./*/{qn('p:nvPr')}/{qn('p:ph')}")
        if mark is not None and mark.get("type") in _FOOTERS:
            continue
        box = own if own is not None else framed.get(xml_id) if xml_id is not None else None
        written = any((text.text or "").strip() for text in node.iter(qn("a:t")))
        if written or xml_id in titles or xml_id == recipe.picture_xml_id:
            if box is not None:
                content.append(box)
                if xml_id in titles:
                    title_boxes.append(box)
            continue
        if (box is None or mark is not None or xml_id in addressed or node.tag == GRAPHIC_FRAME
                or box[2] >= FULL_SPAN_SHARE * width or box[3] >= FULL_SPAN_SHARE * height):
            continue
        decor.append((node, box))
    # Декор у заголовка — оформление слайда, даже если касается и снятой зоны.
    decor = [(node, box) for node, box in decor if not touches(box, title_boxes)]

    lines = [(node, box) for node, box in decor if _is_line(node, box)]
    shapes = [(node, box) for node, box in decor if not _is_line(node, box)]
    gone: list[tuple[Any, _Box]] = []
    reach: list[_Box] = []

    def orphaned(node: Any, box: _Box) -> bool:
        """Конец у снятой зоны — линия ведёт в пустоту. Конец только у снятой линии — каскад,
        если ни один конец не у текста: ствол развилки между заполненными зонами остаётся."""
        ends = _ends(node, box)
        if any(touches(end, dropped) for end in ends):
            return True
        return any(touches(end, reach) for end in ends) and not any(
            touches(end, content) for end in ends
        )

    while more := [
        (node, box) for node, box in lines
        if all(node is not taken for taken, _ in gone) and orphaned(node, box)
    ]:
        gone += more
        reach += [box for _, box in more]
    line_ends = [end for node, box in gone for end in _ends(node, box)]
    plates = [
        (node, box) for node, box in shapes
        if (touches(box, dropped) or touches(box, line_ends)) and not touches(box, content)
    ]
    inner = [
        node for node, box in shapes
        if all(node is not plate for plate, _ in plates)
        and any(_within(box, plate_box) for _, plate_box in plates)
    ]
    for node in [node for node, _ in gone + plates] + inner:
        _remove(node)
