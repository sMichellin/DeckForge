"""Составные компоненты вместо OOXML SmartArt. Change (21) `smartart-icons`.

python-pptx не создаёт diagram-part, а инъекция готового XML даёт нередактируемый объект.
Собираем из автофигур и коннекторов: каждый элемент — отдельная редактируемая фигура,
что выполняет C3 строже, чем настоящий SmartArt (§10). Фигуры собраны в группу, чтобы
компонент двигался в PowerPoint целиком.

Раскладка — `layout.diagram`, кегль — из `fit_report`, цвета — ссылками на тему.
"""

from __future__ import annotations

from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu, Pt

from deckforge.domain.base import BBox
from deckforge.domain.enums import ColorRef, TextRole
from deckforge.domain.rules import contrast_ratio
from deckforge.domain.slide import SmartArtBlock
from deckforge.domain.template import ComponentKind, TemplateManifest
from deckforge.domain.units import EMU_PER_PT
from deckforge.layout.by_design import DesignRules
from deckforge.layout.diagram import diagram_geometry
from deckforge.rendering.theme_binding import apply_theme_color, theme_font_token

#: Толщина коннектора в долях кегля подписей: вес линии следует за весом текста.
LINK_WEIGHT = 0.1


def text_on(fill: ColorRef, manifest: TemplateManifest) -> ColorRef:
    """Цвет текста на заливке: `dk1` или `lt1` — что контрастнее с ней в этой теме."""
    colors = manifest.theme.colors
    background = colors.get(fill)
    dark = contrast_ratio(colors.get(ColorRef.DK1), background)
    light = contrast_ratio(colors.get(ColorRef.LT1), background)
    return ColorRef.DK1 if dark >= light else ColorRef.LT1


def node_colors(block: SmartArtBlock, default: ColorRef = ColorRef.ACCENT1) -> list[ColorRef]:
    """Заливки узлов: слоты из IR, иначе `default` — акцент по роли дизайн-системы (DG3)."""
    refs = block.color_refs or [default]
    return [refs[i % len(refs)] for i in range(len(block.items))]


def add_smartart(
    slide: object,
    block: SmartArtBlock,
    manifest: TemplateManifest,
    *,
    size_pt: float,
    text_color: ColorRef | None,
    design: DesignRules | None = None,
    fill: ColorRef = ColorRef.ACCENT1,
) -> object:
    """Группа фигур компонента. `text_color` — цвет свободного текста макета: им подписаны
    элементы шкалы времени и нарисованы коннекторы, потому что они лежат на фоне слайда.

    `design` — плитка из каталога дизайн-системы (та же, что у вписывания), `fill` —
    заливка узлов, когда IR слотов не назвал."""
    box = block.bbox
    if box is None:
        raise ValueError(f"компонент {block.block_id} без координат")
    tile = design.tile() if design is not None else manifest.component(ComponentKind.TILE)
    geometry = diagram_geometry(block.pattern, len(block.items), box, tile)
    body = manifest.typography(TextRole.BODY)
    on_background = text_color or (body.color_ref if body else None) or ColorRef.DK1
    font_token = theme_font_token(body.font_ref) if body else None
    bold = body.bold if body else None

    group = slide.shapes.add_group_shape()  # type: ignore[attr-defined]
    group.name = f"Компонент {block.pattern.value}"
    shapes = group.shapes

    # Коннекторы первыми: ось шкалы времени проходит под маркерами.
    for link in geometry.links:
        connector = shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Emu(link.x1), Emu(link.y1), Emu(link.x2), Emu(link.y2)
        )
        _drop_style(connector)
        apply_theme_color(connector.line, on_background)
        connector.line.width = Emu(round(size_pt * EMU_PER_PT * LINK_WEIGHT))
        if geometry.arrows:
            connector._element.spPr.find(qn("a:ln")).append(
                parse_xml(f'<a:tailEnd {nsdecls("a")} type="triangle"/>')
            )

    preset = MSO_SHAPE.OVAL if geometry.round_nodes else MSO_SHAPE.ROUNDED_RECTANGLE
    for text, node, node_fill in zip(
        block.items, geometry.nodes, node_colors(block, fill), strict=True
    ):
        shape = shapes.add_shape(preset, *_emu(node))
        _drop_style(shape)
        apply_theme_color(shape.fill, node_fill)
        shape.line.fill.background()
        if geometry.text_inside:
            _write(shape.text_frame, text, size_pt, text_on(node_fill, manifest), font_token,
                   bold, MSO_ANCHOR.MIDDLE)

    if not geometry.text_inside:
        for text, label in zip(block.items, geometry.labels, strict=True):
            textbox = shapes.add_textbox(*_emu(label))
            _write(textbox.text_frame, text, size_pt, on_background, font_token, bold,
                   MSO_ANCHOR.TOP)
    return group


def _emu(box: BBox) -> tuple[Emu, Emu, Emu, Emu]:
    return Emu(box.x), Emu(box.y), Emu(box.cx), Emu(box.cy)


def _drop_style(shape: object) -> None:
    """Стиль автофигуры тянет из темы тень и обводку — всё нужное задаём явно."""
    element = shape._element  # type: ignore[attr-defined]
    style = element.find(qn("p:style"))
    if style is not None:
        element.remove(style)


def _write(
    frame: object,
    text: str,
    size_pt: float,
    color: ColorRef,
    font_token: str | None,
    bold: bool | None,
    anchor: MSO_ANCHOR,
) -> None:
    frame.word_wrap = True  # type: ignore[attr-defined]
    frame.auto_size = MSO_AUTO_SIZE.NONE  # type: ignore[attr-defined]
    frame.vertical_anchor = anchor  # type: ignore[attr-defined]
    paragraph = frame.paragraphs[0]  # type: ignore[attr-defined]
    paragraph.text = text
    paragraph.alignment = PP_ALIGN.CENTER
    for run in paragraph.runs:
        run.font.size = Pt(size_pt)
        apply_theme_color(run.font, color)
        if font_token is not None:
            run.font.name = font_token
        if bold is not None:
            run.font.bold = bold
