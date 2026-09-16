"""Вставка изображений с `fit: cover/contain` без искажения пропорций. Change (13)/(24).

Картинка всегда вставляется целиком и с исходными пропорциями: `cover` обрезает лишнее
через `crop_*` (обрезка обратима в PowerPoint), `contain` уменьшает рамку по центру.
"""

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx.util import Emu

from deckforge.domain.enums import ImageFit
from deckforge.domain.slide import ImageBlock


def cover_crop(img_w: int, img_h: int, box_w: int, box_h: int) -> tuple[float, float, float, float]:
    """Доли обрезки (слева, сверху, справа, снизу), чтобы картинка заполнила рамку."""
    img_ratio, box_ratio = img_w / img_h, box_w / box_h
    if img_ratio > box_ratio:
        side = (1 - box_ratio / img_ratio) / 2
        return (side, 0.0, side, 0.0)
    side = (1 - img_ratio / box_ratio) / 2
    return (0.0, side, 0.0, side)


def contain_box(
    img_w: int, img_h: int, x: int, y: int, cx: int, cy: int
) -> tuple[int, int, int, int]:
    """Наибольший прямоугольник с пропорциями картинки внутри рамки, по центру."""
    scale = min(cx / img_w, cy / img_h)
    width, height = round(img_w * scale), round(img_h * scale)
    return (x + (cx - width) // 2, y + (cy - height) // 2, width, height)


def add_image(slide: object, block: ImageBlock, image_path: Path) -> object:
    """Кладёт картинку в рамку блока. Координаты обязательны — их проверил `validate`."""
    box = block.bbox
    if box is None:
        raise ValueError(f"картинка {block.block_id} без координат")
    with Image.open(image_path) as image:
        img_w, img_h = image.size

    shapes = slide.shapes  # type: ignore[attr-defined]
    if block.fit is ImageFit.CONTAIN:
        x, y, cx, cy = contain_box(img_w, img_h, box.x, box.y, box.cx, box.cy)
        picture = shapes.add_picture(str(image_path), Emu(x), Emu(y), Emu(cx), Emu(cy))
    else:
        picture = shapes.add_picture(
            str(image_path), Emu(box.x), Emu(box.y), Emu(box.cx), Emu(box.cy)
        )
        left, top, right, bottom = cover_crop(img_w, img_h, box.cx, box.cy)
        picture.crop_left, picture.crop_top = left, top
        picture.crop_right, picture.crop_bottom = right, bottom

    # python-pptx ставит в описание имя файла — это утечка пути, а не альтернативный текст.
    picture._element.nvPicPr.cNvPr.set("descr", block.alt_text or "")
    return picture
