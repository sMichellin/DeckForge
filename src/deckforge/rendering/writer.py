"""`SlideIR` → .pptx. Changes (13) `pptx-writer`, (14) диаграммы, таблицы, KPI,
(21) составные компоненты и иконки.

Инварианты рендерера (ADR-002):
* слайд создаётся **только** на макете из шаблона — по части пакета, а не по позиции
  в `prs.slide_layouts` (там только первый мастер и другой порядок, см. proposal change 13);
* текст кладётся в плейсхолдеры, а не в новые текстбоксы, где плейсхолдер есть;
* цвет задаётся ссылкой на тему (`MSO_THEME_COLOR`), не литеральным RGB;
* кегль берётся из `manifest.typography_scale` через `fit_report`, не подбирается на глаз.
Любой литерал RGB/шрифта/EMU в этом модуле — ошибка линтера `lint_no_template_constants`.
"""

from __future__ import annotations

import io
import re
import zipfile
from collections.abc import Collection
from pathlib import Path

from lxml import etree
from PIL import Image, UnidentifiedImageError
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from deckforge.designsystem import DesignSystem
from deckforge.designsystem.contrast import TextClass, readable_ref
from deckforge.domain.base import BBox
from deckforge.domain.content import ContentPackage, Dataset
from deckforge.domain.enums import ColorRef, ImageSource, ListStyle, TextRole
from deckforge.domain.rules import TEXT_SLOTS, readable_text_ref
from deckforge.domain.slide import (
    Block,
    BulletItem,
    BulletsBlock,
    CalloutBlock,
    ChartBlock,
    DeckIR,
    FitResult,
    IconBlock,
    ImageBlock,
    KpiBlock,
    QuoteBlock,
    SlideIR,
    SmartArtBlock,
    TableBlock,
    TextBlock,
)
from deckforge.domain.template import (
    BulletStyle,
    ComponentKind,
    LayoutSpec,
    TemplateManifest,
)
from deckforge.domain.units import EMU_PER_PT
from deckforge.layout.by_design import DesignRules
from deckforge.layout.diagram import SUPPORTED_PATTERNS
from deckforge.layout.errors import LayoutFitError
from deckforge.layout.fitting import fit_block, fit_table, table_row_heights
from deckforge.layout.fonts import FontLibrary
from deckforge.layout.lists import draws_icons, icon_list_geometry, list_style
from deckforge.layout.tabular import dataset_bullets, table_cells, table_has_header
from deckforge.parsing.ooxml.layouts import resolve_placeholders
from deckforge.parsing.package import TemplatePackage
from deckforge.rendering.boxed import add_boxed
from deckforge.rendering.charts import add_chart, chart_problem
from deckforge.rendering.icons import add_icon, draw_icon, icon_nodes
from deckforge.rendering.images import add_image
from deckforge.rendering.recipe_slide import clone_recipe
from deckforge.rendering.smartart import add_smartart
from deckforge.rendering.tables import add_table, template_table_style
from deckforge.rendering.theme_binding import (
    apply_theme_color,
    scheme_token,
    theme_font_token,
)

#: Плейсхолдеры, в которые текст не кладётся: там ждут картинку, диаграмму, таблицу.
_NON_TEXT_PLACEHOLDERS = frozenset({"PIC", "CHART", "TBL", "MEDIA", "CLIPART", "DGM", "SLDIMG"})

_COLOR_REFS = {ref.value: ref for ref in ColorRef}
#: Элементы свойств текста: цвет внутри них — цвет букв, а не заливки фигуры.
_TEXT_PROPERTIES = frozenset({"defRPr", "rPr", "endParaRPr"})

_CONTENT_TYPE_POTX = "presentationml.template.main+xml"
_CONTENT_TYPE_PPTX = "presentationml.presentation.main+xml"
_PARAGRAPH_BREAK = re.compile(r"\r\n|[\n\r]")


class WriterError(ValueError):
    """IR нельзя записать, не нарушив инвариантов. Файл при этом не создаётся."""


def _open_template(path: Path) -> object:
    """python-pptx не открывает .potx: подменяем content type главной части в памяти."""
    data = path.read_bytes()
    with zipfile.ZipFile(io.BytesIO(data)) as zin:
        types = zin.read("[Content_Types].xml").decode("utf-8")
        if _CONTENT_TYPE_POTX not in types:
            return Presentation(io.BytesIO(data))
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                payload = zin.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    payload = types.replace(_CONTENT_TYPE_POTX, _CONTENT_TYPE_PPTX).encode("utf-8")
                zout.writestr(item, payload)
    buffer.seek(0)
    return Presentation(buffer)


def _drop_sample_slides(prs: object, keep: list[object]) -> None:
    """Слайды-примеры шаблона в колоду не попадают; части без связей python-pptx не пишет.

    Произвольные показы и разделы ссылаются на слайды-примеры: без них ссылки повисли бы,
    и PowerPoint предложил бы «восстановить» файл. К новой колоде они отношения не имеют.
    """
    presentation = prs.part._element  # type: ignore[attr-defined]
    for node in presentation.findall(qn("p:custShowLst")):
        presentation.remove(node)
    for ext in presentation.findall(f"{qn('p:extLst')}/{qn('p:ext')}"):
        if any(etree.QName(child).localname == "sectionLst" for child in ext):
            ext.getparent().remove(ext)

    # Слайды колоды уже в файле рядом с примерами: копии примеров по рецепту и обычные
    # удалять нельзя (changes `recipe-slide-in-the-writer`, `slide-order-follows-the-deck`).
    kept = {slide.part.partname for slide in keep}  # type: ignore[attr-defined]
    slide_ids = prs.slides._sldIdLst  # type: ignore[attr-defined]
    for slide_id in list(slide_ids):
        part = prs.part.related_part(slide_id.rId)  # type: ignore[attr-defined]
        if part.partname in kept:
            continue
        prs.part.drop_rel(slide_id.rId)  # type: ignore[attr-defined]
        slide_ids.remove(slide_id)


def _lines_in(block: object) -> bool:
    """Есть ли что писать в зону рецепта: пустой блок оставил бы там текст шаблона."""
    if isinstance(block, BulletsBlock):
        return any(item.text.strip() for item in block.items)
    if isinstance(block, TextBlock):
        return bool(block.text.strip())
    return False


def _text_rank(placeholder: object) -> int:
    """Основной текст важнее подзаголовка, подзаголовок — заголовка: у заголовка часто акцент."""
    kind = str(placeholder.placeholder_format.type).split(" ")[0]  # type: ignore[attr-defined]
    return {"BODY": 0, "OBJECT": 0, "SUBTITLE": 1, "TITLE": 2, "CENTER_TITLE": 2}.get(kind, 3)


def _layout_text_color(layout: object) -> ColorRef | None:
    """Цвет текста макета — ссылкой на тему из его плейсхолдеров, основной текст первым.

    Тёмный макет в шаблонах кейса тёмный за счёт фона, а не `clrMap`: светлый цвет задан
    плейсхолдерам. Свободный текст (текстбокс, подпись KPI, подписи диаграммы) берёт этот цвет,
    иначе цвет роли из типошкалы дал бы тёмные буквы на тёмном фоне. Заголовок — последним:
    на части макетов VK Education он акцентный, а текст — нет.
    """
    placeholders = sorted(layout.placeholders, key=_text_rank)  # type: ignore[attr-defined]
    for placeholder in placeholders:
        for node in placeholder._element.iter(qn("a:schemeClr")):
            ancestors = {etree.QName(a).localname for a in node.iterancestors()}
            ref = _COLOR_REFS.get(node.get("val") or "")
            if ref is not None and ancestors & _TEXT_PROPERTIES:
                return ref
    return None


def _paragraphs(block: TextBlock | BulletsBlock) -> list[tuple[str, int]]:
    if isinstance(block, TextBlock):
        return [(text, 0) for text in _PARAGRAPH_BREAK.split(block.text)]
    return [(item.text, item.level) for item in block.items]


class SlideValidator:
    """Инварианты IR до записи — общие для pptx и html: оба формата показывают одно и то же."""

    def __init__(
        self, manifest: TemplateManifest, recipes: Collection[str] | None = None
    ) -> None:
        self.manifest = manifest
        #: Имена композиций шаблона. Каталог даёт тот, у кого он на руках: `domain`
        #: о дизайн-системе не знает (ADR-003), поэтому существование рецепта
        #: проверяется здесь, а не в `SlideIR.by_recipe`. Без каталога проверка молчит —
        #: html зовёт валидатор без него и о композициях не судит.
        self.recipes = recipes

    def validate(self, slide: SlideIR, content: ContentPackage | None = None) -> None:
        """Инварианты §4.4: макет есть в манифесте, placeholder_idx существует,
        координаты внутри полей, цвета — только `color_ref`, текст вписан."""
        problems = self.problems(slide, content)
        if problems:
            raise WriterError("\n".join(problems))

    def problems(self, slide: SlideIR, content: ContentPackage | None) -> list[str]:
        layout = self.manifest.layout(slide.layout_id)
        if layout is None:
            return [f"{slide.slide_id}: макета {slide.layout_id} нет в манифесте"]
        if not slide.blocks:
            # Такой слайд — пустой фон шаблона; живой прогон показал, что без проверки
            # он молча попадает в колоду.
            return [f"{slide.slide_id}: слайд без блоков"]
        if slide.recipe_id and self.recipes is not None and slide.recipe_id not in self.recipes:
            # Выдуманный моделью рецепт прошёл весь конвейер и упал в писателе `KeyError`
            # через 138 секунд прогона: слайд по рецепту, которого нет, молча уходил
            # собираться на пустом макете. Утечка контракта обязана называть себя.
            return [
                f"{slide.slide_id}: рецепта {slide.recipe_id} нет в каталоге композиций шаблона"
            ]
        out: list[str] = []
        if all(isinstance(block, ImageBlock) for block in slide.blocks):
            out.append(f"{slide.slide_id}: слайд из одних картинок нарушает C3")
        taken: set[int] = set()
        for block in slide.blocks:
            where = f"{slide.slide_id}/{block.block_id}"
            # Блок слайда по рецепту стоит в зоне шаблона: рамку ему дал автор, кегль —
            # его же `rPr`. Ни координат, ни отчёта о вписывании у такого блока нет
            # и быть не должно (change `recipe-slide-in-the-writer`). Спрашивается весь
            # слайд, а не отдельный блок: смешанный слайд по рецепту не собран, и его
            # блоки вне зон обязаны пройти обычные проверки (`SlideIR.by_recipe`, RG3).
            if slide.by_recipe:
                if not _lines_in(block):
                    out.append(f"{where}: блок в зоне рецепта без текста")
                continue
            idx = getattr(block, "placeholder_idx", None)
            if idx is not None:
                if idx in taken:
                    out.append(f"{where}: плейсхолдер idx={idx} уже занят другим блоком")
                taken.add(idx)
            if isinstance(block, TextBlock | BulletsBlock):
                out += self._text_problems(where, block, layout, slide)
            elif isinstance(block, ImageBlock):
                out += self._image_problems(where, block, content)
            elif isinstance(block, ChartBlock):
                out += self._chart_problems(where, block, content)
            elif isinstance(block, TableBlock):
                out += self._box_problems(where, block.bbox, block.type)
                try:
                    table_cells(block, _dataset(content, block.dataset_ref))
                except LayoutFitError as exc:
                    out.append(f"{where}: {exc}")
                out += self._fit_problems(where, block.block_id, slide)
            elif isinstance(block, KpiBlock):
                out += self._box_problems(where, block.bbox, block.type)
                out += self._fit_problems(where, block.block_id, slide)
            elif isinstance(block, SmartArtBlock):
                out += self._box_problems(where, block.bbox, block.type)
                if block.pattern not in SUPPORTED_PATTERNS:
                    out.append(f"{where}: паттерн {block.pattern.value} не строится, "
                               "а в буллеты не заменён")
                if not all(item.strip() for item in block.items):
                    out.append(f"{where}: пустой элемент компонента")
                out += self._fit_problems(where, block.block_id, slide)
            elif isinstance(block, IconBlock):
                out += self._box_problems(where, block.bbox, block.type)
                if icon_nodes(block.query) is None:
                    out.append(f"{where}: иконки «{block.query}» нет в Lucide")
            elif isinstance(block, QuoteBlock | CalloutBlock):
                # Цитата и callout — свободные блоки (DG4): место им даёт решатель,
                # кегль — вписывание по дизайн-системе. Без того и другого не пишутся.
                out += self._box_problems(where, block.bbox, block.type)
                out += self._fit_problems(where, block.block_id, slide)
        return out

    def _box_problems(self, where: str, box: BBox | None, kind: str) -> list[str]:
        if box is None:
            return [f"{where}: {kind} без координат"]
        if not self.manifest.content_bbox.contains(box):
            return [f"{where}: координаты выходят за поля шаблона"]
        return []

    def _fit_problems(self, where: str, block_id: str, slide: SlideIR) -> list[str]:
        fit = slide.fit_report.get(block_id)
        ladder = self.manifest.size_ladder_pt
        if fit is None:
            return [f"{where}: нет записи в fit_report — вписывание не выполнялось"]
        out = []
        if fit.overflow:
            out.append(f"{where}: переполнение, стратегия {fit.strategy}")
        if ladder and fit.final_size_pt not in ladder:
            out.append(f"{where}: кегль {fit.final_size_pt:g} вне шкалы шаблона {ladder}")
        return out

    def _chart_problems(
        self, where: str, block: ChartBlock, content: ContentPackage | None
    ) -> list[str]:
        out = self._box_problems(where, block.bbox, block.type)
        dataset = _dataset(content, block.dataset_ref)
        if dataset is None:
            out.append(f"{where}: датасета {block.dataset_ref} нет в ContentPackage")
        elif (problem := chart_problem(block.chart_type, dataset)) is not None:
            out.append(f"{where}: диаграмму не построить — {problem}")
        return out

    def _text_problems(
        self, where: str, block: TextBlock | BulletsBlock, layout: LayoutSpec, slide: SlideIR
    ) -> list[str]:
        out: list[str] = []
        texts = [block.text] if isinstance(block, TextBlock) else [i.text for i in block.items]
        if not any(text.strip() for text in texts):
            out.append(f"{where}: пустой текст — на слайде осталось бы «Введите текст»")
        if block.placeholder_idx is not None and block.bbox is not None:
            out.append(f"{where}: заданы и плейсхолдер, и координаты — координаты потерялись бы")
        if block.placeholder_idx is not None:
            placeholder = layout.placeholder(block.placeholder_idx)
            if placeholder is None:
                out.append(
                    f"{where}: в макете {layout.layout_id} нет плейсхолдера "
                    f"idx={block.placeholder_idx}"
                )
            elif placeholder.ph_type in _NON_TEXT_PLACEHOLDERS:
                out.append(f"{where}: текст в плейсхолдер {placeholder.ph_type}")
        elif block.bbox is None:
            out.append(f"{where}: нет ни плейсхолдера, ни координат")
        elif not self.manifest.content_bbox.contains(block.bbox):
            out.append(f"{where}: координаты выходят за поля шаблона")

        return out + self._fit_problems(where, block.block_id, slide)

    def _image_problems(
        self, where: str, block: ImageBlock, content: ContentPackage | None
    ) -> list[str]:
        out: list[str] = []
        if block.bbox is None:
            out.append(f"{where}: картинка без координат")
        elif not self.manifest.slide_size.bbox.contains(block.bbox):
            # Картинке можно во весь слайд — «фото на весь слайд» обычный макет, — но не за край.
            out.append(f"{where}: картинка выходит за слайд")
        if block.source is ImageSource.GENERATED:
            out.append(f"{where}: сгенерированные картинки добавит change (24) image-generation")
        elif content is None or content.asset(block.asset_ref or "") is None:
            out.append(f"{where}: ассета {block.asset_ref} нет в ContentPackage")
        else:
            asset = content.asset(block.asset_ref or "")
            if asset is not None and not Path(asset.path).is_file():
                out.append(f"{where}: файл ассета не найден: {asset.path}")
            elif asset is not None:
                try:
                    with Image.open(asset.path) as image:
                        image.verify()
                except (UnidentifiedImageError, OSError):
                    out.append(f"{where}: ассет {asset.path} не картинка")
        return out


class SlideDegrader:
    """Цепочки §15 — общие для pptx и html: «диаграмма → таблица → буллеты»,
    «составной компонент → буллеты», неизвестная иконка убирается."""

    def __init__(self, manifest: TemplateManifest, fonts: FontLibrary | None = None) -> None:
        self.manifest = manifest
        self.fonts = fonts
        #: Что подменено и почему — для аудита и интерфейса.
        self.degradations: list[str] = []

    def degrade(self, slide: SlideIR, content: ContentPackage | None) -> SlideIR:
        """Цепочка §15: диаграмма, которую не построить, → таблица → буллеты.

        Подменённый блок вписывается сразу, иначе писатель не знал бы, влезает ли он.
        Каждая подмена записывается в `degradations`: молча данные не меняют вид.
        """
        layout = self.manifest.layout(slide.layout_id)
        if layout is None:
            return slide
        blocks: list[Block] = []
        report = dict(slide.fit_report)
        for block in slide.blocks:
            replaced: tuple[Block, FitResult] | None = None
            where = f"{slide.slide_id}/{block.block_id}"
            if isinstance(block, ChartBlock) and block.bbox is not None:
                dataset = _dataset(content, block.dataset_ref)
                reason = chart_problem(block.chart_type, dataset) if dataset else None
                if dataset is not None and reason is not None:
                    table = TableBlock(block_id=block.block_id, dataset_ref=block.dataset_ref,
                                       **_coords(block.bbox))
                    try:
                        fit = fit_table(table, block.bbox, self.manifest, dataset=dataset,
                                        fonts=self.fonts)
                    except LayoutFitError:
                        # Из этих данных не собрать и таблицу: оставляем диаграмму,
                        # валидация назовёт причину.
                        fit = None
                    if fit is not None:
                        replaced = (table, fit)
                        self.degradations.append(f"{where}: диаграмма → таблица ({reason})")
                        bullets = dataset_bullets(dataset)
                        if fit.overflow and bullets:
                            replaced = self._as_bullets(block.block_id, block.bbox, bullets,
                                                        layout)
                            self.degradations.append(f"{where}: таблица → буллеты (не влезла)")
            elif isinstance(block, TableBlock) and block.bbox is not None:
                table_fit = report.get(block.block_id)
                dataset = _dataset(content, block.dataset_ref)
                items = (
                    dataset_bullets(dataset) if dataset is not None and not block.rows
                    else [
                        " — ".join(f"{h}: {v}" if h else v
                                   for h, v in zip(block.header, row, strict=False))
                        if block.header else " — ".join(row)
                        for row in block.rows
                    ]
                )
                if table_fit is not None and table_fit.overflow and items:
                    replaced = self._as_bullets(block.block_id, block.bbox, items, layout)
                    self.degradations.append(f"{where}: таблица → буллеты (не влезла)")
            elif isinstance(block, SmartArtBlock) and block.bbox is not None:
                smartart_fit = report.get(block.block_id)
                reason = (
                    "паттерн не поддерживается" if block.pattern not in SUPPORTED_PATTERNS
                    else "не влез" if smartart_fit is not None and smartart_fit.overflow
                    else None
                )
                if reason is not None:
                    replaced = self._as_bullets(block.block_id, block.bbox, block.items, layout)
                    self.degradations.append(
                        f"{where}: smartart {block.pattern.value} → буллеты ({reason})"
                    )
            elif isinstance(block, BulletsBlock) and draws_icons(block):
                missing = [item.icon or "—" for item in block.items
                           if not item.icon or icon_nodes(item.icon) is None]
                if missing:
                    # Иконочный список с дырой выглядит сломанным: у одних пунктов знак,
                    # у других пустота. Весь список — маркером шаблона, и это названо.
                    plain = block.model_copy(update={
                        "style": ListStyle.BULLETED,
                        "items": [item.model_copy(update={"icon": None}) for item in block.items],
                    })
                    replaced = (plain, fit_block(plain, layout, self.manifest, fonts=self.fonts))
                    self.degradations.append(
                        f"{where}: иконочный список → маркированный "
                        f"(иконок нет в наборе: {', '.join(missing)})"
                    )
            elif isinstance(block, IconBlock) and icon_nodes(block.query) is None:
                # Иконка — украшение: без неё слайд цел, а ошибка записи потеряла бы колоду.
                self.degradations.append(
                    f"{where}: иконки «{block.query}» нет в Lucide — блок убран"
                )
                continue
            if replaced is None:
                blocks.append(block)
            else:
                blocks.append(replaced[0])
                report[block.block_id] = replaced[1]
        return slide.model_copy(update={"blocks": blocks, "fit_report": report})

    def _as_bullets(
        self, block_id: str, box: BBox, items: list[str], layout: LayoutSpec
    ) -> tuple[Block, FitResult]:
        bullets = BulletsBlock(block_id=block_id, items=[BulletItem(text=t) for t in items],
                               **_coords(box))
        return bullets, fit_block(bullets, layout, self.manifest, fonts=self.fonts)


class PptxWriter:
    def __init__(
        self,
        template_path: Path,
        manifest: TemplateManifest,
        fonts: FontLibrary | None = None,
        design_system: DesignSystem | None = None,
    ) -> None:
        self.template_path = template_path
        self.manifest = manifest
        #: Шрифты для вписывания подменённых блоков (цепочка деградации диаграммы).
        self.fonts = fonts
        #: Ответы дизайн-системы (DG3): акцент по роли, кегли числа, плитка, полоса
        #: цитаты. Те же, что видело вписывание, — иначе записанное разойдётся с вписанным.
        self.design = DesignRules(manifest, design_system)
        #: Каталог композиций шаблона: по нему writer клонирует слайд-пример
        #: (change `recipe-slide-in-the-writer`).
        self.recipes = {
            recipe.recipe_id: recipe
            for recipe in (design_system.recipes if design_system else [])
        }
        #: Что было подменено при последней записи и почему — для аудита и интерфейса.
        self.degradations: list[str] = []
        self.validator = SlideValidator(manifest, self.recipes)

    def validate(self, slide: SlideIR, content: ContentPackage | None = None) -> None:
        self.validator.validate(slide, content)

    # --- запись --------------------------------------------------------------

    def write(
        self, deck: DeckIR, out_path: Path, content: ContentPackage | None = None
    ) -> Path:
        if deck.template_id != self.manifest.template_id:
            raise WriterError(
                f"template_id колоды {deck.template_id} не совпадает с манифестом "
                f"{self.manifest.template_id}"
            )
        degrader = SlideDegrader(self.manifest, self.fonts)
        slides = [degrader.degrade(slide, content) for slide in deck.slides]
        self.degradations = degrader.degradations
        problems = [p for slide in slides for p in self.validator.problems(slide, content)]
        if problems:
            raise WriterError("\n".join(problems))

        prs = _open_template(self.template_path)
        layouts = self._layouts_by_id(prs)
        table_style = template_table_style(prs)
        # Один проход в порядке колоды: каждый слайд встаёт в конец `sldIdLst`, так что
        # порядок в файле = порядок IR. Примеры удаляются после прохода — рецепту есть
        # что копировать, а новые части не получают имён удалённых примеров
        # (changes `recipe-slide-in-the-writer`, `slide-order-follows-the-deck`).
        written: list[object] = []
        for slide in slides:
            # Тот же предикат, что у проверки и у вписывания. Рецепт в каталоге есть:
            # его отсутствие проверка уже назвала и до записи не пустила. Проверка
            # на `None` рядом — сужение типа для mypy, а не второе условие: непустой
            # `recipe_id` входит в сам предикат.
            recipe_id = slide.recipe_id
            if slide.by_recipe and recipe_id is not None:
                written.append(clone_recipe(prs, self.recipes[recipe_id], slide))
            else:
                written.append(
                    self._render_slide(prs, layouts[slide.layout_id], slide, content, table_style)
                )
        _drop_sample_slides(prs, keep=written)

        out_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(out_path))  # type: ignore[attr-defined]
        return out_path

    def _layouts_by_id(self, prs: object) -> dict[str, object]:
        """`layout_id` → макет python-pptx через имя части пакета.

        Порядок `LayoutSpec.index` восстанавливается тем же обходом, что в парсере; имя макета
        сверяется, чтобы расхождение обходов или манифест от другого файла не положили слайд
        на чужой макет молча.
        """
        with TemplatePackage(self.template_path) as pkg:
            parts: list[tuple[str, str]] = []
            for master in pkg.master_parts():
                master_xml = pkg.read(master)
                for part in pkg.layout_parts(master):
                    if pkg.has(part) and resolve_placeholders(pkg.read(part), master_xml):
                        parts.append((part, pkg.layout_name(part)))

        by_part = {
            str(layout.part.partname).lstrip("/"): layout
            for master in prs.slide_masters  # type: ignore[attr-defined]
            for layout in master.slide_layouts
        }
        out: dict[str, object] = {}
        for spec in self.manifest.layouts:
            if spec.index >= len(parts):
                raise WriterError(f"макет {spec.layout_id}: в шаблоне нет макета №{spec.index}")
            part, name = parts[spec.index]
            if name.strip() == spec.name and part not in by_part:
                raise WriterError(
                    f"макет {spec.layout_id} «{spec.name}» есть в связях мастера, "
                    "но не в его списке макетов — PowerPoint его не покажет"
                )
            if name.strip() != spec.name:
                raise WriterError(
                    f"макет {spec.layout_id} «{spec.name}» не найден в шаблоне: на его месте "
                    f"«{name.strip()}» — манифест снят не с этого файла?"
                )
            out[spec.layout_id] = by_part[part]
        return out

    def _render_slide(
        self,
        prs: object,
        layout: object,
        slide_ir: SlideIR,
        content: ContentPackage | None,
        table_style: str | None,
    ) -> object:
        slide = prs.slides.add_slide(layout)  # type: ignore[attr-defined]
        text_color = _layout_text_color(layout) or self._readable_on_background(slide_ir)
        #: Фон, по которому дизайн-система выбирает видимый акцент (DG3).
        background = self.design.background_hex(slide_ir.layout_id)
        used: set[int] = set()
        for block in slide_ir.blocks:
            if isinstance(block, TextBlock | BulletsBlock):
                size_pt = self._fit_of(slide_ir, block).final_size_pt
                if block.placeholder_idx is not None:
                    self._fill_placeholder(slide, layout, block, size_pt, background)
                    used.add(block.placeholder_idx)
                elif isinstance(block, BulletsBlock) and draws_icons(block):
                    self._add_icon_list(slide, block, size_pt, text_color, background)
                else:
                    self._add_textbox(slide, block, size_pt, text_color, background)
            elif isinstance(block, ImageBlock) and content is not None:
                asset = content.asset(block.asset_ref or "")
                if asset is not None:
                    add_image(slide, block, Path(asset.path))
            elif isinstance(block, ChartBlock):
                dataset = _dataset(content, block.dataset_ref)
                if dataset is not None:
                    add_chart(slide, block, dataset, self.manifest, text_color=text_color)
            elif isinstance(block, TableBlock):
                dataset = _dataset(content, block.dataset_ref)
                cells = table_cells(block, dataset)
                size_pt = self._fit_of(slide_ir, block).final_size_pt
                body = self.manifest.typography(TextRole.BODY)
                add_table(
                    slide, block, cells, size_pt=size_pt, style_id=table_style,
                    row_heights=table_row_heights(
                        block, block.bbox, self.manifest, size_pt,  # type: ignore[arg-type]
                        dataset=dataset, fonts=self.fonts,
                    ),
                    font_token=theme_font_token(body.font_ref) if body else None,
                    has_header=table_has_header(block),
                )
            elif isinstance(block, KpiBlock):
                self._add_kpi(
                    slide, block, self._fit_of(slide_ir, block).final_size_pt, text_color,
                    background,
                )
            elif isinstance(block, SmartArtBlock):
                add_smartart(
                    slide, block, self.manifest,
                    size_pt=self._fit_of(slide_ir, block).final_size_pt,
                    text_color=text_color,
                    design=self.design,
                    fill=self.design.block_accent(background),
                )
            elif isinstance(block, IconBlock):
                add_icon(slide, block, default=self.design.block_accent(background))
            elif isinstance(block, QuoteBlock | CalloutBlock):
                add_boxed(
                    slide, block, self._fit_of(slide_ir, block), self.manifest,
                    self.design, text_color=text_color, background_hex=background,
                )

        # Пустой плейсхолдер в PowerPoint показывает «Введите текст» — такой слайд выглядит
        # недоделанным.
        for placeholder in list(slide.placeholders):
            if placeholder.placeholder_format.idx not in used:
                placeholder._element.getparent().remove(placeholder._element)

        if slide_ir.speaker_note:
            slide.notes_slide.notes_text_frame.text = slide_ir.speaker_note
        return slide

    def _fit_of(self, slide_ir: SlideIR, block: Block) -> FitResult:
        """Замер блока, а если его нет — кегль роли из типошкалы шаблона.

        Колода из десяти слайдов лучше отсутствия колоды (`degradation-is-not-an-error`):
        пропущенная запись о вписывании — повод написать блок кеглем его роли и назвать
        подмену, а не уронить прогон на последней стадии. Ни одного `fit_report[...]`
        по ключу в писателе не остаётся.
        """
        fit = slide_ir.fit_report.get(block.block_id)
        if fit is not None:
            return fit
        role = getattr(block, "role", None) or TextRole.BODY
        step = self.manifest.typography(role) or self.manifest.typography(TextRole.BODY)
        ladder = self.manifest.size_ladder_pt
        size = step.size_pt if step is not None else (ladder[len(ladder) // 2] if ladder else None)
        if size is None:
            raise WriterError(
                f"{slide_ir.slide_id}/{block.block_id}: нет ни замера, ни типошкалы шаблона"
            )
        self.degradations.append(
            f"{slide_ir.slide_id}/{block.block_id}: замера нет — кегль роли {size:g} pt"
        )
        return FitResult(final_size_pt=size)

    def _readable_on_background(self, slide_ir: SlideIR) -> ColorRef | None:
        """Цвет свободного текста, когда макет не назвал его ни в одном плейсхолдере.

        Свободная фигура в OOXML не наследует цвет ни от кого: родителя-плейсхолдера
        у неё нет. Без явной ссылки PowerPoint рисует её чёрной, и на тёмном макете
        это не «цвет по умолчанию», а невидимый текст — ровно то, что случилось
        с заголовками 19.09.

        До этого места доходят только макеты, у которых цвет не объявлен нигде:
        решение автора шаблона всегда главнее измерения. Фона нет в манифесте (он снят
        прежним парсером) — выбирать не из чего, и тогда цвет остаётся за типошкалой,
        как было.

        Порог — правилом `designsystem.contrast`, тем же, что у аудита (change
        `one-contrast-rule`): класс обычного текста, 4,5. Цвет один на весь свободный
        текст слайда — заголовок, подписи диаграммы, показатели, — поэтому берётся класс
        с самым строгим минимумом, а не самый мягкий. Выбор — из пары «тёмный/светлый»
        темы: акценты для текста не предназначены.

        Порог не берёт ни один слот — это дефект шаблона, и `readable_ref` честно
        отвечает `None`. Оставить текст чёрным по умолчанию ещё хуже, поэтому ставится
        лучший из слотов, как до change.
        """
        layout = self.manifest.layout(slide_ir.layout_id)
        background = layout.background if layout is not None else None
        if background is None:
            return None
        chosen = readable_ref(
            self.manifest.theme, background.color_hex, TextClass.BODY, slots=TEXT_SLOTS
        )
        return chosen or readable_text_ref(self.manifest, background.color_hex)

    def _fill_placeholder(
        self,
        slide: object,
        layout: object,
        block: TextBlock | BulletsBlock,
        size_pt: float,
        background: str | None = None,
    ) -> None:
        idx = block.placeholder_idx
        try:
            shape = slide.placeholders[idx]  # type: ignore[attr-defined]
        except KeyError:
            # Дату, футер и номер слайда python-pptx с макета не копирует — клонируем сами.
            source = next(
                (
                    ph for ph in layout.placeholders  # type: ignore[attr-defined]
                    if ph.placeholder_format.idx == idx
                ),
                None,
            )
            if source is None:
                # Парсер видит плейсхолдеры и внутри групп, python-pptx — только верхний уровень.
                raise WriterError(
                    f"плейсхолдер idx={idx} есть в манифесте, но недоступен в макете "
                    "(вложен в группу?)"
                ) from None
            slide.shapes.clone_placeholder(source)  # type: ignore[attr-defined]
            shape = slide.placeholders[idx]  # type: ignore[attr-defined]
        # Гарнитура не задаётся: её наследует плейсхолдер шаблона.
        color = block.color_ref if isinstance(block, TextBlock) else None
        self._write_paragraphs(shape.text_frame, _paragraphs(block), size_pt, color, None, None)
        # Маркер наследуется не всегда: все три шаблона кейса гасят первый уровень тела
        # (`buNone`) и маркируют список со второго, поэтому пункты в плейсхолдере выходили
        # абзацами (прогон VK Education 6b1d9e82b612). Манифест знает, каким знаком и с каким
        # выносом шаблон рисует список, — ставим его явно, как и в свободном блоке.
        if isinstance(block, BulletsBlock):
            self._mark_list(shape.text_frame, block, size_pt, background)

    def _add_textbox(
        self,
        slide: object,
        block: TextBlock | BulletsBlock,
        size_pt: float,
        text_color: ColorRef | None,
        background: str | None = None,
    ) -> None:
        box: BBox = block.bbox  # type: ignore[assignment]
        step = self.manifest.typography(block.role)
        shape = slide.shapes.add_textbox(  # type: ignore[attr-defined]
            Emu(box.x), Emu(box.y), Emu(box.cx), Emu(box.cy)
        )
        shape.text_frame.word_wrap = True
        # Рамку свободному блоку считает решатель: он делит свободную площадь слайда,
        # и её верхний край — граница области, а не решение о том, где начинается текст.
        # Прижатый к нему текст оставлял под собой восемь сантиметров пустоты
        # (прогон 2ac85990b2f2, s04: занято 16 % при норме 25–75).
        shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        # Текстбокс ничего не наследует: гарнитура и цвет — ссылками на тему из типошкалы.
        color = block.color_ref if isinstance(block, TextBlock) else None
        self._write_paragraphs(
            shape.text_frame,
            _paragraphs(block),
            size_pt,
            color or text_color or (step.color_ref if step else None),
            theme_font_token(step.font_ref) if step else None,
            step.bold if step else None,
        )
        # Маркер списка текстбокс тоже не наследует: без него тезисы читаются абзацами
        # (прогон 2ac85990b2f2). Знак, гарнитура, цвет и вынос — из мастера шаблона.
        if isinstance(block, BulletsBlock):
            self._mark_list(shape.text_frame, block, size_pt, background)

    def _mark_list(
        self, text_frame: object, block: BulletsBlock, size_pt: float, background: str | None
    ) -> None:
        """Знак пунктов по стилю списка (DG3): номер — автонумерация цветом акцента по роли
        ДС, иначе маркер шаблона. Иконочный список сюда попадает только в плейсхолдере —
        там иконку к строке не привязать, и он рисуется маркером шаблона."""
        if list_style(block) is ListStyle.NUMBERED:
            _apply_numbering(
                text_frame, self.manifest, size_pt,
                self.design.number_ink(background, size_pt=size_pt),
            )
        elif self.manifest.bullet_levels:
            _apply_bullets(text_frame, self.manifest, size_pt)

    def _add_icon_list(
        self,
        slide: object,
        block: BulletsBlock,
        size_pt: float,
        text_color: ColorRef | None,
        background: str | None,
    ) -> None:
        """Иконочный список (DG3): иконка набора проекта у каждого пункта, текст — одна
        рамка справа от колонки иконок. Иконки и текст — в одной группе, как цитата:
        блок двигается целиком. Цвет иконок — акцент по роли ДС (порог графики)."""
        box: BBox = block.bbox  # type: ignore[assignment]
        parts = icon_list_geometry(
            block, box, size_pt, self.manifest, self.design, fonts=self.fonts
        )
        step = self.manifest.typography(block.role)
        accent = self.design.block_accent(background)
        group = slide.shapes.add_group_shape()  # type: ignore[attr-defined]
        group.name = "Иконочный список"
        for item, icon_box in zip(block.items, parts.icons, strict=True):
            draw_icon(group.shapes, item.icon or "", icon_box, accent)
        frame_box = parts.frame
        shape = group.shapes.add_textbox(
            Emu(frame_box.x), Emu(frame_box.y), Emu(frame_box.cx), Emu(frame_box.cy)
        )
        shape.text_frame.word_wrap = True
        shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        # Уровни не переносятся: иконка стоит у левого края колонки, и сдвиг пункта
        # оторвал бы его от своей иконки.
        self._write_paragraphs(
            shape.text_frame,
            [(item.text, 0) for item in block.items],
            size_pt,
            text_color or (step.color_ref if step else None),
            theme_font_token(step.font_ref) if step else None,
            step.bold if step else None,
        )
        for paragraph in shape.text_frame.paragraphs:
            _clear_marker(paragraph)

    def _add_kpi(
        self,
        slide: object,
        block: KpiBlock,
        size_pt: float,
        text_color: ColorRef | None,
        background: str | None = None,
    ) -> None:
        """Колонка на показатель: значение и подпись — ссылками на тему, кегли из шкалы.

        Кегль подписи и цвет значения — по дизайн-системе (DG3): подпись — мелкий кегль
        числа шаблона, значение — акцент по роли. Дизайн-система не ответила — как было:
        подпись по отношению из компонента (DS4), значение — `accent1`."""
        box: BBox = block.bbox  # type: ignore[assignment]
        value_step = self.manifest.typography(TextRole.SUBTITLE) or self.manifest.typography(
            TextRole.BODY
        )
        label_step = self.manifest.typography(TextRole.CAPTION) or self.manifest.typography(
            TextRole.BODY
        )
        sizes = self.design.kpi_sizes()
        # Отношение кегля подписи к кеглю значения шаблон показывает сам — на своём
        # показателе из примеров (DS4). Нет такого компонента — кегль роли, как было.
        label_pt = (
            sizes.label_for(size_pt, self.manifest)
            if sizes is not None
            else _label_size_pt(self.manifest, size_pt, label_step, size_pt)
        )
        accent = self.design.block_accent(background, size_pt=size_pt)
        width = box.cx // len(block.items)
        for i, item in enumerate(block.items):
            shape = slide.shapes.add_textbox(  # type: ignore[attr-defined]
                Emu(box.x + i * width), Emu(box.y), Emu(width), Emu(box.cy)
            )
            frame = shape.text_frame
            frame.word_wrap = True
            # Колонка показателя получает всю высоту рамки, а занимает две строки:
            # прижатые к верхнему краю, они оставляли под собой пустое поле
            # (прогон f0eb600a3ad7: ни одного слайда в норме плотности).
            frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            frame.clear()
            value = frame.paragraphs[0]
            value.text = item.value
            _style_runs(
                value, size_pt, item.color_ref or accent,
                theme_font_token(value_step.font_ref) if value_step else None,
                value_step.bold if value_step else None,
            )
            label = frame.add_paragraph()
            label.text = item.label
            _style_runs(
                label,
                label_pt,
                text_color or (label_step.color_ref if label_step else None),
                theme_font_token(label_step.font_ref) if label_step else None,
                None,
            )

    @staticmethod
    def _write_paragraphs(
        text_frame: object,
        paragraphs: list[tuple[str, int]],
        size_pt: float,
        color: object,
        font_token: str | None,
        bold: bool | None,
    ) -> None:
        text_frame.clear()  # type: ignore[attr-defined]
        for i, (text, level) in enumerate(paragraphs):
            paragraph = (
                text_frame.paragraphs[0] if i == 0  # type: ignore[attr-defined]
                else text_frame.add_paragraph()  # type: ignore[attr-defined]
            )
            paragraph.text = text
            paragraph.level = level
            _style_runs(paragraph, size_pt, color, font_token, bold)


def _apply_bullets(text_frame: object, manifest: TemplateManifest, size_pt: float) -> None:
    """Ставит маркер шаблона каждому абзацу свободного списка — по уровню пункта.

    python-pptx маркеров не умеет, поэтому свойства абзаца пишутся в XML напрямую —
    теми же узлами, которыми их записал бы PowerPoint: `buChar`, `buFont`, `buClr`.
    """
    declared = len(manifest.bullet_levels)
    for paragraph in text_frame.paragraphs:  # type: ignore[attr-defined]
        level = int(getattr(paragraph, "level", 0) or 0)
        bullet = manifest.bullet_for(level)
        if bullet is None:
            continue
        # Уровни глубже объявленных в шаблоне: маркер берётся ближайший сверху,
        # а сдвиг добавляется — иначе вложенность, видимая в IR, на слайде исчезнет.
        _write_bullet(paragraph, bullet, size_pt, max(0, level - (declared - 1)))


def _write_bullet(
    paragraph: object, bullet: BulletStyle, size_pt: float, extra_levels: int
) -> None:
    """Маркер и вынос одного абзаца.

    Отступы шаблон задаёт не всегда. Когда их нет, вынос считается от кегля — это
    типографская арифметика, а не константа шаблона: маркер отодвигается на ширину
    примерно одного знака, иначе он налезает на текст. Уровень, не объявленный
    в шаблоне, сдвигается от объявленного на тот же вынос — иначе вложенность,
    видимая в IR, на слайде исчезнет.
    """
    indent = bullet.indent_emu or -round(size_pt * EMU_PER_PT)
    margin = bullet.margin_left_emu or -indent
    extra = extra_levels * -indent
    props = (
        paragraph._pPr  # type: ignore[attr-defined]
        if paragraph._pPr is not None  # type: ignore[attr-defined]
        else paragraph._p.get_or_add_pPr()  # type: ignore[attr-defined]
    )
    props.set("marL", str(margin + extra))
    props.set("indent", str(indent))
    for tag in ("a:buClr", "a:buFont", "a:buChar", "a:buNone", "a:buAutoNum"):
        for node in props.findall(qn(tag)):
            props.remove(node)
    if bullet.color_ref is not None:
        clr = props.makeelement(qn("a:buClr"), {})
        scheme = clr.makeelement(qn("a:schemeClr"), {"val": scheme_token(bullet.color_ref)})
        clr.append(scheme)
        props.append(clr)
    if bullet.font:
        props.append(props.makeelement(qn("a:buFont"), {"typeface": bullet.font}))
    props.append(props.makeelement(qn("a:buChar"), {"char": bullet.char}))


def _apply_numbering(
    text_frame: object, manifest: TemplateManifest, size_pt: float, color: ColorRef | None
) -> None:
    """Нумерованный список: автонумерация PowerPoint (`a:buAutoNum`), вынос — маркера
    шаблона этого уровня, цвет номера — слот темы (`number_ink`). Номер — свойство абзаца,
    а не текст: пункт остаётся тем, что написала композиция, а нумерация — живой."""
    declared = len(manifest.bullet_levels)
    for paragraph in text_frame.paragraphs:  # type: ignore[attr-defined]
        level = int(getattr(paragraph, "level", 0) or 0)
        bullet = manifest.bullet_for(level)
        extra = max(0, level - (declared - 1)) if declared else level
        indent = (bullet.indent_emu if bullet else 0) or -round(size_pt * EMU_PER_PT)
        margin = (bullet.margin_left_emu if bullet else 0) or -indent
        props = _clear_marker(paragraph)
        props.set("marL", str(margin + extra * -indent))  # type: ignore[attr-defined]
        props.set("indent", str(indent))  # type: ignore[attr-defined]
        nodes = []
        if color is not None:
            clr = props.makeelement(qn("a:buClr"), {})  # type: ignore[attr-defined]
            clr.append(clr.makeelement(qn("a:schemeClr"), {"val": scheme_token(color)}))
            nodes.append(clr)
        nodes.append(
            props.makeelement(qn("a:buAutoNum"), {"type": "arabicPeriod"})  # type: ignore[attr-defined]
        )
        _insert_marker(props, nodes)


def _clear_marker(paragraph: object) -> object:
    """Свойства абзаца без знака пункта: прежний маркер, номер и их цвет убираются."""
    props = (
        paragraph._pPr  # type: ignore[attr-defined]
        if paragraph._pPr is not None  # type: ignore[attr-defined]
        else paragraph._p.get_or_add_pPr()  # type: ignore[attr-defined]
    )
    for tag in ("a:buClr", "a:buFont", "a:buChar", "a:buNone", "a:buAutoNum"):
        for node in props.findall(qn(tag)):
            props.remove(node)
    return props


#: Узлы, перед которыми по схеме `a:pPr` стоят знак пункта и его цвет.
_AFTER_MARKER = ("a:tabLst", "a:defRPr", "a:extLst")


def _insert_marker(props: object, nodes: list[object]) -> None:
    """Знак пункта на своё место в `a:pPr`: до `tabLst`, `defRPr` и `extLst` — таков
    порядок схемы, и PowerPoint иначе считает файл повреждённым."""
    anchor = None
    for tag in _AFTER_MARKER:
        anchor = props.find(qn(tag))  # type: ignore[attr-defined]
        if anchor is not None:
            break
    for node in nodes:
        if anchor is not None:
            anchor.addprevious(node)
        else:
            props.append(node)  # type: ignore[attr-defined]


def _label_size_pt(
    manifest: TemplateManifest,
    value_pt: float,
    label_step: object,
    fallback_pt: float,
) -> float:
    """Кегль подписи показателя: по отношению, которое шаблон показал на своём примере.

    Задача B3 просила крупное значение над мелкой подписью, но насколько мелкой —
    вопрос дизайна, а не вкуса. Компонент `kpi` из примеров шаблона отвечает на него
    числом: у VK WorkSpace 66 pt над 20 pt, у VK Education 36 над 16.

    Итог обязательно ложится на ступень шкалы шаблона (правило 6): отношение —
    это пропорция, а не разрешение писать любым кеглем.
    """
    component = manifest.component(ComponentKind.KPI)
    sizes = component.text_sizes_pt if component else []
    own = getattr(label_step, "size_pt", None) or fallback_pt
    if len(sizes) < 2 or sizes[0] <= 0:
        return float(own)
    wanted = value_pt * (sizes[1] / sizes[0])
    ladder = manifest.size_ladder_pt
    if not ladder:
        return float(own)
    return min(ladder, key=lambda step: abs(step - wanted))


def _style_runs(
    paragraph: object, size_pt: float, color: object, font_token: str | None, bold: bool | None
) -> None:
    for run in paragraph.runs:  # type: ignore[attr-defined]
        # Кегль явно: рендер обязан совпасть с расчётом `layout`, а наследованный
        # кегль плейсхолдера может отличаться от кегля роли.
        run.font.size = Pt(size_pt)
        if color is not None:
            apply_theme_color(run.font, color)  # type: ignore[arg-type]
        if font_token is not None:
            run.font.name = font_token
        if bold is not None:
            run.font.bold = bold


def _dataset(content: ContentPackage | None, dataset_ref: str | None) -> Dataset | None:
    if content is None or not dataset_ref:
        return None
    return content.dataset(dataset_ref)


def _coords(box: BBox) -> dict[str, int]:
    return {"x": box.x, "y": box.y, "cx": box.cx, "cy": box.cy}
