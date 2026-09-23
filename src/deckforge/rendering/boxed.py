"""Цитата и callout нативными объектами pptx. Change `draw-by-the-design-system` (DG3).

Полоса — автофигура-прямоугольник с заливкой слотом темы, текст — текстовая рамка:
оба редактируются в PowerPoint, растра нет (C3). Фигуры собраны в группу, чтобы блок
двигался целиком, — как составные компоненты (`rendering/smartart.py`).

Раскладка — `layout.boxed`, кегль — из `fit_report`, акцент и шаг — из дизайн-системы
(`layout.by_design`). Своих чисел, цветов и гарнитур здесь нет (гейт C6).
"""

from __future__ import annotations

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from deckforge.domain.base import BBox
from deckforge.domain.enums import ColorRef, TextRole
from deckforge.domain.slide import FitResult, QuoteBlock
from deckforge.domain.template import TemplateManifest
from deckforge.layout.boxed import BoxedBlock, geometry, paragraphs, style_of
from deckforge.layout.by_design import DesignRules
from deckforge.rendering.theme_binding import apply_theme_color, theme_font_token


def boxed_accent(
    block: BoxedBlock, design: DesignRules, background_hex: str | None
) -> ColorRef:
    """Слот полосы: из IR, иначе акцент по роли дизайн-системы (DG3).

    `accent_ref` у цитаты и callout необязателен (DG4): модель называет цвет, только когда
    ей есть что сказать, а «какой у этого шаблона акцент» знает дизайн-система.
    """
    if block.accent_ref is not None:
        return block.accent_ref
    if isinstance(block, QuoteBlock):
        return design.accent(background_hex)
    return design.callout_accent(block.tone, background_hex)


def add_boxed(
    slide: object,
    block: BoxedBlock,
    fit: FitResult,
    manifest: TemplateManifest,
    design: DesignRules,
    *,
    text_color: ColorRef | None,
    background_hex: str | None,
) -> object:
    """Группа «полоса + текст». `text_color` — цвет свободного текста слайда: цитата
    и callout лежат на фоне слайда, как любой свободный текст."""
    box = block.bbox
    if box is None:
        raise ValueError(f"блок {block.block_id} ({block.type}) без координат")
    style = style_of(block, design)
    parts = geometry(box, style, fit.required_cy_emu)
    step = manifest.typography(style.role) or manifest.typography(TextRole.BODY)
    accent = boxed_accent(block, design, background_hex)
    ink = text_color or (step.color_ref if step else None)
    font_token = theme_font_token(step.font_ref) if step else None

    group = slide.shapes.add_group_shape()  # type: ignore[attr-defined]
    group.name = "Цитата" if isinstance(block, QuoteBlock) else f"Callout {block.tone.value}"
    bar = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, *_emu(parts.bar))
    _drop_style(bar)
    apply_theme_color(bar.fill, accent)
    bar.line.fill.background()

    textbox = group.shapes.add_textbox(*_emu(parts.frame))
    frame = textbox.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    # Рамку дал решатель, текста в ней несколько строк: по середине, вровень с полосой.
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.clear()
    for index, paragraph in enumerate(
        paragraphs(block, style, fit.final_size_pt, bold=bool(step and step.bold))
    ):
        target = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        target.text = paragraph.text
        color = (
            design.accent_ink(
                accent, background_hex, size_pt=paragraph.size_pt, bold=paragraph.bold
            )
            or ink
            if paragraph.label
            else ink
        )
        for run in target.runs:
            run.font.size = Pt(paragraph.size_pt)
            run.font.bold = paragraph.bold
            if color is not None:
                apply_theme_color(run.font, color)
            if font_token is not None:
                run.font.name = font_token
    return group


def _emu(box: BBox) -> tuple[Emu, Emu, Emu, Emu]:
    return Emu(box.x), Emu(box.y), Emu(box.cx), Emu(box.cy)


def _drop_style(shape: object) -> None:
    """Стиль автофигуры тянет из темы тень и обводку — у полосы их быть не должно."""
    element = shape._element  # type: ignore[attr-defined]
    style = element.find(qn("p:style"))
    if style is not None:
        element.remove(style)
