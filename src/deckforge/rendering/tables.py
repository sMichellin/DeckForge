"""Нативные таблицы со стилем из шаблона. Change (14) `native-charts-tables`.

Стиль таблицы — стиль по умолчанию самого шаблона (`def` в `ppt/tableStyles.xml`): так таблица
наследует оформление, а не получает вид, придуманный кодом. Ячейки считает `layout.tabular`,
чтобы записанное совпало с тем, что было вписано.
"""

from __future__ import annotations

from lxml import etree
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from deckforge.domain.slide import TableBlock
from deckforge.rendering.units import size_hundredths

_TABLE_STYLES_REL = "/tableStyles"


def template_table_style(prs: object) -> str | None:
    """GUID стиля таблицы по умолчанию из пакета шаблона; нет части или атрибута — `None`."""
    for rel in prs.part.rels.values():  # type: ignore[attr-defined]
        if rel.reltype.endswith(_TABLE_STYLES_REL) and not rel.is_external:
            root = etree.fromstring(rel.target_part.blob)
            style = root.get("def")
            return str(style) if style else None
    return None


def add_table(
    slide: object,
    block: TableBlock,
    cells: list[list[str]],
    *,
    size_pt: float,
    style_id: str | None,
    row_heights: list[int] | None = None,
    font_token: str | None = None,
    has_header: bool = True,
) -> object:
    """Таблица в рамке блока.

    `row_heights` — высоты из `layout.fitting.table_row_heights`: без них python-pptx делит
    рамку поровну, и строка с переносами вылезает за рамку. `font_token` — шрифт темы, которым
    ячейки были измерены.
    """
    box = block.bbox
    if box is None:
        raise ValueError(f"таблица {block.block_id} без координат")
    rows, columns = len(cells), max(len(row) for row in cells)
    frame = slide.shapes.add_table(  # type: ignore[attr-defined]
        rows, columns, Emu(box.x), Emu(box.y), Emu(box.cx), Emu(box.cy)
    )
    table = frame.table
    table.first_row = block.first_row_header and has_header
    if row_heights is not None:
        for row, height in zip(table.rows, row_heights, strict=True):
            row.height = Emu(height)
        frame.height = Emu(sum(row_heights))
    table.horz_banding = block.banding
    if style_id:
        table._tbl.tblPr.find(qn("a:tableStyleId")).text = style_id

    for r, row in enumerate(cells):
        for c in range(columns):
            text_frame = table.cell(r, c).text_frame
            text_frame.text = row[c] if c < len(row) else ""
            for paragraph in text_frame.paragraphs:
                # Пустая ячейка прогонов не имеет: высоту строки PowerPoint берёт из
                # endParaRPr, и без кегля там она считалась бы от 18 pt по умолчанию.
                paragraph._p.get_or_add_endParaRPr().set("sz", size_hundredths(size_pt))
                for run in paragraph.runs:
                    # Кегль — вписанный `layout`; цвет и жирность шапки — от стиля шаблона.
                    run.font.size = Pt(size_pt)
                    if font_token is not None:
                        run.font.name = font_token
    return frame
