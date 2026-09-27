"""Проверки соответствия шаблону (§5.1). Change (15) `audit-deterministic`.

Проверки этой группы отвечают на вопрос «этот слайд собран по правилам загруженного
шаблона?». Эталон — `TemplateManifest`, собственных представлений о «правильном»
у проверок нет: ни цвета, ни кегля, ни гарнитуры в коде не встречается (гейт C6).

Две проверки здесь работают не так, как звучат, и это следствие ADR-002.

**Цвет.** В `SlideIR` цвет выразим только ссылкой на слот темы (`ColorRef`), литерал
`#RRGGBB` запрещён схемой. Поэтому «цвет не из палитры» на уровне IR сработать не может
в принципе: любая ссылка по построению принадлежит палитре. Проверка ловит то, что
на этом уровне ещё осмысленно — ссылку на слот, которого в теме нет, и две одинаковые
ссылки на соседних сериях диаграммы. Полная проверка цвета делается по готовому файлу.

**Гарнитура.** Имени гарнитуры в IR тоже нет, а гарнитура плейсхолдера принадлежит
шаблону по построению. Чужой шрифт появляется только при записи файла — там проверка
его и ищет, если файл передан. По одному IR остаётся считать **число** гарнитур: это
уже свойство колоды, а не шаблона.
"""

from __future__ import annotations

from collections.abc import Iterable
from contextlib import suppress
from typing import Any

from pptx.oxml.ns import qn

from deckforge.audit.findings import make_finding
from deckforge.audit.geometry import (
    FULL_BLEED_SHARE,
    block_bbox,
    block_text,
    carries_text,
    covers,
    layout_of,
    placeholder_of,
    self_positioned_blocks,
)
from deckforge.audit.recipes import catalogue, recipe_layout_part
from deckforge.audit.registry import CheckContext, CheckUnavailable, check
from deckforge.designsystem.contrast import (
    TextClass,
    comfort_ratio,
    required_ratio,
    text_classes,
)
from deckforge.designsystem.models import Recipe, TypeLevel
from deckforge.domain.audit import Finding
from deckforge.domain.base import BBox
from deckforge.domain.content import ContentPackage
from deckforge.domain.enums import AutoFix, ColorRef, Severity, TextRole
from deckforge.domain.rules import contrast_ratio
from deckforge.domain.slide import Block, BulletsBlock, ChartBlock, KpiBlock, SlideIR, TableBlock
from deckforge.domain.template import LayoutSpec, TemplateManifest
from deckforge.layout.errors import LayoutFitError
from deckforge.layout.tabular import table_cells
from deckforge.rendering.theme_binding import font_family_for_token

#: Доля своей площади, начиная с которой фигура макета считается лежащей в области
#: контента, — то есть подложкой под содержание, а не знаком в полях. Величина
#: безразмерная и от шаблона не зависит: это про приём, а не про бренд.
IN_CONTENT_SHARE = 0.5


def template_fonts(manifest: TemplateManifest) -> set[str]:
    """Гарнитуры, которые шаблон действительно использует.

    Это не только пара из темы. На шаблонах кейса тема заявляет Arial, а плейсхолдеры
    набраны Play: сверка с одной темой дала бы находку на каждом слайде — то есть
    проверку, которая врёт на эталоне и обесценивает весь отчёт.
    """
    fonts = {
        family
        for family in (manifest.theme.fonts.major_latin, manifest.theme.fonts.minor_latin)
        if family
    }
    fonts |= {
        placeholder.font_family
        for layout in manifest.layouts
        for placeholder in layout.placeholders
        if placeholder.font_family
    }
    # Слайды-примеры — третий и самый честный источник: там видно, чем шаблон набран
    # на самом деле. У VK Tech Play стоит на 666 фигурах примеров, а тема зовёт Arial
    # (change `design-system-from-examples`, DS2).
    fonts |= {usage.family for usage in manifest.usage.fonts if usage.family}
    return fonts


def _fonts_in_file(path: object) -> dict[str, list[str]]:
    """Гарнитуры, которыми набран готовый файл: `slide_id-номер` → список гарнитур.

    По `SlideIR` эту проверку сделать нельзя: имени гарнитуры в IR нет вовсе, а кегль
    и шрифт приезжают из плейсхолдера шаблона — то есть по построению «из шаблона».
    Чужая гарнитура может появиться только при записи файла, там её и надо искать.
    """
    try:
        from pptx import Presentation

        presentation = Presentation(str(path))
    except Exception:
        return {}

    found: dict[str, list[str]] = {}
    for number, slide in enumerate(presentation.slides, start=1):
        families: list[str] = []
        for shape in slide.shapes:
            frame = getattr(shape, "text_frame", None)
            if frame is None:
                continue
            for paragraph in frame.paragraphs:
                for run in paragraph.runs:
                    name = run.font.name
                    if name:
                        families.append(name)
        found[str(number)] = families
    return found


@check(id="template.font_not_in_theme", deterministic=True, severity=Severity.ERROR,
       title="Шрифт не из шаблона или гарнитур больше двух")
def font_not_in_theme(ctx: CheckContext) -> Iterable[Finding]:
    """Шрифт не из шаблона или гарнитур больше двух."""
    manifest = ctx.manifest
    known = template_fonts(manifest)
    max_families = int(ctx.param("max_families", 2))
    used: set[str] = set()

    # Гарнитуры плейсхолдеров, в которые лягут блоки: они принадлежат шаблону
    # по построению, но их **количество** — уже свойство колоды.
    for slide in ctx.deck.slides:
        layout = layout_of(slide, manifest)
        for block in slide.blocks:
            placeholder = placeholder_of(block, layout)
            family = placeholder.font_family if placeholder is not None else None
            if family is None:
                role = getattr(block, "role", None)
                step = manifest.typography(role) if isinstance(role, TextRole) else None
                family = manifest.theme.fonts.get(step.font_ref) if step is not None else None
            if family:
                used.add(family)

    slide_ids = [slide.slide_id for slide in ctx.deck.slides]
    if ctx.deck_path is not None:
        for page, families in sorted(_fonts_in_file(ctx.deck_path).items()):
            index = int(page) - 1
            slide_id = slide_ids[index] if 0 <= index < len(slide_ids) else None
            for written in sorted(set(families)):
                # В файле `typeface` бывает ссылкой на тему («+mn-lt»), а не названием
                # гарнитуры. Без разрешения ссылки проверка выдаёт ошибку на каждом
                # прогоне и заодно считает ссылку и сам шрифт двумя разными гарнитурами.
                family = font_family_for_token(written, manifest) or written
                used.add(family)
                if family in known:
                    continue
                yield make_finding(
                    check_id="template.font_not_in_theme",
                    slide_id=slide_id,
                    reason=f"font:{family}",
                    message=(
                        f"Гарнитура «{family}» в шаблоне не встречается: "
                        + ", ".join(sorted(known))
                    ),
                    evidence={"family": family, "known": ", ".join(sorted(known))},
                )

    if len(used) > max_families:
        yield make_finding(
            check_id="template.font_not_in_theme",
            reason="too_many_families",
            message=(
                f"В колоде {len(used)} гарнитуры при допустимых {max_families}: "
                + ", ".join(sorted(used))
            ),
            evidence={"families": ", ".join(sorted(used))},
        )


def template_sizes(manifest: TemplateManifest) -> list[float]:
    """Кегли шаблона: шкала плюс фактические кегли плейсхолдеров.

    Шкала — четыре ступени по ролям, а плейсхолдеры живут своими размерами: заголовок
    раздела и заголовок слайда набраны по-разному, хотя роль у них одна. Сверять только
    со шкалой значит ругаться на кегль, который шаблон сам и задал.
    """
    sizes = set(manifest.size_ladder_pt)
    sizes |= {
        placeholder.size_pt
        for layout in manifest.layouts
        for placeholder in layout.placeholders
        if placeholder.size_pt
    }
    return sorted(sizes, reverse=True)


@check(id="template.size_not_in_scale", deterministic=True, severity=Severity.WARNING,
       title="Кегль не из типографической шкалы шаблона")
def size_not_in_scale(ctx: CheckContext) -> Iterable[Finding]:
    """Кегль не из типографической шкалы шаблона."""
    allowed = template_sizes(ctx.manifest)
    if not allowed:
        return
    for slide in ctx.deck.slides:
        layout = layout_of(slide, ctx.manifest)
        for block in slide.blocks:
            size = getattr(block, "size_pt", None)
            if size is None:
                # Кегль мог быть выбран вписыванием (change 12) — тогда он в fit_report.
                measured = slide.fit_report.get(block.block_id)
                size = measured.final_size_pt if measured is not None else None
            if size is None:
                continue
            if any(abs(size - step) < 0.01 for step in allowed):
                continue
            yield make_finding(
                check_id="template.size_not_in_scale",
                slide_id=slide.slide_id,
                block_id=block.block_id,
                bbox=block_bbox(block, layout),
                reason=f"size:{size:g}",
                message=(
                    f"Кегль {size:g} pt в шаблоне не встречается: "
                    + ", ".join(f"{step:g}" for step in allowed[:8])
                ),
                evidence={"size_pt": f"{size:g}"},
            )


@check(id="template.color_not_in_palette", deterministic=True, severity=Severity.ERROR,
       auto_fix=AutoFix.MAP_TO_NEAREST_THEME_COLOR, title="Цвет не из палитры шаблона")
def color_not_in_palette(ctx: CheckContext) -> Iterable[Finding]:
    """Цвет не из палитры шаблона."""
    colors = ctx.manifest.theme.colors
    known = set(colors.model_dump())

    for slide in ctx.deck.slides:
        layout = layout_of(slide, ctx.manifest)
        for block in slide.blocks:
            refs: list[ColorRef] = []
            if isinstance(block, ChartBlock):
                refs = list(block.series_color_refs)
            elif isinstance(block, KpiBlock):
                refs = [item.color_ref for item in block.items if item.color_ref is not None]
            else:
                own = getattr(block, "color_ref", None)
                refs = [own] if isinstance(own, ColorRef) else []

            for ref in refs:
                if ref.value not in known:
                    yield make_finding(
                        check_id="template.color_not_in_palette",
                        slide_id=slide.slide_id,
                        block_id=block.block_id,
                        bbox=block_bbox(block, layout),
                        reason=f"unknown:{ref.value}",
                        message=f"Слота {ref.value} нет в палитре шаблона",
                        evidence={"color_ref": ref.value},
                    )

            # Две серии одного цвета на диаграмме неразличимы — это тот же дефект
            # палитры, только проявленный внутри одного блока.
            if isinstance(block, ChartBlock) and len(set(refs)) != len(refs):
                duplicates = sorted({ref.value for ref in refs if refs.count(ref) > 1})
                yield make_finding(
                    check_id="template.color_not_in_palette",
                    slide_id=slide.slide_id,
                    block_id=block.block_id,
                    bbox=block_bbox(block, layout),
                    reason="duplicate_series",
                    message=(
                        "Серии диаграммы окрашены одним цветом и неразличимы: "
                        + ", ".join(duplicates)
                    ),
                    evidence={"duplicates": ", ".join(duplicates)},
                )


def _layout_parts_in_file(ctx: CheckContext) -> list[str] | None:
    """Имена частей макетов, на которых собраны слайды готового файла.

    Адресоваться к макету по индексу нельзя: `LayoutSpec.index` сквозной по мастерам,
    а макеты без пригодных плейсхолдеров пропущены. На шаблонах кейса `slide_layouts[index]`
    промахивается в 34 случаях из 37 — поэтому сверка идёт по `part_name`.
    """
    path = ctx.deck_path
    if path is None:
        return None
    try:
        from pptx import Presentation

        presentation = Presentation(str(path))
    except Exception:
        return None
    return [str(slide.slide_layout.part.partname).lstrip("/") for slide in presentation.slides]


@check(id="template.layout_not_from_template", deterministic=True, severity=Severity.ERROR,
       title="Слайд собран не на макете из шаблона")
def layout_not_from_template(ctx: CheckContext) -> Iterable[Finding]:
    """Слайд собран не на макете из шаблона.

    Слайд по рецепту (`SlideIR.by_recipe`) — копия слайда-примера, и лежит он на макете
    примера, а не на том, что выбрал план. Его часть макета сверяется с макетом примера
    рецепта (RG27): совпала — молчит, не совпала — писатель склонировал не тот пример.
    """
    parts_in_file = _layout_parts_in_file(ctx)
    if parts_in_file is not None:
        recipes = catalogue(ctx)
        for number, slide in enumerate(ctx.deck.slides):
            if number >= len(parts_in_file):
                break
            if slide.by_recipe:
                yield from _recipe_layout_mismatch(ctx, slide, parts_in_file[number], recipes)
                continue
            declared = layout_of(slide, ctx.manifest)
            actual = parts_in_file[number]
            if declared is None or declared.part_name.lstrip("/") == actual:
                continue
            yield make_finding(
                check_id="template.layout_not_from_template",
                slide_id=slide.slide_id,
                reason=f"part:{actual}",
                message=(
                    f"Слайд собран на макете {actual}, а план объявлял "
                    f"{declared.part_name} ({declared.layout_id})"
                ),
                evidence={"expected": declared.part_name, "actual": actual},
            )

    for slide in ctx.deck.slides:
        layout = layout_of(slide, ctx.manifest)
        if layout is None:
            yield make_finding(
                check_id="template.layout_not_from_template",
                slide_id=slide.slide_id,
                reason=f"layout:{slide.layout_id}",
                message=(
                    f"Макета {slide.layout_id} нет в шаблоне: доступны "
                    + ", ".join(item.layout_id for item in ctx.manifest.layouts)
                ),
                evidence={"layout_id": slide.layout_id},
            )
            continue

        for block in slide.blocks:
            idx = getattr(block, "placeholder_idx", None)
            if idx is None or layout.placeholder(idx) is not None:
                continue
            yield make_finding(
                check_id="template.layout_not_from_template",
                slide_id=slide.slide_id,
                block_id=block.block_id,
                reason=f"placeholder:{idx}",
                message=(
                    f"Блок {block.block_id} ссылается на плейсхолдер {idx}, "
                    f"которого нет в макете {layout.layout_id}"
                ),
                evidence={"layout_id": layout.layout_id, "placeholder_idx": str(idx)},
            )


def _recipe_layout_mismatch(
    ctx: CheckContext, slide: SlideIR, actual: str, recipes: dict[str, Recipe]
) -> Iterable[Finding]:
    """Слайд по рецепту лежит не на макете своего примера.

    Рецепта нет в каталоге — это вопрос `template.recipe_not_in_catalogue`; макет примера
    неизвестен — сверять не с чем, а гадать хуже, чем промолчать.
    """
    recipe = recipes.get(slide.recipe_id or "")
    expected = recipe_layout_part(recipe, ctx.manifest) if recipe is not None else None
    if recipe is None or expected is None or expected == actual:
        return
    yield make_finding(
        check_id="template.layout_not_from_template",
        slide_id=slide.slide_id,
        reason=f"recipe:{recipe.recipe_id}:{actual}",
        message=(
            f"Слайд собран по рецепту {recipe.recipe_id} на макете {actual}, "
            f"а пример рецепта лежит на {expected}"
        ),
        evidence={"recipe_id": recipe.recipe_id, "expected": expected, "actual": actual},
    )


@check(id="template.decor_moved", deterministic=True, severity=Severity.WARNING,
       title="Логотип или колонтитул сдвинуты с положенного места")
def decor_moved(ctx: CheckContext) -> Iterable[Finding]:
    """Логотип или колонтитул сдвинуты с положенного места.

    Проверка отвечает на два вопроса по отдельности, и прежде путала их (C12:
    19 предупреждений из 35 на каждом прогоне VK WorkSpace, ни одного про логотип).

    **Что защищаем — знак, а не подложку.** `LayoutSpec.shapes` по своему описанию
    «фон, фотографии, декор»: на трёх шаблонах кейса 13 из 34, 36 из 76 и 28 из 28
    таких фигур лежат внутри области контента и занимают до 76 % слайда. Это подложка
    под содержание — решатель ставит текст ровно туда, и иначе на VK WorkSpace ставить
    его будет некуда. Знак шаблон держит **в полях**, поэтому защищается фигура,
    которая в область контента заходит меньше чем наполовину. Логотип защищён всегда:
    парсер отличает его от прочих картинок мастера по построению (мелкий, у края).

    **Что значит «накрыл» — скрыл, а не задел краем.** Порог — доля площади знака,
    как `min_overlap_ratio` у `layout.overlap`. Прежний `tolerance_emu` был длиной
    и сравнивался с площадью в EMU², то есть с нулём.
    """
    min_ratio = ctx.param("min_cover_ratio", 0.5)
    slide_box = ctx.manifest.slide_size.bbox
    content = ctx.manifest.content_bbox
    logo = ctx.manifest.decor.logo

    for slide in ctx.deck.slides:
        layout = layout_of(slide, ctx.manifest)

        decor_boxes: list[tuple[str, BBox]] = []
        if logo is not None:
            decor_boxes.append(("логотип шаблона", logo.bbox))
        if layout is not None:
            decor_boxes += [
                (f"декор макета {shape.shape_id}", shape.bbox)
                for shape in layout.shapes
                # Фон во весь слайд лежит под контентом по замыслу, а не по ошибке.
                if not covers(shape.bbox, slide_box, FULL_BLEED_SHARE)
                and content.intersection_area(shape.bbox) < IN_CONTENT_SHARE * shape.bbox.area
            ]
        if not decor_boxes:
            continue

        # Плейсхолдер, наложенный на декор, — так нарисован сам шаблон. Наша вина
        # начинается там, где мы поставили блок своими координатами.
        for block, bbox in self_positioned_blocks(slide, ctx.manifest):
            # Сдвинуть декор из IR нельзя — он приезжает с мастера (ADR-002).
            # Испортить его можно единственным способом: накрыть своим блоком.
            for label, decor_box in decor_boxes:
                if decor_box.area <= 0:
                    continue
                ratio = bbox.intersection_area(decor_box) / decor_box.area
                if ratio < min_ratio:
                    continue
                yield make_finding(
                    check_id="template.decor_moved",
                    slide_id=slide.slide_id,
                    block_id=block.block_id,
                    bbox=bbox,
                    reason=f"covered:{label}",
                    message=f"Блок {block.block_id} накрывает {label} на {ratio:.0%} площади",
                    evidence={"ratio": f"{ratio:.3f}", "decor": label},
                )


#: Какой параметр проверки задаёт минимум какого класса текста. Класс выбирает слой
#: `designsystem.contrast`, значение — `configs/audit_checks.yaml`. У подписи минимум
#: тот же, что у обычного текста: отличается она только запасом.
CONTRAST_PARAM: dict[TextClass, str] = {
    TextClass.BODY: "body",
    TextClass.CAPTION: "body",
    TextClass.LARGE: "large",
    TextClass.GRAPHICS: "graphics",
}


def contrast_thresholds(ctx: CheckContext) -> dict[TextClass, float]:
    """Минимум контраста по классу текста: значение из конфига, по умолчанию — слоя ДС.

    Значения по умолчанию берутся у `designsystem.contrast.required_ratio`, а не пишутся
    здесь второй раз: если параметр пропал из YAML, проверка съезжает на правило слоя,
    а не на собственное число.
    """
    return {kind: ctx.param(CONTRAST_PARAM[kind], required_ratio(kind)) for kind in TextClass}


def block_text_classes(block: object, manifest: TemplateManifest) -> tuple[TextClass, TextClass]:
    """Два ответа слоя ДС о блоке: чей минимум он обязан взять и чей запас ему нужен.

    Само правило — «минимум по кеглю, запас по роли» — живёт в слое:
    `designsystem.contrast.text_classes`. Им же считает страница дизайн-системы, иначе
    она показала бы подписи 18 pt минимум 4,5, а аудит спрашивал бы 3,0
    (change `contrast-minimum-by-size`). Здесь только то, чего слой знать не может:
    кегль и начертание блока.

    Кегль — свой у блока, иначе ступени шкалы по роли; начертание — всегда ступени:
    в IR полужирного нет, его задаёт шаблон.
    """
    role = getattr(block, "role", None)
    role = role if isinstance(role, TextRole) else None
    step = manifest.typography(role) if role is not None else None
    size = getattr(block, "size_pt", None)
    if size is None and step is not None:
        size = step.size_pt
    bold = bool(step is not None and step.bold)
    return text_classes(size, bold=bold, role=role)


@check(id="template.contrast_below_wcag", deterministic=True, severity=Severity.ERROR,
       title="Контраст текста к фону ниже порога своего класса")
def contrast_below_wcag(ctx: CheckContext) -> Iterable[Finding]:
    """Контраст текста к фону ниже порога своего класса.

    **Порог — по классу текста, правилом `designsystem.contrast`** (change
    `one-contrast-rule`). Один свод правил на страницу дизайн-системы, вёрстку и аудит:
    крупный текст (от 18 pt или от 14 pt полужирным) — 3,0, обычный — 4,5. Подпись
    берёт минимум своего кегля, а сверх него у неё есть комфортный порог слоя (7,0).
    Подпись между минимумом и комфортом — находка `info`, а не ошибка: минимум взят,
    не хватает запаса. Иначе каждый шаблон с серой подписью давал бы ошибку на каждом
    слайде, как `font_not_in_theme` до C3.

    Фон берётся **макета**, а не темы. Разница не теоретическая: 19.09 колода из
    двенадцати заголовков цвета `dk1` на макете, залитом `dk1`, прошла аудит без единой
    находки — контраст 1:1 сравнивался со светлым слотом темы и выходил 21:1.

    Проверяются блоки, у которых цвет задан **явно**: это наш выбор, и отвечаем за него
    мы. Блок без `color_ref` наследует цвет плейсхолдера шаблона, и спрашивать за него
    с генератора нельзя — шаблон не нарушает сам себя (тот же довод, что у
    `self_positioned_blocks`). Свободный текст без `color_ref` цвет получает при записи
    тем же правилом слоя (`contrast.readable_ref`, класс обычного текста), поэтому
    неразличимым он становится только на шаблоне, где порог не берёт ни один слот.
    """
    manifest = ctx.manifest
    colors = manifest.theme.colors
    thresholds = contrast_thresholds(ctx)

    for slide in ctx.deck.slides:
        layout = layout_of(slide, ctx.manifest)
        background, background_source = _background_of(layout, manifest)
        for block in slide.blocks:
            if not carries_text(block):
                continue
            ref = getattr(block, "color_ref", None)
            if not isinstance(ref, ColorRef):
                continue
            kind, role_kind = block_text_classes(block, manifest)
            required = thresholds[kind]
            # Запас слой даёт не каждому классу: у кого его нет, у того комфорт — это
            # минимум из конфига, и промежутка для `info` не остаётся.
            has_margin = comfort_ratio(role_kind) > required_ratio(role_kind)
            comfort = max(required, comfort_ratio(role_kind)) if has_margin else required

            foreground = colors.get(ref)
            ratio = contrast_ratio(foreground, background)
            if ratio >= comfort:
                continue
            evidence = {
                "foreground": foreground,
                "background": background,
                "background_source": background_source,
                "ratio": f"{ratio:.2f}",
                "text_class": role_kind.value,
                "minimum_class": kind.value,
                "required": f"{required:.1f}",
                "comfort": f"{comfort:.1f}",
            }
            if ratio >= required:
                # Минимум взят, запаса нет: замечание, а не нарушение.
                yield make_finding(
                    check_id="template.contrast_below_wcag",
                    slide_id=slide.slide_id,
                    block_id=block.block_id,
                    bbox=block_bbox(block, layout),
                    reason=f"contrast-tight:{ref.value}",
                    message=(
                        f"Контраст подписи блока {block.block_id} к фону {ratio:.1f}:1 — "
                        f"минимум {required:.1f}:1 взят, но без запаса {comfort:.1f}:1"
                    ),
                    evidence=evidence,
                    severity=Severity.INFO,
                )
                continue
            yield make_finding(
                check_id="template.contrast_below_wcag",
                slide_id=slide.slide_id,
                block_id=block.block_id,
                bbox=block_bbox(block, layout),
                reason=f"contrast:{ref.value}",
                message=(
                    f"Контраст текста блока {block.block_id} к фону {ratio:.1f}:1 "
                    f"при требуемых {required:.1f}:1"
                ),
                evidence=evidence,
            )


def _background_of(layout: LayoutSpec | None, manifest: TemplateManifest) -> tuple[str, str]:
    """Фон, по которому читается текст слайда, и честное объяснение, откуда он взят.

    Манифест, снятый парсером до change (24), фона не описывает. Подставлять вместо него
    светлый слот темы молча нельзя — именно это и дало ложное «нарушений нет», — поэтому
    источник едет в `evidence` каждой находки.
    """
    background = layout.background if layout is not None else None
    if background is None:
        return (
            manifest.theme.colors.get(ColorRef.LT1),
            "тема lt1: манифест снят парсером без разбора фона",
        )
    if background.is_image:
        return (
            background.color_hex,
            f"{background.source}: усреднённый цвет подложки, вердикт приблизителен",
        )
    return background.color_hex, background.source


#: Сколько знаков чужой фразы показать в находке: достаточно, чтобы узнать её глазами.
SAMPLE_SNIPPET_CHARS = 80


def _block_pieces(block: Block, content: ContentPackage | None) -> list[str]:
    """Куски текста блока так, как writer раскладывает их по абзацам и ячейкам.

    Показатель — значение и подпись отдельными абзацами; таблица — по ячейке, числа
    датасета — в том виде, в каком их пишет writer (`table_cells`, «2,5»).
    """
    pieces = [block_text(block)]
    if isinstance(block, BulletsBlock):
        pieces += [item.text for item in block.items]
    elif isinstance(block, KpiBlock):
        pieces += [part for item in block.items for part in (item.value, item.label)]
    elif isinstance(block, TableBlock):
        dataset = content.dataset(block.dataset_ref) if content and block.dataset_ref else None
        # Таблицу нечем заполнить — writer её и не запишет.
        with suppress(LayoutFitError):
            pieces += [cell for row in table_cells(block, dataset) for cell in row]
    return pieces


def _our_lines(ctx: CheckContext) -> set[str]:
    """Всё, что колода написала сама, построчно и без краевых пробелов.

    Набор общий на колоду, а не на слайд, и от порядка слайдов не зависит только он.
    Страницу выбирает `sample_text_left` — по номеру слайда в IR, и в смешанной колоде
    она может оказаться чужой. Пока writer (поток B) не исправил порядок, на чужой
    странице возможна ложная находка: её текст writer пишет не только из блоков,
    разобранных в `_block_pieces` (SmartArt, callout, деградации диаграмм и таблиц).
    """
    lines: set[str] = set()
    for slide in ctx.deck.slides:
        for block in slide.blocks:
            for piece in _block_pieces(block, ctx.content):
                lines.add(piece.strip())
                lines.update(line.strip() for line in piece.splitlines())
    lines.discard("")
    return lines


def _written_paragraphs(shape: Any) -> list[str]:
    """Строки фигуры, набранные прогонами `a:r`.

    Мягкий перенос `a:br` делит абзац на строки — так же, как перевод строки делит
    наш текст в `_our_lines`. Поля `a:fld` (номер слайда, дата) не в счёт: их текст
    подставляет PowerPoint, а не автор примера.
    """
    lines: list[str] = []
    for paragraph in shape.iter(qn("a:p")):
        text = "".join(
            "\n" if node.tag == qn("a:br") else node.findtext(qn("a:t")) or ""
            for node in paragraph
            if node.tag in (qn("a:r"), qn("a:br"))
        )
        lines.extend(line.strip() for line in text.splitlines() if line.strip())
    return lines


#: Типы плейсхолдеров колонтитулов (`p:ph/@type`): их текст — служебный, а не текст
#: примера, даже набранный вручную («Конфиденциально»).
FOOTER_PLACEHOLDERS = frozenset({"ftr", "dt", "sldNum", "hdr"})


def _text_shapes(tree: Any) -> Iterable[tuple[int | None, str, Any]]:
    """Фигуры слайда с текстом на любой глубине групп, кроме колонтитулов:
    `cNvPr id`, имя, узел."""
    for shape in tree.iter(qn("p:sp"), qn("p:graphicFrame")):
        placeholder = shape.find(f"./*/{qn('p:nvPr')}/{qn('p:ph')}")
        if placeholder is not None and placeholder.get("type") in FOOTER_PLACEHOLDERS:
            continue
        props = shape.find(f"./*/{qn('p:cNvPr')}")
        raw = props.get("id") if props is not None else None
        name = props.get("name", "") if props is not None else ""
        yield (int(raw) if raw and raw.isdigit() else None), name, shape


@check(id="template.sample_text_left", deterministic=True, severity=Severity.ERROR,
       title="На слайде по рецепту остался текст слайда-примера шаблона")
def sample_text_left(ctx: CheckContext) -> Iterable[Finding]:
    """На слайде по рецепту остался текст слайда-примера шаблона.

    Слайд по рецепту — копия примера целиком, с его текстом (change
    `recipe-slide-in-the-writer`). Writer заменяет текст зон нашим, а лишние зоны
    удаляет или стирает. Всё, что после этого осталось на слайде и чего колода
    не писала, — фраза автора шаблона: «Имя Фамилия», «Описание преимущества».

    Исходного текста примера в манифесте нет сознательно, и он не нужен: на копии
    примера других источников текста, кроме примера и колоды, не бывает. Поэтому
    эталон — строки `SlideIR` колоды, а не файл шаблона.

    Страница файла для слайда по рецепту выбирается по номеру слайда в IR, как
    у соседних проверок по готовому файлу. Writer кладёт слайды по рецепту в начало
    файла (`rendering/writer.py`, `write()`: копии рецептов до остальных), поэтому
    в смешанной колоде страница по номеру может оказаться чужой: проверка тогда
    пропустит нарушителя, назовёт не тот `slide_id` или даст ложную находку на тексте
    чужого слайда. Чинит порядок writer (поток B), а не аудит.
    """
    path = ctx.deck_path
    if path is None:
        raise CheckUnavailable("файла колоды ещё нет: текст примера виден только в .pptx")
    if not ctx.manifest.examples:
        raise CheckUnavailable("в шаблоне нет слайдов-примеров: рецептов не бывает")
    recipe_slides = [
        (number, slide) for number, slide in enumerate(ctx.deck.slides) if slide.recipe_id
    ]
    if not recipe_slides:
        raise CheckUnavailable("ни одного слайда по рецепту: текста примера взяться неоткуда")
    try:
        from pptx import Presentation

        pages = list(Presentation(str(path)).slides)
    except Exception as error:
        raise CheckUnavailable(f"файл колоды не открылся: {type(error).__name__}") from error
    if len(pages) < len(ctx.deck.slides):
        raise CheckUnavailable(
            f"в файле {len(pages)} слайдов, а в IR {len(ctx.deck.slides)}: "
            "страницу слайда по рецепту не найти"
        )

    ours = _our_lines(ctx)
    for number, slide in recipe_slides:
        for xml_id, name, shape in _text_shapes(pages[number].shapes._spTree):
            foreign = [text for text in _written_paragraphs(shape) if text not in ours]
            if not foreign:
                continue
            text = " / ".join(foreign)
            yield make_finding(
                check_id="template.sample_text_left",
                slide_id=slide.slide_id,
                reason=f"shape:{xml_id}",
                message=(
                    f"На слайде {slide.slide_id} (рецепт {slide.recipe_id}) в фигуре "
                    f"«{name}» остался текст примера: «{text[:SAMPLE_SNIPPET_CHARS]}»"
                ),
                evidence={
                    "recipe_id": slide.recipe_id or "",
                    "xml_id": str(xml_id) if xml_id is not None else "нет",
                    "shape": name,
                    "text": text[:SAMPLE_SNIPPET_CHARS],
                },
            )


def _catalogue(ctx: CheckContext) -> set[str]:
    """Идентификаторы рецептов шаблона — из дизайн-системы контекста, а нет её — тем же
    `derive`, которым каталог строит `parse` (`audit.recipes.catalogue`, RG27).
    Считается один раз на вызов проверки.
    """
    return set(catalogue(ctx))


@check(id="template.recipe_not_in_catalogue", deterministic=True, severity=Severity.ERROR,
       title="Слайд назван рецептом, которого нет в каталоге шаблона")
def recipe_not_in_catalogue(ctx: CheckContext) -> Iterable[Finding]:
    """Слайд назван рецептом, которого нет в каталоге композиций шаблона.

    Такой `recipe_id` — выдумка модели, а не решение каталога: 24.09 она прошла весь
    конвейер и упала в писателе `KeyError` (change `recipe-is-not-the-models-word`).
    Слайд без `recipe_id` здесь не в счёт — это вопрос `template.slide_without_recipe`.
    """
    catalogue = _catalogue(ctx)
    for slide in ctx.deck.slides:
        if not slide.recipe_id or slide.recipe_id in catalogue:
            continue
        yield make_finding(
            check_id="template.recipe_not_in_catalogue",
            slide_id=slide.slide_id,
            reason=f"recipe:{slide.recipe_id}",
            message=(
                f"Слайд {slide.slide_id} назван рецептом {slide.recipe_id}, а в каталоге "
                f"шаблона его нет (рецептов в каталоге: {len(catalogue)})"
            ),
            evidence={"recipe_id": slide.recipe_id, "catalogue_size": str(len(catalogue))},
        )


@check(id="template.slide_without_recipe", deterministic=True, severity=Severity.INFO,
       title="Слайд собран не по рецепту, хотя у шаблона есть каталог композиций")
def slide_without_recipe(ctx: CheckContext) -> Iterable[Finding]:
    """Слайд собран не по рецепту, хотя у шаблона есть каталог композиций.

    Оговорка, а не ошибка: слайд по макету корректен, но идёт мимо дизайн-системы
    шаблона. «По рецепту» — единый предикат `SlideIR.by_recipe`: смешанный слайд
    (часть блоков вне зон) писатель тоже собирает не по рецепту. Каталог пуст —
    сравнивать не с чем, это пропуск, а не «прошла».
    """
    catalogue = _catalogue(ctx)
    if not catalogue:
        raise CheckUnavailable("у шаблона нет каталога композиций: рецептов не бывает")
    for slide in ctx.deck.slides:
        if slide.by_recipe:
            continue
        yield make_finding(
            check_id="template.slide_without_recipe",
            slide_id=slide.slide_id,
            reason="not_by_recipe",
            message=(
                f"Слайд {slide.slide_id} собран по макету {slide.layout_id}, а не по рецепту: "
                f"в каталоге шаблона {len(catalogue)} рецептов"
            ),
            evidence={"recipe_id": slide.recipe_id or "", "layout_id": slide.layout_id},
        )


_Frame = tuple[float, float, float, float]

_GROUP = qn("p:grpSp")
_CONNECTOR = qn("p:cxnSp")
_SHAPE = qn("p:sp")
_PLACED_TAGS = (_SHAPE, qn("p:pic"), _CONNECTOR, qn("p:graphicFrame"))


def _xfrm_pair(xfrm: Any, off: str, ext: str) -> _Frame | None:
    start = xfrm.find(qn(off)) if xfrm is not None else None
    size = xfrm.find(qn(ext)) if xfrm is not None else None
    if start is None or size is None:
        return None
    return (float(start.get("x", 0)), float(start.get("y", 0)),
            float(size.get("cx", 0)), float(size.get("cy", 0)))


def _node_xfrm(node: Any) -> Any:
    if node.tag == _GROUP:
        return node.find(f"{qn('p:grpSpPr')}/{qn('a:xfrm')}")
    found = node.find(f"{qn('p:spPr')}/{qn('a:xfrm')}")
    return found if found is not None else node.find(qn("p:xfrm"))


def _shapes_on_page(
    parent: Any, scale: _Frame = (0, 0, 1, 1), origin: tuple[float, float] = (0, 0)
) -> Iterable[tuple[Any, _Frame | None]]:
    """Фигуры страницы на любой глубине групп с рамкой в координатах слайда.

    Свой обход, а не писателя: проверка меряет файл, а не решение писателя. Линии разбор
    примеров не хранит (`cx = 0` он отбрасывает) — их геометрия есть только в файле.
    """
    ox, oy, sx, sy = scale
    for node in parent:
        if node.tag not in (_GROUP, *_PLACED_TAGS):
            continue
        xfrm = _node_xfrm(node)
        own = _xfrm_pair(xfrm, "a:off", "a:ext")
        frame = None if own is None else (
            ox + (own[0] - origin[0]) * sx, oy + (own[1] - origin[1]) * sy,
            own[2] * sx, own[3] * sy,
        )
        if node.tag != _GROUP:
            yield node, frame
            continue
        child = _xfrm_pair(xfrm, "a:chOff", "a:chExt")
        if frame is None or child is None or not child[2] or not child[3]:
            yield from _shapes_on_page(node, scale, origin)
        else:
            yield from _shapes_on_page(
                node, (frame[0], frame[1], frame[2] / child[2], frame[3] / child[3]),
                (child[0], child[1]),
            )


def _apart(a: _Frame, b: _Frame) -> float:
    """Зазор между рамками; 0 — касаются или пересекаются."""
    return max(b[0] - (a[0] + a[2]), a[0] - (b[0] + b[2]),
               b[1] - (a[1] + a[3]), a[1] - (b[1] + b[3]), 0)


def _line_ends(node: Any, frame: _Frame) -> tuple[_Frame, _Frame]:
    """Концы линии — углы рамки с учётом `flipH`/`flipV`."""
    xfrm = _node_xfrm(node)
    x0, y0, x1, y1 = frame[0], frame[1], frame[0] + frame[2], frame[1] + frame[3]
    if xfrm is not None and xfrm.get("flipH") == "1":
        x0, x1 = x1, x0
    if xfrm is not None and xfrm.get("flipV") == "1":
        y0, y1 = y1, y0
    return (x0, y0, 0, 0), (x1, y1, 0, 0)


def _has_text(node: Any) -> bool:
    return any((text.text or "").strip() for text in node.iter(qn("a:t")))


def _shape_id(node: Any) -> str:
    props = node.find(f"./*/{qn('p:cNvPr')}")
    return props.get("id", "") if props is not None else ""


def _spans_the_slide(frame: _Frame, size: tuple[float, float]) -> bool:
    """Подложка или полоса во всю ширину или высоту — оформление слайда, не декор зоны."""
    return frame[2] >= FULL_BLEED_SHARE * size[0] or frame[3] >= FULL_BLEED_SHARE * size[1]


@check(id="template.decor_leads_nowhere", deterministic=True, severity=Severity.WARNING,
       title="Стрелка или рамка примера ведёт к месту без текста")
def decor_leads_nowhere(ctx: CheckContext) -> Iterable[Finding]:
    """Стрелка или плашка примера ведёт к месту, где в файле нет текста (RG52).

    Education s06, рецепт `ex013`: зоны схемы нашим текстом не заполнились, писатель их снял,
    а стрелки и плашка остались — чертёж без подписей при нуле ошибок аудита. Проверка меряет
    **готовый файл**, а не решение писателя (его правило она не импортирует — иначе замер
    подтверждал бы сам себя): геометрия зон — из каталога, фигур — со страницы, с группами.

    Пустое место — рамка зоны рецепта, у которой в файле нет фигуры или в фигуре нет текста.
    Находка — линия (`p:cxnSp` или фигура, у которой меньшая сторона не больше `line_share`
    большей), конец которой в пределах `touch_share` ширины слайда от пустого места и не у фигуры
    с текстом; и плашка `p:sp` без текста, которая касается пустого места и ни одной фигуры
    с текстом. Не декор примера: плейсхолдеры, фигуры, которые адресует рецепт, подложки во всю
    сторону слайда (`FULL_BLEED_SHARE`); колонтитулы — не текст. Картинка рецепта — содержимое.
    Заголовок — не пустое место: его писатель не снимает никогда.

    Граница: только слайды по рецепту и только рамки зон рецепта. На слайде по макету декор —
    оформление макета, и зон, к которым он «ведёт», нет. Страница — по номеру слайда в IR,
    как у `template.sample_text_left`.
    """
    path = ctx.deck_path
    if path is None:
        raise CheckUnavailable("файла колоды ещё нет: декор примера виден только в .pptx")
    if not ctx.manifest.examples:
        raise CheckUnavailable("в шаблоне нет слайдов-примеров: рецептов не бывает")
    recipes = catalogue(ctx)
    recipe_slides = [
        (number, slide, recipes[slide.recipe_id])
        for number, slide in enumerate(ctx.deck.slides)
        if slide.recipe_id and slide.recipe_id in recipes
    ]
    if not recipe_slides:
        raise CheckUnavailable("ни одного слайда по рецепту из каталога: декора примера нет")
    try:
        from pptx import Presentation

        prs = Presentation(str(path))
        pages = list(prs.slides)
        size = (float(prs.slide_width or 0), float(prs.slide_height or 0))
    except Exception as error:
        raise CheckUnavailable(f"файл колоды не открылся: {type(error).__name__}") from error
    if len(pages) < len(ctx.deck.slides):
        raise CheckUnavailable(
            f"в файле {len(pages)} слайдов, а в IR {len(ctx.deck.slides)}: "
            "страницу слайда по рецепту не найти"
        )
    tolerance = ctx.param("touch_share", 0.01) * size[0]
    line_share = ctx.param("line_share", 0.02)
    for number, slide, recipe in recipe_slides:
        yield from _decor_findings(slide, recipe, pages[number], size, tolerance, line_share)


def _decor_findings(
    slide: SlideIR,
    recipe: Recipe,
    page: Any,
    size: tuple[float, float],
    tolerance: float,
    line_share: float,
) -> Iterable[Finding]:
    placed = list(_shapes_on_page(page.shapes._spTree))
    texts: list[_Frame] = []
    written: set[str] = set()
    for node, frame in placed:
        mark = node.find(f"./*/{qn('p:nvPr')}/{qn('p:ph')}")
        if mark is not None and mark.get("type") in FOOTER_PLACEHOLDERS:
            continue
        if _has_text(node):
            written.add(_shape_id(node))
        elif _shape_id(node) != str(recipe.picture_xml_id):
            continue
        if frame is not None:
            texts.append(frame)
    empty = [
        (zone.zone_id, (float(zone.x), float(zone.y), float(zone.cx), float(zone.cy)))
        for zone in recipe.zones
        if zone.x is not None and zone.y is not None and zone.cx and zone.cy
        and zone.role is not TypeLevel.SLIDE_TITLE and str(zone.xml_id) not in written
    ]
    if not empty:
        return
    addressed = {str(zone.xml_id) for zone in recipe.zones}
    addressed.update(str(xml_id) for row in recipe.repeat_xml_ids for xml_id in row)
    addressed.add(str(recipe.picture_xml_id))

    def near(frame: _Frame, targets: list[_Frame]) -> bool:
        return any(_apart(frame, target) <= tolerance for target in targets)

    for node, frame in placed:
        xml_id = _shape_id(node)
        if (frame is None or _has_text(node) or xml_id in addressed
                or node.find(f"./*/{qn('p:nvPr')}/{qn('p:ph')}") is not None
                or _spans_the_slide(frame, size)):
            continue
        small, big = sorted((frame[2], frame[3]))
        if node.tag == _CONNECTOR or (big > 0 and small <= line_share * big):
            what, ends = "линия", [e for e in _line_ends(node, frame) if not near(e, texts)]
        elif node.tag == _SHAPE and not near(frame, texts):
            what, ends = "плашка", [frame]
        else:
            continue
        zone_id = next(
            (zone_id for end in ends for zone_id, zone in empty if near(end, [zone])), None
        )
        if zone_id is None:
            continue
        props = node.find(f"./*/{qn('p:cNvPr')}")
        name = props.get("name", "") if props is not None else ""
        yield make_finding(
            check_id="template.decor_leads_nowhere",
            slide_id=slide.slide_id,
            reason=f"shape:{xml_id}",
            message=(
                f"На слайде {slide.slide_id} (рецепт {recipe.recipe_id}) {what} «{name}» "
                f"ведёт к зоне {zone_id}, где в файле нет текста"
            ),
            evidence={
                "recipe_id": recipe.recipe_id,
                "xml_id": xml_id or "нет",
                "shape": name,
                "zone_id": zone_id,
            },
        )
